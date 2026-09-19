"""Independent worksheet trace, identity, content and package byte audit."""

import io
import json
import zipfile

from lxml import etree
from openpyxl import load_workbook

from tests.codex_native_pdf.trace import digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def workbook_reads(calls, target):
    buffers, complete, observed = {}, {}, set()
    revisions = [item["sha256"] for item in target["history"]]
    structure_ops = {
        "add_worksheets",
        "rename_worksheet",
        "reorder_worksheets",
        "delete_worksheets",
    }
    mutations = []
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_workbook":
            require(
                args["asset_id"] == target["asset_id"], "Wrong workbook read identity"
            )
            revision = args.get("revision")
            require(
                revision in revisions and revision == result["inspected_revision"],
                "Unpinned workbook structure read",
            )
            key = (
                revision,
                args.get("workbook_view", "structure"),
                result["text_sha256"],
            )
            start, end = result["excerpt_char_range"]
            if start == 0:
                buffers[key] = ""
            require(
                start == args.get("text_offset", 0) == len(buffers.get(key, "")),
                "Noncontiguous workbook read",
            )
            buffers[key] += result["text_excerpt"]
            require(end == len(buffers[key]), "Workbook excerpt range mismatch")
            if result["next_text_offset"] is not None:
                continue
            require(
                digest(buffers[key].encode()) == key[2], "Workbook read hash mismatch"
            )
            record = json.loads(buffers[key])
            require(
                record["asset_id"] == target["asset_id"]
                and record["revision"] == revision,
                "Workbook read changed binding",
            )
            if key[1] == "references":
                require(
                    isinstance(record["references"], list),
                    "Missing reference inventory",
                )
                complete[revision] = record
                observed.add(revision)
        elif args["op"] in structure_ops or (
            args["op"] == "update"
            and any(edit["sheet"] == "Review" for edit in args["edits"])
        ):
            require(
                args["asset_id"] == target["asset_id"]
                and args["expected_revision"] in observed,
                "Mutation precedes complete workbook reference read",
            )
            require(
                args["expected_revision"] == revisions[len(mutations) + 1]
                and result["asset"]["revision"] == revisions[len(mutations) + 2],
                "Mutation differs from the exact revision chain",
            )
            mutations.append(args["op"])
    require(
        mutations
        == [
            "add_worksheets",
            "update",
            "rename_worksheet",
            "reorder_worksheets",
            "delete_worksheets",
        ],
        "Wrong worksheet lifecycle",
    )
    require(set(revisions[1:]) <= complete.keys(), "Missing before/after workbook read")
    return complete


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def validate_workbook_history(calls, target, directory):
    require(len(target["history"]) == 7, "Unexpected worksheet history")
    records = workbook_reads(calls, target)
    expected_orders = [
        ["Sheet1"],
        ["Sheet1"],
        ["Sheet1", "Review", "Temporary"],
        ["Sheet1", "Review", "Temporary"],
        ["資料 O'Brien", "Review", "Temporary"],
        ["Review", "資料 O'Brien", "Temporary"],
        ["Review", "資料 O'Brien"],
    ]
    original_key, table_part, previous_parts = None, None, None
    for index, revision in enumerate(target["history"]):
        data = (directory / "revisions" / revision["sha256"]).read_bytes()
        require(
            digest(data) == revision["sha256"], "Stored workbook revision hash mismatch"
        )
        package = parts(data)
        workbook = load_workbook(io.BytesIO(data))
        try:
            require(
                workbook.sheetnames == expected_orders[index],
                "Wrong native sheet order",
            )
            values = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
            if index:
                values[1][1] = "008"
            name = "Sheet1" if index < 4 else "資料 O'Brien"
            sheet = workbook[name]
            require(
                sheet.max_row == 3 and sheet.max_column == 5,
                "Wrong source table extent",
            )
            for row, expected in zip(sheet.iter_rows(), values, strict=True):
                require(
                    [cell.value for cell in row] == expected
                    and all(cell.data_type == "s" for cell in row),
                    "Original literal table changed",
                )
            if index >= 3:
                review = workbook["Review"]
                require(
                    review["A1"].value == "Current count"
                    and review["B1"].data_type == "f",
                    "Missing native Review formula",
                )
                require(
                    review["B1"].value
                    == ("=Sheet1!B2" if index < 4 else "='資料 O''Brien'!B2"),
                    "Native formula was not rewritten exactly",
                )
            if index:
                record = records[revision["sha256"]]
                require(
                    [item["name"] for item in record["worksheets"]]
                    == workbook.sheetnames,
                    "Tool structure differs from native file",
                )
                current = next(
                    item["key"] for item in record["worksheets"] if item["name"] == name
                )
                original_key = original_key or current
                require(
                    current == original_key,
                    "Renaming/reordering changed sheet identity",
                )
                table_part = table_part or current["part"]
                registry = etree.fromstring(package[record["workbook_part"]])
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                native_sheets = registry.findall("s:sheets/s:sheet", ns)
                require(
                    [(n.get("name"), n.get("sheetId")) for n in native_sheets]
                    == [
                        (item["name"], item["key"]["sheet_id"])
                        for item in record["worksheets"]
                    ],
                    "Native sheetId differs from read-back",
                )
                if index >= 2:
                    require(
                        package[table_part] == previous_parts[table_part],
                        "Structure operation changed original table XML",
                    )
                    unchanged = set(previous_parts) - set(
                        revision["result"]["changed_parts"]
                    )
                    require(
                        all(
                            package.get(path) == previous_parts[path]
                            for path in unchanged
                        ),
                        "Unreported package change",
                    )
        finally:
            workbook.close()
        previous_parts = package
    last = records[target["revision"]]
    require(
        any(item["text"] == "'資料 O''Brien'!B2" for item in last["references"]),
        "Final reference inventory lacks rewritten formula",
    )
    require(
        last["operation_result"]["changes"][0]["secure_erasure"] is False,
        "Missing detached-part disclosure",
    )
