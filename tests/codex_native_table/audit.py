"""Independent actual trace and stored A2T/native snapshot correspondence checks."""

import io
import json
import zipfile

from openpyxl import load_workbook

from tests.codex_native_pdf.artifacts import load_asset
from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def workspace_reads(calls):
    buffers, observed, records, edits, mutations = {}, set(), {}, [], []
    frozen_reads, verified = set(), set()
    for call in calls:
        if call["tool"] == "table_data":
            args = call["arguments"]
            require(args["operation"] == "update_cell", "Unexpected A2T mutation")
            require(
                any(key[0] == args["table_id"] for key in observed),
                "A2T edit precedes complete workspace read",
            )
            edits.append(args)
            continue
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_table_workspace":
            key = (result["table_id"], result["table_sha256"], result["text_sha256"])
            start, end = result["excerpt_char_range"]
            if start == 0:
                buffers[key] = ""
            require(
                start == args.get("text_offset", 0) == len(buffers.get(key, ""))
                and end == start + len(result["text_excerpt"]),
                "Noncontiguous workspace read",
            )
            require(
                args["table_id"] == key[0]
                and (not start or args.get("table_sha256") == key[1]),
                "Workspace read lacks exact hash pin",
            )
            buffers[key] += result["text_excerpt"]
            if result["next_text_offset"] is not None:
                continue
            require(
                digest(buffers[key].encode()) == key[2], "Workspace read hash mismatch"
            )
            record = json.loads(buffers[key])
            require(
                digest(canonical(record["table"])) == key[1],
                "Workspace table hash mismatch",
            )
            observed.add(key[:2])
            records[key[:2]] = record
            if args.get("workspace_reference"):
                frozen_reads.add(canonical(args["workspace_reference"]))
        elif args["op"] == "verify" and result.get("valid"):
            verified.add(canonical(args["reference"]))
        elif args["op"] in {"apply_table_workspace", "create_workbook_from_table"}:
            key = (args["table_id"], args["expected_table_sha256"])
            require(
                key in observed,
                "Native mutation precedes complete pinned workspace read",
            )
            require(
                result["table_sha256"]
                == key[1]
                == result["workspace_reference"]["revision"],
                "Native mutation snapshot hash mismatch",
            )
            mutations.append((args, result))
            if args["op"] == "create_workbook_from_table":
                reference = canonical(args.get("workspace_reference"))
                require(
                    reference in frozen_reads and reference in verified,
                    "Creation precedes complete frozen snapshot read and proof",
                )
    require(
        len(edits) == 1 and len(mutations) == 2,
        "Missing or unexpected bridge mutations",
    )
    return records, edits[0], mutations


def validate_table_workflow(calls, workspace, target):
    records, edit, mutations = workspace_reads(calls)
    require(
        [args["op"] for args, _ in mutations]
        == ["apply_table_workspace", "create_workbook_from_table"],
        "Wrong bridge operation order",
    )
    args, applied = mutations[0]
    require(
        args["asset_id"] == target["asset_id"]
        and args["expected_revision"] == target["history"][0]["sha256"]
        and applied["asset"]["revision"] == target["revision"],
        "Wrong native revision transition",
    )
    record = records[(args["table_id"], args["expected_table_sha256"])]
    table = record["table"]
    binding = table["native_binding"]
    require(
        binding["source"]["asset_id"] == target["asset_id"]
        and binding["source"]["revision"] == target["history"][0]["sha256"],
        "A2T binding advanced or changed",
    )
    require(
        edit["table_id"] == table["id"]
        and edit["row_id"] == table["row_ids"][1]
        and edit["column_name"] == "B"
        and edit["value"] == {"kind": "string", "value": "008"},
        "Wrong typed cell edit",
    )
    original = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
    values = [row.copy() for row in original]
    values[1][1] = "008"
    require(
        table["rows"]
        == [
            dict(
                zip(
                    "ABCDE",
                    [{"kind": "string", "value": value} for value in row],
                    strict=True,
                )
            )
            for row in values
        ],
        "A2T values differ from exact native values",
    )
    require(
        table["row_ids"] == binding["row_ids"]
        and [c["name"] for c in table["columns"]]
        == binding["columns"]
        == list("ABCDE"),
        "Source correspondence changed",
    )
    require(len(record["source_cells"]) == 15, "Incomplete original cell evidence")
    for index, cell in enumerate(record["source_cells"]):
        row, col = divmod(index, 5)
        source = cell["source"]
        evidence = source["evidence"]
        core = {k: v for k, v in source.items() if k != "evidence"}
        require(
            cell["row_id"] == table["row_ids"][row]
            and cell["column_name"] == "ABCDE"[col]
            and source["cell"] == f"{'ABCDE'[col]}{row + 1}"
            and source["kind"] == "string"
            and source["value"] == original[row][col],
            "Incorrect source cell mapping",
        )
        require(
            evidence["asset_id"] == target["asset_id"]
            and evidence["revision"] == binding["source"]["revision"]
            and evidence["value_sha256"] == digest(canonical(core)),
            "Invalid source cell reference",
        )
    for _request, result in mutations:
        ref = result["workspace_reference"]
        asset = load_asset(workspace, ref["asset_id"])
        path = (
            workspace
            / "data"
            / "native-assets"
            / ref["asset_id"]
            / "revisions"
            / ref["revision"]
        )
        require(
            asset["format"] == "a2t"
            and digest(path.read_bytes()) == ref["revision"]
            and json.loads(path.read_bytes()) == table,
            "Frozen A2T snapshot mismatch",
        )
    creation, created = mutations[1]
    require(
        creation["workspace_reference"] == applied["workspace_reference"],
        "Independent creation did not use frozen input",
    )
    independent = load_asset(workspace, created["asset"]["asset_id"])
    require(
        independent["asset_id"] != target["asset_id"]
        and len(independent["history"]) == 1,
        "Independent workbook identity/history mismatch",
    )
    final_bytes = (workspace / "independent.xlsx").read_bytes()
    require(
        digest(final_bytes) == independent["revision"],
        "Independent publication mismatch",
    )
    book = load_workbook(io.BytesIO(final_bytes))
    try:
        require(
            book.sheetnames == ["Data"]
            and [list(row) for row in book["Data"].values] == values
            and all(cell.data_type == "s" for row in book["Data"] for cell in row),
            "Independent literal workbook mismatch",
        )
    finally:
        book.close()
    event = target["history"][-1]["result"]["changes"][-1]
    require(
        event["workspace_reference"] == applied["workspace_reference"]
        and event["edited_cell_count"] == 1,
        "Native history lost A2T lineage",
    )
    directory = workspace / "data" / "native-assets" / target["asset_id"] / "revisions"
    with (
        zipfile.ZipFile(directory / target["history"][0]["sha256"]) as before,
        zipfile.ZipFile(directory / target["revision"]) as after,
    ):
        for name in before.namelist():
            if name not in {
                "xl/worksheets/sheet1.xml",
                "xl/workbook.xml",
                "xl/sharedStrings.xml",
            }:
                require(
                    before.read(name) == after.read(name),
                    "Unrelated native part changed",
                )
