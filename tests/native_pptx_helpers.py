"""Real PresentationML fixtures with text styles, groups, tables, media and notes."""

from __future__ import annotations

import base64
import hashlib
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from src.domain.native_pptx import NativePptxRunLocator, NativePptxTextEdit
from src.infrastructure.native_pptx import NativePresentation
from tests.native_workbook_helpers import _replace


def build_presentation() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shape.name = "Styled text"
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text, run.font.bold, run.font.size = "原始", True, Pt(24)
    run.hyperlink.address = "https://example.com/evidence"
    run = paragraph.add_run()
    run.text, run.font.italic = " unchanged", True
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(3), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "Header"
    table.cell(1, 0).text = "Value"
    table.cell(1, 1).text = "12"
    table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.bold = True
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1))
    child.name, child.text = "Grouped text", "Nested"
    pixel = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
    )
    slide.shapes.add_picture(io.BytesIO(pixel), Inches(7), Inches(1), Inches(1))
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Keep chart", [1, 2])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(5),
        Inches(3),
        Inches(3),
        Inches(2),
        chart_data,
    )
    slide.notes_slide.notes_text_frame.text = "備註 Keep notes"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    ).text = "Other slide"
    output = io.BytesIO()
    presentation.save(output)
    return _replace(
        output.getvalue(),
        {"customXml/preserve.xml": b"<foreign exact='true'> unchanged </foreign>"},
    )


def find_shape(data: bytes, name: str) -> dict:
    return next(
        record
        for record in NativePresentation().iter_shapes(data)
        if record["name"] == name
    )


def edit_run(
    record: dict, text: str, *, paragraph=0, run=0, row=None, column=None
) -> NativePptxTextEdit:
    paragraphs = (
        record["paragraphs"]
        if row is None
        else record["table"]["rows"][row]["cells"][column]["paragraphs"]
    )
    item = next(
        item for item in paragraphs[paragraph]["items"] if item.get("run") == run
    )
    locator = NativePptxRunLocator(
        **record["locator"], paragraph=paragraph, run=run, row=row, column=column
    )
    return NativePptxTextEdit(
        locator=locator,
        expected_text_sha256=hashlib.sha256(item["text"].encode("utf-8")).hexdigest(),
        text=text,
    )
