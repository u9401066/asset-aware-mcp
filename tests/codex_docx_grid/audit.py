"""Audit actual tool order, native packages, source pixels and complete previews."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from tests.codex_docx_structure.audit import validate_docx, validate_wiki
from tests.codex_docx_structure.fonts import environment
from tests.codex_docx_structure.renders import validate_renders
from tests.codex_native_pdf.artifacts import load_asset, read_json, validate_source
from tests.codex_native_pdf.trace import canonical, complete_records, digest, payload
from tests.codex_pdf.trace import call_failed, require, tool_errors
from tests.codex_pptx_tables.audit import validate_images

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def xml(node):
    return etree.tostring(node, method="c14n", exclusive=True)


def calls_from(events):
    require(any(e["type"] == "turn.completed" for e in events), "Incomplete Codex turn")
    calls = []
    for event in events:
        if event["type"] != "item.completed":
            continue
        item = event["item"]
        require(
            item["type"] in {"agent_message", "reasoning", "plan", "mcp_tool_call"},
            "Non-MCP action",
        )
        if item["type"] != "mcp_tool_call":
            continue
        require(
            item["server"] == "asset_aware_under_test"
            and item["tool"] == "document"
            and item["arguments"]["op"] == "native",
            "Unexpected tool/server",
        )
        if not call_failed(item):
            calls.append(item)
    operations = {c["arguments"]["native_request"]["op"] for c in calls}
    required = {
        "contract",
        "register",
        "render_pdf_page",
        "read_pdf_page",
        "create_docx",
        "read_docx",
        "read_docx_table",
        "update_docx_table_grid",
        "render_docx_page",
        "verify",
        "publish",
        "export_wiki",
        "history",
    }
    require(required <= operations, "Missing native grid workflow operations")
    require(
        operations
        <= required
        | {
            "schema",
            "contract_details",
            "inspect",
            "read_pdf",
            "list",
            "read_docx_block",
        },
        "Unexpected mutation/workflow",
    )
    return calls


def validate_receipt(record, target, revision):
    history = next(
        stage for stage in reversed(target["history"]) if stage["sha256"] == revision
    )
    require(
        "operation_result" in record
        and record["operation_result"] == history.get("result"),
        "Full table receipt differs from latest matching history",
    )


def validate_reads(calls, workspace, target, source):
    buffers, completed, tables, refs, proofs = {}, set(), {}, set(), set()
    page_records = complete_records(calls)
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        op = args["op"]
        if op == "read_docx":
            refs.update(canonical(b["evidence"]) for b in result["blocks"])
        if op in {"read_docx", "read_docx_table"}:
            revision = result["inspected_revision"]
            require(args.get("revision") == revision, "Unpinned Word read")
            section = result["dfm" if op == "read_docx" else "table"]
            key = (op, args["asset_id"], revision, section["text_sha256"])
            start, end = section["excerpt_char_range"]
            require(
                start == args.get("text_offset", 0)
                and end - start == len(section["text_excerpt"]),
                "Read range mismatch",
            )
            if start == 0:
                buffers[key] = ""
            require(len(buffers.get(key, "")) == start, "Missing Word read chunk")
            buffers[key] += section["text_excerpt"]
            if section["next_text_offset"] is not None:
                require(section["next_text_offset"] == end, "Bad Word continuation")
                continue
            text = buffers.pop(key)
            require(
                len(text) == section["text_length"]
                and digest(text.encode()) == key[-1],
                "Word read hash mismatch",
            )
            completed.add(key[:3])
            if op == "read_docx_table":
                record = json.loads(text)
                validate_receipt(record, target, revision)
                require(
                    record["evidence"] == args["docx_table_reference"],
                    "Table reference override",
                )
                reference = canonical(record["evidence"])
                require(reference in refs, "Table read lacks actual block reference")
                path = (
                    workspace
                    / "data/native-assets"
                    / target["asset_id"]
                    / "revisions"
                    / revision
                )
                with ZipFile(path) as archive:
                    native = etree.fromstring(archive.read("word/document.xml")).find(
                        W + "body/" + W + "tbl"
                    )
                require(
                    xml(native) == xml(etree.fromstring(record["native_xml"].encode())),
                    "Full table XML differs from pinned package",
                )
                require(
                    record["rows"] == 4 and record["columns"] == 5,
                    "Table read dimensions differ",
                )
                tables[reference] = record
        if op == "update_docx_table_grid":
            ref = canonical(args["docx_table_grid"]["reference"])
            require(ref in tables, "Grid mutation precedes full table read")
            require(
                ("read_docx", args["asset_id"], args["expected_revision"]) in completed,
                "Grid mutation precedes full DFM read",
            )
            refs.add(canonical(result["review_request"]["docx_table_reference"]))
        if op == "verify":
            ref = canonical(args["reference"])
            require(
                result.get("valid") is True and (ref in refs or ref in page_records),
                "Invalid/unseen proof",
            )
            if args["reference"]["asset_id"] == source["asset_id"]:
                proofs.add("source")
            elif args["reference"]["asset_id"] == target["asset_id"]:
                require(
                    not result["is_current_managed_revision"],
                    "Expected historical Word proof",
                )
                proofs.add("historical")
    require(
        not buffers and proofs == {"source", "historical"}, "Incomplete reads/proofs"
    )
    for stage in target["history"]:
        for op in ("read_docx", "read_docx_table"):
            require(
                (op, target["asset_id"], stage["sha256"]) in completed,
                "Missing complete revision read",
            )
    return len(tables)


def validate_history(workspace, target):
    history = target["history"]
    require(
        len(history) == 3 and target["source"] is None,
        "Expected create/grid/restore history",
    )
    root = workspace / "data/native-assets" / target["asset_id"] / "revisions"
    initial = root / history[0]["sha256"]
    require(validate_docx(initial) == ("007", False), "Initial transcription differs")
    with ZipFile(initial) as archive:
        original = {name: archive.read(name) for name in archive.namelist()}
    baseline = Document(initial)
    texts = [[c.text for c in row.cells] for row in baseline.tables[0].rows]
    paragraphs = [
        xml(p._p)
        for row in baseline.tables[0].rows
        for c in row.cells
        for p in c.paragraphs
    ]
    seen = set()
    for i, stage in enumerate(history[1:], 1):
        path = root / stage["sha256"]
        document = Document(path)
        table = document.tables[0]
        require(
            [[c.text for c in row.cells] for row in table.rows] == texts,
            "Native text/merge drift",
        )
        require(
            [xml(p._p) for row in table.rows for c in row.cells for p in c.paragraphs]
            == paragraphs,
            "Rich paragraph formatting drift",
        )
        expected_widths = [1500, 1700, 1500, 1300, 1500] if i == 1 else [1500] * 5
        require(
            [c.width.twips for c in table.columns] == expected_widths,
            "Grid width request not applied",
        )
        with ZipFile(path) as archive:
            require(set(archive.namelist()) == set(original), "Package inventory drift")
            require(
                all(
                    archive.read(k) == v
                    for k, v in original.items()
                    if k != "word/document.xml"
                ),
                "Unrelated package part drift",
            )
        changes = stage["result"]["changes"]
        require(
            len(changes) == 1 and changes[0]["operation"] == "update_docx_table_grid",
            "Wrong history operation",
        )
        seen.update(edit["op"] for edit in changes[0]["edits"])
        if i == 1:
            require(
                any(e.get("moved_blocks", 0) > 0 for e in changes[0]["edits"]),
                "Merge did not move actual content",
            )
    require(
        seen == {"insert", "delete", "resize", "merge", "split"},
        "Grid CRUD coverage incomplete",
    )
    require(
        validate_docx(root / target["revision"]) == ("007", False),
        "Final transcription/format differs",
    )
    return history[1]["sha256"]


def audit(output):
    expected, final = (
        read_json(output / "expected.json"),
        read_json(output / "last-message.txt"),
    )
    workspace = output / "workspace"
    events = [
        json.loads(line)
        for line in (output / "events.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    calls = calls_from(events)
    source, target = (
        load_asset(workspace, final["source_asset_id"]),
        load_asset(workspace, final["docx_asset_id"]),
    )
    validate_source(workspace, expected, source)
    images = validate_images(calls, workspace, source)
    tables = validate_reads(calls, workspace, target, source)
    intermediate = validate_history(workspace, target)
    require(
        digest((workspace / "verified.docx").read_bytes()) == target["revision"],
        "Published file differs",
    )
    validate_wiki(workspace, target)
    require(
        isinstance(final.get("limitations"), list) and final["limitations"],
        "Missing limitations",
    )
    with environment(expected["font_environment"]):
        renders = validate_renders(
            workspace, target, calls, final, {intermediate}, require_cjk=True
        )
    return {
        "passed": True,
        "tool_calls": len(calls),
        "tool_errors": tool_errors(events),
        "source_images": images,
        "complete_table_revisions": tables,
        "managed_revisions": len(target["history"]),
        "renders": renders,
        "limitations": final["limitations"],
    }


def write_audit(output):
    try:
        result = audit(output)
    except (ValueError, KeyError, OSError, TypeError, AssertionError) as exc:
        result = {"passed": False, "error": str(exc)}
    (output / "audit.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    result = write_audit(parser.parse_args().output)
    print(json.dumps(result, ensure_ascii=False))
    raise SystemExit(0 if result["passed"] else 1)
