"""A supplied Word table with deliberately clipped rows and missing repeat headers."""

from __future__ import annotations

import io
import os
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

HEADERS = ["Study 007 µg", "Sample / Reading"]
ROWS = 14


def row_lines(index):
    return [
        f"ROW {index:02d}",
        "Count 007",
        "Reading -0.50 mg/L",
        f"END {index:02d} / 1,234.50",
    ]


def clipped_document():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Twips(8500), Twips(7200)
    section.top_margin = section.bottom_margin = Twips(360)
    section.left_margin = section.right_margin = Twips(500)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(12)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    table = document.add_table(rows=ROWS + 2, cols=2)
    table.autofit = False
    for column in table.columns:
        column.width = Twips(3750)
    for row, title in zip(table.rows[:2], HEADERS, strict=True):
        cell = row.cells[0].merge(row.cells[1])
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("1122AA")
        row.height_rule, row.height = WD_ROW_HEIGHT_RULE.AT_LEAST, Pt(22)
    for i, row in enumerate(table.rows[2:], 1):
        row.height_rule, row.height = WD_ROW_HEIGHT_RULE.EXACTLY, Pt(13)
        for j, cell in enumerate(row.cells):
            run = cell.paragraphs[0].add_run(
                "\n".join(row_lines(i))
                if j == 0
                else f"Sample {i:02d}\nµg\nValue\nCONFIRMED {i:02d}"
            )
            run.font.name, run.font.size = "Arial", Pt(12)
            if j == 0:
                run.font.color.rgb = RGBColor.from_string("008000")
            for border in ["top", "left", "bottom", "right"]:
                pr = cell._tc.get_or_add_tcPr()
                edges = pr.find(qn("w:tcBorders"))
                if edges is None:
                    edges = OxmlElement("w:tcBorders")
                    pr.append(edges)
                side = OxmlElement("w:" + border)
                side.set(qn("w:val"), "single")
                side.set(qn("w:sz"), "4")
                edges.append(side)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def independent_pdf(data):
    binary = os.environ.get("LIBREOFFICE_BIN") or "libreoffice"
    with tempfile.TemporaryDirectory(prefix="word-layout-audit-") as temp:
        root = Path(temp)
        source = root / "independent.docx"
        source.write_bytes(data)
        profile = root / "profile"
        (profile / "user").mkdir(parents=True)
        (profile / "user/registrymodifications.xcu").write_text(
            "<?xml version='1.0'?><oor:items xmlns:oor='http://openoffice.org/2001/registry'><item oor:path='/org.openoffice.Office.Common/Misc'><prop oor:name='UseSystemFileDialog' oor:op='fuse'><value>false</value></prop></item></oor:items>",
            encoding="utf-8",
        )
        subprocess.run(
            [
                binary,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf:writer_pdf_Export",
                "--outdir",
                temp,
                str(source),
            ],
            check=True,
            timeout=60,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        assert source.read_bytes() == data
        return (root / "independent.pdf").read_bytes()


def inspect_pages(data, corrected):
    pdf_bytes = independent_pdf(data)
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as pdf:
        texts = [page.get_text() for page in pdf]
        if corrected:
            assert 2 <= len(pdf) <= 6
            for text in texts:
                assert all(header in text for header in HEADERS)
            for i in range(1, ROWS + 1):
                owners = [j for j, text in enumerate(texts) if f"ROW {i:02d}" in text]
                assert len(owners) == 1
                text = texts[owners[0]]
                assert all(line in text for line in row_lines(i))
                assert f"CONFIRMED {i:02d}" in text
        else:
            assert len(pdf) == 1
            assert not all(f"END {i:02d}" in texts[0] for i in range(1, ROWS + 1))
        return pdf_bytes, texts
