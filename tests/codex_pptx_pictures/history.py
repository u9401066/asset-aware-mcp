"""Independent package and shape checks for the fixed live picture scenario."""

from __future__ import annotations

import zipfile

from lxml import etree
from pptx import Presentation

from tests.codex_native_pdf.trace import canonical, digest
from tests.codex_pdf.trace import require

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"p": P_NS, "a": A_NS}


def xml(node):
    return etree.tostring(node, method="c14n")


def validate_stage(source, path, expected, first_mapping=None):
    baseline, actual = Presentation(source), Presentation(path)
    require(len(actual.slides) == len(baseline.slides), "Slide inventory changed")
    for index, (before, after) in enumerate(
        zip(baseline.slides, actual.slides, strict=True)
    ):
        extra = len(expected) if index == 0 else 0
        require(len(after.shapes) == len(before.shapes) + extra, "Wrong shape count")
        for old, new in zip(before.shapes, after.shapes, strict=False):
            require(xml(old.element) == xml(new.element), "Original shape changed")
    added = list(actual.slides[0].shapes)[len(baseline.slides[0].shapes) :]
    require([s.image.blob for s in added] == expected, "Wrong picture bytes or order")
    if first_mapping is not None:
        old = etree.fromstring(first_mapping)
        new = etree.fromstring(xml(added[0].element))
        new.find("p:blipFill/a:blip", NS).set(
            f"{{{R_NS}}}embed",
            old.find("p:blipFill/a:blip", NS).get(f"{{{R_NS}}}embed"),
        )
        require(xml(old) == xml(new), "Replacement changed picture mapping")
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path) as output:
        allowed = {
            "[Content_Types].xml",
            "ppt/slides/slide1.xml",
            "ppt/slides/_rels/slide1.xml.rels",
        }
        for name in original.namelist():
            if name not in allowed:
                require(
                    original.read(name) == output.read(name), "Untouched part changed"
                )
    return xml(added[0].element)


def validate_history(workspace, deck, child, image_assets, records):
    root = workspace / "data" / "native-assets" / deck["asset_id"] / "revisions"
    require(len(deck["history"]) == 4, "Unexpected deck history")
    require(
        deck["history"][0]["sha256"]
        == digest((workspace / "source.pptx").read_bytes()),
        "Wrong initial deck",
    )
    original, replacement = [
        (workspace / n).read_bytes() for n in ("original.png", "replacement.png")
    ]
    stages, mapping = {}, None
    for revision, (operation, expected) in zip(
        deck["history"][1:],
        [
            ("add_picture", [original, original]),
            ("replace_picture", [replacement, original]),
            ("delete_shape", [replacement]),
        ],
        strict=True,
    ):
        changes = revision["result"]["changes"]
        require(
            changes and all(c["operation"] == operation for c in changes),
            "Wrong history operation",
        )
        stages[operation] = revision
        current = validate_stage(
            workspace / "source.pptx", root / revision["sha256"], expected, mapping
        )
        mapping = mapping or current
    validate_image_lineage(stages, image_assets)
    validate_extracted(workspace, child, original, records)
    require(
        digest((workspace / "verified.pptx").read_bytes()) == deck["revision"],
        "Published PPTX differs",
    )


def validate_image_lineage(stages, image_assets):
    copies = stages["add_picture"]["result"]["changes"]
    require(
        len(copies) == 2 and copies[0]["media_part"] == copies[1]["media_part"],
        "Initial pictures must share media",
    )
    for operation, asset in [
        ("add_picture", image_assets[0]),
        ("replace_picture", image_assets[1]),
    ]:
        for change in stages[operation]["result"]["changes"]:
            ref = change["source_reference"]
            require(
                (ref["asset_id"], ref["revision"])
                == (asset["asset_id"], asset["revision"]),
                "Image creation lineage mismatch",
            )


def validate_extracted(workspace, child, original, records):
    require(len(child["history"]) == 1, "Extracted image was mutated")
    origin = child["history"][0]["result"]["changes"][0]
    require(
        origin["operation"] == "extract_picture"
        and origin["media_sha256"] == digest(original),
        "Extraction lineage mismatch",
    )
    require(
        canonical(origin["source_reference"]) in records,
        "Extraction lacks full shape readback",
    )
    require(
        (workspace / "extracted.png").read_bytes() == original
        and child["revision"] == digest(original),
        "Extracted bytes differ",
    )


def validate_references(calls, records, deck):
    from tests.codex_native_pdf.trace import payload

    available, historical = set(), False
    for call in calls:
        args = call["arguments"]["native_request"]
        if args["op"] == "read_pptx_shape":
            shape = payload(call)["shape"]
            key = canonical(shape["evidence"])
            if shape["next_text_offset"] is None and key in records:
                available.add(key)
        refs = args.get("pptx_shape_refs", []) + [
            e["reference"] for e in args.get("pptx_picture_edits", [])
        ]
        require(
            all(canonical(ref) in available for ref in refs),
            "Mutation lacks prior complete shape readback",
        )
        if args["op"] == "verify":
            ref, proof = args["reference"], payload(call)
            historical |= (
                canonical(ref) in available
                and ref["asset_id"] == deck["asset_id"]
                and proof.get("valid") is True
                and proof.get("is_current_managed_revision") is False
            )
    require(historical, "Historical deck shape proof missing")
