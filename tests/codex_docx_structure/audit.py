"""Independent audit of scanned pixels, Word structures, revisions and wiki bytes."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from tests.codex_docx_structure.trace import calls_from, validate_reads
from tests.codex_native_pdf.artifacts import load_asset, read_json, validate_source
from tests.codex_native_pdf.trace import canonical, complete_records, digest
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require, tool_errors
from tests.codex_pptx_tables.audit import validate_images


def validate_docx(path):
    document = Document(path)
    require(len(document.tables) in {1, 2}, "Wrong native table count")
    require(document.paragraphs[0].text == "研究 007 µg", "Heading text differs")
    heading = document.paragraphs[0].runs[0]
    require(
        heading.bold
        and heading.font.name == "Arial"
        and heading.font.size.pt == 11.5
        and str(heading.font.color.rgb) == "1122AA",
        "Heading format differs",
    )
    table = document.tables[0]
    require(len(table.rows) == 4 and len(table.columns) == 5, "Grid size differs")
    require(
        [c.width.twips for c in table.columns] == [1500] * 5, "Column widths differ"
    )
    require([r.height.twips for r in table.rows] == [300] * 4, "Row heights differ")
    require(table.autofit is False, "Table is not fixed layout")
    require(
        table.rows[0]._tr.find("./" + qn("w:trPr") + "/" + qn("w:tblHeader"))
        is not None,
        "Repeated header missing",
    )
    require(
        all(table.cell(0, c)._tc is table.cell(0, 0)._tc for c in range(5)),
        "Title merge missing",
    )
    expected = [["Scanned inventory"] * 5, list(COLUMNS), *map(list, PAGE_ROWS[0])]
    actual = [[cell.text for cell in row.cells] for row in table.rows]
    require(actual[2][1] in {"007", "008"}, "Leading zero or edit lost")
    expected[2][1] = actual[2][1]
    require(actual == expected, "Exact scanned transcription differs")
    for row in table.rows:
        for cell in row.cells:
            require(
                all(
                    run.font.size.pt == 11.5 and run.font.name == "Arial"
                    for p in cell.paragraphs
                    for run in p.runs
                ),
                "Cell run formatting differs",
            )
    temporary = len(document.tables) == 2
    require(
        [p.text for p in document.paragraphs]
        == (["研究 007 µg", "Temporary 007"] if temporary else ["研究 007 µg"]),
        "Unexpected body paragraph",
    )
    if temporary:
        require(
            document.tables[1].cell(0, 0).text == "Discard me", "Temporary table lost"
        )
    return actual[2][1], temporary


def validate_history(workspace, asset):
    root = workspace / "data" / "native-assets" / asset["asset_id"] / "revisions"
    require(asset["source"] is None and not asset["archived"], "Not independent DOCX")
    history = asset["history"]
    require(len(history) >= 5, "Missing DOCX history stages")
    states, previous = [], None
    for stage in history:
        path = root / stage["sha256"]
        states.append(validate_docx(path))
        with ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        if previous:
            changes = (stage.get("result") or {}).get("changes", [])
            if any(
                c.get("operation") in {"insert_blocks", "delete_blocks"}
                for c in changes
            ):
                require(
                    parts.keys() == previous.keys(), "Structural part inventory drift"
                )
                require(
                    all(
                        parts[k] == v
                        for k, v in previous.items()
                        if k != "word/document.xml"
                    ),
                    "Untouched part bytes drift",
                )
        previous = parts
    require(states[0] == states[-1] == ("007", False), "Initial/final table differs")
    require(
        ("008", False) in states and ("007", True) in states,
        "Temporary text edit or insertion missing",
    )
    return len(states)


def validate_wiki(workspace, asset):
    manifests = [
        p
        for p in (workspace / "native-wiki").rglob("manifest.json")
        if read_json(p).get("asset_id") == asset["asset_id"]
    ]
    require(len(manifests) == 1, "Expected one DOCX wiki snapshot")
    root = manifests[0].parent
    manifest = read_json(manifests[0])
    require(manifest["projection"] == "docx-blocks-v1", "Wrong wiki projection")
    require(
        (manifest["asset_id"], manifest["revision"])
        == (asset["asset_id"], asset["revision"]),
        "Wiki identity differs",
    )
    for name, info in manifest["files"].items():
        require(Path(name).name == name, "Unsafe wiki filename")
        raw = (root / name).read_bytes()
        require(
            digest(raw) == info["sha256"] and len(raw) == info["size_bytes"],
            "Wiki artifact bytes differ",
        )
    require(
        (root / manifest["source_attachment"]).read_bytes()
        == (workspace / "verified.docx").read_bytes(),
        "Wiki source bytes differ",
    )
    records = [
        json.loads(line)
        for line in (root / "records.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    require(len(records) == 2, "Wiki does not cover final heading and table")
    for record in records:
        core = {
            k: record[k]
            for k in ("schema_version", "block_id", "kind", "locator", "representation")
        }
        require(
            digest(canonical(core)) == record["evidence"]["value_sha256"],
            "Wiki block hash differs",
        )
        require(
            "[[" in (root / record["note"]).read_text(encoding="utf-8"),
            "Wiki links absent",
        )
    with ZipFile(workspace / "verified.docx") as archive:
        require(
            set(manifest["part_attachments"]) == set(archive.namelist()),
            "Wiki part coverage differs",
        )
        for name, attachment in manifest["part_attachments"].items():
            require(
                (root / attachment["attachment"]).read_bytes() == archive.read(name),
                "Wiki original part differs",
            )


def audit(output):
    expected = read_json(output / "expected.json")
    final = read_json(output / "last-message.txt")
    workspace = output / "workspace"
    events = [
        json.loads(line)
        for line in (output / "events.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    calls = calls_from(events)
    source = load_asset(workspace, final["source_asset_id"])
    target = load_asset(workspace, final["docx_asset_id"])
    validate_source(workspace, expected, source)
    images = validate_images(calls, workspace, source)
    reads = validate_reads(calls, target, source, complete_records(calls))
    revisions = validate_history(workspace, target)
    published = workspace / "verified.docx"
    require(
        digest(published.read_bytes()) == target["revision"], "Published hash differs"
    )
    require(validate_docx(published) == ("007", False), "Published content differs")
    validate_wiki(workspace, target)
    require(
        isinstance(final.get("limitations"), list) and final["limitations"],
        "Review limitations missing",
    )
    return {
        "passed": True,
        "tool_calls": len(calls),
        "tool_errors": tool_errors(events),
        "source_images": images,
        "complete_dfm_revisions": reads,
        "managed_revisions": revisions,
        "limitations": final["limitations"],
    }


def write_audit(output):
    try:
        result = audit(output)
    except (ValueError, KeyError, OSError, TypeError, AssertionError) as exc:
        result = {"passed": False, "error": str(exc)}
    (output / "audit.json").write_text(
        json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    options = parser.parse_args()
    result = write_audit(options.output)
    print(json.dumps(result, ensure_ascii=False))
    raise SystemExit(0 if result["passed"] else 1)
