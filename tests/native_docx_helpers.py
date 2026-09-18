"""Real Word fixtures for native version/DFM integration regressions."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from PIL import Image

from src.application.native_document_service import NativeDocumentService
from src.application.native_docx_bridge import NativeDocxBridge
from src.domain.native_assets import NativeDocumentRequest
from src.infrastructure.native_asset_store import FileNativeAssetRepository
from src.infrastructure.native_docx_workspace import FileNativeDocxWorkspaces
from src.infrastructure.native_spreadsheet import SpreadsheetFileAdapter


def build_docx() -> bytes:
    document = Document()
    document.add_heading("研究文件", 1)
    paragraph = document.add_paragraph()
    paragraph.add_run("原始段落").bold = True
    paragraph.add_run(" with italic context").italic = True
    document.add_paragraph("Unchanged paragraph")
    table = document.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Drug"
    table.cell(1, 1).text = "Old value"
    document.sections[0].header.paragraphs[0].text = "保留頁首"
    document.sections[0].footer.paragraphs[0].text = "保留頁尾"
    picture = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(picture, format="PNG")
    picture.seek(0)
    document.add_picture(picture)
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def replace_parts(data: bytes, replacements: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w") as target,
    ):
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info)))
        for name in replacements.keys() - set(source.namelist()):
            target.writestr(name, replacements[name])
    return output.getvalue()


def call(service: NativeDocumentService, **fields: Any) -> dict[str, Any]:
    return service.execute(NativeDocumentRequest.model_validate(fields))


def read_dfm(service: NativeDocumentService, asset: dict[str, Any]) -> str:
    assert service.docx is not None
    return service.docx.read(
        service.repository.read(asset["asset_id"], asset["revision"]),
        asset["asset_id"],
        asset["revision"],
    ).dfm_text


@pytest.fixture
def native_docx(tmp_path: Path) -> tuple[NativeDocumentService, dict[str, Any], Path]:
    source = tmp_path / "研究.docx"
    source.write_bytes(build_docx())
    service = NativeDocumentService(
        FileNativeAssetRepository(tmp_path / "native"),
        SpreadsheetFileAdapter(),
        docx=NativeDocxBridge(FileNativeDocxWorkspaces()),
    )
    asset = call(service, op="register", source_path=str(source))["asset"]
    return service, asset, source
