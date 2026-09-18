"""Actual SDK2 image asset -> picture CRUD -> extraction/evidence/wiki/writeback."""

from __future__ import annotations

import base64
import hashlib
import json
import os
import sys
from pathlib import Path

import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation

from tests.integration.test_native_pptx_shape_stdio import native, read_complete
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_picture_helpers import raster, request


async def preview(client, asset, locator, expected_bytes):
    result = await client.call_tool(
        "document",
        {
            "op": "native",
            "native_request": {
                "op": "read_pptx_picture",
                "asset_id": asset["asset_id"],
                "revision": asset["revision"],
                "pptx_locator": locator,
                "render_size": 512,
            },
        },
    )
    metadata = json.loads("".join(b.text for b in result.content if b.type == "text"))
    images = [b for b in result.content if b.type == "image"]
    assert len(images) == 1 and images[0].mime_type == "image/png"
    raw = base64.b64decode(images[0].data, validate=True)
    assert hashlib.sha256(raw).hexdigest() == metadata["image_sha256"]
    assert metadata["image"]["sha256"] == hashlib.sha256(expected_bytes).hexdigest()
    assert "embedded_raster_only" in metadata["preview_scope"]
    return metadata


async def extract_and_verify(client, current, locator, image_asset):
    extracted = await native(
        client,
        op="extract_pptx_picture",
        asset_id=current["asset_id"],
        revision=current["revision"],
        pptx_locator=locator,
    )
    child = extracted["asset"]
    assert child["revision"] == image_asset["revision"]
    proof = await native(client, op="verify", reference=child["file_reference"])
    assert proof["valid"] and proof["verification_scope"] == "immutable_file_bytes"


async def exercise_pictures(
    client, asset, image_asset, replacement, data, image, new_image
):
    item = request(data, image).model_dump()
    item["image"] = image_asset["file_reference"]
    added = await native(
        client,
        op="add_pptx_pictures",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        pptx_pictures=[item, item],
    )
    assert added["success"] and not added["source_written"]
    current = added["asset"]
    locators = [c["locator"] for c in added["operation_result"]["changes"]]
    first = await read_complete(client, current, locators[0])
    await preview(client, current, locators[0], image)
    await extract_and_verify(client, current, locators[0], image_asset)
    replaced = await native(
        client,
        op="replace_pptx_pictures",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        pptx_picture_edits=[
            {"reference": first["evidence"], "image": replacement["file_reference"]}
        ],
    )
    current = replaced["asset"]
    await preview(client, current, locators[0], new_image)
    await preview(client, current, locators[1], image)
    proof = await native(client, op="verify", reference=first["evidence"])
    assert proof["valid"] and not proof["is_current_managed_revision"]
    second = await read_complete(client, current, locators[1])
    deleted = await native(
        client,
        op="delete_pptx_shapes",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        pptx_shape_refs=[second["evidence"]],
    )
    assert deleted["success"]
    return deleted["asset"]


async def publish_and_writeback(client, current, original, source, target):
    published = await native(
        client,
        op="publish",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        output_path=str(target),
    )
    assert published["success"] and target.is_file()
    wiki = await native(
        client,
        op="export_wiki",
        asset_id=current["asset_id"],
        output_dir=str(target.parent / "wiki"),
    )
    manifest = json.loads((Path(wiki["output_dir"]) / "manifest.json").read_text())
    assert manifest["projection"] == "pptx-shapes-v1"
    old = source.read_bytes()
    written = await native(
        client,
        op="writeback",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        expected_source_sha256=original["source"]["sha256"],
    )
    assert Path(written["backup_path"]).read_bytes() == old
    assert source.read_bytes() == target.read_bytes()


@pytest.mark.timeout(120)
async def test_picture_asset_crud_over_sdk2(tmp_path):
    source, image_path, replacement_path = [
        tmp_path / name for name in ("source.pptx", "image.png", "replacement.jpg")
    ]
    data, image, new_image = build_presentation(), raster(), raster("JPEG", "blue")
    for path, raw in [
        (source, data),
        (image_path, image),
        (replacement_path, new_image),
    ]:
        path.write_bytes(raw)
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
            "ASSET_AWARE_MCP_TOOL_SURFACE": "balanced",
        },
    )
    async with Client(stdio_client(params)) as client:
        contract = await native(client, op="contract", for_op="add_pptx_pictures")
        assert "read_pptx_picture" in contract["formats"]["pptx"]
        assets = [
            (await native(client, op="register", source_path=str(path)))["asset"]
            for path in (source, image_path, replacement_path)
        ]
        current = await exercise_pictures(client, *assets, data, image, new_image)
        assert source.read_bytes() == data
        target = tmp_path / "published.pptx"
        await publish_and_writeback(client, current, assets[0], source, target)
        shape = Presentation(target).slides[0].shapes[-1]
        assert shape.image.blob == new_image
        assert (
            image_path.read_bytes() == image
            and replacement_path.read_bytes() == new_image
        )
