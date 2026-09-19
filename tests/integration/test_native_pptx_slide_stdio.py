"""SDK2 native slide CRUD, historical shapes, paged layouts and guarded writeback."""

import os
import sys

import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation

from tests.integration.test_native_pptx_picture_stdio import publish_and_writeback
from tests.integration.test_native_pptx_shape_stdio import native, read_complete
from tests.native_pptx_helpers import build_presentation


@pytest.mark.timeout(120)
async def test_slide_structure_evidence_and_writeback_sdk2(tmp_path):
    data = build_presentation()
    source = tmp_path / "source.pptx"
    source.write_bytes(data)
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
            "ASSET_AWARE_MCP_TOOL_SURFACE": "balanced",
        },
    )
    async with Client(stdio_client(params)) as client:
        contract = await native(client, op="contract", for_op="add_pptx_slides")
        assert "read_pptx_layouts" in contract["formats"]["pptx"]
        asset = (await native(client, op="register", source_path=str(source)))["asset"]
        listing = await native(
            client,
            op="read_pptx",
            asset_id=asset["asset_id"],
            revision=asset["revision"],
        )
        old = await read_complete(client, asset, listing["shapes"][0]["locator"])
        layouts, offset = [], 0
        while True:
            result = await native(
                client,
                op="read_pptx_layouts",
                asset_id=asset["asset_id"],
                revision=asset["revision"],
                offset=offset,
                limit=3,
            )
            assert result["inspected_revision"] == asset["revision"]
            layouts.extend(result["layouts"])
            offset = result["next_offset"]
            if offset is None:
                break
        assert len(layouts) == len({item["part"] for item in layouts}) == 11
        layout = next(item for item in layouts if item["type"] == "blank")
        args = {
            "op": "add_pptx_slides",
            "asset_id": asset["asset_id"],
            "expected_revision": asset["revision"],
            "pptx_slide_insert": {
                "index": 1,
                "slides": [
                    {
                        "layout_part": layout["part"],
                        "textboxes": [
                            {
                                "paragraphs": [
                                    [{"text": "SDK2 new slide 007 µg", "bold": True}]
                                ]
                            }
                        ],
                    }
                ],
            },
        }
        added = await native(client, **args)
        assert added["success"] and not added["source_written"]
        assert not (await native(client, **args))["success"]
        current = added["asset"]
        listing = await native(client, **added["review_request"])
        assert listing["metadata"]["slide_count"] == 3
        new_key = added["operation_result"]["changes"][0]["slides"][0]
        new_locator = next(
            s["locator"]
            for s in listing["shapes"]
            if s["locator"]["slide_id"] == new_key["slide_id"]
        )
        new_record = await read_complete(client, current, new_locator)
        order = list(reversed(listing["slides"]))
        reordered = await native(
            client,
            op="reorder_pptx_slides",
            asset_id=asset["asset_id"],
            expected_revision=current["revision"],
            pptx_slide_order=order,
        )
        assert reordered["success"]
        assert (await native(client, **reordered["review_request"]))["slides"] == order
        deleted = await native(
            client,
            op="delete_pptx_slides",
            asset_id=asset["asset_id"],
            expected_revision=reordered["asset"]["revision"],
            pptx_slide_keys=[new_key],
        )
        assert deleted["success"]
        remaining = await native(client, **deleted["review_request"])
        assert remaining["metadata"]["slide_count"] == 2
        assert all(s["slide_id"] != new_key["slide_id"] for s in remaining["slides"])
        for record in (old, new_record):
            proof = await native(client, op="verify", reference=record["evidence"])
            assert proof["valid"] and not proof["is_current_managed_revision"]
        assert source.read_bytes() == data
        target = tmp_path / "published.pptx"
        await publish_and_writeback(client, deleted["asset"], asset, source, target)
    deck = Presentation(target)
    assert len(deck.slides) == 2
    assert deck.slides[0].shapes[0].text == "Other slide"
    assert deck.slides[1].notes_slide.notes_text_frame.text == "備註 Keep notes"
