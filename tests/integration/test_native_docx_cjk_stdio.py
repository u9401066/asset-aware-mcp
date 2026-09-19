"""Same DOCX bytes before/after CJK fallback correction over actual SDK2 images."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client
from PIL import Image

from tests.codex_docx_structure.fonts import environment, snapshot
from tests.codex_docx_structure.renders import replay
from tests.integration.test_native_docx_stdio_e2e import _native


@pytest.mark.skipif(
    not os.environ.get("NATIVE_DOCX_FONT_FIXTURE"),
    reason="Set NATIVE_DOCX_FONT_FIXTURE to a prepared private Linux fixture",
)
@pytest.mark.timeout(150)
async def test_same_source_cjk_correction_over_sdk2(tmp_path):
    fonts = snapshot(Path(os.environ["NATIVE_DOCX_FONT_FIXTURE"]))
    source = tmp_path / "source.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("研究 007 µg")
    run.font.name, run.font.size, run.bold = "Arial", Pt(11.5), True
    doc.save(source)
    original, mtime = source.read_bytes(), source.stat().st_mtime_ns
    images = []
    for enabled in (False, True):
        with environment(fonts, cjk=enabled):
            params = StdioServerParameters(
                command=sys.executable,
                args=["-m", "src.server"],
                env={
                    **os.environ,
                    "DATA_DIR": str(tmp_path / str(enabled)),
                    "ENABLE_LIGHTRAG": "false",
                    "ASSET_AWARE_DISABLE_DOTENV": "true",
                },
            )
            async with Client(stdio_client(params)) as client:
                asset = (await _native(client, op="register", source_path=str(source)))[
                    "asset"
                ]
                response = await client.call_tool(
                    "document",
                    {
                        "op": "native",
                        "native_request": {
                            "op": "render_docx_page",
                            "asset_id": asset["asset_id"],
                            "revision": asset["revision"],
                            "docx_page_index": 0,
                            "render_size": 1400,
                        },
                    },
                )
                result = json.loads(
                    "".join(
                        item.text for item in response.content if item.type == "text"
                    )
                )
                result = result.get("result", result)
                assert result["success"], result
                assert (
                    result["inspected_revision"] == hashlib.sha256(original).hexdigest()
                )
                blocks = [item for item in response.content if item.type == "image"]
                assert len(blocks) == 1
                png = base64.b64decode(blocks[0].data)
                assert hashlib.sha256(png).hexdigest() == result["image_sha256"]
                count, width, height, pixels = replay(
                    original, 0, 1400, require_cjk=enabled
                )
                with Image.open(io.BytesIO(png)) as image:
                    assert count == 1 and image.size == (width, height)
                    assert image.convert("RGB").tobytes() == pixels
                if not enabled:
                    with pytest.raises(ValueError, match="Chinese heading"):
                        replay(original, 0, 1400, require_cjk=True)
                images.append(png)
                assert (
                    source.read_bytes() == original
                    and source.stat().st_mtime_ns == mtime
                )
    assert images[0] != images[1]
