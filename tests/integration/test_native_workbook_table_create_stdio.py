"""SDK2 creates a workbook and native Table with immutable source-cell evidence."""

import io
import os
import sys

import openpyxl
import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client

from tests.integration.test_native_pptx_shape_stdio import native
from tests.integration.test_native_workbook_stdio import read_workbook


@pytest.mark.timeout(120)
async def test_independent_table_creation_and_history_over_sdk2(tmp_path):
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
            "ASSET_AWARE_MCP_TOOL_SURFACE": "balanced",
        },
    )
    async with Client(stdio_client(params)) as client:
        contract = await native(client, op="contract", for_op="add_workbook_table")
        assert contract["workbook_table_creation_enabled"]
        assert "add_workbook_table" in contract["formats"]["xlsx"]
        edits = [
            {"sheet": "Sheet1", "cell": cell, "kind": "string", "value": value}
            for cell, value in {
                "A1": "Code",
                "B1": "Length",
                "A2": "007",
                "A3": "0010",
            }.items()
        ]
        asset = (
            await native(
                client,
                op="create",
                workbook={
                    "name": "independent.xlsx",
                    "sheets": ["Sheet1"],
                    "edits": edits,
                },
            )
        )["asset"]
        before = await read_workbook(client, asset)
        assert before["tables"] == []
        old = (
            await native(
                client,
                op="read_cell",
                asset_id=asset["asset_id"],
                revision=asset["revision"],
                sheet="Sheet1",
                cell="A2",
            )
        )["cell"]["evidence"]
        args = {
            "op": "add_workbook_table",
            "asset_id": asset["asset_id"],
            "expected_revision": asset["revision"],
            "table_create": {
                "worksheet": before["worksheets"][0]["key"],
                "ref": "A1:B4",
                "name": "Inventory",
                "totals_row": True,
                "columns": [
                    {"name": "Code", "totals": {"kind": "label", "value": "Reviewed"}},
                    {
                        "name": "Length",
                        "calculated": {
                            "formula": "=LEN([@Code])",
                            "policy": "require_matching",
                        },
                        "totals": {"kind": "function", "value": "sum"},
                    },
                ],
            },
        }
        result = await native(client, **args)
        assert result["success"] and not result["source_written"]
        assert result["asset"]["revision_count"] == 2
        current = await read_workbook(client, result["asset"])
        change = current["operation_result"]["changes"][0]
        assert change["operation"] == "add_workbook_table"
        assert change["created_table"]["part"] == current["tables"][0]["part"]
        assert change["cells"]["B2"]["before"]["kind"] == "blank"
        proof = await native(client, op="verify", reference=old)
        assert proof["valid"] and not proof["is_current_managed_revision"]
        assert not (await native(client, **args))["success"]
        assert await read_workbook(client, asset) == before
        path = tmp_path / "result.xlsx"
        assert (
            await native(
                client,
                op="publish",
                asset_id=asset["asset_id"],
                expected_revision=result["asset"]["revision"],
                output_path=str(path),
            )
        )["success"]
    book = openpyxl.load_workbook(io.BytesIO(path.read_bytes()))
    try:
        sheet = book["Sheet1"]
        assert sheet["A2"].value == "007" and sheet["A2"].data_type == "s"
        assert sheet["B3"].value == "=LEN([@Code])"
        assert sheet["B4"].value == "=SUBTOTAL(109,[Length])"
        assert sheet.tables["Inventory"].ref == "A1:B4"
        assert sheet.tables["Inventory"].autoFilter.ref == "A1:B3"
    finally:
        book.close()
