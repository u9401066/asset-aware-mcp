"""Independent Calc import, native ODS edits, formula results and actual PDF layout."""

import hashlib
import json
import os
from pathlib import Path

import openpyxl
import pymupdf
import pytest
import xlsxwriter

from src.domain.native_ods import NativeODSCreate
from src.infrastructure.native_odf_package import q
from src.infrastructure.native_ods import NativeODS, create_native_ods
from src.infrastructure.native_office_process import (
    office_binary,
    prepare_profile,
    run_office,
)
from tests.native_ods_helpers import edit, locator

pytestmark = [
    pytest.mark.skipif(
        os.environ.get("NATIVE_ODS_RENDER_TEST") != "1",
        reason="Set NATIVE_ODS_RENDER_TEST=1 with LibreOffice Calc installed",
    ),
    pytest.mark.timeout(180),
]


def convert(source, destination, output_format):
    destination.mkdir()
    profile = prepare_profile(destination)
    log = run_office(
        [
            os.environ.get("NATIVE_ODS_CALC_BIN") or office_binary("Calc"),
            profile,
            "--headless",
            "--norestore",
            "--convert-to",
            output_format,
            "--outdir",
            str(destination),
            str(source),
        ],
        destination,
        45,
    )
    suffix = output_format.split(":", 1)[0]
    result = destination / (source.stem + "." + suffix)
    assert result.is_file(), log
    return result


def test_real_calc_creation_native_read_edit_and_render(tmp_path):
    source = tmp_path / "original.xlsx"
    with xlsxwriter.Workbook(source) as book:
        sheet = book.add_worksheet("Sheet1")
        red = book.add_format({"font_color": "#FF0000", "font_size": 20, "bold": True})
        number = book.add_format({"num_format": "0.00", "font_color": "#0000FF"})
        sheet.write("A1", "OLD TITLE", red)
        sheet.write_number("A2", 3.43, number)
        sheet.write_formula("B2", "=A2*2", number, 6.86)
        sheet.write_formula("B3", "=A2>4", None, False)
        sheet.write_formula("B4", '=IF(A2>4,"high","low")', None, "low")
        sheet.write("A4", "UNCHANGED")
        sheet.set_column("A:B", 25)
        sheet.print_area("A1:B4")
    original_ods = convert(source, tmp_path / "source-ods", "ods")
    original = original_ods.read_bytes()
    mtime = original_ods.stat().st_mtime_ns
    book = NativeODS(original)
    old_style = book.read_cell(locator())["attributes"][q("table", "style-name")]
    changed, receipt = book.edit([edit(value="NEW TITLE"), edit(1, 0, "5.75", "float")])
    edited = tmp_path / "edited.ods"
    edited.write_bytes(changed)
    assert (
        NativeODS(changed).read_cell(locator())["attributes"][q("table", "style-name")]
        == old_style
    )
    assert (
        original_ods.read_bytes() == original
        and original_ods.stat().st_mtime_ns == mtime
    )
    for name, content in book.package.parts.items():
        if name != "content.xml":
            assert NativeODS(changed).package.parts[name] == content
    independent = convert(edited, tmp_path / "readback", "xlsx")
    readback = openpyxl.load_workbook(independent, data_only=True)
    assert readback.active["A1"].value == "NEW TITLE"
    assert readback.active["A2"].value == 5.75
    assert readback.active["B2"].value == 11.5
    assert readback.active["B3"].value is True
    assert readback.active["B4"].value == "high"
    assert readback.active["A2"].number_format == "0.00"
    assert readback.active["A1"].font.bold
    assert readback.active["A1"].font.color.rgb[-6:] == "FF0000"
    assert readback.active["A4"].value == "UNCHANGED"
    pdf = convert(edited, tmp_path / "render", "pdf:calc_pdf_Export")
    original_pdf = convert(
        original_ods, tmp_path / "before-render", "pdf:calc_pdf_Export"
    )
    with pymupdf.open(pdf) as rendered:
        assert len(rendered) == 1
        text = rendered[0].get_text()
        assert all(
            item in text for item in ("NEW TITLE", "5.75", "11.50", "UNCHANGED", "high")
        )
        spans = [
            span
            for block in rendered[0].get_text("dict")["blocks"]
            if "lines" in block
            for line in block["lines"]
            for span in line["spans"]
        ]
        title = next(span for span in spans if "NEW TITLE" in span["text"])
        assert title["color"] == 0xFF0000 and title["size"] >= 19
        rendered[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(
            tmp_path / "edited-preview.png"
        )
        with pymupdf.open(original_pdf) as old_render:
            old_rect = old_render[0].search_for("UNCHANGED")[0]
            assert rendered[0].search_for("UNCHANGED")[0] == old_rect
            assert (
                rendered[0].get_pixmap(clip=old_rect).samples
                == old_render[0].get_pixmap(clip=old_rect).samples
            )
            old_render[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(
                tmp_path / "original-preview.png"
            )
    assert receipt.changed_parts == ["content.xml"]
    assert receipt.repairs == ["invalidated_typed_formula_caches"]
    assert (
        original_ods.read_bytes() == original
        and original_ods.stat().st_mtime_ns == mtime
    )
    if output := os.environ.get("NATIVE_ODS_ARTIFACTS"):
        root = Path(output)
        root.mkdir(parents=True, exist_ok=False)
        files = [
            original_ods,
            edited,
            independent,
            pdf,
            original_pdf,
            tmp_path / "edited-preview.png",
            tmp_path / "original-preview.png",
        ]
        hashes = {}
        for file in files:
            data = file.read_bytes()
            (root / file.name).write_bytes(data)
            hashes[file.name] = hashlib.sha256(data).hexdigest()
        (root / "receipt.json").write_text(receipt.model_dump_json(indent=2))
        (root / "proof.json").write_text(
            json.dumps(
                {
                    "sha256": hashes,
                    "renderer": os.environ.get("NATIVE_ODS_CALC_BIN")
                    or office_binary("Calc"),
                    "source_unchanged": True,
                    "independent_formula_value": readback.active["B2"].value,
                },
                indent=2,
            )
        )


def test_native_created_ods_is_readable_by_calc_including_uncached_formula(tmp_path):
    original = create_native_ods(NativeODSCreate())
    data, _ = NativeODS(original).edit(
        [
            edit(value="literal =1+1"),
            edit(1, 0, "3.43", "float"),
            edit(1, 1, "=[.A2]*2", "formula"),
            edit(2, 0, True, "boolean"),
        ]
    )
    source = tmp_path / "created.ods"
    source.write_bytes(data)
    result = convert(source, tmp_path / "converted", "xlsx")
    readback = openpyxl.load_workbook(result, data_only=True)
    assert readback.active["A1"].value == "literal =1+1"
    assert readback.active["A2"].value == 3.43
    assert readback.active["B2"].value == 6.86
    assert readback.active["A3"].value is True
