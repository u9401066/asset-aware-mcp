"""Actual SDK2 discovery, full table paging and atomic native Word grid mutation."""

from __future__ import annotations

import hashlib
import json
import os
import sys
from functools import partial

import pytest
from docx import Document
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client

from tests.integration.test_native_docx_stdio_e2e import _native
from tests.native_docx_structure_helpers import creation


async def complete(native, fields):
    chunks, offset, digest = [], 0, None
    while True:
        response = await native(**fields, text_offset=offset, text_limit=2500)
        assert response["success"], response
        page = response["table"]
        assert digest in (None, page["text_sha256"])
        digest = page["text_sha256"]
        assert page["excerpt_char_range"][0] == offset
        chunks.append(page["text_excerpt"])
        offset = page["next_text_offset"]
        if offset is None:
            break
    text = "".join(chunks)
    assert hashlib.sha256(text.encode()).hexdigest() == digest
    return json.loads(text)


@pytest.mark.timeout(90)
async def test_docx_grid_over_sdk2(tmp_path):
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
        },
    )
    async with Client(stdio_client(params)) as client:
        native = partial(_native, client)
        contract = await native(op="contract", for_op="update_docx_table_grid")
        assert contract["docx_table_grid_enabled"]
        created = await native(op="create_docx", docx_create=creation().model_dump())
        asset = created["asset"]
        blocks = await native(**created["review_request"], limit=100)
        reference = next(
            b["evidence"] for b in blocks["blocks"] if b["type"] == "table"
        )
        read = {
            "op": "read_docx_table",
            "asset_id": asset["asset_id"],
            "revision": asset["revision"],
            "docx_table_reference": reference,
        }
        before = await complete(native, read)
        edits = [
            {"op": "insert", "axis": "row", "index": 2, "sizes_twips": [500]},
            {"op": "insert", "axis": "column", "index": 1, "sizes_twips": [600]},
            {"op": "resize", "axis": "column", "index": 2, "sizes_twips": [2000]},
            {
                "op": "merge",
                "row": 2,
                "column": 1,
                "end_row": 2,
                "end_column": 2,
                "content_policy": "require_empty",
            },
            {"op": "split", "row": 2, "column": 1},
            {"op": "delete", "axis": "row", "index": 1, "count": 1},
        ]
        result = await native(
            op="update_docx_table_grid",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_table_grid={"reference": reference, "edits": edits},
        )
        assert result["success"], result
        current = await complete(native, result["review_request"])
        assert [current["rows"], current["columns"]] == [3, 4]
        assert current["column_widths_twips"] == [1200, 600, 2000, 2400]
        assert await complete(native, read) == before
        stale = await native(
            op="update_docx_table_grid",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_table_grid={"reference": reference, "edits": edits},
        )
        assert not stale["success"]
        proof = await native(op="verify", reference=reference)
        assert proof["valid"] and not proof["is_current_managed_revision"]
        path = tmp_path / "grid.docx"
        published = await native(
            op="publish",
            asset_id=asset["asset_id"],
            expected_revision=result["asset"]["revision"],
            output_path=str(path),
        )
        assert published["success"], published
        document = Document(path)
        assert (
            document.tables[0].cell(1, 0).text
            == document.tables[0].cell(2, 0).text
            == "007"
        )
        assert document.tables[0].cell(2, 2).text == "1,234.50"

        literal = "長資料" * 1200
        large = await native(
            op="update_docx_table_grid",
            asset_id=asset["asset_id"],
            expected_revision=result["asset"]["revision"],
            docx_table_grid={
                "reference": result["review_request"]["docx_table_reference"],
                "edits": [
                    {
                        "op": "insert",
                        "axis": "row",
                        "index": 3,
                        "sizes_twips": [500],
                        "cells": [
                            [
                                {"paragraphs": [{"runs": [{"text": literal}]}]},
                                {},
                                {},
                                {},
                            ]
                        ],
                    }
                ],
            },
        )
        assert large["success"], large
        assert len(json.dumps(large, ensure_ascii=False)) < 10000
        full = await complete(native, large["review_request"])
        request = full["operation_result"]["changes"][0]["edits"][0]["request"]
        assert request["cells"][0][0]["paragraphs"][0]["runs"][0]["text"] == literal
