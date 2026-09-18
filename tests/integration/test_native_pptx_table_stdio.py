"""Real SDK2 table creation, complete readback, cell edit, delete, wiki and writeback."""

from __future__ import annotations

import os
import sys

import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation

from tests.integration.test_native_pptx_picture_stdio import publish_and_writeback
from tests.integration.test_native_pptx_shape_stdio import native, read_complete
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_table_helpers import table_addition


async def edit_first_table(client, current, locator, first):
    run = first["table"]["rows"][1]["cells"][1]["paragraphs"][0]["items"][0]
    edited = await native(
        client,
        op="update_pptx",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        pptx_edits=[
            {
                "locator": {
                    **locator,
                    "row": 1,
                    "column": 1,
                    "paragraph": 0,
                    "run": 0,
                },
                "expected_text_sha256": run["text_sha256"],
                "text": "008",
            }
        ],
    )
    current = edited["asset"]
    updated = await read_complete(client, current, locator)
    assert (
        updated["table"]["rows"][1]["cells"][1]["paragraphs"][0]["items"][0]["text"]
        == "008"
    )
    proof = await native(client, op="verify", reference=first["evidence"])
    assert proof["valid"] and not proof["is_current_managed_revision"]
    return current


async def mutate_tables(client, asset, data):
    item = table_addition(data).model_dump()
    added = await native(
        client,
        op="add_pptx_tables",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        pptx_tables=[item, item],
    )
    assert added["success"] and not added["source_written"]
    current = added["asset"]
    locators = [c["locator"] for c in added["operation_result"]["changes"]]
    first = await read_complete(client, current, locators[0])
    current = await edit_first_table(client, current, locators[0], first)
    second = await read_complete(client, current, locators[1])
    deleted = await native(
        client,
        op="delete_pptx_shapes",
        asset_id=current["asset_id"],
        expected_revision=current["revision"],
        pptx_shape_refs=[second["evidence"]],
    )
    assert deleted["success"]
    return deleted["asset"]


@pytest.mark.timeout(120)
async def test_native_table_crud_over_sdk2(tmp_path):
    data = build_presentation()
    source = tmp_path / "source.pptx"
    source.write_bytes(data)
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
            "ASSET_AWARE_MCP_TOOL_SURFACE": "balanced",
        },
    )
    async with Client(stdio_client(params)) as client:
        contract = await native(client, op="contract", for_op="add_pptx_tables")
        assert "add_pptx_tables" in contract["formats"]["pptx"]
        asset = (await native(client, op="register", source_path=str(source)))["asset"]
        current = await mutate_tables(client, asset, data)
        target = tmp_path / "published.pptx"
        await publish_and_writeback(client, current, asset, source, target)
    shape = Presentation(target).slides[0].shapes[-1]
    assert shape.table.cell(1, 1).text == "008"
    assert shape.table.cell(0, 0).is_merge_origin
    assert shape.table.cell(2, 2).text == "=SUM(A1:A2)"
