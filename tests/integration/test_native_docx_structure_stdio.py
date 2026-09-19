"""Real SDK 2 creation and reference-bound body CRUD, with independent DOCX reads."""

from __future__ import annotations

import hashlib
import os
import sys
from functools import partial
from pathlib import Path

import pytest
from docx import Document
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client

from tests.integration.test_native_docx_stdio_e2e import _native, _read_complete_dfm
from tests.native_docx_structure_helpers import creation, paragraph


@pytest.mark.timeout(90)
async def test_native_docx_structure_over_sdk2_stdio(tmp_path: Path) -> None:
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
        },
    )
    async with Client(stdio_client(params)) as client:
        native = partial(_native, client)
        contract = await native(op="contract", for_op="create_docx")
        assert contract["docx_structure_enabled"]
        created = await native(op="create_docx", docx_create=creation().model_dump())
        assert created["success"] and not created["source_written"]
        asset = created["asset"]
        assert asset["source"] is None
        original_dfm = await _read_complete_dfm(native, asset)
        read = await native(**created["review_request"], limit=100)
        table = next(b for b in read["blocks"] if b["type"] == "table")
        added = await native(
            op="add_docx_blocks",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_insert={
                "position": "after",
                "anchor": table["evidence"],
                "blocks": [paragraph("Disposable 007", italic=True)],
            },
        )
        assert added["success"], added
        current = added["asset"]
        read = await native(**added["review_request"], limit=100)
        temp = next(b for b in read["blocks"] if b["preview"] == "Disposable 007")
        removed = await native(
            op="delete_docx_blocks",
            asset_id=current["asset_id"],
            expected_revision=current["revision"],
            docx_block_refs=[temp["evidence"]],
        )
        assert removed["success"], removed
        stale = await native(
            op="delete_docx_blocks",
            asset_id=current["asset_id"],
            expected_revision=current["revision"],
            docx_block_refs=[temp["evidence"]],
        )
        assert not stale["success"]
        restored = removed["asset"]
        text = await _read_complete_dfm(native, restored)
        assert "Disposable" not in text
        edited = await native(
            op="update_docx",
            asset_id=restored["asset_id"],
            expected_revision=restored["revision"],
            docx_edit={"dfm_text": text.replace("1,234.50", "1,234.75")},
        )
        assert edited["success"], edited
        final = edited["asset"]
        published = tmp_path / "研究.docx"
        result = await native(
            op="publish",
            asset_id=final["asset_id"],
            expected_revision=final["revision"],
            output_path=str(published),
        )
        assert result["success"], result
        assert hashlib.sha256(published.read_bytes()).hexdigest() == final["revision"]
        document = Document(published)
        assert document.tables[0].cell(1, 0).text == "007"
        assert document.tables[0].cell(2, 1).text == "1,234.75"
        assert document.tables[0].cell(1, 0)._tc is document.tables[0].cell(2, 0)._tc
        assert document.paragraphs[0].runs[0].font.size.pt == 11.5
        assert await _read_complete_dfm(native, asset) == original_dfm
        for reference in (table["evidence"], temp["evidence"]):
            verified = await native(op="verify", reference=reference)
            assert verified["valid"] and not verified["is_current_managed_revision"]
        wiki = await native(
            op="export_wiki",
            asset_id=final["asset_id"],
            revision=final["revision"],
            output_dir=str(tmp_path / "wiki"),
        )
        assert wiki["success"], wiki
        records = Path(wiki["output_dir"]) / "records.jsonl"
        assert "1,234.75" in records.read_text(encoding="utf-8")
