"""Real SDK2 native presentation creation, evidence, edits and explicit publication."""

from __future__ import annotations

import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

import pytest
from mcp import Client, StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation

from tests.integration.test_native_asset_stdio_e2e import _unwrap
from tests.native_pptx_helpers import build_presentation, edit_run


@pytest.mark.timeout(60)
async def test_native_pptx_sdk2(tmp_path: Path) -> None:
    source = tmp_path / "original.pptx"
    source.write_bytes(build_presentation())
    params = StdioServerParameters(
        command=sys.executable,
        args=["-m", "src.server"],
        env={
            **os.environ,
            "DATA_DIR": str(tmp_path / "data"),
            "ENABLE_LIGHTRAG": "false",
            "ASSET_AWARE_DISABLE_DOTENV": "true",
            "ASSET_AWARE_MCP_TOOL_SURFACE": "balanced",
        },
    )
    async with Client(stdio_client(params)) as client:

        async def native(**request: Any) -> dict[str, Any]:
            result = _unwrap(
                await client.call_tool(
                    "document", {"op": "native", "native_request": request}
                )
            )
            assert "response_truncated" not in result
            return result

        contract = await native(op="contract", for_op="update_pptx")
        assert "read_pptx_shape" in contract["formats"]["pptx"]
        created = await native(
            op="create_pptx",
            presentation={
                "name": "new.pptx",
                "slides": [
                    {
                        "textboxes": [
                            {"paragraphs": [[{"text": "Created", "bold": True}]]}
                        ]
                    }
                ],
            },
        )
        assert created["success"] and created["asset"]["source"] is None
        registered = await native(op="register", source_path=str(source))
        asset = registered["asset"]
        overview = await native(op="read_pptx", asset_id=asset["asset_id"], limit=1)
        locator = overview["shapes"][0]["locator"]
        text = ""
        offset = 0
        digest = None
        while True:
            page = await native(
                op="read_pptx_shape",
                asset_id=asset["asset_id"],
                revision=asset["revision"],
                pptx_locator=locator,
                text_offset=offset,
                text_limit=1000,
            )
            shape = page["shape"]
            digest = digest or shape["text_sha256"]
            assert shape["text_sha256"] == digest
            text += shape["text_excerpt"]
            if shape["next_text_offset"] is None:
                break
            offset = shape["next_text_offset"]
        assert hashlib.sha256(text.encode("utf-8")).hexdigest() == digest
        record = json.loads(text)
        assert (await native(op="verify", reference=record["evidence"]))["valid"]
        updated = await native(
            op="update_pptx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            pptx_edits=[edit_run(record, "SDK2 changed").model_dump()],
        )
        assert updated["success"] and updated["source_written"] is False
        current = updated["asset"]
        assert (await native(op="verify", reference=record["evidence"]))["valid"]
        exported = await native(
            op="export_wiki",
            asset_id=asset["asset_id"],
            output_dir=str(tmp_path / "wiki"),
        )
        assert exported["success"] and exported["shape_count"] > 5
        records = [
            json.loads(line)
            for line in (Path(exported["output_dir"]) / "records.jsonl")
            .read_text(encoding="utf-8")
            .splitlines()
        ]
        assert (await native(op="verify", reference=records[0]["evidence"]))["valid"]
        target = tmp_path / "published.pptx"
        assert (
            await native(
                op="publish",
                asset_id=asset["asset_id"],
                expected_revision=current["revision"],
                output_path=str(target),
            )
        )["success"]
        assert hashlib.sha256(target.read_bytes()).hexdigest() == current["revision"]
        assert Presentation(target).slides[0].shapes[0].text == "SDK2 changed unchanged"
        assert (
            await native(
                op="writeback",
                asset_id=asset["asset_id"],
                expected_revision=current["revision"],
                expected_source_sha256=asset["source"]["sha256"],
            )
        )["source_written"]
        assert source.read_bytes() == target.read_bytes()
        rejected = await native(
            op="update_pptx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            pptx_edits=[edit_run(record, "stale").model_dump()],
        )
        assert rejected["success"] is False
