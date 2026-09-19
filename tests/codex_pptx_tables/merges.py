"""Independently audit paragraph migration, split behavior and every live revision."""

import io
import json
from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require
from tests.codex_pptx_tables.derivations import complete_page
from tests.codex_pptx_tables.grids import unchanged_package

TEMP = ["TEMP", "000", "+0.00", "µg", "尾"]
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def xml(node):
    return etree.tostring(node, method="c14n", exclusive=True, with_tail=False)


def shape(data):
    deck = Presentation(io.BytesIO(data))
    require(
        len(deck.slides) == 1 and len(deck.slides[0].shapes) == 1,
        "Merge stage shape count",
    )
    return deck.slides[0].shapes[0]


def check_merge_stage(data, stage):
    frame = shape(data)
    table = frame.table
    rows = 5 if stage <= 3 else 4
    expected = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
    ]
    if stage <= 3:
        expected.append(TEMP if stage == 1 else ["\n".join(TEMP), "", "", "", ""])
    require(
        [r.height for r in table.rows] == [550000] * rows
        and [c.width for c in table.columns] == [1200000] * 5,
        "Merge grid dimensions differ",
    )
    require(
        (frame.left, frame.top, frame.width, frame.height)
        == (100000, 100000, 6000000, rows * 550000),
        "Merge frame differs",
    )
    require(
        [[table.cell(r, c).text for c in range(5)] for r in range(rows)] == expected,
        "Merge stage exact text differs",
    )
    for r in range(rows):
        merged = (r == 0 and stage != 5) or (r == 4 and stage == 2)
        for c in range(5):
            cell = table.cell(r, c)
            require(
                cell.is_merge_origin == (merged and c == 0)
                and cell.is_spanned == (merged and c > 0),
                "Merge stage topology differs",
            )
        if merged:
            require(
                table.cell(r, 0).span_width == 5 and table.cell(r, 0).span_height == 1,
                "Wrong merge span",
            )
    if stage == 1:
        runs = [table.cell(4, c).text_frame.paragraphs[0].runs[0] for c in range(5)]
        require(
            all(run.font.size == Pt(14) for run in runs)
            and runs[1].font.bold
            and runs[2].font.italic,
            "Temporary rich formatting missing",
        )


def check_merge_delta(before, after, stage):
    old, new = shape(before), shape(after)
    for r in range(4):
        for c in range(5):
            cells = []
            for table in (old.table, new.table):
                node = deepcopy(table.cell(r, c)._tc)
                for name in ("rowSpan", "gridSpan", "hMerge", "vMerge"):
                    node.attrib.pop(name, None)
                cells.append(xml(node))
            require(cells[0] == cells[1], "Original scanned cell XML changed")
    stripped = []
    for frame in (old, new):
        node = deepcopy(frame.element)
        node.find("p:xfrm/a:ext", NS).set("cy", "checked")
        for row in node.findall(".//a:tbl/a:tr", NS):
            row.getparent().remove(row)
        stripped.append(xml(node))
    require(stripped[0] == stripped[1], "Table properties or frame metadata changed")
    if stage in {2, 3}:
        old_cells = [old.table.cell(4, c)._tc for c in range(5)]
        new_cells = [new.table.cell(4, c)._tc for c in range(5)]
        expected = [
            xml(p)
            for cell in (old_cells if stage == 2 else old_cells[:1])
            for p in cell.findall("a:txBody/a:p", NS)
        ]
        require(
            [xml(p) for p in new_cells[0].findall("a:txBody/a:p", NS)] == expected,
            "Migrated rich paragraph XML differs",
        )
        for before_cell, after_cell in zip(old_cells, new_cells, strict=True):
            require(
                xml(before_cell.find("a:tcPr", NS))
                == xml(after_cell.find("a:tcPr", NS)),
                "Cell formatting changed during merge",
            )


def validate_merges(workspace, deck, calls, *, preceding_grid=False):
    root = workspace / "data" / "native-assets" / deck["asset_id"] / "revisions"
    available, chunks, outputs = set(), {}, []
    count, skip = 0, 5 if preceding_grid else 0
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_pptx_shape":
            record = result["shape"]
            key = canonical(record["evidence"])
            if complete_page(record, key, chunks, digest_field="text_sha256"):
                available.add(key)
        if args["op"] != "update_pptx_table_grid":
            continue
        count += 1
        if count <= skip:
            continue
        ref = args["pptx_table_grid"]["reference"]
        require(canonical(ref) in available, "Merge mutation lacks full prior readback")
        require(
            ref["asset_id"] == args["asset_id"] == deck["asset_id"]
            and ref["revision"] == args["expected_revision"],
            "Merge input reference mismatch",
        )
        revision = result["asset"]["revision"]
        before, after = (
            (root / ref["revision"]).read_bytes(),
            (root / revision).read_bytes(),
        )
        require(
            digest(before) == ref["revision"] and digest(after) == revision,
            "Merge revision digest mismatch",
        )
        check_merge_stage(after, count - skip)
        check_merge_delta(before, after, count - skip)
        unchanged_package(before, after, ref["locator"]["part"])
        outputs.append((revision, ref["locator"]))
    require(count == skip + 6, "Missing or extra successful merge stages")
    for revision, locator in outputs:
        require(
            any(
                json.loads(key)["revision"] == revision
                and json.loads(key)["locator"] == locator
                for key in available
            ),
            "Merge result lacks complete readback",
        )
    return {
        "mutations": 6,
        "intermediate_revisions_checked": len(outputs),
        "paragraph_xml_verified": True,
        "split_redistributes_text": False,
    }
