"""Independent intermediate-grid and chronological reference audit for Codex runs."""

import io
import json
import zipfile

from lxml import etree
from pptx import Presentation

from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require
from tests.codex_pptx_tables.derivations import complete_page

TEMP = ["TEMP", "000", "", "+0.00", "tmp", "N/A"]


def check_stage(data, stage):
    deck = Presentation(io.BytesIO(data))
    require(
        len(deck.slides) == 1 and len(deck.slides[0].shapes) == 1,
        "Grid stage shape count",
    )
    shape = deck.slides[0].shapes[0]
    table = shape.table
    values = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
    ]
    widths, heights = [1200000] * 5, [550000] * 4
    if stage < 5:
        widths.insert(2, 450000 if stage >= 3 else 300000)
        for row in values:
            row.insert(2, "")
    if stage in {2, 3}:
        heights.insert(2, 600000 if stage == 3 else 400000)
        values.insert(2, TEMP)
    require(
        [c.width for c in table.columns] == widths
        and [r.height for r in table.rows] == heights,
        "Intermediate grid dimensions differ",
    )
    require(
        (shape.left, shape.top, shape.width, shape.height)
        == (100000, 100000, sum(widths), sum(heights)),
        "Intermediate frame differs",
    )
    require(
        [
            [table.cell(r, c).text for c in range(len(widths))]
            for r in range(len(heights))
        ]
        == values,
        "Intermediate grid content differs",
    )
    require(
        table.cell(0, 0).span_width == len(widths)
        and table.cell(0, 0).span_height == 1,
        "Intermediate title merge differs",
    )
    require(
        all(table.cell(0, c).is_spanned for c in range(1, len(widths))),
        "Incomplete title coverage",
    )
    require(
        all(
            not table.cell(r, c).is_spanned and not table.cell(r, c).is_merge_origin
            for r in range(1, len(heights))
            for c in range(len(widths))
        ),
        "Unexpected data merge",
    )


def unchanged_package(before, after, part):
    with (
        zipfile.ZipFile(io.BytesIO(before)) as a,
        zipfile.ZipFile(io.BytesIO(after)) as b,
    ):
        require(a.namelist() == b.namelist(), "Grid package inventory differs")
        for name in a.namelist():
            if name != part:
                require(a.read(name) == b.read(name), "Grid changed unrelated part")
        left, right = etree.fromstring(a.read(part)), etree.fromstring(b.read(part))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        for root in (left, right):
            shapes = root.findall("p:cSld/p:spTree/p:graphicFrame", ns)
            require(len(shapes) == 1, "Expected one native table frame")
            shapes[0].getparent().remove(shapes[0])
        require(
            etree.tostring(left, method="c14n") == etree.tostring(right, method="c14n"),
            "Grid changed surrounding XML",
        )


def surviving_cells(before, after, stage):
    old = Presentation(io.BytesIO(before)).slides[0].shapes[0].table
    new = Presentation(io.BytesIO(after)).slides[0].shapes[0].table
    for r in range(len(old.rows)):
        for c in range(len(old.columns)):
            if (stage == 4 and r == 2) or (stage == 5 and c == 2):
                continue
            nr = r + int(stage == 2 and r >= 2) - int(stage == 4 and r > 2)
            nc = c + int(stage == 1 and c >= 2) - int(stage == 5 and c > 2)
            cells = []
            for cell in (old.cell(r, c), new.cell(nr, nc)):
                node = etree.fromstring(etree.tostring(cell._tc))
                for name in ("rowSpan", "gridSpan", "hMerge", "vMerge"):
                    node.attrib.pop(name, None)
                cells.append(etree.tostring(node, method="c14n", exclusive=True))
            require(
                cells[0] == cells[1], "Surviving cell content or formatting changed"
            )


def validate_grids(workspace, deck, calls):
    root = workspace / "data" / "native-assets" / deck["asset_id"] / "revisions"
    available, chunks, outputs, count = set(), {}, [], 0
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_pptx_shape":
            shape = result["shape"]
            key = canonical(shape["evidence"])
            if complete_page(shape, key, chunks, digest_field="text_sha256"):
                available.add(key)
        if args["op"] != "update_pptx_table_grid":
            continue
        ref = args["pptx_table_grid"]["reference"]
        require(
            canonical(ref) in available, "Grid edit lacks prior complete shape read"
        )
        require(
            ref["asset_id"] == args["asset_id"] == deck["asset_id"]
            and ref["revision"] == args["expected_revision"],
            "Grid input reference mismatch",
        )
        before = (root / ref["revision"]).read_bytes()
        revision = result["asset"]["revision"]
        after = (root / revision).read_bytes()
        require(
            digest(before) == ref["revision"] and digest(after) == revision,
            "Grid revision digest mismatch",
        )
        count += 1
        check_stage(after, count)
        unchanged_package(before, after, ref["locator"]["part"])
        surviving_cells(before, after, count)
        outputs.append((revision, ref["locator"]))
    require(count == 5, "Missing or unexpected successful grid mutation stages")
    for revision, locator in outputs:
        require(
            any(
                json.loads(key)["revision"] == revision
                and json.loads(key)["locator"] == locator
                for key in available
            ),
            "Grid result lacks complete shape readback",
        )
    return {
        "mutations": count,
        "intermediate_revisions_checked": len(outputs),
        "full_readbacks": True,
        "unchanged_package_parts": True,
    }
