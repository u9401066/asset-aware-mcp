"""Independent ZIP/PresentationML and trace audit of live slide structure mutations."""

import io
import json
import posixpath
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.trace import require
from tests.codex_pptx_tables.derivations import complete_page
from tests.codex_pptx_tables.tables import validate_table

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
MAIN = "ppt/presentation.xml"
RELS = "ppt/_rels/presentation.xml.rels"
APP = "docProps/app.xml"
OPS = {"add_pptx_slides", "reorder_pptx_slides", "delete_pptx_slides"}
TEXTS = ["TEMP SLIDE 007 µg", "SECOND +0.00"]


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def slide_keys(data):
    package = parts(data)
    relationships = {
        r.get("Id"): r.get("Target") for r in etree.fromstring(package[RELS])
    }
    return [
        {
            "slide_id": node.get("id"),
            "part": posixpath.normpath(
                posixpath.join("ppt", relationships[node.get("{" + NS["r"] + "}id")])
            ),
        }
        for node in etree.fromstring(package[MAIN]).findall("p:sldIdLst/p:sldId", NS)
    ]


def xml(node):
    return etree.tostring(node, method="c14n", exclusive=True)


def check_slide_stage(before, after, stage, original, created):
    expected = (
        [original, *created]
        if stage == 1
        else [created[1], original, created[0]]
        if stage == 2
        else [original]
    )
    require(slide_keys(after) == expected, "Slide identity/order mismatch")
    old, new = parts(before), parts(after)
    exceptions = {MAIN, RELS, APP}
    if stage == 1:
        exceptions.add("[Content_Types].xml")
        added = {s["part"] for s in created}
        added |= {
            p.rsplit("/", 1)[0] + "/_rels/" + p.rsplit("/", 1)[1] + ".rels"
            for p in added
        }
        require(
            set(new) - set(old) == added and set(old) <= set(new),
            "Unexpected new slide parts",
        )
    else:
        require(set(new) == set(old), "Slide operation lost or added unplanned parts")
    if stage == 1:
        types = [etree.fromstring(raw["[Content_Types].xml"]) for raw in (old, new)]
        selected = [
            node
            for node in types[1]
            if node.get("PartName", "").lstrip("/") in {s["part"] for s in created}
        ]
        require(
            len(selected) == 2
            and all(
                node.get("ContentType")
                == "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
                for node in selected
            ),
            "New slide content types differ",
        )
        for node in selected:
            types[1].remove(node)
        require(xml(types[0]) == xml(types[1]), "Unrelated content type changed")
    for name, raw in old.items():
        if name not in exceptions:
            require(new[name] == raw, "Slide operation changed unrelated part")
    roots = [etree.fromstring(raw[MAIN]) for raw in (old, new)]
    for root in roots:
        root.remove(root.find("p:sldIdLst", NS))
    require(
        xml(roots[0]) == xml(roots[1]), "Slide operation changed presentation metadata"
    )
    if stage == 2:
        require(new[RELS] == old[RELS], "Reorder changed relationships")
    else:
        rel_roots = [etree.fromstring(raw[RELS]) for raw in (old, new)]
        larger, smaller = (
            (rel_roots[1], rel_roots[0]) if stage == 1 else (rel_roots[0], rel_roots[1])
        )
        current = {s["part"].removeprefix("ppt/") for s in created}
        removed = [r for r in larger if r.get("Target") in current]
        require(len(removed) == 2, "Wrong slide relationship delta")
        for relation in removed:
            larger.remove(relation)
        require(xml(larger) == xml(smaller), "Unrelated relationship changed")
    if APP in old:
        app_roots = [etree.fromstring(raw[APP]) for raw in (old, new)]
        namespace = (
            "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
        )
        for name, value in (("Slides", len(expected)), ("Notes", 0)):
            actual = app_roots[1].find(f"{{{namespace}}}{name}")
            if actual is not None:
                require(actual.text == str(value), "Slide count property stale")
            for root in app_roots:
                node = root.find(f"{{{namespace}}}{name}")
                if node is not None:
                    node.text = "checked"
        require(
            xml(app_roots[0]) == xml(app_roots[1]),
            "Unrelated document properties changed",
        )
    deck = Presentation(io.BytesIO(after))
    for index, key in enumerate(created):
        if stage == 3:
            continue
        slide = next(s for s in deck.slides if str(s.slide_id) == key["slide_id"])
        require(
            len(slide.shapes) == 1 and slide.slide_layout.name == "Blank",
            "Unexpected temporary slide contents/layout",
        )
        shape = slide.shapes[0]
        require(shape.text == TEXTS[index], "Temporary slide exact text mismatch")
        require(
            (shape.left, shape.top, shape.width, shape.height)
            == (100000, 100000, 6000000, 550000),
            "Temporary textbox geometry mismatch",
        )
        paragraphs = shape.text_frame.paragraphs
        require(
            len(paragraphs) == 1 and len(paragraphs[0].runs) == 1,
            "Unexpected paragraph/run structure",
        )
        font = paragraphs[0].runs[0].font
        require(
            font.size == Pt(14)
            and font.bold == (index == 0)
            and font.italic == (index == 1),
            "Temporary slide rich formatting differs",
        )


def validate_slides(workspace, deck, calls):
    root = workspace / "data" / "native-assets" / deck["asset_id"] / "revisions"
    listings, layouts, available, chunks, outputs = {}, {}, set(), {}, []
    original, created = None, None
    historical_deleted = False
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_pptx" and result["asset_id"] == deck["asset_id"]:
            revision = result["inspected_revision"]
            coverage = listings.setdefault(revision, {})
            for index, key in enumerate(result["slides"], args.get("offset", 0)):
                coverage[index] = key
        if args["op"] == "read_pptx_layouts" and result["asset_id"] == deck["asset_id"]:
            layouts.setdefault(result["inspected_revision"], []).extend(
                result["layouts"]
            )
        if args["op"] == "read_pptx_shape":
            record = result["shape"]
            key = canonical(record["evidence"])
            if complete_page(record, key, chunks, digest_field="text_sha256"):
                available.add(key)
        if args["op"] == "verify" and len(outputs) == 3 and created:
            reference = args["reference"]
            historical_deleted |= (
                reference.get("locator", {}).get("slide_id")
                in {s["slide_id"] for s in created}
                and result.get("valid") is True
                and result.get("is_current_managed_revision") is False
            )
        if args["op"] not in OPS:
            continue
        before_revision = args["expected_revision"]
        revision = result["asset"]["revision"]
        require(args["asset_id"] == deck["asset_id"], "Wrong slide target asset")
        before, after = (
            (root / before_revision).read_bytes(),
            (root / revision).read_bytes(),
        )
        require(
            digest(before) == before_revision and digest(after) == revision,
            "Slide revision digest differs",
        )
        require(
            listings.get(before_revision) == dict(enumerate(slide_keys(before))),
            "Slide mutation lacks complete current listing",
        )
        stage = len(outputs) + 1
        require(
            args["op"]
            == ["add_pptx_slides", "reorder_pptx_slides", "delete_pptx_slides"][
                stage - 1
            ],
            "Unexpected slide mutation sequence",
        )
        if stage == 1:
            validate_table(io.BytesIO(before))
            original = slide_keys(before)[0]
            created = slide_keys(after)[1:]
            require(len(created) == 2, "Expected two inserted temporary slides")
            advertised = layouts.get(before_revision, [])
            require(
                all(
                    any(
                        layout["part"] == item["layout_part"]
                        and layout["type"] == "blank"
                        for layout in advertised
                    )
                    for item in args["pptx_slide_insert"]["slides"]
                ),
                "Slide insertion lacks current layout discovery",
            )
        check_slide_stage(before, after, stage, original, created)
        outputs.append(revision)
    require(len(outputs) == 3, "Missing slide stages")
    for revision in outputs:
        require(
            listings.get(revision)
            == dict(enumerate(slide_keys((root / revision).read_bytes()))),
            "Slide result lacks full listing",
        )
    require(created is not None, "Missing created slides")
    for key in created:
        require(
            any(
                json.loads(record)["revision"] == outputs[0]
                and json.loads(record)["locator"]["slide_id"] == key["slide_id"]
                for record in available
            ),
            "Temporary slide lacks full shape readback",
        )
    require(historical_deleted, "Deleted slide historical proof missing")
    for revision in outputs[1:]:
        require(
            any(
                json.loads(record)["revision"] == revision
                and json.loads(record)["locator"]["slide_id"] == original["slide_id"]
                for record in available
            ),
            "Original table lacks complete structural readback",
        )
    return {
        "mutations": 3,
        "intermediate_revisions_checked": 3,
        "full_slide_listings": True,
        "original_parts_preserved": True,
        "temporary_shapes_fully_read": True,
    }
