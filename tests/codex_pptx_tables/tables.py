"""Independent native table geometry, text and revision checks for the live corpus."""

from pptx import Presentation

from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def validate_table(path, *, expected_count=1, changed=False, check_text=True):
    deck = Presentation(path)
    require(
        len(deck.slides) == 1 and len(deck.slides[0].shapes) == expected_count,
        "Wrong slide/table count",
    )
    expected = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
    ]
    if changed:
        expected[2][1] = "008"
    for index, shape in enumerate(deck.slides[0].shapes):
        table = shape.table
        require(
            (shape.left, shape.top, shape.width, shape.height)
            == (100000, 100000, 6000000, 2200000),
            "Wrong table rectangle",
        )
        require(
            [c.width for c in table.columns] == [1200000] * 5
            and [r.height for r in table.rows] == [550000] * 4,
            "Wrong table grid",
        )
        actual = [[table.cell(r, c).text for c in range(5)] for r in range(4)]
        wanted = [row[:] for row in expected]
        if changed and index > 0:
            wanted[2][1] = PAGE_ROWS[0][0][1]
        if check_text:
            require(actual == wanted, "Exact scanned table transcription mismatch")
        require(
            table.cell(0, 0).is_merge_origin
            and table.cell(0, 0).span_width == 5
            and table.cell(0, 0).span_height == 1,
            "Merged title missing",
        )
        require(
            all(table.cell(0, c).is_spanned for c in range(1, 5)),
            "Merged title coverage wrong",
        )
        require(
            all(not table.cell(r, c).is_spanned for r in range(1, 4) for c in range(5)),
            "Unexpected data merge",
        )


def exact(path, *, changed=False):
    try:
        validate_table(path, expected_count=2, changed=changed)
    except ValueError:
        return False
    return True


def validate_revisions(root, deck, *, grid=False, slides=False):
    stages = [r for r in deck["history"] if r.get("result") and r["result"]["changes"]]
    if grid:
        stages = [
            r
            for r in stages
            if not any(
                c.get("operation") == "edit_table_grid" for c in r["result"]["changes"]
            )
        ]
    if slides:
        operations = {"add_slides", "reorder_slides", "delete_slides_retaining_parts"}
        stages = [
            r
            for r in stages
            if not any(c.get("operation") in operations for c in r["result"]["changes"])
        ]
    require(len(stages) >= 4, "Missing add/edit/restore/delete stages")
    require(
        all(c.get("operation") == "add_table" for c in stages[0]["result"]["changes"]),
        "First mutation was not native table creation",
    )
    first_exact = exact(root / stages[0]["sha256"])
    temporary = restored = False
    for stage in stages[:-1]:
        path = root / stage["sha256"]
        validate_table(path, expected_count=2, check_text=False)
        if temporary and exact(path):
            restored = True
        temporary |= exact(path, changed=True)
    require(temporary and restored, "Temporary cell edit or exact restoration missing")
    validate_table(root / stages[-1]["sha256"])
    return first_exact
