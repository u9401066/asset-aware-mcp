"""Independent grid trace ordering, coordinate/value and package-member audit."""

import io
import json
import zipfile

from openpyxl import load_workbook

from tests.codex_native_pdf.trace import digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def complete_workbook_reads(calls, target):
    buffers, observed, changes = {}, set(), []
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_workbook":
            require(
                args["asset_id"] == target["asset_id"], "Wrong grid workbook identity"
            )
            revision = args.get("revision")
            if revision is None:
                # Exploratory reads are allowed but cannot authorize a mutation.
                continue
            require(
                revision == result["inspected_revision"], "Unpinned grid workbook read"
            )
            key = (
                revision,
                args.get("workbook_view", "structure"),
                result["text_sha256"],
            )
            start, end = result["excerpt_char_range"]
            if start == 0:
                buffers[key] = ""
            require(
                start == args.get("text_offset", 0) == len(buffers.get(key, "")),
                "Noncontiguous grid workbook read",
            )
            buffers[key] += result["text_excerpt"]
            require(end == len(buffers[key]), "Grid workbook excerpt range mismatch")
            if result["next_text_offset"] is not None:
                continue
            require(
                digest(buffers[key].encode()) == key[2], "Grid workbook hash mismatch"
            )
            record = json.loads(buffers[key])
            require(
                record["asset_id"] == target["asset_id"]
                and record["revision"] == revision,
                "Grid workbook read changed binding",
            )
            if key[1] == "references":
                require(
                    isinstance(record["references"], list),
                    "Missing grid reference inventory",
                )
                stage = len(changes) + 1
                require(
                    stage < len(target["history"]), "Read after unexpected grid stage"
                )
                history = next(
                    item
                    for item in reversed(target["history"][: stage + 1])
                    if item["sha256"] == revision
                )
                if history.get("result") is not None:
                    require(
                        record.get("operation_result") == history["result"],
                        "Grid receipt read differs from stored history",
                    )
                observed.add((stage, revision))
        elif args["op"] == "update_worksheet_grid":
            require(
                args["asset_id"] == target["asset_id"]
                and (len(changes) + 1, args["expected_revision"]) in observed,
                "Grid mutation precedes complete reference read",
            )
            index = len(changes) + 1
            require(index + 1 < len(target["history"]), "Extra grid mutation")
            require(
                args["expected_revision"] == target["history"][index]["sha256"]
                and result["asset"]["revision"]
                == target["history"][index + 1]["sha256"],
                "Grid mutation differs from revision chain",
            )
            changes.append(args)
    require(
        (3, target["revision"]) in observed,
        "Final grid receipt was not completely read",
    )
    require(len(changes) == 2, "Missing grid insert/delete mutations")
    for index, request in enumerate(changes):
        grid = request["worksheet_grid"]
        require(
            grid["worksheet"] == {"sheet_id": "1", "part": "xl/worksheets/sheet1.xml"},
            "Wrong grid worksheet key",
        )
        operation = "insert" if index == 0 else "delete"
        actual = []
        for edit in grid["edits"]:
            normalized = {"count": 1, **edit}
            for key, default in (
                ("inherit_format", "before"),
                ("merged_anchor", "preserve"),
                ("collapsed_objects", "preserve_size"),
            ):
                if normalized.get(key) in {None, default}:
                    normalized.pop(key, None)
            actual.append(normalized)
        require(
            actual
            == [
                {"axis": "row", "operation": operation, "at": 2, "count": 1},
                {"axis": "column", "operation": operation, "at": 3, "count": 1},
            ],
            "Unexpected grid transformation",
        )


def validate_values(data, index):
    book = load_workbook(io.BytesIO(data))
    try:
        require(book.sheetnames == ["Sheet1"], "Wrong grid worksheet list")
        sheet = book["Sheet1"]
        values = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
        if index:
            values[1][1] = "008"
        expected = {}
        for row_index, row in enumerate(values, 1):
            for col_index, value in enumerate(row, 1):
                new_row = row_index + int(index == 2 and row_index >= 2)
                new_col = col_index + int(index == 2 and col_index >= 3)
                expected[new_row, new_col] = value
        require(
            (sheet.max_row, sheet.max_column) == ((4, 6) if index == 2 else (3, 5)),
            "Wrong grid table bounds",
        )
        for row in sheet.iter_rows():
            for cell in row:
                expected_value = expected.get((cell.row, cell.column))
                require(
                    cell.value == expected_value, "Native grid literal or blank differs"
                )
                if expected_value is not None:
                    require(cell.data_type == "s", "Native grid changed literal type")
    finally:
        book.close()


def validate_grid_history(calls, target, directory):
    require(len(target["history"]) == 4, "Unexpected grid workbook history")
    complete_workbook_reads(calls, target)
    previous = None
    for index, item in enumerate(target["history"]):
        data = (directory / "revisions" / item["sha256"]).read_bytes()
        require(digest(data) == item["sha256"], "Grid stored revision hash mismatch")
        validate_values(data, index)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        if index >= 2:
            result = item["result"]
            require(
                result["changes"][0]["operation"] == "update_worksheet_grid",
                "Missing stored grid receipt",
            )
            require(
                "complete_grid_xml_read_back" in result["checks"],
                "Missing grid readback check",
            )
            require(
                set(parts) == set(previous),
                "Grid unexpectedly added or removed package members",
            )
            for name in parts:
                if name not in result["changed_parts"]:
                    require(
                        parts[name] == previous[name],
                        "Grid changed an undeclared package member",
                    )
                if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                    require(
                        parts[name] == previous[name],
                        "Grid changed an unrelated package member",
                    )
            require(
                parts["xl/styles.xml"] == previous["xl/styles.xml"],
                "Grid changed source styles",
            )
        previous = parts
