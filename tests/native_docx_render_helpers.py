"""Two native Word pages with visible colors, header/footer and editable table."""

import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def paged_docx():
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Revision-pinned Word review"
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for index, color in enumerate(("FF0000", "0000FF")):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.page_break_before = index > 0
        run = paragraph.add_run(f"WORD PAGE {index + 1}")
        run.font.size = Pt(32)
        run.font.color.rgb = RGBColor.from_string(color)
    table = document.add_table(2, 2)
    table.cell(0, 0).text = "Count"
    table.cell(0, 1).text = "Reading"
    table.cell(1, 0).text = "007"
    table.cell(1, 1).text = "-0.50 mg/L"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
