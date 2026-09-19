"""Audit real MCP actions, immutable revisions and measurable clipping corrections."""

import base64
import io
import json

import openpyxl

from tests.codex_native_pdf.artifacts import load_asset, read_json
from tests.codex_native_pdf.trace import (
    complete_records,
    digest,
    payload,
    validate_image_pixels,
)
from tests.codex_pdf.trace import call_failed, require, tool_errors
from tests.codex_table_grid.reads import read_events
from tests.codex_workbook_rendition.audit import receipts
from tests.native_workbook_helpers import _parts
from tests.native_workbook_layout_helpers import check_corrected_pdf


def calls_from(events):
    require(any(e["type"] == "turn.completed" for e in events), "Incomplete Codex turn")
    calls = []
    allowed = {
        "contract",
        "schema",
        "inspect",
        "register",
        "read_workbook",
        "read_cell",
        "read_worksheet_layout",
        "update_worksheet_layout",
        "create_workbook_rendition",
        "read_rendition",
        "read_pdf",
        "read_pdf_page",
        "render_pdf_page",
        "verify",
        "publish",
        "export_wiki",
        "history",
    }
    for event in events:
        if event["type"] != "item.completed":
            continue
        item = event["item"]
        require(
            item["type"] in {"agent_message", "reasoning", "plan", "mcp_tool_call"},
            "Non-MCP action",
        )
        if item["type"] != "mcp_tool_call":
            continue
        require(
            item["server"] == "asset_aware_under_test"
            and item["tool"] == "document"
            and item["arguments"]["op"] == "native",
            "Wrong MCP tool",
        )
        require(
            item["arguments"]["native_request"]["op"] in allowed, "Unexpected operation"
        )
        if not call_failed(item):
            calls.append(item)
    require(
        allowed - {"schema", "inspect"}
        <= {c["arguments"]["native_request"]["op"] for c in calls},
        "Incomplete layout workflow",
    )
    return calls


def layouts(calls):
    buffers, complete = {}, {}
    for call in calls:
        args = call["arguments"]["native_request"]
        if args["op"] != "read_worksheet_layout":
            continue
        result = payload(call)
        key = (
            args["asset_id"],
            args["revision"],
            args["worksheet_key"]["part"],
            result["text_sha256"],
        )
        require(result["inspected_revision"] == key[1], "Wrong layout revision")
        start, end = result["excerpt_char_range"]
        if start == 0:
            require(key not in buffers, "Restarted unfinished layout read")
            buffers[key] = ""
        require(
            start == args.get("text_offset", 0) == len(buffers.get(key, "")),
            "Noncontiguous layout",
        )
        buffers[key] += result["text_excerpt"]
        require(end == len(buffers[key]), "Wrong layout offset")
        if result["next_text_offset"] is None:
            text = buffers.pop(key)
            require(digest(text.encode()) == key[3], "Layout hash mismatch")
            record = json.loads(text)
            require(
                record["revision"] == key[1]
                and record["worksheet"]["key"] == args["worksheet_key"],
                "Layout identity mismatch",
            )
            complete[key[:3]] = record
    require(not buffers, "Unfinished layout reads")
    return complete


def audit(output):
    expected, final = (
        read_json(output / "expected.json"),
        read_json(output / "last-message.txt"),
    )
    workspace = output / "workspace"
    events = [
        json.loads(line) for line in (output / "events.jsonl").read_text().splitlines()
    ]
    calls = calls_from(events)
    source = load_asset(workspace, final["source_asset_id"])
    original, updated = (
        (workspace / "source.xlsx").read_bytes(),
        (workspace / "verified.xlsx").read_bytes(),
    )
    require(
        digest(original) == expected["source_sha256"]
        and (workspace / "source.xlsx").stat().st_mtime_ns
        == expected["source_mtime_ns"],
        "Human source changed",
    )
    require(
        digest(updated) == source["revision"] and len(source["history"]) == 4,
        "Wrong workbook history",
    )
    old, new = _parts(original), _parts(updated)
    require(set(old) == set(new), "Workbook package inventory changed")
    require(
        {p for p in old if old[p] != new[p]}
        <= {
            "xl/workbook.xml",
            "xl/worksheets/sheet1.xml",
            "xl/worksheets/sheet3.xml",
            "xl/worksheets/sheet4.xml",
        },
        "Unrelated package part changed",
    )
    before, after = (
        openpyxl.load_workbook(io.BytesIO(original)),
        openpyxl.load_workbook(io.BytesIO(updated)),
    )
    try:
        require(before.sheetnames == after.sheetnames, "Sheet identities changed")
        for sheet in before:
            require(
                sheet.sheet_state == after[sheet.title].sheet_state,
                "Sheet visibility changed",
            )
            for row in sheet:
                for cell in row:
                    value = after[sheet.title][cell.coordinate]
                    require(
                        (cell.value, cell.style_id, cell.data_type)
                        == (value.value, value.style_id, value.data_type),
                        "Cell content or style changed",
                    )
    finally:
        before.close()
        after.close()
    records = layouts(calls)
    updates = [
        c
        for c in calls
        if c["arguments"]["native_request"]["op"] == "update_worksheet_layout"
    ]
    require(len(updates) == 3, "Wrong layout update count")
    for call, history in zip(updates, source["history"][1:], strict=True):
        args, result = call["arguments"]["native_request"], payload(call)
        part = args["worksheet_layout"]["worksheet"]["part"]
        previous = records[(source["asset_id"], args["expected_revision"], part)]
        current = records[(source["asset_id"], result["asset"]["revision"], part)]
        require(
            current["operation_result"] == history["result"], "Wrong layout receipt"
        )
        change = history["result"]["changes"][0]
        require(
            change["before"] == previous["layout"]
            and change["after"] == current["layout"],
            "Incomplete dimension review",
        )
    require(
        {r[3] for r in read_events(calls) if r[1].get("workbook_view") == "references"}
        >= {digest(original), source["revision"]},
        "Missing workbook references",
    )
    pdfs = [
        load_asset(workspace, identity) for identity in final["rendition_asset_ids"]
    ]
    require(
        len(pdfs) == 2 and len({p["asset_id"] for p in pdfs}) == 2,
        "Expected independent before/after PDFs",
    )
    reports, pages = receipts(calls), complete_records(calls)
    expected_pages, pdf_bytes = set(), []
    for pdf, name, revision in zip(
        pdfs,
        ("baseline.pdf", "corrected.pdf"),
        (digest(original), source["revision"]),
        strict=True,
    ):
        require(len(pdf["history"]) == 1 and pdf["source"] is None, "PDF was mutated")
        report = reports[(pdf["asset_id"], pdf["revision"])]
        require(report == pdf["history"][0]["result"], "Incomplete rendition receipt")
        change = report["changes"][0]
        require(
            change["source_reference"]["asset_id"] == source["asset_id"]
            and change["source_reference"]["revision"] == revision,
            "Wrong rendition source",
        )
        require(
            change["requested"]["mode"] == "whole_sheet"
            and change["requested"]["calculation"] == "recalculate",
            "Wrong rendition policies",
        )
        data = (workspace / name).read_bytes()
        require(digest(data) == pdf["revision"], "Wrong published PDF")
        pdf_bytes.append(data)
        expected_pages.update(
            (pdf["asset_id"], pdf["revision"], index) for index in range(4)
        )
    pixel_proof = check_corrected_pdf(*pdf_bytes)
    require(
        expected_pages
        <= {
            (
                r["evidence"]["asset_id"],
                r["evidence"]["revision"],
                r["locator"]["page_index"],
            )
            for r in pages.values()
        },
        "Missing complete page records",
    )
    seen, images, historical, verified, frozen_replay = set(), 0, False, set(), False
    last_update = calls.index(updates[-1])
    for index, call in enumerate(calls):
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "verify" and result.get("valid") is True:
            ref = args["reference"]
            verified.add((ref["asset_id"], ref["revision"]))
            historical |= (
                ref["asset_id"] == source["asset_id"]
                and ref["revision"] == digest(original)
                and result.get("is_current_managed_revision") is False
            )
        if args["op"] == "render_pdf_page":
            blocks = [b for b in call["result"]["content"] if b["type"] == "image"]
            require(len(blocks) == 1, "Missing actual PNG")
            raw = base64.b64decode(blocks[0]["data"], validate=True)
            require(digest(raw) == result["image_sha256"], "Wrong PNG identity")
            validate_image_pixels(workspace, result, args, raw)
            seen.add(
                (args["asset_id"], args["revision"], args["pdf_locator"]["page_index"])
            )
            frozen_replay |= (
                index > last_update and args["asset_id"] == pdfs[0]["asset_id"]
            )
            images += 1
    require(
        expected_pages <= seen and images >= 9 and frozen_replay,
        "Incomplete image review or historical replay",
    )
    require(
        historical and {(p["asset_id"], p["revision"]) for p in pdfs} <= verified,
        "Missing historical/file verification",
    )
    manifests = list((workspace / "wiki").glob("native-*/manifest.json"))
    require(len(manifests) == 1, "Wrong Wiki inventory")
    manifest, folder = read_json(manifests[0]), manifests[0].parent
    require(manifest["revision"] == pdfs[-1]["revision"], "Wrong Wiki revision")
    require(
        (folder / manifest["rendition"]["source_attachment"]).read_bytes() == updated,
        "Wiki source differs",
    )
    require(
        read_json(folder / "rendition.json") == pdfs[-1]["history"][0]["result"],
        "Wiki receipt differs",
    )
    for name, entry in manifest["files"].items():
        data = (folder / name).read_bytes()
        require(
            digest(data) == entry["sha256"] and len(data) == entry["size_bytes"],
            "Wiki file differs",
        )
    return {
        "passed": True,
        "successful_calls": len(calls),
        "tool_errors": tool_errors(events),
        "images": images,
        "title_pixel_proof": pixel_proof,
        "observations": final.get("observations"),
        "limitations": final.get("limitations"),
    }


def write_audit(output):
    try:
        report = audit(output)
    except (AssertionError, KeyError, ValueError, OSError) as exc:
        report = {"passed": False, "error": str(exc)}
    (output / "audit.json").write_text(json.dumps(report, indent=2, ensure_ascii=False))
    return report
