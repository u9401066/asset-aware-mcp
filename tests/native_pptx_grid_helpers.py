"""Independent PowerPoint fixtures and immutable table references for grid edits."""

import hashlib
import io

from pptx import Presentation

from src.application.native_pptx_operations import attach_pptx_evidence
from src.domain.native_pptx_grid import NativePptxTableGridEdit
from src.infrastructure.native_pptx import NativePresentation
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_table_helpers import table_addition


def grid_fixture(**kwargs):
    adapter = NativePresentation()
    data = build_presentation()
    item = table_addition(data, **kwargs)
    data, result = adapter.add_tables(data, [item])
    from src.domain.native_pptx import NativePptxShapeLocator

    locator = NativePptxShapeLocator(**result.changes[0]["locator"])
    return data, reference(data, locator), item


def reference(data, locator):
    record = NativePresentation().read_shape(data, locator)
    attach_pptx_evidence(record, "file_" + "a" * 32, hashlib.sha256(data).hexdigest())
    return record["evidence"]


def edit(data, ref, *edits):
    return NativePresentation().edit_table_grid(
        data, NativePptxTableGridEdit(reference=ref, edits=list(edits))
    )


def table(data):
    return Presentation(io.BytesIO(data)).slides[0].shapes[-1].table


def texts(data):
    target = table(data)
    return [
        [target.cell(r, c).text for c in range(len(target.columns))]
        for r in range(len(target.rows))
    ]
