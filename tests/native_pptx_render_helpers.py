"""Distinct slide backgrounds expose hidden-slide and reorder mapping errors."""

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches


def colored_deck() -> bytes:
    presentation = Presentation()
    for index, color in enumerate(((255, 0, 0), (0, 255, 0), (0, 0, 255))):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*color)
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(6), Inches(1)
        ).text = f"SLIDE {index + 1}"
        if index == 1:
            slide._element.set("show", "0")
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
