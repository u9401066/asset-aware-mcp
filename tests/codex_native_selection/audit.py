"""Independent trace/package/selection checks without the production selector code."""

import argparse
import io
import json
from pathlib import Path

from openpyxl import load_workbook

from tests.codex_native_pdf.artifacts import load_asset, read_json, validate_source
from tests.codex_native_pdf.trace import canonical, complete_records, digest, payload
from tests.codex_native_selection.records import selected_records
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import call_failed, require, tool_errors
from tests.codex_pptx_tables.audit import validate_images
from tests.codex_pptx_tables.derivations import complete_page


def validate_ledger_trace(calls, ledger):
    current = {**ledger, "events": []}
    chunks, observed, mutations, proofs = {}, set(), 0, 0
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        if args["op"] == "read_derivations":
            key = result["derivations_sha256"]
            require(
                key == digest(canonical(current)),
                "Ledger read differs from trace state",
            )
            if complete_page(result, key, chunks, digest_field="derivations_sha256"):
                observed.add(key)
        elif args["op"] == "record_derivation":
            before = digest(canonical(current))
            require(
                before in observed and args["expected_derivations_sha256"] == before,
                "Assertion lacks complete pinned ledger read",
            )
            claim = {"supersedes": None, **args["derivation"]}
            claim["review"] = {
                "semantic_accuracy": "not_checked",
                "rendered_layout": "not_checked",
                "formula_results": "not_checked",
                "notes": "",
                **claim.get("review", {}),
            }
            require(
                claim == ledger["events"][0]["derivation"],
                "Stored assertion differs from request",
            )
            current["events"].append(ledger["events"][0])
            mutations += 1
            require(
                result["derivations_sha256"] == digest(canonical(current)),
                "Append result hash mismatch",
            )
        elif args["op"] == "verify_derivation":
            require(
                result["active"]
                and result["references_valid"]
                and result["derivation_id"] == ledger["events"][0]["derivation_id"],
                "Invalid derivation proof",
            )
            proofs += 1
    require(
        mutations == 1 and proofs >= 1 and digest(canonical(ledger)) in observed,
        "Missing complete ledger/proof workflow",
    )


def audit(output):
    expected, final = (
        read_json(output / "expected.json"),
        read_json(output / "last-message.txt"),
    )
    workspace = output / "workspace"
    events = [
        json.loads(line)
        for line in (output / "events.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    require(any(e["type"] == "turn.completed" for e in events), "Incomplete Codex turn")
    calls, table_calls, ordered_calls = [], [], []
    for event in events:
        if event["type"] != "item.completed":
            continue
        item = event["item"]
        require(
            item["type"] in {"agent_message", "reasoning", "plan", "mcp_tool_call"},
            "Non-MCP action",
        )
        if item["type"] != "mcp_tool_call":
            continue
        require(
            item["server"] == "asset_aware_under_test"
            and (
                (item["tool"] == "document" and item["arguments"]["op"] == "native")
                or (expected.get("tables") and item["tool"] == "table_data")
            ),
            "Unexpected server/tool",
        )
        if not call_failed(item):
            ordered_calls.append(item)
            (table_calls if item["tool"] == "table_data" else calls).append(item)
    ops = {c["arguments"]["native_request"]["op"] for c in calls}
    require(
        {
            "contract",
            "create",
            "read_cell",
            "read_selection",
            "record_derivation",
            "verify_derivation",
            "apply_table_workspace" if expected.get("tables") else "update",
            "verify",
            "publish",
            "export_wiki",
            "history",
        }
        <= ops,
        "Incomplete workflow",
    )
    require("writeback" not in ops, "Unexpected writeback")
    source, target = (
        load_asset(workspace, final[key])
        for key in ("source_asset_id", "table_asset_id")
    )
    validate_source(workspace, expected, source)
    images = validate_images(calls, workspace, source)
    pages = complete_records(calls)
    records = selected_records(calls, pages)
    directory = workspace / "data" / "native-assets" / target["asset_id"]
    if expected.get("worksheets"):
        from tests.codex_native_workbook.audit import validate_workbook_history

        validate_workbook_history(calls, target, directory)
    else:
        require(len(target["history"]) == 2, "Unexpected workbook history")
        for index, item in enumerate(target["history"]):
            with_bytes = (directory / "revisions" / item["sha256"]).read_bytes()
            workbook = load_workbook(io.BytesIO(with_bytes))
            try:
                require(workbook.sheetnames == ["Sheet1"], "Wrong worksheet")
                values = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
                if index:
                    values[1][1] = "008"
                sheet = workbook["Sheet1"]
                require(
                    sheet.max_row == 3 and sheet.max_column == 5, "Wrong table bounds"
                )
                for row, expected_row in zip(sheet.iter_rows(), values, strict=True):
                    require(
                        [c.value for c in row] == expected_row
                        and all(c.data_type == "s" for c in row),
                        "Literal native values differ",
                    )
            finally:
                workbook.close()
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Published workbook mismatch",
    )
    ledger = read_json(directory / "derivations.json")
    require(len(ledger["events"]) == 1, "Unexpected assertion count")
    event = ledger["events"][0]
    claim = event["derivation"]
    require(
        event["derivation_id"] == digest(canonical(claim)), "Assertion hash mismatch"
    )
    require(
        canonical(claim["target"]) in records
        and claim["target"]["selector"]["pointer"] == "/value",
        "Missing precise target",
    )
    require(
        claim["target"]["revision"] == target["history"][0]["sha256"],
        "Assertion migrated to new revision",
    )
    require(
        len(claim["sources"]) == 1
        and canonical(claim["sources"][0]) in pages
        and claim["sources"][0]["asset_id"] == source["asset_id"]
        and claim["sources"][0]["revision"] == source["revision"]
        and claim["sources"][0]["locator"]["page_index"] == 0,
        "Wrong source endpoint",
    )
    require(
        claim["review"]["rendered_layout"] == "not_checked",
        "Unsupported layout verdict",
    )
    require(
        any(
            c["arguments"]["native_request"]["op"] == "verify"
            and payload(c).get("valid")
            and payload(c).get("is_current_managed_revision") is False
            and c["arguments"]["native_request"]["reference"] == claim["target"]
            for c in calls
        ),
        "Historical selected-value proof missing",
    )
    validate_ledger_trace(calls, ledger)
    validate_wikis(workspace, target, source, ledger, records)
    if expected.get("tables"):
        from tests.codex_native_table.audit import validate_table_workflow

        validate_table_workflow(ordered_calls, workspace, target)
    return {
        "passed": True,
        "mcp_calls": len(calls) + len(table_calls),
        "tool_errors": tool_errors(events),
        "images": images,
        "selection_records": len(records),
        "agent_reported_limitations": final.get("limitations", []),
        "scope": "Synthetic scan page to literal XLSX table; precise historical value evidence and two revision-specific wikis. No OCR or Excel rendering guarantee.",
        "worksheet_structure_evaluated": bool(expected.get("worksheets")),
        "native_table_workspaces_evaluated": bool(expected.get("tables")),
    }


def validate_wikis(workspace, target, source, ledger, records):
    manifests = list((workspace / "native-wiki").rglob("manifest.json"))
    require(len(manifests) == 2, "Expected historical and current wikis")
    active_count = 0
    for path in manifests:
        manifest = read_json(path)
        for name, metadata in manifest["files"].items():
            require(Path(name).name == name, "Unsafe wiki attachment path")
            data = (path.parent / name).read_bytes()
            require(
                digest(data) == metadata["sha256"]
                and len(data) == metadata["size_bytes"],
                "Wiki artifact hash mismatch",
            )
        info = manifest["derivations"]
        require(
            info["ledger_sha256"] == digest(canonical(ledger))
            and read_json(path.parent / info["ledger_file"]) == ledger,
            "Wiki ledger mismatch",
        )
        ids = info["active_ids_for_revision"]
        if ids:
            active_count += 1
            require(
                ids == [ledger["events"][0]["derivation_id"]]
                and manifest["revision"] != target["revision"],
                "Wrong active assertion revision",
            )
            require(len(info["selection_records"]) == 1, "Missing selection sidecar")
            for selected in info["selection_records"].values():
                require(
                    read_json(path.parent / selected["record_file"])
                    == records[canonical(selected["reference"])],
                    "Wiki selection record differs",
                )
            attachment = next(iter(info["source_attachments"].values()))
            require(
                digest((path.parent / attachment["attachment"]).read_bytes())
                == source["revision"],
                "Changed source attachment",
            )
        else:
            require(
                manifest["revision"] == target["revision"]
                and "selection_records" not in info,
                "Unexpected inherited selection",
            )
    require(active_count == 1, "Assertion inherited or lost")


def write_audit(output):
    try:
        report = audit(output)
    except Exception as exc:
        report = {"passed": False, "error": str(exc)}
    (output / "audit.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    report = write_audit(parser.parse_args().output.resolve())
    print(json.dumps(report, indent=2, ensure_ascii=False))
    raise SystemExit(0 if report["passed"] else 1)
