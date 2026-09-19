"""Independent ZIP/openpyxl and immutable Wiki checks for newly created Tables."""

import io
from pathlib import Path

from openpyxl import load_workbook

from tests.codex_native_pdf.artifacts import read_json
from tests.codex_native_pdf.trace import digest
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require
from tests.native_workbook_helpers import _parts


def check_packages(workspace, target):
    history = target["history"]
    require(
        len(history) == 2 and target["source"] is None,
        "Expected independent creation and one Table commit",
    )
    previous = None
    for index, entry in enumerate(history):
        data = (
            workspace
            / "data/native-assets"
            / target["asset_id"]
            / "revisions"
            / entry["sha256"]
        ).read_bytes()
        parts = _parts(data)
        book = load_workbook(io.BytesIO(data))
        try:
            require(book.sheetnames == ["Sheet1"], "Unexpected worksheet inventory")
            sheet = book["Sheet1"]
            for row, truth in zip(
                sheet.iter_rows(max_row=3, max_col=5),
                [COLUMNS, *PAGE_ROWS[0]],
                strict=True,
            ):
                require(
                    tuple(cell.value for cell in row) == truth
                    and all(cell.data_type == "s" for cell in row),
                    "Literal scan values/types differ",
                )
            require(sheet["F1"].value == "CountLength", "Calculated header differs")
            if index == 0:
                require(
                    not sheet.tables
                    and all(sheet[f"F{r}"].value is None for r in (2, 3, 4)),
                    "Baseline already has Table/calculations",
                )
            else:
                check_table(sheet, entry)
        finally:
            book.close()
        if previous is not None:
            require(
                parts["xl/styles.xml"] == previous["xl/styles.xml"],
                "Native cell styles changed",
            )
            require(
                all(
                    parts[name] == raw
                    for name, raw in previous.items()
                    if name not in entry["result"]["changed_parts"]
                ),
                "Untouched package bytes changed",
            )
        previous = parts
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Published workbook differs",
    )


def check_table(sheet, entry):
    require(set(sheet.tables) == {"Inventory"}, "Wrong native Table inventory")
    table = sheet.tables["Inventory"]
    require(
        table.ref == "A1:F4"
        and table.autoFilter.ref == "A1:F3"
        and table.totalsRowCount == 1,
        "Table/filter/totals range differs",
    )
    require(
        [c.id for c in table.tableColumns] == list(range(1, 7))
        and [c.name for c in table.tableColumns] == [*COLUMNS, "CountLength"],
        "Native column identities differ",
    )
    require(
        table.tableStyleInfo.name == "TableStyleMedium2"
        and table.tableStyleInfo.showRowStripes,
        "Requested Table style differs",
    )
    require(
        table.tableColumns[5].calculatedColumnFormula.text == "LEN([@Count])"
        and table.tableColumns[5].totalsRowFunction == "sum",
        "Formula metadata differs",
    )
    require(
        [sheet[f"F{r}"].value for r in (2, 3, 4)]
        == ["=LEN([@Count])", "=LEN([@Count])", "=SUBTOTAL(109,[CountLength])"],
        "Formula cells differ",
    )
    require(
        sheet["A4"].value == "Reviewed"
        and all(sheet.cell(4, c).value is None for c in (2, 3, 4, 5)),
        "Totals replaced ordinary data",
    )
    change = entry["result"]["changes"][0]
    require(
        change["operation"] == "add_workbook_table"
        and change["created_table"]["ref"] == "A1:F4",
        "Creation receipt missing",
    )
    require(set(change["cells"]) == {"A4", "F2", "F3", "F4"}, "Unexpected cell writes")
    require(
        all(cell["before"]["kind"] == "blank" for cell in change["cells"].values()),
        "Creation before-value receipt differs",
    )


def check_wikis(workspace, target):
    manifests = list((workspace / "native-wiki").rglob("manifest.json"))
    require(len(manifests) == 2, "Expected baseline/current Wiki snapshots")
    revisions = set()
    for path in manifests:
        manifest = read_json(path)
        require(manifest["asset_id"] == target["asset_id"], "Wrong Wiki asset")
        revisions.add(manifest["revision"])
        for name, metadata in manifest["files"].items():
            require(Path(name).name == name, "Unsafe Wiki attachment")
            data = (path.parent / name).read_bytes()
            require(
                digest(data) == metadata["sha256"]
                and len(data) == metadata["size_bytes"],
                "Wiki file hash differs",
            )
        require(
            any(
                digest((path.parent / name).read_bytes()) == manifest["revision"]
                for name in manifest["files"]
            ),
            "Missing exact native Wiki attachment",
        )
    require(
        revisions == {h["sha256"] for h in target["history"]}, "Wrong Wiki revisions"
    )
