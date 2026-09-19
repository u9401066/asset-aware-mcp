"""Independent four-page source with inherited sections and an obsolete definition."""

import hashlib
import io
import zipfile
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION_START
from lxml import etree

from tests.native_docx_stories_helpers import HEADER_PART, HEADER_TEXT, story_document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
TYPES = "{http://schemas.openxmlformats.org/package/2006/content-types}"
DOC_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
NEW_HEADER = "word/section-b/header.xml"
NEW_FOOTER = "word/section-b/footer.xml"
OBSOLETE = "word/obsolete-header.xml"
FOOTER_TEXT = "Section B verified / 007 µg"
CORRECTED = HEADER_TEXT.replace("SOURCE ", "SECTION B VERIFIED ", 1)


def source_document():
    doc = Document(io.BytesIO(story_document()))
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.different_first_page_header_footer = False
    doc.add_paragraph("Section C / keep original header and footer")
    buf = io.BytesIO()
    doc.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    old = next(
        name
        for name, raw in parts.items()
        if name.endswith(".xml") and HEADER_TEXT.encode() in raw
    )
    parts[HEADER_PART] = parts.pop(old)
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    for node in rels:
        if node.get("Target") == old.removeprefix("word/"):
            node.set("Target", "running/report.xml")
    etree.SubElement(
        rels,
        REL + "Relationship",
        Id="rIdObsolete",
        Type=DOC_REL + "header",
        Target="obsolete-header.xml",
    )
    types = etree.fromstring(parts["[Content_Types].xml"])
    for node in types:
        if node.get("PartName") == "/" + old:
            node.set("PartName", "/" + HEADER_PART)
    etree.SubElement(
        types,
        TYPES + "Override",
        PartName="/" + OBSOLETE,
        ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
    )
    root = etree.Element(W + "hdr", nsmap={"w": W[1:-1]})
    etree.SubElement(
        etree.SubElement(etree.SubElement(root, W + "p"), W + "r"), W + "t"
    ).text = "OBSOLETE DRAFT 099"
    parts[OBSOLETE] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
    parts["word/_rels/document.xml.rels"] = etree.tostring(
        rels, encoding="UTF-8", xml_declaration=True
    )
    parts["[Content_Types].xml"] = etree.tostring(
        types, encoding="UTF-8", xml_declaration=True
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, raw in parts.items():
            z.writestr(name, raw)
    return out.getvalue()


def binding(section, part, kind="header"):
    return {
        "op": "bind",
        "section_index": section,
        "story_kind": kind,
        "variant": "default",
        "part": part,
    }


def structure_edits(header, obsolete, footer):
    return [
        {
            "op": "clone",
            "source_part": HEADER_PART,
            "source_part_sha256": header["raw_part_sha256"],
            "part": NEW_HEADER,
        },
        {
            "op": "create",
            "part": NEW_FOOTER,
            "story_kind": "footer",
            "blocks": [
                {
                    "kind": "paragraph",
                    "runs": [
                        {"text": FOOTER_TEXT, "font_name": "Arial", "font_size_pt": 9.0}
                    ],
                }
            ],
        },
        binding(1, NEW_HEADER),
        binding(2, HEADER_PART),
        binding(1, NEW_FOOTER, "footer"),
        binding(2, footer, "footer"),
        {
            "op": "delete",
            "part": OBSOLETE,
            "expected_part_sha256": obsolete["raw_part_sha256"],
        },
    ]


def text_edit(record):
    node = next(n for n in record["text_nodes"] if n["text"] == HEADER_TEXT)
    return {
        "part": NEW_HEADER,
        "shared_scope": "all_sections_using_part",
        "edits": [{"op": "set_text", "path": node["path"], "text": CORRECTED}],
    }


def inspect_pages(data, stage):
    import pymupdf

    from tests.native_docx_layout_helpers import independent_pdf

    pdf_data = independent_pdf(data)
    with pymupdf.open(stream=pdf_data, filetype="pdf") as pdf:
        texts = [p.get_text() for p in pdf]
    assert len(texts) == 4, texts
    assert "FIRST PAGE ONLY" in texts[0] and "Page " not in texts[0]
    for index in [1, 3]:
        assert "SOURCE 007" in texts[index] and f"Page {index + 1}" in texts[index]
        assert (
            "SECTION B VERIFIED" not in texts[index] and FOOTER_TEXT not in texts[index]
        )
    assert ("SECTION B VERIFIED 007" if stage == 2 else "SOURCE 007") in texts[2]
    assert (FOOTER_TEXT if stage else "Page 3") in texts[2]
    for text in texts[1:]:
        assert "END-HEADER" in text and "KEEP-ITALIC" in text and "-0.50 mg/L" in text
    assert all("OBSOLETE DRAFT" not in t for t in texts)
    return pdf_data, texts


def assert_native_stages(original, intermediate, final):
    def parts(data):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return {n: z.read(n) for n in z.namelist()}

    before, mid, after = map(parts, [original, intermediate, final])
    assert mid[NEW_HEADER] == before[HEADER_PART]
    root = etree.fromstring(mid[NEW_HEADER])
    expected = deepcopy(root)
    next(n for n in expected.iter(W + "t") if n.text == HEADER_TEXT).text = CORRECTED
    assert etree.tostring(expected, method="c14n", exclusive=True) == etree.tostring(
        etree.fromstring(after[NEW_HEADER]), method="c14n", exclusive=True
    )
    assert (
        mid.keys()
        == after.keys()
        == (before.keys() - {OBSOLETE}) | {NEW_HEADER, NEW_FOOTER}
    )
    assert all(mid[n] == after[n] for n in mid if n != NEW_HEADER)
    allowed = {
        "word/document.xml",
        "word/_rels/document.xml.rels",
        "[Content_Types].xml",
        OBSOLETE,
    }
    assert all(after[n] == raw for n, raw in before.items() if n not in allowed)
    assert HEADER_TEXT.encode() in after[HEADER_PART]
    left, right = (
        etree.fromstring(before["word/document.xml"]),
        etree.fromstring(after["word/document.xml"]),
    )
    for root in [left, right]:
        for n in list(root.iter()):
            if n.tag in {W + "headerReference", W + "footerReference"}:
                n.getparent().remove(n)
    assert etree.tostring(left, method="c14n", exclusive=True) == etree.tostring(
        right, method="c14n", exclusive=True
    )
    old_types = {
        n.get("PartName"): dict(n.attrib)
        for n in etree.fromstring(before["[Content_Types].xml"])
        if n.tag == TYPES + "Override"
    }
    new_types = {
        n.get("PartName"): dict(n.attrib)
        for n in etree.fromstring(after["[Content_Types].xml"])
        if n.tag == TYPES + "Override"
    }
    assert all(
        new_types[n] == value for n, value in old_types.items() if n != "/" + OBSOLETE
    )
    assert new_types.keys() == (old_types.keys() - {"/" + OBSOLETE}) | {
        "/" + NEW_HEADER,
        "/" + NEW_FOOTER,
    }
    footer = etree.fromstring(after[NEW_FOOTER])
    assert len(footer) == 1 and footer[0].tag == W + "p"
    assert "".join(n.text or "" for n in footer.iter(W + "t")) == FOOTER_TEXT
    assert footer.find(".//" + W + "rFonts").get(W + "ascii") == "Arial"
    assert footer.find(".//" + W + "sz").get(W + "val") == "18"
    doc = Document(io.BytesIO(final))
    assert doc.sections[2].header.part.partname == doc.sections[0].header.part.partname
    assert doc.sections[1].header.part.partname != doc.sections[0].header.part.partname
    assert doc.sections[2].footer.part.partname == doc.sections[0].footer.part.partname
    assert FOOTER_TEXT in doc.sections[1].footer.paragraphs[0].text
    return {
        "source_sha256": hashlib.sha256(original).hexdigest(),
        "page_scope": "only section B changes",
    }
