"""Independent openpyxl and ZIP verification of Table membership and frozen assets."""

import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from tests.codex_native_pdf.artifacts import read_json
from tests.codex_native_pdf.trace import digest
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def check_packages(workspace, target, expected):
    history = target["history"]
    require(
        len(history) == 3,
        "Expected registration, transcription and one structural commit",
    )
    require(
        history[0]["sha256"] == expected["template_sha256"], "Wrong template revision"
    )
    template = workspace / "template.xlsx"
    require(
        digest(template.read_bytes()) == expected["template_sha256"]
        and template.stat().st_mtime_ns == expected["template_mtime_ns"],
        "Template source changed",
    )
    original_parts = None
    for index, entry in enumerate(history):
        data = (
            workspace
            / "data/native-assets"
            / target["asset_id"]
            / "revisions"
            / entry["sha256"]
        ).read_bytes()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        if original_parts is None:
            original_parts = parts
        else:
            require(
                parts["xl/styles.xml"] == original_parts["xl/styles.xml"],
                "Native styles changed",
            )
        if not index:
            continue
        book = load_workbook(io.BytesIO(data))
        try:
            sheet = book["Sheet1"]
            values = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
            if index == 2:
                values[1][1] = "008"
                values.append(["Added", "000", "0.0", "mg/L", "manual"])
            require(
                sheet.max_row == len(values)
                and sheet.max_column == (7 if index == 2 else 6),
                "Wrong native table dimensions",
            )
            for row, truth in zip(sheet.iter_rows(max_col=5), values, strict=True):
                require(
                    [cell.value for cell in row] == truth
                    and all(cell.data_type == "s" for cell in row),
                    "Literal scan/manual values or types differ",
                )
            native = sheet.tables["Inventory"]
            require(
                native.ref
                == native.autoFilter.ref
                == ("A1:G4" if index == 2 else "A1:F3"),
                "Table/filter membership differs",
            )
            require(
                [col.id for col in native.tableColumns]
                == list(range(1, 8 if index == 2 else 7)),
                "Native table column identity changed",
            )
            require(sheet["F1"].value == "CountLength", "Calculated header changed")
            for row in range(2, len(values) + 1):
                require(
                    sheet[f"F{row}"].value == "=LEN([[#This Row],Count])",
                    "Calculated formula differs",
                )
            if index == 2:
                require(
                    [sheet[f"G{row}"].value for row in range(1, 5)]
                    == ["Column7", "checked", "checked", "checked"],
                    "Generated header/manual values differ",
                )
        finally:
            book.close()
        if index == 2:
            previous = history[1]["sha256"]
            with zipfile.ZipFile(
                workspace
                / "data/native-assets"
                / target["asset_id"]
                / "revisions"
                / previous
            ) as archive:
                changed = set(entry["result"]["changed_parts"])
                require(
                    all(
                        parts[name] == archive.read(name)
                        for name in archive.namelist()
                        if name not in changed
                    ),
                    "Unmodified package member differs",
                )
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Published workbook mismatch",
    )


def check_wikis(workspace, target):
    manifests = list((workspace / "native-wiki").rglob("manifest.json"))
    require(len(manifests) == 2, "Expected baseline/current Wiki snapshots")
    revisions = set()
    for path in manifests:
        manifest = read_json(path)
        require(manifest["asset_id"] == target["asset_id"], "Wrong Wiki asset")
        revisions.add(manifest["revision"])
        for name, metadata in manifest["files"].items():
            require(Path(name).name == name, "Unsafe Wiki attachment name")
            data = (path.parent / name).read_bytes()
            require(
                digest(data) == metadata["sha256"]
                and len(data) == metadata["size_bytes"],
                "Wiki file differs from manifest",
            )
        require(
            any(
                digest((path.parent / name).read_bytes()) == manifest["revision"]
                for name in manifest["files"]
            ),
            "Wiki lacks the exact native file",
        )
    require(
        revisions == {h["sha256"] for h in target["history"][1:]},
        "Wrong Wiki revisions",
    )
