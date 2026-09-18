"""Reject unsafe native shape changes and preserve actual connected presentations."""

from __future__ import annotations

import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR

from src.infrastructure.native_ooxml import NativeOOXMLPackage, xml_bytes
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import NS, P_NS, NativePptxPackage
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_shape_helpers import addition, reference
from tests.native_workbook_helpers import _replace


@pytest.fixture
def data():
    return build_presentation()


def test_connected_shape_and_connector_can_be_deleted_together(data):
    presentation = Presentation(io.BytesIO(data))
    slide = presentation.slides[0]
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 100000, 100000)
    connector.name = "Dependent connector"
    connector.begin_connect(slide.shapes[0], 0)
    part = NativePptxPackage(data).slides[0]["part"]
    linked = _replace(data, {part: xml_bytes(slide.element)})
    adapter = NativePresentation()
    records = {record["name"]: record for record in adapter.iter_shapes(linked)}
    target = reference(linked, records["Styled text"])
    connection = reference(linked, records["Dependent connector"])
    with pytest.raises(ValueError, match="Surviving shape reference"):
        adapter.delete_shapes(linked, [target])
    updated, result = adapter.delete_shapes(linked, [connection, target])
    reopened = Presentation(io.BytesIO(updated))
    assert len(reopened.slides[0].shapes) == len(slide.shapes) - 2
    assert result.changed_parts == [part]
    before, after = NativePptxPackage(linked), NativePptxPackage(updated)
    assert all(
        before.parts[name] == after.parts[name] for name in before.parts if name != part
    )


@pytest.mark.parametrize("kind", ["ext", "chExt"])
def test_zero_extent_groups_are_rejected_without_reformatting(data, kind):
    package, request = NativePptxPackage(data), addition(data, grouped=True)
    part = request.container.part
    root = package.xml(part)
    root.find(f".//p:grpSp/p:grpSpPr/a:xfrm/a:{kind}", NS).set("cx", "0")
    with pytest.raises(ValueError, match="nonzero extents"):
        NativePresentation().add_shapes(
            _replace(data, {part: xml_bytes(root)}), [request]
        )


def test_shape_id_exhaustion_fails(data):
    package, request = NativePptxPackage(data), addition(data)
    part = request.container.part
    root = package.xml(part)
    root.find("p:cSld/p:spTree/p:nvGrpSpPr/p:cNvPr", NS).set("id", "4294967295")
    with pytest.raises(ValueError, match="exhausted"):
        NativePresentation().add_shapes(
            _replace(data, {part: xml_bytes(root)}), [request]
        )


def test_insert_before_extension_list_preserves_extension(data):
    package, request = NativePptxPackage(data), addition(data)
    root = package.xml(request.container.part)
    tree = root.find("p:cSld/p:spTree", NS)
    extension = etree.SubElement(tree, f"{{{P_NS}}}extLst")
    etree.SubElement(extension, f"{{{P_NS}}}ext", uri="test-preserve")
    extended = _replace(data, {request.container.part: xml_bytes(root)})
    updated, _ = NativePresentation().add_shapes(extended, [request])
    final = (
        NativePptxPackage(updated)
        .xml(request.container.part)
        .find("p:cSld/p:spTree", NS)
    )
    assert final[-1].tag == f"{{{P_NS}}}extLst"
    assert final[-2].find("p:txBody/a:p/a:r/a:t", NS).text == "新增 µ"


@pytest.mark.parametrize("operation", ["add", "delete"])
@pytest.mark.parametrize("guard", ["signed", "protected"])
def test_protected_packages_reject_both_structure_operations(data, operation, guard):
    package = NativePptxPackage(data)
    if guard == "signed":
        data = _replace(data, {"_xmlsignatures/test.xml": b"<signature/>"})
    else:
        root = package.presentation
        etree.SubElement(root, f"{{{P_NS}}}modifyVerifier")
        data = _replace(data, {package.main_part: xml_bytes(root)})
    adapter = NativePresentation()
    with pytest.raises(ValueError, match=r"signed|Protected"):
        if operation == "add":
            adapter.add_shapes(data, [addition(data)])
        else:
            adapter.delete_shapes(
                data, [reference(data, next(adapter.iter_shapes(data)))]
            )


@pytest.mark.parametrize("operation", ["add", "delete"])
@pytest.mark.parametrize("damage", ["unrelated_part", "unrelated_xml"])
def test_corrupted_serialization_is_rejected(data, monkeypatch, operation, damage):
    original = NativeOOXMLPackage.replace

    def corrupt(package, replacements, removed=None):
        candidate = original(package, replacements, removed)
        if damage == "unrelated_part":
            return _replace(candidate, {"customXml/preserve.xml": b"<changed/>"})
        part = next(iter(replacements))
        root = NativePptxPackage(candidate).xml(part)
        root.set("show", "0")
        return _replace(candidate, {part: xml_bytes(root)})

    monkeypatch.setattr(NativePptxPackage, "replace", corrupt)
    adapter = NativePresentation()
    with pytest.raises(ValueError, match="outside"):
        if operation == "add":
            adapter.add_shapes(data, [addition(data)])
        else:
            adapter.delete_shapes(
                data, [reference(data, next(adapter.iter_shapes(data)))]
            )


def test_addition_revalidates_aggregate_shape_limit(data, monkeypatch):
    package = NativePptxPackage(data)
    list(package.shapes())
    monkeypatch.setattr(
        "src.infrastructure.native_pptx_package.MAX_PPTX_COMPONENTS",
        package.component_count,
    )
    with pytest.raises(ValueError, match="shape limit"):
        NativePresentation().add_shapes(data, [addition(data)])
