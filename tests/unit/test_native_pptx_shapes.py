"""Shape-tree edits preserve real package features and reject broken dependencies."""

from __future__ import annotations

import hashlib
import io
import json

import pytest
from lxml import etree
from pptx import Presentation

from src.domain.native_pptx import NativePptxShapeLocator, shape_representation_sha256
from src.infrastructure.native_ooxml import xml_bytes
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import A_NS, P_NS, NativePptxPackage
from tests.native_pptx_helpers import build_presentation, edit_run
from tests.native_pptx_shape_helpers import addition, reference
from tests.native_workbook_helpers import _replace


@pytest.fixture
def data():
    return build_presentation()


def assert_other_parts(before, after, changed):
    old, new = NativePptxPackage(before), NativePptxPackage(after)
    assert set(old.parts) == set(new.parts)
    assert all(
        old.parts[name] == new.parts[name] for name in old.parts if name not in changed
    )
    return old, new


@pytest.mark.parametrize(
    "region,grouped", [("slide", False), ("slide", True), ("notes", False)]
)
def test_add_read_edit_delete_textboxes_preserves_package(data, region, grouped):
    adapter, request = (
        NativePresentation(),
        addition(data, region=region, grouped=grouped),
    )
    before_records = list(adapter.iter_shapes(data))
    updated, changes = adapter.add_shapes(data, [request])
    locator = NativePptxShapeLocator(**changes.changes[0]["locator"])
    record = adapter.read_shape(updated, locator)
    assert [item["text"] for item in record["paragraphs"][0]["items"]] == [
        "新增 µ",
        " preserved",
    ]
    assert 'b="1"' in record["paragraphs"][0]["items"][0]["xml"]
    assert 'sz="2400"' in record["paragraphs"][0]["items"][0]["xml"]
    assert_other_parts(data, updated, changes.changed_parts)
    assert len(list(adapter.iter_shapes(updated))) == len(before_records) + 1
    for previous in before_records:
        if (
            grouped
            and previous["locator"]["shape_id"] == request.container.group_shape_id
        ):
            continue
        assert (
            adapter.read_shape(updated, NativePptxShapeLocator(**previous["locator"]))
            == previous
        )
    Presentation(io.BytesIO(updated))
    updated, _ = adapter.edit(updated, [edit_run(record, "Edited new shape")])
    record = adapter.read_shape(updated, locator)
    assert record["paragraphs"][0]["items"][0]["text"] == "Edited new shape"
    restored, removed = adapter.delete_shapes(updated, [reference(updated, record)])
    assert removed.changes[0]["removed_components"] == 1
    original, final = assert_other_parts(data, restored, changes.changed_parts)
    for part in changes.changed_parts:
        assert etree.tostring(original.xml(part), method="c14n") == etree.tostring(
            final.xml(part), method="c14n"
        )
    Presentation(io.BytesIO(restored))


@pytest.mark.parametrize("kind", ["pic", "graphicFrame", "grpSp"])
def test_delete_existing_shapes_retains_related_media_and_charts(data, kind):
    adapter = NativePresentation()
    record = next(r for r in adapter.iter_shapes(data) if r["kind"] == kind)
    updated, result = adapter.delete_shapes(data, [reference(data, record)])
    assert_other_parts(data, updated, result.changed_parts)
    assert result.changes[0]["removed_components"] == (2 if kind == "grpSp" else 1)
    assert not any(
        r["locator"] == record["locator"] for r in adapter.iter_shapes(updated)
    )
    Presentation(io.BytesIO(updated))


def test_delete_requires_existing_exact_nonoverlapping_references(data):
    adapter = NativePresentation()
    records = list(adapter.iter_shapes(data))
    group = next(record for record in records if record["kind"] == "grpSp")
    child = next(record for record in records if record["group_path"])
    group_ref, child_ref = reference(data, group), reference(data, child)
    for refs, error in [
        ([group_ref, group_ref], "Duplicate"),
        ([group_ref, child_ref], "Overlapping"),
        ([group_ref.model_copy(update={"value_sha256": "0" * 64})], "Stale"),
    ]:
        with pytest.raises(ValueError, match=error):
            adapter.delete_shapes(data, refs)


@pytest.mark.parametrize(
    "element,attribute",
    [
        (f"{{{P_NS}}}spTgt", "spid"),
        (f"{{{P_NS}}}bldP", "spid"),
        (f"{{{A_NS}}}stCxn", "id"),
        (f"{{{A_NS}}}endCxn", "id"),
    ],
)
def test_surviving_shape_references_block_deletion(data, element, attribute):
    adapter, package = NativePresentation(), NativePptxPackage(data)
    record = next(r for r in adapter.iter_shapes(data) if r["name"] == "Styled text")
    part = record["locator"]["part"]
    root = package.xml(part)
    etree.SubElement(root, element, {attribute: record["locator"]["shape_id"]})
    linked = _replace(data, {part: xml_bytes(root)})
    ref = reference(
        linked, adapter.read_shape(linked, NativePptxShapeLocator(**record["locator"]))
    )
    with pytest.raises(ValueError, match="Surviving shape reference"):
        adapter.delete_shapes(linked, [ref])


def test_additions_do_not_adopt_dangling_shape_ids(data):
    package = NativePptxPackage(data)
    part = package.slides[0]["part"]
    root = package.xml(part)
    etree.SubElement(root, f"{{{P_NS}}}spTgt", spid="200")
    linked = _replace(data, {part: xml_bytes(root)})
    _, result = NativePresentation().add_shapes(
        linked, [addition(linked), addition(linked)]
    )
    assert [item["locator"]["shape_id"] for item in result.changes] == ["201", "202"]


def test_shape_v1_hash_remains_identical(data):
    record = next(NativePresentation().iter_shapes(data))
    legacy = {**record, "schema_version": "native-pptx-shape-v1"}
    expected = hashlib.sha256(
        json.dumps(
            legacy, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode()
    ).hexdigest()
    assert shape_representation_sha256(record) == expected
