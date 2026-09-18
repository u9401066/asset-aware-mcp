"""Native edits must preserve real workbook features and reject stale sources."""

from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import load_workbook

from src.application.native_document_service import NativeDocumentService
from src.domain.native_assets import (
    NativeCellEdit,
    NativeDocumentRequest,
    NativeWorkbookCreate,
)
from src.infrastructure.native_asset_store import FileNativeAssetRepository
from src.infrastructure.native_ooxml import SHEET_NS, NativeOOXMLPackage
from src.infrastructure.native_spreadsheet import (
    NativeSpreadsheet,
    SpreadsheetFileAdapter,
)
from tests.native_workbook_helpers import _call, _edit, _parts, _replace, build_workbook


@pytest.fixture
def workbook() -> bytes:
    return build_workbook()


def test_cell_update_preserves_all_unrelated_native_parts(workbook: bytes) -> None:
    before = _parts(workbook)
    result, report = NativeSpreadsheet(workbook).edit(
        [
            _edit("A1", "Changed"),
            _edit("A2", 42, "number"),
            _edit("E2", "=literal string"),
            _edit("F2", "=A2*3", "formula"),
            _edit("G2", True, "boolean"),
            _edit("A3", None, "blank"),
        ]
    )
    after = _parts(result)
    assert before.keys() == after.keys()
    assert report.changed_parts == [
        "xl/sharedStrings.xml",
        "xl/workbook.xml",
        "xl/worksheets/sheet1.xml",
    ]
    for name in before.keys() - set(report.changed_parts):
        assert before[name] == after[name], name
    assert any("charts/" in name for name in before)
    assert any("comments" in name for name in before)
    assert "xl/styles.xml" in before
    cells = {
        c["cell"]: c for c in NativeSpreadsheet(result).inspect(sheet="Data")["cells"]
    }
    assert cells["A1"]["value"] == "Changed"
    assert cells["A2"]["value"] == 42
    assert cells["A2"]["style_index"] == report.changes[1]["before"]["style_index"]
    assert cells["E2"]["kind"] == "string"
    assert cells["F2"]["kind"] == "formula"
    assert cells["F2"]["cached_value"] is None
    assert cells["B2"]["cached_value_verified"] is False
    assert cells["A3"]["kind"] == "blank"
    assert "formula_recalculation_requested" in report.repairs
    assert "rendered_layout" in report.review_required
    independently_read = load_workbook(io.BytesIO(result), data_only=False)
    sheet = independently_read["Data"]
    assert sheet["A1"].value == "Changed"
    assert sheet["A1"].font.bold is True
    assert sheet["A2"].number_format == "$0.00"
    assert sheet["A2"].value == 42
    assert sheet["E2"].data_type == "s"
    assert sheet["F2"].data_type == "f"
    assert sheet["A1"].comment.text == "Keep this comment"
    assert sheet["C3"].hyperlink.target == "https://example.com"
    assert "A5:D6" in sheet.merged_cells


@pytest.mark.parametrize(
    "edit, reason",
    [
        (_edit("D2", "flattened"), "Rich text"),
        (_edit("B5", "wrong"), "anchor"),
        (_edit("A1", "overwrite", sheet="Other"), "Protected"),
    ],
)
def test_unsupported_edit_fails_without_changing_source(
    workbook: bytes, edit: NativeCellEdit, reason: str
) -> None:
    parser = NativeSpreadsheet(workbook)
    with pytest.raises(ValueError, match=reason):
        parser.edit([edit])
    assert parser.package.original == workbook


def test_noop_is_whole_file_byte_identical(workbook: bytes) -> None:
    result, report = NativeSpreadsheet(workbook).edit([_edit("A1", "Original")])
    assert result == workbook
    assert report.changed_parts == []


def test_signed_package_cannot_be_silently_invalidated(workbook: bytes) -> None:
    signed = _replace(workbook, {"_xmlsignatures/sig1.xml": b"<signature/>"})
    with pytest.raises(ValueError, match="signed"):
        NativeSpreadsheet(signed).edit([_edit("A1", "Changed")])


def test_array_formula_follower_is_not_treated_as_ordinary_cell(
    workbook: bytes,
) -> None:
    parts = _parts(workbook)
    xml = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    formula = xml.find(f".//{{{SHEET_NS}}}f")
    assert formula is not None
    formula.set("t", "array")
    formula.set("ref", "B2:B4")
    data = _replace(workbook, {"xl/worksheets/sheet1.xml": etree.tostring(xml)})
    with pytest.raises(ValueError, match="range-aware"):
        NativeSpreadsheet(data).edit([_edit("B3", 5, "number")])


@pytest.mark.parametrize("address", ["A0", "a1", "XFE1", "A1048577", "A1:B2", "../A1"])
def test_cell_bounds(address: str) -> None:
    with pytest.raises(ValueError):
        _edit(address, "x")


@pytest.mark.parametrize(
    "value, kind",
    [
        (True, "number"),
        ("4", "number"),
        (float("inf"), "number"),
        (1, "boolean"),
        ("SUM(A1)", "formula"),
        ("\x00", "string"),
    ],
)
def test_typed_cell_values(value: object, kind: str) -> None:
    with pytest.raises(ValueError):
        _edit("A1", value, kind)


def test_dtd_is_rejected_without_entity_resolution(workbook: bytes) -> None:
    data = _replace(
        workbook,
        {
            "xl/workbook.xml": b'<!DOCTYPE workbook [<!ENTITY x SYSTEM "file:///etc/passwd">]><workbook xmlns="'
            + SHEET_NS.encode()
            + b'">&x;</workbook>'
        },
    )
    with pytest.raises(ValueError, match="DTD"):
        NativeSpreadsheet(data)


def test_duplicate_package_members_are_rejected() -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        with pytest.warns(UserWarning):
            archive.writestr("[Content_Types].xml", b"<Types/>")
    with pytest.raises(ValueError, match="Ambiguous"):
        NativeOOXMLPackage(output.getvalue())


@pytest.fixture
def service(tmp_path: Path) -> NativeDocumentService:
    return NativeDocumentService(
        FileNativeAssetRepository(tmp_path / "store"), SpreadsheetFileAdapter()
    )


def test_register_update_history_publish_writeback_and_archive(
    service: NativeDocumentService, workbook: bytes, tmp_path: Path
) -> None:
    source = tmp_path / "original.xlsx"
    source.write_bytes(workbook)
    asset = _call(service, op="register", source_path=str(source))["asset"]
    identity, original_revision = asset["asset_id"], asset["revision"]
    updated = _call(
        service,
        op="update",
        asset_id=identity,
        expected_revision=original_revision,
        edits=[{"sheet": "Data", "cell": "A1", "value": "Agent edit"}],
    )
    revision = updated["asset"]["revision"]
    assert updated["asset"]["asset_id"] == identity
    assert revision != original_revision
    assert source.read_bytes() == workbook
    original = _call(
        service, op="inspect", asset_id=identity, revision=original_revision
    )["content"]["cells"][0]
    current = _call(service, op="inspect", asset_id=identity)["content"]["cells"][0]
    assert original["value"] == "Original"
    assert current["value"] == "Agent edit"
    assert original["evidence"]["locator"] == current["evidence"]["locator"]
    assert original["evidence"]["value_sha256"] != current["evidence"]["value_sha256"]
    assert len(_call(service, op="history", asset_id=identity)["history"]) == 2
    with pytest.raises(ValueError, match="stale"):
        _call(
            service,
            op="update",
            asset_id=identity,
            expected_revision=original_revision,
            edits=[{"sheet": "Data", "cell": "A1", "value": "Lost update"}],
        )
    exported = tmp_path / "published.xlsx"
    _call(
        service,
        op="publish",
        asset_id=identity,
        expected_revision=revision,
        output_path=str(exported),
    )
    with pytest.raises(ValueError, match="new output"):
        _call(
            service,
            op="publish",
            asset_id=identity,
            expected_revision=revision,
            output_path=str(exported),
        )
    saved = _call(
        service,
        op="writeback",
        asset_id=identity,
        expected_revision=revision,
        expected_source_sha256=original_revision,
    )
    assert source.read_bytes() == exported.read_bytes()
    assert Path(saved["backup_path"]).read_bytes() == workbook
    _call(service, op="archive", asset_id=identity, expected_revision=revision)
    assert source.exists()
    assert _call(service, op="inspect", asset_id=identity)["asset"]["archived"] is True
    with pytest.raises(ValueError, match="Archived"):
        _call(
            service,
            op="update",
            asset_id=identity,
            expected_revision=revision,
            edits=[{"sheet": "Data", "cell": "A1", "value": "No"}],
        )


def test_stale_external_write_is_rejected_even_with_same_mtime(
    service: NativeDocumentService, workbook: bytes, tmp_path: Path
) -> None:
    source = tmp_path / "original.xlsx"
    source.write_bytes(workbook)
    asset = _call(service, op="register", source_path=str(source))["asset"]
    identity, revision = asset["asset_id"], asset["revision"]
    stamp = source.stat()
    external = bytearray(workbook)
    external[-1] ^= 1
    source.write_bytes(external)
    os.utime(source, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
    with pytest.raises(ValueError, match="externally"):
        _call(
            service,
            op="writeback",
            asset_id=identity,
            expected_revision=revision,
            expected_source_sha256=revision,
        )
    assert source.read_bytes() == external


def test_create_without_any_pdf_and_keep_formula_looking_text_literal(
    service: NativeDocumentService,
) -> None:
    asset = _call(
        service,
        op="create",
        workbook={
            "name": "new.xlsx",
            "sheets": ["Budget"],
            "edits": [
                {"sheet": "Budget", "cell": "A1", "value": "=1+1"},
                {"sheet": "Budget", "cell": "A2", "kind": "formula", "value": "=1+1"},
            ],
        },
    )["asset"]
    assert asset["source"] is None
    cells = _call(service, op="inspect", asset_id=asset["asset_id"])["content"]["cells"]
    assert cells[0]["kind"] == "string"
    assert cells[1]["kind"] == "formula"
    assert cells[1]["cached_value_verified"] is False


def test_unknown_format_is_addressable_without_fabricated_editor(
    service: NativeDocumentService, tmp_path: Path
) -> None:
    source = tmp_path / "handed-over.custom"
    source.write_bytes(b"binary\x00payload")
    asset = _call(service, op="register", source_path=str(source))["asset"]
    assert not asset["capabilities"]["edit_cells"]
    assert (
        _call(service, op="inspect", asset_id=asset["asset_id"])["content"][
            "representation"
        ]
        == "opaque_binary"
    )


def test_request_does_not_ignore_irrelevant_fields() -> None:
    with pytest.raises(ValueError, match="not used"):
        NativeDocumentRequest(op="contract", source_path="ignored")
    with pytest.raises(ValueError, match="unique"):
        NativeWorkbookCreate(sheets=["Sheet", "sheet"])


def test_table_header_changes_need_metadata_aware_operations(workbook: bytes) -> None:
    with pytest.raises(ValueError, match="Table headers"):
        NativeSpreadsheet(workbook).edit([_edit("A8", "Renamed")])
    result, _ = NativeSpreadsheet(workbook).edit([_edit("A9", 22, "number")])
    assert load_workbook(io.BytesIO(result))["Data"]["A9"].value == 22


def test_inconsistent_row_locator_is_rejected(workbook: bytes) -> None:
    data = _replace(
        workbook,
        {
            "xl/worksheets/sheet1.xml": _parts(workbook)[
                "xl/worksheets/sheet1.xml"
            ].replace(b'r="A2"', b'r="A4"', 1)
        },
    )
    with pytest.raises(ValueError, match="locator disagrees"):
        NativeSpreadsheet(data).edit([_edit("A4", "Wrong row")])


def test_extremely_large_integer_fails_as_validation_error() -> None:
    with pytest.raises(ValueError, match="finite"):
        _edit("A1", 10**1000, "number")


def test_failed_metadata_commit_keeps_current_revision(
    service: NativeDocumentService,
    workbook: bytes,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    path.write_bytes(workbook)
    asset = _call(service, op="register", source_path=str(path))["asset"]

    def fail(_asset: object) -> None:
        raise OSError("simulated metadata write failure")

    monkeypatch.setattr(service.repository, "_save", fail)
    with pytest.raises(OSError, match="metadata"):
        _call(
            service,
            op="update",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            edits=[{"sheet": "Data", "cell": "A1", "value": "New"}],
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    assert path.read_bytes() == workbook


def test_writeback_reports_partial_success_when_metadata_fails(
    service: NativeDocumentService,
    workbook: bytes,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "source.xlsx"
    path.write_bytes(workbook)
    asset = _call(service, op="register", source_path=str(path))["asset"]
    updated = _call(
        service,
        op="update",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        edits=[{"sheet": "Data", "cell": "A1", "value": "New"}],
    )["asset"]

    def fail(_asset: object) -> None:
        raise OSError("simulated metadata write failure")

    monkeypatch.setattr(service.repository, "_save", fail)
    result = _call(
        service,
        op="writeback",
        asset_id=asset["asset_id"],
        expected_revision=updated["revision"],
        expected_source_sha256=asset["revision"],
    )
    assert result["success"] is False
    assert result["source_written"] is True
    assert result["reconciliation_required"] is True
    assert Path(result["backup_path"]).read_bytes() == workbook
    assert load_workbook(path)["Data"]["A1"].value == "New"
