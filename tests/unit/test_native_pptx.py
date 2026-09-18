"""Native presentations preserve untouched structures under precise run edits."""

from __future__ import annotations

import io

import pytest
from lxml import etree
from pptx import Presentation

from src.domain.native_pptx import NativePptxShapeLocator, NativePresentationCreate
from src.infrastructure.native_ooxml import REL_NS
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import NS, P_NS, NativePptxPackage
from tests.native_pptx_helpers import build_presentation, edit_run, find_shape
from tests.native_workbook_helpers import _replace


@pytest.fixture
def presentation():
    return build_presentation()


def test_decomposition_exposes_runs_groups_tables_notes_and_native_xml(presentation):
    adapter = NativePresentation()
    info = adapter.inspect(presentation)
    assert info["slide_count"] == 2
    records = list(adapter.iter_shapes(presentation))
    assert info["shape_count"] == len(records)
    styled = find_shape(presentation, "Styled text")
    assert [item["text"] for item in styled["paragraphs"][0]["items"]] == [
        "原始",
        " unchanged",
    ]
    assert 'b="1"' in styled["xml"] and "hlinkClick" in styled["xml"]
    grouped = find_shape(presentation, "Grouped text")
    assert len(grouped["group_path"]) == 1
    assert "local_EMU" in grouped["coordinate_scope"]
    assert any(
        record["locator"]["region"] == "notes" and "備註" in record["xml"]
        for record in records
    )
    assert any("table" in record for record in records)
    assert any(record["kind"] == "pic" for record in records)
    assert (
        adapter.read_shape(presentation, NativePptxShapeLocator(**styled["locator"]))
        == styled
    )


def test_scoped_edit_preserves_styles_other_runs_and_all_unedited_parts(presentation):
    adapter = NativePresentation()
    record = find_shape(presentation, "Styled text")
    updated, report = adapter.edit(presentation, [edit_run(record, "修訂文字")])
    assert report.changed_parts == [record["locator"]["part"]]
    before, after = NativePptxPackage(presentation), NativePptxPackage(updated)
    assert before.parts.keys() == after.parts.keys()
    for name in before.parts:
        if name not in report.changed_parts:
            assert before.parts[name] == after.parts[name], name
    reopened = Presentation(io.BytesIO(updated))
    shape = next(
        shape for shape in reopened.slides[0].shapes if shape.name == "Styled text"
    )
    runs = shape.text_frame.paragraphs[0].runs
    assert runs[0].text == "修訂文字" and runs[0].font.bold is True
    assert runs[0].font.size.pt == 24
    assert runs[0].hyperlink.address == "https://example.com/evidence"
    assert runs[1].text == " unchanged" and runs[1].font.italic is True
    assert "rendered_layout" in report.review_required
    assert report.repairs == []


def test_table_group_and_notes_edits_reopen_without_structural_changes(presentation):
    adapter = NativePresentation()
    records = list(adapter.iter_shapes(presentation))
    table = next(record for record in records if "table" in record)
    notes = next(
        record
        for record in records
        if record["locator"]["region"] == "notes" and "備註" in record["xml"]
    )
    group = find_shape(presentation, "Grouped text")
    edits = [
        edit_run(table, "19", row=1, column=1),
        edit_run(notes, "Changed note"),
        edit_run(group, "Nested update"),
    ]
    updated, report = adapter.edit(presentation, edits)
    assert len(report.changed_parts) == 2
    reopened = Presentation(io.BytesIO(updated))
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "Changed note"
    table_shape = next(shape for shape in reopened.slides[0].shapes if shape.has_table)
    assert table_shape.table.cell(1, 1).text == "19"
    assert (
        table_shape.table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.bold is True
    )
    assert "Nested update" in find_shape(updated, "Grouped text")["xml"]


def test_noop_is_byte_identical_and_stale_or_duplicate_runs_fail(presentation):
    adapter = NativePresentation()
    record = find_shape(presentation, "Styled text")
    edit = edit_run(record, "原始")
    updated, report = adapter.edit(presentation, [edit])
    assert updated == presentation and report.changed_parts == []
    with pytest.raises(ValueError, match="Duplicate"):
        adapter.edit(presentation, [edit, edit])
    stale = edit.model_copy(update={"expected_text_sha256": "0" * 64})
    with pytest.raises(ValueError, match="Stale"):
        adapter.edit(presentation, [stale])


@pytest.mark.parametrize(
    "mutation", ["part", "slide_id", "shape_id", "region", "paragraph", "run", "row"]
)
def test_invalid_or_mismatched_locators_never_guess_a_target(presentation, mutation):
    edit = edit_run(find_shape(presentation, "Styled text"), "Changed")
    values = {
        "part": "ppt/slides/slide2.xml",
        "slide_id": "9999",
        "shape_id": "9999",
        "region": "notes",
        "paragraph": 99,
        "run": 99,
        "row": 1,
    }
    locator = edit.locator.model_copy(update={mutation: values[mutation]})
    with pytest.raises(ValueError):
        NativePresentation().edit(
            presentation, [edit.model_copy(update={"locator": locator})]
        )


@pytest.mark.parametrize(
    "kind", ["signature", "protection", "duplicate_shape", "duplicate_slide"]
)
def test_protection_signatures_and_ambiguous_identity_fail_closed(presentation, kind):
    package = NativePptxPackage(presentation)
    part = package.main_part
    root = package.xml(part)
    if kind == "signature":
        part = "_rels/.rels"
        root = package.xml(part)
        etree.SubElement(
            root,
            f"{{{REL_NS}}}Relationship",
            Id="relocatedSignature",
            Type=f"{REL_NS}/digital-signature/origin",
            Target="signatures/custom-origin",
        )
    elif kind == "protection":
        etree.SubElement(root, f"{{{P_NS}}}modifyVerifier")
    elif kind == "duplicate_slide":
        ids = root.findall("p:sldIdLst/p:sldId", NS)
        ids[1].set("id", ids[0].get("id"))
    else:
        part = package.slides[0]["part"]
        root = package.xml(part)
        ids = root.findall(".//p:cNvPr", NS)
        ids[1].set("id", ids[0].get("id"))
    changed = _replace(presentation, {part: etree.tostring(root)})
    edit = edit_run(find_shape(presentation, "Styled text"), "Changed")
    with pytest.raises(ValueError):
        NativePresentation().edit(changed, [edit])


def test_slide_identity_follows_relationships_after_reordering(presentation):
    package = NativePptxPackage(presentation)
    root = package.xml(package.main_part)
    slide_list = root.find("p:sldIdLst", NS)
    first = slide_list[0]
    slide_list.remove(first)
    slide_list.append(first)
    reordered = _replace(presentation, {package.main_part: etree.tostring(root)})
    before = find_shape(presentation, "Styled text")
    after = find_shape(reordered, "Styled text")
    assert after["locator"] == before["locator"]
    updated, _ = NativePresentation().edit(reordered, [edit_run(after, "Reordered")])
    assert "Reordered" in Presentation(io.BytesIO(updated)).slides[1].shapes[0].text


def test_independent_native_creation_preserves_explicit_runs_and_dimensions():
    request = NativePresentationCreate.model_validate(
        {
            "name": "new.pptx",
            "slides": [
                {
                    "textboxes": [
                        {
                            "paragraphs": [
                                [{"text": "Bold", "bold": True}, {"text": " normal"}],
                                [{"text": "Second"}],
                            ]
                        }
                    ],
                    "notes": "Notes",
                }
            ],
        }
    )
    data = NativePresentation().create(request)
    presentation = Presentation(io.BytesIO(data))
    assert presentation.slide_width == request.width
    shape = presentation.slides[0].shapes[0]
    assert shape.text_frame.paragraphs[0].runs[0].font.bold is True
    assert shape.text == "Bold normal\nSecond"
    assert presentation.slides[0].notes_slide.notes_text_frame.text == "Notes"


def test_zero_root_identity_and_field_break_runs_are_retained(presentation):
    package = NativePptxPackage(presentation)
    part = package.slides[0]["part"]
    root = package.xml(part)
    root.find("p:cSld/p:spTree/p:nvGrpSpPr/p:cNvPr", NS).set("id", "0")
    paragraph = root.find("p:cSld/p:spTree/p:sp/p:txBody/a:p", NS)
    field = etree.Element(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}fld",
        id="keep",
        type="slidenum",
    )
    etree.SubElement(
        field, "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
    ).text = "1"
    paragraph.insert(1, field)
    paragraph.insert(
        2, etree.Element("{http://schemas.openxmlformats.org/drawingml/2006/main}br")
    )
    data = _replace(presentation, {part: etree.tostring(root)})
    record = find_shape(data, "Styled text")
    assert [item["kind"] for item in record["paragraphs"][0]["items"]] == [
        "r",
        "fld",
        "br",
        "r",
    ]
    updated, _ = NativePresentation().edit(
        data, [edit_run(record, "Last regular run", run=1)]
    )
    changed = find_shape(updated, "Styled text")
    assert [item["text"] for item in changed["paragraphs"][0]["items"]] == [
        "原始",
        "1",
        "\n",
        "Last regular run",
    ]


def test_preservation_guard_detects_mutation_inside_a_changed_part(
    presentation, monkeypatch
):
    from src.infrastructure.native_ooxml import NativeOOXMLPackage

    original_replace = NativeOOXMLPackage.replace

    def corrupt(package, replacements, removed=None):
        data = original_replace(package, replacements, removed)
        part = next(iter(replacements))
        root = NativePptxPackage(data).xml(part)
        root.set("show", "0")
        return _replace(data, {part: etree.tostring(root)})

    monkeypatch.setattr(NativePptxPackage, "replace", corrupt)
    with pytest.raises(ValueError, match="outside the requested"):
        NativePresentation().edit(
            presentation, [edit_run(find_shape(presentation, "Styled text"), "Changed")]
        )
