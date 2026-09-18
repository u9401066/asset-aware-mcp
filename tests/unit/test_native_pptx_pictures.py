"""Pictures preserve source bytes, package contents and shared-image mappings."""

from __future__ import annotations

import io

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation

from src.domain.native_pptx import NativePptxShapeLocator
from src.domain.native_pptx_picture import NativePptxPictureReplace
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import NS, NativePptxPackage
from src.infrastructure.native_pptx_pictures import (
    add_pictures,
    read_picture,
    replace_pictures,
)
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_picture_helpers import file_ref, raster, request, sources
from tests.native_pptx_shape_helpers import reference


@pytest.fixture
def data():
    return build_presentation()


def unchanged_parts(before, after, changed):
    original, actual = NativePptxPackage(before), NativePptxPackage(after)
    assert all(
        actual.parts[k] == v for k, v in original.parts.items() if k not in changed
    )
    assert set(original.parts) <= set(actual.parts)


@pytest.mark.parametrize("format", ["PNG", "JPEG"])
@pytest.mark.parametrize("fit", ["contain", "cover", "stretch"])
def test_add_picture_exact_bytes_fit_and_readback(data, format, fit):
    image = raster(format)
    updated, report = add_pictures(
        data, [request(data, image, fit=fit)], sources(image)
    )
    locator = NativePptxShapeLocator(**report.changes[0]["locator"])
    read = read_picture(updated, locator, 512)
    assert read["image_bytes"] == image
    assert report.changes[0]["source_reference"] == file_ref(image).model_dump()
    with Image.open(io.BytesIO(read["image_png"])) as preview:
        assert preview.size == (80, 40)
    shape = Presentation(io.BytesIO(updated)).slides[0].shapes[-1]
    assert shape.image.blob == image and shape.name == "Evidence picture"
    assert shape.left == 914400
    assert shape.height == (914400 if fit == "contain" else 1828800)
    assert shape.top == (1371600 if fit == "contain" else 914400)
    assert shape.crop_left == (0.25 if fit == "cover" else 0)
    assert shape.crop_right == shape.crop_left
    unchanged_parts(data, updated, report.changed_parts)


@pytest.mark.parametrize("region,grouped", [("notes", False), ("slide", True)])
def test_add_picture_in_existing_container_preserves_transforms(data, region, grouped):
    image = raster()
    item = request(data, image, region=region, grouped=grouped)
    before = NativePptxPackage(data).xml(item.container.part)
    updated, report = add_pictures(data, [item], sources(image))
    after = NativePptxPackage(updated).xml(item.container.part)
    assert [etree.tostring(e) for e in before.findall(".//p:grpSpPr", NS)] == [
        etree.tostring(e) for e in after.findall(".//p:grpSpPr", NS)
    ]
    locator = NativePptxShapeLocator(**report.changes[0]["locator"])
    assert read_picture(updated, locator)["image_bytes"] == image
    unchanged_parts(data, updated, report.changed_parts)


def test_replacement_isolates_shared_media_and_keeps_mapping(data):
    old, new = raster(), raster("JPEG", "blue", (40, 80))
    initial, added = add_pictures(
        data, [request(data, old, fit="cover")] * 2, sources(old)
    )
    locators = [NativePptxShapeLocator(**c["locator"]) for c in added.changes]
    assert added.changes[0]["media_part"] == added.changes[1]["media_part"]
    first = NativePresentation().read_shape(initial, locators[0])
    edit = NativePptxPictureReplace(
        reference=reference(initial, first), image=file_ref(new)
    )
    updated, report = replace_pictures(initial, [edit], sources(new))
    assert read_picture(updated, locators[0])["image_bytes"] == new
    assert read_picture(updated, locators[1])["image_bytes"] == old
    shapes = Presentation(io.BytesIO(updated)).slides[0].shapes
    assert shapes[-2].image.blob == new and shapes[-1].image.blob == old
    assert shapes[-2].crop_left == shapes[-1].crop_left == 0.25
    assert report.changes[0]["mapping"] == "preserve_existing"
    unchanged_parts(initial, updated, report.changed_parts)
    current = NativePresentation().read_shape(updated, locators[0])
    deleted, removed = NativePresentation().delete_shapes(
        updated, [reference(updated, current)]
    )
    unchanged_parts(updated, deleted, removed.changed_parts)
    assert NativePptxPackage(deleted).parts[added.changes[0]["media_part"]] == old
    assert NativePptxPackage(deleted).parts[report.changes[0]["media_part"]] == new


def test_stale_and_duplicate_picture_replacement(data):
    image = raster()
    initial, report = add_pictures(data, [request(data, image)], sources(image))
    locator = NativePptxShapeLocator(**report.changes[0]["locator"])
    ref = reference(initial, NativePresentation().read_shape(initial, locator))
    edit = NativePptxPictureReplace(reference=ref, image=file_ref(image))
    with pytest.raises(ValueError, match="Duplicate"):
        replace_pictures(initial, [edit, edit], sources(image))
    for field in ["revision", "value_sha256"]:
        stale = edit.model_copy(
            update={"reference": ref.model_copy(update={field: "0" * 64})}
        )
        with pytest.raises(ValueError, match="Stale"):
            replace_pictures(initial, [stale], sources(image))
