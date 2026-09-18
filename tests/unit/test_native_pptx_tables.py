"""Independent table shape/grid/text/merge readback and existing CRUD preservation."""

from __future__ import annotations

import hashlib
import io

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from src.domain.native_pptx import NativePptxShapeLocator, NativePptxTextEdit
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import NS, NativePptxPackage
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_shape_helpers import reference
from tests.native_pptx_table_helpers import table_addition


def independent_shape(data, region, grouped):
    slide = Presentation(io.BytesIO(data)).slides[0]
    shapes = slide.notes_slide.shapes if region == "notes" else slide.shapes
    if grouped:
        shapes = next(s for s in shapes if hasattr(s, "shapes")).shapes
    return shapes[-1]


@pytest.mark.parametrize(
    "region,grouped", [("slide", False), ("notes", False), ("slide", True)]
)
@pytest.mark.parametrize("merge", ["horizontal", "vertical", "rectangle", None])
def test_table_grid_styles_merges_and_untouched_parts(region, grouped, merge):
    data = build_presentation()
    item = table_addition(data, region=region, grouped=grouped, merge=merge)
    updated, result = NativePresentation().add_tables(data, [item])
    before, after = NativePptxPackage(data), NativePptxPackage(updated)
    assert set(before.parts) == set(after.parts)
    for part in before.parts:
        if part != item.container.part:
            assert before.parts[part] == after.parts[part]
    shape = independent_shape(updated, region, grouped)
    table = shape.table
    assert [c.width for c in table.columns] == item.table.column_widths
    assert [r.height for r in table.rows] == item.table.row_heights
    assert (shape.left, shape.top, shape.width, shape.height) == (
        100000,
        200000,
        5000000,
        1800000,
    )
    assert shape.name == "原生表格" and table.first_col and table.vert_banding
    anchor = table.cell(0, 0)
    assert anchor.text == "樣本 & 結果" and str(anchor.fill.fore_color.rgb) == "204060"
    assert anchor.vertical_anchor == MSO_ANCHOR.TOP and anchor.margin_left == 10000
    assert anchor.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert str(anchor.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFEEDD"
    assert table.cell(1, 2).text == "-0.50\nmg/L"
    assert table.cell(2, 2).text == "=SUM(A1:A2)"
    if merge:
        spec = item.table.merges[0]
        assert anchor.is_merge_origin
        assert (anchor.span_height, anchor.span_width) == (
            spec.end_row + 1,
            spec.end_column + 1,
        )
        assert table.cell(spec.end_row, spec.end_column).is_spanned
    assert result.changes[0]["operation"] == "add_table"


def test_added_table_runs_edit_and_delete_without_resaving_source():
    original = build_presentation()
    backend = NativePresentation()
    data, result = backend.add_tables(original, [table_addition(original)])
    locator = NativePptxShapeLocator(**result.changes[0]["locator"])
    edit = NativePptxTextEdit(
        locator={
            **locator.model_dump(),
            "row": 1,
            "column": 1,
            "paragraph": 0,
            "run": 0,
        },
        expected_text_sha256=hashlib.sha256(b"007").hexdigest(),
        text="008",
    )
    edited, _ = backend.edit(data, [edit])
    table = independent_shape(edited, "slide", False).table
    assert table.cell(1, 1).text == "008" and table.cell(0, 0).is_merge_origin
    ref = reference(edited, backend.read_shape(edited, locator))
    deleted, _ = backend.delete_shapes(edited, [ref])
    before, after = NativePptxPackage(original), NativePptxPackage(deleted)
    assert before.parts.keys() == after.parts.keys()
    for name, raw in before.parts.items():
        if name != locator.part:
            assert after.parts[name] == raw
    assert len(Presentation(io.BytesIO(deleted)).slides[0].shapes) == len(
        Presentation(io.BytesIO(original)).slides[0].shapes
    )
    span_edit = edit.model_copy(
        update={"locator": edit.locator.model_copy(update={"row": 0, "column": 1})}
    )
    with pytest.raises(ValueError, match="anchor"):
        backend.edit(data, [span_edit])


def test_destination_custom_default_style_is_preserved():
    from src.infrastructure.native_ooxml import xml_bytes
    from tests.native_workbook_helpers import _replace

    data = build_presentation()
    package = NativePptxPackage(data)
    styles = package.xml("ppt/tableStyles.xml")
    style_id = "{01234567-89AB-CDEF-0123-456789ABCDEF}"
    styles.set("def", style_id)
    data = _replace(data, {"ppt/tableStyles.xml": xml_bytes(styles)})
    updated, result = NativePresentation().add_tables(data, [table_addition(data)])
    package = NativePptxPackage(updated)
    locator = NativePptxShapeLocator(**result.changes[0]["locator"])
    assert (
        package.locate(locator).findtext(
            "a:graphic/a:graphicData/a:tbl/a:tblPr/a:tableStyleId", namespaces=NS
        )
        == style_id
    )
    assert package.parts["ppt/tableStyles.xml"] == xml_bytes(styles)
    assert result.changes[0]["table_style_id"] == style_id
