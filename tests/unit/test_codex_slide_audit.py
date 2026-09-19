"""Independent slide audit rejects drift in content, format, IDs and retained XML."""

import io

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Pt

from tests.codex_pptx_tables.slides import (
    APP,
    MAIN,
    NS,
    TEXTS,
    check_slide_stage,
    parts,
    slide_keys,
)
from tests.native_workbook_helpers import _replace


def save(deck):
    output = io.BytesIO()
    deck.save(output)
    data = output.getvalue()
    node = etree.fromstring(parts(data)[APP])
    node.find(
        "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Slides"
    ).text = str(len(deck.slides))
    return _replace(data, {APP: etree.tostring(node)})


def fixture():
    deck = Presentation()
    original = deck.slides.add_slide(deck.slide_layouts[6])
    original.shapes.add_textbox(100000, 100000, 6000000, 550000).text = "Keep original"
    before = save(deck)
    for index, text in enumerate(TEXTS):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(100000, 100000, 6000000, 550000)
        shape.text = text
        font = shape.text_frame.paragraphs[0].runs[0].font
        font.size, font.bold, font.italic = Pt(14), index == 0, index == 1
    return before, deck


@pytest.mark.parametrize(
    "failure",
    [None, "zero", "style", "geometry", "source", "metadata", "content_type", "count"],
)
def test_independent_slide_insertion_audit(failure):
    before, deck = fixture()
    if failure == "zero":
        deck.slides[1].shapes[0].text_frame.paragraphs[0].runs[
            0
        ].text = "TEMP SLIDE 7 µg"
    elif failure == "style":
        deck.slides[1].shapes[0].text_frame.paragraphs[0].runs[0].font.bold = False
    elif failure == "geometry":
        deck.slides[1].shapes[0].left += 1
    elif failure == "source":
        deck.slides[0].shapes[0].text = "Changed original"
    after = save(deck)
    original = slide_keys(before)[0]
    created = slide_keys(after)[1:]
    package = parts(after)
    if failure == "metadata":
        root = etree.fromstring(package[MAIN])
        root.set("showSpecialPlsOnTitleSld", "0")
        after = _replace(after, {MAIN: etree.tostring(root)})
    elif failure == "content_type":
        root = etree.fromstring(package["[Content_Types].xml"])
        root[0].set("ContentType", "application/octet-stream")
        after = _replace(after, {"[Content_Types].xml": etree.tostring(root)})
    elif failure == "count":
        root = etree.fromstring(package[APP])
        root.find(
            "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Slides"
        ).text = "2"
        after = _replace(after, {APP: etree.tostring(root)})
    if failure:
        with pytest.raises(ValueError):
            check_slide_stage(before, after, 1, original, created)
    else:
        check_slide_stage(before, after, 1, original, created)


def test_independent_slide_reorder_audit_rejects_wrong_identity_order():
    before, deck = fixture()
    added = save(deck)
    original, created = slide_keys(before)[0], slide_keys(added)[1:]
    root = etree.fromstring(parts(added)[MAIN])
    slides = root.find("p:sldIdLst", NS)
    nodes = list(slides)
    for n in nodes:
        slides.remove(n)
    for index in (2, 0, 1):
        slides.append(nodes[index])
    reordered = _replace(added, {MAIN: etree.tostring(root)})
    check_slide_stage(added, reordered, 2, original, created)
    with pytest.raises(ValueError, match="identity/order"):
        check_slide_stage(added, added, 2, original, created)
