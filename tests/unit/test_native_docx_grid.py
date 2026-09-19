"""Independent Word/XML checks for grid topology and native content survival."""

from __future__ import annotations

import hashlib
import io
from copy import deepcopy

import pytest
from docx import Document
from lxml import etree

from src.domain.native_docx_grid import NativeDocxTableGridEdit
from src.infrastructure.native_docx_grid import edit_table, read_table
from src.infrastructure.native_docx_grid_model import W, WordGrid
from src.infrastructure.native_docx_structure import NativeDocxStructure, canonical
from src.infrastructure.native_docx_workspace import MAIN_PART, checked_docx
from src.infrastructure.native_ooxml import xml_bytes
from tests.native_docx_structure_helpers import creation

CHAIN = [
    {
        "source_part": MAIN_PART,
        "source_story": "body",
        "source_element": "w:tbl",
        "table_index": 0,
        "source_order": 1,
    }
]


def source():
    return NativeDocxStructure().create(creation())


def request(data, *edits):
    return NativeDocxTableGridEdit.model_validate(
        {
            "reference": {
                "schema_version": "native-docx-block-ref-v1",
                "asset_id": "file_" + "a" * 32,
                "revision": hashlib.sha256(data).hexdigest(),
                "locator": {"block_id": "t001", "part": MAIN_PART},
                "value_sha256": "b" * 64,
                "verification_scope": "immutable_native_representation",
            },
            "edits": list(edits),
        }
    )


def transform(data, callback):
    package = checked_docx(data)
    root = package.xml(MAIN_PART)
    callback(root.find(W + "body")[1])
    return package.replace({MAIN_PART: xml_bytes(root)})


def run(data, *edits, chain=None):
    return edit_table(data, chain or CHAIN, request(data, *edits))


def test_read_distinguishes_physical_cells_and_merged_coverage():
    record = read_table(source(), CHAIN)
    assert (record["rows"], record["columns"]) == (3, 3)
    assert record["repeat_header_rows"] == [0]
    assert record["regions"][0]["col_span"] == 3
    assert record["regions"][1]["row_span"] == 2
    assert len(record["regions"][1]["physical_cells"]) == 2
    assert "vMerge" in record["native_xml"] and "szCs" in record["native_xml"]


def test_insert_inside_merges_preserves_rich_native_content_and_package():
    before = source()
    updated, result = run(
        before,
        {"op": "insert", "axis": "row", "index": 2, "sizes_twips": [420]},
        {"op": "insert", "axis": "column", "index": 1, "sizes_twips": [700]},
    )
    document = Document(io.BytesIO(updated))
    table = document.tables[0]
    assert len(table.rows) == 4 and len(table.columns) == 4
    assert table.cell(3, 0).text == "007"
    assert table.cell(0, 3).text == "Header"
    assert table.cell(1, 2).text == "-0.50\nSecond"
    assert table.cell(1, 2).paragraphs[0].runs[0].italic
    assert table.rows[2].height.twips == 420
    assert [c.width.twips for c in table.columns] == [1200, 700, 1800, 2400]
    old, new = checked_docx(before), checked_docx(updated)
    assert old.parts.keys() == new.parts.keys()
    assert all(new.parts[k] == v for k, v in old.parts.items() if k != MAIN_PART)
    assert result.changes[0]["after_rows_columns"] == [4, 4]


def test_delete_promotes_anchor_and_keeps_surviving_content():
    updated, result = run(
        source(), {"op": "delete", "axis": "row", "index": 1, "count": 1}
    )
    table = Document(io.BytesIO(updated)).tables[0]
    assert table.cell(1, 0).text == "007"
    assert table.cell(1, 0).paragraphs[0].runs[0].font.size.pt == 12
    assert table.cell(1, 1).text == "1,234.50"
    assert result.changes[0]["edits"][0]["promoted_merge_anchors"] == 1
    assert table.cell(1, 0)._tc.tcPr.find(W + "vMerge") is None


def test_delete_column_shrinks_horizontal_merge_and_preserves_anchor():
    updated, _ = run(
        source(), {"op": "delete", "axis": "column", "index": 0, "count": 1}
    )
    table = Document(io.BytesIO(updated)).tables[0]
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == table.cell(0, 1).text == "Header"
    assert table.cell(1, 0).text == "-0.50\nSecond"


def test_merge_append_and_split_keep_rich_paragraphs_and_do_not_guess_content():
    updated, result = run(
        source(),
        {
            "op": "merge",
            "row": 1,
            "column": 1,
            "end_row": 2,
            "end_column": 2,
            "content_policy": "append_blocks",
        },
        {"op": "split", "row": 1, "column": 1},
    )
    table = Document(io.BytesIO(updated)).tables[0]
    assert table.cell(1, 1).text == "-0.50\nSecond\nmg/L\n1,234.50\nHIGH"
    assert table.cell(1, 1).paragraphs[0].runs[0].italic
    assert all(table.cell(r, c).text == "" for r, c in [(1, 2), (2, 1), (2, 2)])
    assert result.changes[0]["edits"][0]["moved_blocks"] == 3


def test_insert_typed_cells_and_resize_dimensions():
    updated, _ = run(
        source(),
        {
            "op": "insert",
            "axis": "row",
            "index": 3,
            "sizes_twips": [350],
            "cells": [
                [
                    {"paragraphs": [{"runs": [{"text": "新值 0007", "bold": True}]}]},
                    {},
                    {},
                ]
            ],
        },
        {"op": "resize", "axis": "row", "index": 3, "sizes_twips": [720]},
        {"op": "resize", "axis": "column", "index": 1, "sizes_twips": [2100, 2300]},
    )
    table = Document(io.BytesIO(updated)).tables[0]
    assert (
        table.cell(3, 0).text == "新值 0007"
        and table.cell(3, 0).paragraphs[0].runs[0].bold
    )
    assert table.rows[3].height.twips == 720
    assert table.autofit is False
    assert table.cell(0, 0).width.twips == 5600


@pytest.mark.parametrize(
    "edit,match",
    [
        (
            {
                "op": "merge",
                "row": 1,
                "column": 1,
                "end_row": 2,
                "end_column": 2,
                "content_policy": "require_empty",
            },
            "discard",
        ),
        (
            {
                "op": "merge",
                "row": 1,
                "column": 0,
                "end_row": 1,
                "end_column": 1,
                "content_policy": "append_blocks",
            },
            "completely",
        ),
        (
            {
                "op": "merge",
                "row": 0,
                "column": 0,
                "end_row": 2,
                "end_column": 2,
                "content_policy": "append_blocks",
            },
            "header",
        ),
        ({"op": "split", "row": 2, "column": 0}, "anchor"),
        ({"op": "delete", "axis": "row", "index": 0, "count": 3}, "last axis"),
        ({"op": "insert", "axis": "row", "index": 4, "sizes_twips": [300]}, "bounds"),
        (
            {"op": "resize", "axis": "column", "index": 2, "sizes_twips": [300, 400]},
            "exceeds",
        ),
        (
            {
                "op": "insert",
                "axis": "row",
                "index": 2,
                "sizes_twips": [300],
                "cells": [[{"paragraphs": [{"runs": [{"text": "hidden"}]}]}, {}, {}]],
            },
            "default empty",
        ),
    ],
)
def test_invalid_operations_fail_before_any_source_mutation(edit, match):
    data = source()
    digest = hashlib.sha256(data).hexdigest()
    with pytest.raises(ValueError, match=match):
        run(data, edit)
    assert hashlib.sha256(data).hexdigest() == digest


def omitted_source():
    document = Document()
    document.add_paragraph("intro")
    table = document.add_table(rows=2, cols=5)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.text = str(i)
    row = table.rows[1]._tr
    row.remove(row.tc_lst[-1])
    row.remove(row.tc_lst[0])
    pr = row.get_or_add_trPr()
    etree.SubElement(pr, W + "gridBefore", {W + "val": "1"})
    etree.SubElement(pr, W + "gridAfter", {W + "val": "1"})
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "index,expected", [(0, (2, 1)), (1, (1, 1)), (4, (1, 1)), (5, (1, 2))]
)
def test_insert_preserves_omissions_and_actual_cell_boundaries(index, expected):
    updated, _ = run(
        omitted_source(),
        {"op": "insert", "axis": "column", "index": index, "sizes_twips": [500]},
    )
    record = read_table(updated, CHAIN)
    assert record["row_omissions"][1] == dict(
        zip(("before", "after"), expected, strict=True)
    )
    assert record["columns"] == 6


@pytest.mark.parametrize("index,expected", [(0, (0, 1)), (1, (1, 1)), (4, (1, 0))])
def test_delete_preserves_omission_topology(index, expected):
    updated, _ = run(
        omitted_source(), {"op": "delete", "axis": "column", "index": index, "count": 1}
    )
    assert read_table(updated, CHAIN)["row_omissions"][1] == dict(
        zip(("before", "after"), expected, strict=True)
    )


def test_nested_table_content_survives_merge_and_can_be_edited_directly():
    def add_nested(table):
        cell = table.findall(W + "tr")[1].findall(W + "tc")[2]
        nested = deepcopy(table)
        cell.insert(len(cell) - 1, nested)

    data = transform(source(), add_nested)
    before = checked_docx(data).xml(MAIN_PART).find(".//" + W + "tc/" + W + "tbl")
    updated, _ = run(
        data,
        {
            "op": "merge",
            "row": 1,
            "column": 1,
            "end_row": 1,
            "end_column": 2,
            "content_policy": "append_blocks",
        },
    )
    after = checked_docx(updated).xml(MAIN_PART).find(".//" + W + "tc/" + W + "tbl")
    assert canonical(before) == canonical(after)
    nested_chain = [
        *CHAIN,
        {
            "source_part": MAIN_PART,
            "source_element": "w:tbl",
            "table_index": 0,
            "parent_cell": "1:2",
        },
    ]
    edited, _ = run(
        data,
        {"op": "insert", "axis": "row", "index": 3, "sizes_twips": [500]},
        chain=nested_chain,
    )
    assert read_table(edited, nested_chain)["rows"] == 4
    assert read_table(edited, CHAIN)["rows"] == 3


@pytest.mark.parametrize("tag", ["bookmarkStart", "fldSimple", "tcPrChange"])
def test_structural_dependency_guards_but_resize_retains_native_nodes(tag):
    def add(table):
        etree.SubElement(table.find(".//" + W + "p"), W + tag, {W + "id": "8"})

    data = transform(source(), add)
    with pytest.raises(ValueError, match="dependencies"):
        run(data, {"op": "insert", "axis": "row", "index": 3, "sizes_twips": [500]})
    updated, _ = run(
        data, {"op": "resize", "axis": "row", "index": 2, "sizes_twips": [500]}
    )
    assert checked_docx(updated).xml(MAIN_PART).find(".//" + W + tag) is not None


def test_noop_split_preserves_exact_package_bytes():
    data = source()
    updated, result = run(data, {"op": "split", "row": 2, "column": 2})
    assert updated == data and result.changed_parts == []


def test_mismatched_vertical_span_and_missing_grid_rejected():
    def mismatch(table):
        pr = table.findall(W + "tr")[2].findall(W + "tc")[0].find(W + "tcPr")
        etree.SubElement(pr, W + "gridSpan", {W + "val": "2"})

    with pytest.raises(ValueError, match="mismatched"):
        read_table(transform(source(), mismatch), CHAIN)
    with pytest.raises(ValueError, match="grid"):
        WordGrid.read(etree.Element(W + "tbl"))


@pytest.mark.parametrize("lock", [None, "unlocked", "sdtContentLocked", "binding"])
def test_content_controls_keep_explicit_binding_and_lock_guards(lock):
    package = checked_docx(source())
    root = package.xml(MAIN_PART)
    body = root.find(W + "body")
    table = body[1]
    control = etree.Element(W + "sdt")
    properties = etree.SubElement(control, W + "sdtPr")
    if lock:
        etree.SubElement(
            properties,
            W + ("dataBinding" if lock == "binding" else "lock"),
            {W + "val": lock},
        )
    content = etree.SubElement(control, W + "sdtContent")
    body.replace(table, control)
    content.append(table)
    data = package.replace({MAIN_PART: xml_bytes(root)})
    chain = [{**CHAIN[0], "sdt_index": 0, "source_order": 0}]
    edit = {"op": "resize", "axis": "row", "index": 1, "sizes_twips": [600]}
    if lock in {"binding", "sdtContentLocked"}:
        with pytest.raises(ValueError, match="locked"):
            run(data, edit, chain=chain)
    else:
        updated, _ = run(data, edit, chain=chain)
        assert read_table(updated, chain)["row_heights"][1] == [
            {"value_twips": 600, "rule": "atLeast"}
        ]


def test_resizing_columns_updates_omitted_prefix_suffix_preferred_widths():
    updated, _ = run(
        omitted_source(),
        {
            "op": "resize",
            "axis": "column",
            "index": 0,
            "sizes_twips": [500, 700, 900, 1100, 1300],
        },
    )
    pr = Document(io.BytesIO(updated)).tables[0].rows[1]._tr.trPr
    assert pr.find(W + "wBefore").get(W + "w") == "500"
    assert pr.find(W + "wAfter").get(W + "w") == "1300"
