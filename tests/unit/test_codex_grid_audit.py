"""Independent live-grid checks reject text, geometry, merge and style corruption."""

import io

import pytest
from pptx import Presentation
from pptx.util import Pt

from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pptx_tables.grids import check_stage, surviving_cells
from tests.unit.test_codex_pptx_table_audit import make_presentation


@pytest.mark.parametrize("failure", [None, "leading_zero", "width", "frame", "merge"])
def test_intermediate_scanned_grid_audit(failure):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(4, 6, 100000, 100000, 6300000, 2200000)
    table = shape.table
    for i, width in enumerate([1200000, 1200000, 300000, 1200000, 1200000, 1200000]):
        table.columns[i].width = width
    for row in table.rows:
        row.height = 550000
    table.cell(0, 0).merge(table.cell(0, 4 if failure == "merge" else 5))
    table.cell(0, 0).text = "Scanned inventory"
    for r, values in enumerate([COLUMNS, *PAGE_ROWS[0]], 1):
        for c, value in enumerate(values):
            table.cell(r, c if c < 2 else c + 1).text = value
    if failure == "leading_zero":
        table.cell(2, 1).text = "7"
    elif failure == "width":
        table.columns[2].width += 1
    elif failure == "frame":
        shape.left += 1
    output = io.BytesIO()
    presentation.save(output)
    if failure:
        with pytest.raises(ValueError):
            check_stage(output.getvalue(), 1)
    else:
        check_stage(output.getvalue(), 1)


def test_surviving_cell_audit_detects_format_change_even_when_text_matches():
    presentation = make_presentation()
    shapes = presentation.slides[0].shapes
    shapes[-1].element.getparent().remove(shapes[-1].element)
    before = io.BytesIO()
    presentation.save(before)
    surviving_cells(before.getvalue(), before.getvalue(), 3)
    shapes[0].table.cell(2, 1).text_frame.paragraphs[0].runs[0].font.size = Pt(19)
    after = io.BytesIO()
    presentation.save(after)
    with pytest.raises(ValueError, match="formatting"):
        surviving_cells(before.getvalue(), after.getvalue(), 3)
