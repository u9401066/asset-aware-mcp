"""Slide operations preserve rich decks and leave stable shape evidence addressable."""

import io
from copy import deepcopy

import pytest
from lxml import etree
from pptx import Presentation

from src.domain.native_pptx_slides import NativePptxSlideInsert, NativePptxSlideKey
from src.infrastructure.native_ooxml import (
    DOC_REL_NS,
    REL_NS,
    relationships_path,
    xml_bytes,
)
from src.infrastructure.native_pptx import NativePresentation
from src.infrastructure.native_pptx_package import NS, NativePptxPackage
from tests.native_pptx_helpers import build_presentation
from tests.native_workbook_helpers import _replace


def keys(data):
    return [NativePptxSlideKey(**s) for s in NativePptxPackage(data).slides]


def insert_request(adapter, data, *, index=1, layout_type="blank"):
    layout = next(
        layout for layout in adapter.read_layouts(data) if layout["type"] == layout_type
    )
    return NativePptxSlideInsert(
        index=index,
        slides=[
            {
                "layout_part": layout["part"],
                "textboxes": [
                    {
                        "paragraphs": [
                            [
                                {
                                    "text": "007 µg +0.50 原樣",
                                    "bold": True,
                                    "font_size_pt": 22,
                                }
                            ]
                        ]
                    }
                ],
            }
        ],
    )


def assert_original_parts(before, after, changed):
    old, new = NativePptxPackage(before), NativePptxPackage(after)
    assert set(old.parts) <= set(new.parts)
    for part, raw in old.parts.items():
        if part not in changed:
            assert new.parts[part] == raw, part


@pytest.mark.parametrize("index", [0, 1, 2])
def test_insert_slide_preserves_all_original_shapes_parts_and_native_identities(index):
    adapter, data = NativePresentation(), build_presentation()
    before = NativePptxPackage(data)
    updated, result = adapter.add_slides(
        data, insert_request(adapter, data, index=index)
    )
    after = NativePptxPackage(updated)
    created = result.changes[0]["slides"][0]
    assert after.slides == [*before.slides[:index], created, *before.slides[index:]]
    assert_original_parts(data, updated, result.changed_parts)
    original = {str(r["locator"]): r for r in adapter.iter_shapes(data)}
    for r in adapter.iter_shapes(updated):
        if str(r["locator"]) in original:
            assert r == original[str(r["locator"])]
    deck = Presentation(io.BytesIO(updated))
    shape = deck.slides[index].shapes[0]
    assert shape.text == "007 µg +0.50 原樣"
    assert shape.text_frame.paragraphs[0].runs[0].font.bold is True
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 22
    assert deck.slides[index].slide_layout.name == "Blank"


def test_layout_discovery_and_placeholders_do_not_copy_master_prompt_content():
    adapter, data = NativePresentation(), build_presentation()
    layouts = adapter.read_layouts(data)
    assert len(layouts) == 11
    assert {layout["master_part"] for layout in layouts} == {
        "ppt/slideMasters/slideMaster1.xml"
    }
    request = insert_request(adapter, data, layout_type="title")
    updated, _ = adapter.add_slides(data, request)
    slide = Presentation(io.BytesIO(updated)).slides[1]
    assert len(slide.placeholders) == 2
    assert all(p.text == "" for p in slide.placeholders)
    assert {p.placeholder_format.idx for p in slide.placeholders} == {0, 1}
    assert slide.shapes[-1].text == "007 µg +0.50 原樣"


def test_reorder_keeps_native_ids_media_notes_chart_and_all_shape_records():
    adapter, data = NativePresentation(), build_presentation()
    requested = list(reversed(keys(data)))
    updated, result = adapter.reorder_slides(data, requested)
    assert NativePptxPackage(updated).slides == [k.model_dump() for k in requested]
    assert_original_parts(data, updated, result.changed_parts)
    expected = {str(r["locator"]): r for r in adapter.iter_shapes(data)}
    assert {str(r["locator"]): r for r in adapter.iter_shapes(updated)} == expected
    deck = Presentation(io.BytesIO(updated))
    assert deck.slides[0].shapes[0].text == "Other slide"
    assert deck.slides[1].notes_slide.notes_text_frame.text == "備註 Keep notes"


@pytest.mark.parametrize("delete_all", [False, True])
def test_delete_retains_detached_slide_and_notes_bytes_and_allows_empty_deck(
    delete_all,
):
    adapter, data = NativePresentation(), build_presentation()
    selected = keys(data) if delete_all else keys(data)[:1]
    updated, result = adapter.delete_slides(data, selected)
    assert_original_parts(data, updated, result.changed_parts)
    assert len(Presentation(io.BytesIO(updated)).slides) == (0 if delete_all else 1)
    assert len(NativePptxPackage(updated).parts) == len(NativePptxPackage(data).parts)
    if delete_all:
        restored, _ = adapter.add_slides(
            updated, insert_request(adapter, updated, index=0)
        )
        assert len(Presentation(io.BytesIO(restored)).slides) == 1
        assert_original_parts(
            updated,
            restored,
            [
                NativePptxPackage(data).main_part,
                relationships_path(NativePptxPackage(data).main_part),
                "[Content_Types].xml",
                "docProps/app.xml",
            ],
        )


def test_multiple_insertions_allocate_distinct_parts_and_ids_after_detached_parts():
    adapter, data = NativePresentation(), build_presentation()
    first, result = adapter.add_slides(data, insert_request(adapter, data))
    deleted, _ = adapter.delete_slides(
        first, [NativePptxSlideKey(**result.changes[0]["slides"][0])]
    )
    request = insert_request(adapter, deleted)
    request.slides.append(deepcopy(request.slides[0]))
    updated, result = adapter.add_slides(deleted, request)
    created = result.changes[0]["slides"]
    assert len({s["part"] for s in created}) == 2
    assert "nativeSlide1.xml" not in str(created)
    assert len({s["slide_id"] for s in NativePptxPackage(updated).slides}) == 4


def second_master(data):
    package = NativePptxPackage(data)
    main = package.presentation
    master_part = "ppt/slideMasters/customMaster.xml"
    layout_part = "ppt/slideLayouts/customLayout.xml"
    master = package.xml("ppt/slideMasters/slideMaster1.xml")
    ids = master.find("p:sldLayoutIdLst", NS)
    for item in list(ids)[1:]:
        ids.remove(item)
    ids[0].set(f"{{{DOC_REL_NS}}}id", "rIdLayout")
    master_rels = etree.Element(f"{{{REL_NS}}}Relationships")
    etree.SubElement(
        master_rels,
        f"{{{REL_NS}}}Relationship",
        Id="rIdLayout",
        Type=f"{DOC_REL_NS}/slideLayout",
        Target="../slideLayouts/customLayout.xml",
    )
    etree.SubElement(
        master_rels,
        f"{{{REL_NS}}}Relationship",
        Id="rIdTheme",
        Type=f"{DOC_REL_NS}/theme",
        Target="../theme/theme1.xml",
    )
    layout = package.xml("ppt/slideLayouts/slideLayout1.xml")
    layout.find("p:cSld", NS).set("name", "Second master title")
    layout_rels = etree.Element(f"{{{REL_NS}}}Relationships")
    etree.SubElement(
        layout_rels,
        f"{{{REL_NS}}}Relationship",
        Id="rId1",
        Type=f"{DOC_REL_NS}/slideMaster",
        Target="../slideMasters/customMaster.xml",
    )
    entry = etree.SubElement(
        main.find("p:sldMasterIdLst", NS), f"{{{NS['p']}}}sldMasterId", id="2147483649"
    )
    entry.set(f"{{{DOC_REL_NS}}}id", "rIdCustomMaster")
    rels = package.xml(relationships_path(package.main_part))
    etree.SubElement(
        rels,
        f"{{{REL_NS}}}Relationship",
        Id="rIdCustomMaster",
        Type=f"{DOC_REL_NS}/slideMaster",
        Target="slideMasters/customMaster.xml",
    )
    types = package.xml("[Content_Types].xml")
    for part, kind in ((master_part, "slideMaster"), (layout_part, "slideLayout")):
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/" + part,
            ContentType=f"application/vnd.openxmlformats-officedocument.presentationml.{kind}+xml",
        )
    return _replace(
        data,
        {
            package.main_part: xml_bytes(main),
            relationships_path(package.main_part): xml_bytes(rels),
            "[Content_Types].xml": xml_bytes(types),
            master_part: xml_bytes(master),
            relationships_path(master_part): xml_bytes(master_rels),
            layout_part: xml_bytes(layout),
            relationships_path(layout_part): xml_bytes(layout_rels),
        },
    )


def test_discover_and_insert_using_a_nonfirst_master_without_layout_index_assumptions():
    adapter, data = NativePresentation(), second_master(build_presentation())
    layouts = adapter.read_layouts(data)
    selected = next(
        layout for layout in layouts if layout["name"] == "Second master title"
    )
    request = NativePptxSlideInsert(index=0, slides=[{"layout_part": selected["part"]}])
    updated, result = adapter.add_slides(data, request)
    slide = Presentation(io.BytesIO(updated)).slides[0]
    assert slide.slide_layout.name == "Second master title"
    assert len(slide.placeholders) == 2
    assert_original_parts(data, updated, result.changed_parts)
