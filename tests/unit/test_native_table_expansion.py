"""Explicit membership must coordinate table metadata, native cells and dependencies."""

import hashlib
import io

import openpyxl
import pytest
import xlsxwriter
from lxml import etree

from src.domain.native_grid import NativeGridEdit
from src.infrastructure.native_spreadsheet_reader import NS
from src.infrastructure.native_workbook_grid import NativeWorkbookGrid
from src.infrastructure.native_workbook_structure import NativeWorkbookStructure
from tests.native_grid_helpers import table_workbook
from tests.native_workbook_helpers import _parts, _replace
from tests.unit.test_native_grid_metadata import CACHE
from tests.unit.test_native_grid_named_sources import named_source
from tests.unit.test_native_grid_tables import read, table
from tests.unit.test_native_workbook_grid import request


def insertion(axis="row", at=6, count=1, ref="A2:C5", part="xl/tables/table1.xml"):
    return {
        "axis": axis,
        "operation": "insert",
        "at": at,
        "count": count,
        "expand_tables": [{"part": part, "expected_ref": ref}],
    }


@pytest.mark.parametrize("formula", ["=A3*2", "=[@Input]*2"])
def test_append_rows_extends_data_filters_sort_and_calculated_cells(formula):
    source = table_workbook(formula=formula)
    updated, result = NativeWorkbookGrid().update(source, request(insertion(count=2)))
    root = table(updated)
    assert root.get("ref") == "A2:C7"
    assert root.find("s:autoFilter", NS).get("ref") == "A2:C7"
    assert root.find("s:autoFilter/s:sortState", NS).get("ref") == "A3:C7"
    assert (
        root.find("s:autoFilter/s:sortState/s:sortCondition", NS).get("ref") == "A3:A7"
    )
    for row in (6, 7):
        expected = f"=A{row}*2" if formula == "=A3*2" else "=[[#This Row],Input]*2"
        assert read(updated, f"B{row}")["value"] == expected
        assert (
            read(updated, f"B{row}")["style_index"] == read(source, "B5")["style_index"]
        )
    assert result.changes[0]["generated_table_cells"] == {
        "B6": "calculated",
        "B7": "calculated",
    }
    assert _parts(updated)["xl/styles.xml"] == _parts(source)["xl/styles.xml"]
    assert (
        _parts(updated)["customXml/unchanged.xml"]
        == _parts(source)["customXml/unchanged.xml"]
    )
    book = openpyxl.load_workbook(io.BytesIO(updated))
    assert book["Data"].tables["Table1"].ref == "A2:C7"
    assert [col.id for col in book["Data"].tables["Table1"].tableColumns] == [1, 2, 3]


@pytest.mark.parametrize(
    "at,ids,names,filter_ids,sort_key",
    [
        (
            1,
            [4, 5, 1, 2, 3],
            ["Column1", "Column2", "Input", "Calc", "Label"],
            [2, 4],
            "C3:C5",
        ),
        (
            4,
            [1, 2, 3, 4, 5],
            ["Input", "Calc", "Label", "Column4", "Column5"],
            [0, 2],
            "A3:A5",
        ),
    ],
)
def test_column_edges_preserve_ids_and_sort_key_identity(
    at, ids, names, filter_ids, sort_key
):
    updated, result = NativeWorkbookGrid().update(
        table_workbook(), request(insertion("column", at, 2))
    )
    root = table(updated)
    assert root.get("ref") == "A2:E5"
    cols = root.findall("s:tableColumns/s:tableColumn", NS)
    assert [int(col.get("id")) for col in cols] == ids
    assert [col.get("name") for col in cols] == names
    assert [read(updated, f"{col}2")["value"] for col in "ABCDE"] == names
    auto = root.find("s:autoFilter", NS)
    assert auto.get("ref") == "A2:E5"
    assert [
        int(col.get("colId")) for col in auto.findall("s:filterColumn", NS)
    ] == filter_ids
    assert auto.find("s:sortState", NS).get("ref") == "A3:E5"
    assert auto.find("s:sortState/s:sortCondition", NS).get("ref") == sort_key
    assert read(updated, "A1", "Other")["value"] == "=SUM(Table1[[Input]:[Label]])"
    assert len(result.changes[0]["generated_table_cells"]) == 2


@pytest.mark.parametrize("explicit", [False, True])
def test_insertion_before_totals_extends_filter_even_without_expansion_option(explicit):
    edit = insertion(ref="A2:C6", count=2)
    if not explicit:
        edit.pop("expand_tables")
    updated, _ = NativeWorkbookGrid().update(table_workbook(totals=True), request(edit))
    root = table(updated)
    assert root.get("ref") == "A2:C8"
    assert root.find("s:autoFilter", NS).get("ref") == "A2:C7"
    assert read(updated, "C8")["value"] == "Total"
    assert read(updated, "B8")["value"] == "=SUBTOTAL(109,[Calc])"
    assert read(updated, "B6")["kind"] == "formula"


def test_adjacent_table_is_shifted_but_not_implicitly_expanded():
    output = io.BytesIO()
    with xlsxwriter.Workbook(output, {"in_memory": True}) as book:
        sheet = book.add_worksheet("Data")
        sheet.add_table("A2:C5", {"name": "Left"})
        sheet.add_table("D2:F5", {"name": "Right"})
    updated, _ = NativeWorkbookGrid().update(
        output.getvalue(), request(insertion("column", 4))
    )
    book = openpyxl.load_workbook(io.BytesIO(updated))
    assert book["Data"].tables["Left"].ref == "A2:D5"
    assert book["Data"].tables["Right"].ref == "E2:G5"
    assert [c.id for c in book["Data"].tables["Right"].tableColumns] == [1, 2, 3]


@pytest.mark.parametrize(
    "edit,totals,match",
    [
        (insertion(ref="A2:C4"), False, "stale"),
        (insertion(part="xl/tables/missing.xml"), False, "active table"),
        (insertion(at=4), False, "boundary"),
        (insertion(at=7, ref="A2:C6"), True, "before totals"),
    ],
)
def test_incorrect_expansion_request_fails(edit, totals, match):
    source = table_workbook(totals=totals)
    with pytest.raises(ValueError, match=match):
        NativeWorkbookGrid().update(source, request(edit))


def test_expansion_identity_is_validated_at_each_intermediate_grid():
    source = table_workbook()
    edits = [
        insertion(count=2),
        insertion("column", 4, ref="A2:C7"),
        {"axis": "column", "operation": "insert", "at": 1},
    ]
    updated, result = NativeWorkbookGrid().update(source, request(*edits))
    assert table(updated).get("ref") == "B2:E7"
    assert result.changes[0]["generated_table_cells"] == {
        "C6": "calculated",
        "C7": "calculated",
        "E2": "header",
    }
    assert read(updated, "E2")["value"] == "Column4"
    edits[1]["expand_tables"][0]["expected_ref"] = "A2:C5"
    with pytest.raises(ValueError, match="stale"):
        NativeWorkbookGrid().update(source, request(*edits))


def test_expansion_keeps_pivot_cache_schema_guard_and_refreshes_new_rows():
    source = named_source(("Proxy", "AllData", None))
    updated, result = NativeWorkbookGrid().update(source, request(insertion()))
    assert etree.fromstring(_parts(updated)[CACHE]).get("refreshOnLoad") == "1"
    assert result.changes[0]["edits"][0]["invalidated_named_source_caches"] == [CACHE]
    with pytest.raises(ValueError, match="cache field identities"):
        NativeWorkbookGrid().update(source, request(insertion("column", 4)))
    source = table_workbook()
    root = table(source)
    root.set("tableType", "xml")
    source = _replace(source, {"xl/tables/table1.xml": etree.tostring(root)})
    with pytest.raises(ValueError, match="source field identities"):
        NativeWorkbookGrid().update(source, request(insertion("column", 4)))


def test_deletion_and_duplicate_expansion_are_schema_errors():
    value = insertion()
    value["operation"] = "delete"
    with pytest.raises(ValueError, match="only used for insertion"):
        NativeGridEdit.model_validate(value)
    value = insertion()
    value["expand_tables"] *= 2
    with pytest.raises(ValueError, match="Duplicate"):
        NativeGridEdit.model_validate(value)


def test_headerless_first_row_insertion_and_formula_generation():
    source = table_workbook(formula="=A3*2")
    source, _ = NativeWorkbookGrid().update(
        source, request({"axis": "row", "operation": "delete", "at": 2})
    )
    updated, _ = NativeWorkbookGrid().update(
        source, request(insertion(at=2, ref="A2:C4"))
    )
    assert table(updated).get("ref") == "A2:C5"
    assert read(updated, "B2")["value"] == "=A2*2"
    assert read(updated, "A3")["value"] == 10
    assert table(updated).find("s:autoFilter", NS) is None


def test_array_calculated_column_cannot_be_expanded_as_scalar_formulas():
    source = table_workbook()
    root = table(source)
    root.find("s:tableColumns/s:tableColumn/s:calculatedColumnFormula", NS).set(
        "array", "1"
    )
    source = _replace(source, {"xl/tables/table1.xml": etree.tostring(root)})
    with pytest.raises(ValueError, match="Array calculated"):
        NativeWorkbookGrid().update(source, request(insertion()))


def test_table_inventory_exposes_exact_part_hash_and_complete_non_utf8_definition():
    source = table_workbook()
    root = table(source)
    root.set("comment", "Café")
    raw = etree.tostring(root, encoding="iso-8859-1", xml_declaration=True)
    source = _replace(source, {"xl/tables/table1.xml": raw})
    inventory = NativeWorkbookStructure().read(source, references=True)["tables"]
    assert len(inventory) == 1
    record = inventory[0]
    assert record["worksheet"] == {"sheet_id": "1", "part": "xl/worksheets/sheet1.xml"}
    assert record["part"] == "xl/tables/table1.xml"
    assert (
        record["attributes"]["ref"] == "A2:C5"
        and record["attributes"]["comment"] == "Café"
    )
    assert record["sha256"] == hashlib.sha256(raw).hexdigest()
    parsed = etree.fromstring(record["xml"].encode())
    assert etree.tostring(parsed, method="c14n") == etree.tostring(root, method="c14n")
    assert [c["id"] for c in record["columns"]] == ["1", "2", "3"]
