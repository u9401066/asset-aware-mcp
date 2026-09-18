"""Managed shape CRUD, revision conflicts and immutable evidence snapshots."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation

from src.application.native_wiki_service import NativeWikiService
from src.domain.native_assets import NativeDocumentRequest
from src.infrastructure.native_wiki_publisher import FileNativeWikiPublisher
from tests.native_pptx_helpers import edit_run, find_shape
from tests.native_pptx_shape_helpers import addition, reference
from tests.native_workbook_helpers import _call
from tests.unit.test_native_pptx_operations import managed as managed
from tests.unit.test_native_pptx_operations import read_shape


def shape_request(asset, data, *, deleting=False):
    payload = (
        {
            "pptx_shape_refs": [
                reference(
                    data, find_shape(data, "Styled text"), asset_id=asset["asset_id"]
                ).model_dump()
            ]
        }
        if deleting
        else {"pptx_shapes": [addition(data).model_dump()]}
    )
    return dict(
        op="delete_pptx_shapes" if deleting else "add_pptx_shapes",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        **payload,
    )


def enable_wiki(service, tmp_path):
    service.wiki = NativeWikiService(
        service.repository,
        service.spreadsheets,
        FileNativeWikiPublisher((tmp_path / "store",)),
        presentations=service.presentations,
    )


def test_shape_add_delete_retains_history_wiki_and_source(managed, tmp_path):
    service, asset, source = managed
    enable_wiki(service, tmp_path)
    original = source.read_bytes()
    added = _call(service, **shape_request(asset, original))
    current = added["asset"]
    locator = added["operation_result"]["changes"][0]["locator"]
    record = read_shape(service, current, locator)
    assert "新增 µ" in record["xml"]
    assert (
        _call(service, **added["review_request"])["inspected_revision"]
        == current["revision"]
    )
    old_wiki = _call(
        service,
        op="export_wiki",
        asset_id=asset["asset_id"],
        output_dir=str(tmp_path / "wiki"),
    )
    old_records = Path(old_wiki["output_dir"]) / "records.jsonl"
    snapshot = old_records.read_bytes()
    deleted = _call(
        service,
        op="delete_pptx_shapes",
        asset_id=asset["asset_id"],
        expected_revision=current["revision"],
        pptx_shape_refs=[record["evidence"]],
    )
    assert not deleted["source_written"] and source.read_bytes() == original
    assert not any(
        s["locator"] == locator
        for s in _call(service, **deleted["review_request"])["shapes"]
    )
    proof = _call(service, op="verify", reference=record["evidence"])
    assert proof["valid"] and not proof["is_current_managed_revision"]
    assert read_shape(service, current, locator) == record
    new_wiki = _call(
        service,
        op="export_wiki",
        asset_id=asset["asset_id"],
        output_dir=str(tmp_path / "wiki"),
    )
    assert new_wiki["output_dir"] != old_wiki["output_dir"]
    assert old_records.read_bytes() == snapshot
    assert any(
        json.loads(line)["evidence"] == record["evidence"]
        for line in snapshot.splitlines()
    )
    assert len(_call(service, op="history", asset_id=asset["asset_id"])["history"]) == 3


def test_shape_writeback_preserves_backup_and_detects_source_change(managed):
    service, asset, source = managed
    original = source.read_bytes()
    changed = _call(service, **shape_request(asset, original))["asset"]
    source.write_bytes(original + b"human change")
    request = {
        "op": "writeback",
        "asset_id": asset["asset_id"],
        "expected_revision": changed["revision"],
        "expected_source_sha256": asset["source"]["sha256"],
    }
    with pytest.raises(ValueError):
        _call(service, **request)
    assert source.read_bytes() == original + b"human change"
    source.write_bytes(original)
    # Restore the exact captured identity as well as bytes for this successful path.
    import os

    os.utime(source, ns=(source.stat().st_atime_ns, asset["source"]["mtime_ns"]))
    written = _call(service, **request)
    assert Path(written["backup_path"]).read_bytes() == original
    assert Presentation(source).slides[0].shapes[-1].text.startswith("新增 µ")


@pytest.mark.parametrize("deleting", [False, True])
def test_shape_operations_reject_concurrent_revision(managed, monkeypatch, deleting):
    service, asset, source = managed
    data = source.read_bytes()
    adapter = service.presentations
    method = "delete_shapes" if deleting else "add_shapes"
    operation = getattr(adapter, method)
    competing = None

    def interleave(data, requests):
        nonlocal competing
        candidate = operation(data, requests)
        other, report = adapter.edit(
            data, [edit_run(find_shape(data, "Styled text"), "Concurrent")]
        )
        competing = service.repository.commit(
            asset["asset_id"], asset["revision"], other, report
        ).revision
        return candidate

    monkeypatch.setattr(adapter, method, interleave)
    with pytest.raises(ValueError, match="Stale"):
        _call(service, **shape_request(asset, data, deleting=deleting))
    assert service.repository.load(asset["asset_id"]).revision == competing
    assert source.read_bytes() == data


@pytest.mark.parametrize("deleting", [False, True])
def test_shape_contract_archive_and_stale_guards(managed, deleting):
    service, asset, source = managed
    request = shape_request(asset, source.read_bytes(), deleting=deleting)
    discovery = _call(service, op="contract", for_op=request["op"])
    assert request["op"] in discovery["formats"]["pptx"]
    assert asset["capabilities"][request["op"]]
    with pytest.raises(ValueError, match="stale"):
        _call(service, **{**request, "expected_revision": "0" * 64})
    _call(
        service,
        op="archive",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
    )
    archived = _call(service, op="inspect", asset_id=asset["asset_id"])["asset"]
    assert not archived["capabilities"][request["op"]]
    with pytest.raises(ValueError, match="Archived"):
        _call(service, **request)


@pytest.mark.parametrize(
    "field,value", [("asset_id", "file_" + "0" * 32), ("revision", "0" * 64)]
)
def test_deletion_reference_must_match_asset_and_revision(managed, field, value):
    service, asset, source = managed
    data = source.read_bytes()
    request = shape_request(asset, data, deleting=True)
    request["pptx_shape_refs"][0][field] = value
    with pytest.raises(ValueError, match="different asset or revision"):
        _call(service, **request)
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    assert source.read_bytes() == data


@pytest.mark.parametrize(
    "mutation", ["empty", "oversized", "unrelated", "wrong_part", "unknown_group"]
)
def test_add_shape_invalid_batch_never_commits(managed, mutation):
    service, asset, source = managed
    request = shape_request(asset, source.read_bytes())
    item = request["pptx_shapes"][0]
    if mutation == "empty":
        request["pptx_shapes"] = []
    elif mutation == "oversized":
        request["pptx_shapes"] *= 101
    elif mutation == "unrelated":
        request["sheet"] = "Unused"
    else:
        item = json.loads(json.dumps(item))
        key, value = (
            ("part", "ppt/slides/slide2.xml")
            if mutation == "wrong_part"
            else ("group_shape_id", "999")
        )
        item["container"][key] = value
        request["pptx_shapes"].append(item)
    with pytest.raises(ValueError):
        _call(service, **request)
    assert len(service.repository.load(asset["asset_id"]).history) == 1


@pytest.mark.parametrize("kind", ["runs", "bytes"])
def test_add_shape_aggregate_limits(managed, kind):
    _, asset, source = managed
    request = shape_request(asset, source.read_bytes())
    textbox = request["pptx_shapes"][0]["textbox"]
    run = {"text": "字" * 8000 if kind == "bytes" else "x"}
    textbox["paragraphs"] = [[run] * 100] * (2 if kind == "bytes" else 100)
    if kind == "runs":
        request["pptx_shapes"] *= 3
    with pytest.raises(ValueError, match="20,000 runs or 4 MiB"):
        NativeDocumentRequest.model_validate(request)
