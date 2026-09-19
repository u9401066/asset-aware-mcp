"""Whole-package native grid edits, independent readers and preservation failures."""

import io

import openpyxl
import pytest
from lxml import etree

from src.domain.native_grid import NativeGridUpdate
from src.infrastructure.native_grid_drawings import NS as DNS
from src.infrastructure.native_grid_xml import tag
from src.infrastructure.native_spreadsheet_reader import NS
from src.infrastructure.native_workbook_grid import NativeWorkbookGrid
from tests.native_grid_drawing_helpers import DRAWING_PART, SHEET_PART, drawing_workbook
from tests.native_grid_helpers import table_workbook
from tests.native_workbook_helpers import _parts, _replace, build_workbook


def request(*edits, **kwargs):
    return NativeGridUpdate.model_validate(
        {
            "worksheet": {"sheet_id": "1", "part": SHEET_PART},
            "edits": list(edits) or [{"axis": "row", "operation": "insert", "at": 1}],
            **kwargs,
        }
    )


def root(data, part=SHEET_PART):
    return etree.fromstring(_parts(data)[part])


def test_rich_workbook_moves_cells_merges_tables_comments_charts_and_cross_sheet_formulas():
    source = build_workbook()
    updated, result = NativeWorkbookGrid().update(source, request())
    book = openpyxl.load_workbook(io.BytesIO(updated))
    sheet = book["Data"]
    assert sheet["A2"].value == "Original" and sheet["A2"].font.bold
    assert sheet["A3"].value == 10 and sheet["A3"].number_format == "$0.00"
    assert sheet["B3"].value == "=A3*2"
    assert sheet["D3"].value == "Bold normal"
    assert sheet["A2"].comment.text == "Keep this comment"
    assert sheet["C4"].hyperlink.target == "https://example.com"
    assert str(sheet.merged_cells) == "A6:D7"
    assert sheet.tables["Table1"].ref == "A9:B11"
    assert book["Other"]["B1"].value == "=Data!A3"
    chart = root(updated, "xl/charts/chart1.xml")
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    assert chart.find(".//c:f", ns).text == "Data!$A$3:$A$4"
    assert not chart.findall(".//c:numCache", ns)
    assert (
        root(updated, DRAWING_PART).find("xdr:twoCellAnchor/xdr:from/xdr:row", DNS).text
        == "1"
    )
    for name in (
        "xl/styles.xml",
        "xl/theme/theme1.xml",
        "docProps/app.xml",
        "customXml/item1.xml",
    ):
        assert _parts(source)[name] == _parts(updated)[name]
    assert "complete_grid_xml_read_back" in result.checks
    assert result.changes[0]["edits"][0]["objects"]["comments"][0]["after_cell"] == "A2"


def test_sequential_edits_use_intermediate_grid_and_keep_table_column_identity():
    updated, result = NativeWorkbookGrid().update(
        table_workbook(),
        request(
            {"axis": "column", "operation": "insert", "at": 2},
            {"axis": "row", "operation": "insert", "at": 4},
            {"axis": "column", "operation": "delete", "at": 1},
        ),
    )
    book = openpyxl.load_workbook(io.BytesIO(updated))
    table = book["Data"].tables["Table1"]
    assert table.ref == "A2:C6"
    assert [(column.id, column.name) for column in table.tableColumns] == [
        (4, "Column2"),
        (2, "Calc"),
        (3, "Label"),
    ]
    assert book["Data"]["A2"].value == "Column2"
    assert book["Data"]["B4"].value == "=#REF!*2"
    assert book["Other"]["A2"].value == "=SUM(#REF!)"
    assert book["Other"]["A1"].value == "=SUM(Table1[[Column2]:[Label]])"
    assert len(result.changes[0]["edits"]) == 3


def test_merge_anchor_promotion_preserves_content_and_format():
    updated, _ = NativeWorkbookGrid().update(
        build_workbook(), request({"axis": "row", "operation": "delete", "at": 5})
    )
    sheet = openpyxl.load_workbook(io.BytesIO(updated))["Data"]
    assert str(sheet.merged_cells) == "A5:D5"
    assert sheet["A5"].value == "Merged" and sheet["A5"].font.bold
    assert sheet.tables["Table1"].ref == "A7:B9"


def test_outline_summary_and_cell_watches_follow_deleted_rows():
    source = build_workbook()
    sheet = root(source)
    sheet.find("s:sheetFormatPr", NS).set("outlineLevelRow", "5")
    sheet.find('s:sheetData/s:row[@r="2"]', NS).set("outlineLevel", "5")
    sheet.find('s:sheetData/s:row[@r="3"]', NS).set("outlineLevel", "2")
    watches = etree.SubElement(sheet, tag("cellWatches"))
    etree.SubElement(watches, tag("cellWatch"), r="A2")
    etree.SubElement(watches, tag("cellWatch"), r="A3")
    updated, _ = NativeWorkbookGrid().update(
        _replace(source, {SHEET_PART: etree.tostring(sheet)}),
        request({"axis": "row", "operation": "delete", "at": 2}),
    )
    after = root(updated)
    assert after.find("s:sheetFormatPr", NS).get("outlineLevelRow") == "2"
    assert [
        node.get("r") for node in after.findall("s:cellWatches/s:cellWatch", NS)
    ] == ["A2"]


def test_original_source_bytes_remain_readable_when_late_edit_fails():
    source = build_workbook()
    with pytest.raises(ValueError, match="outside"):
        NativeWorkbookGrid().update(
            source,
            request(
                {"axis": "row", "operation": "insert", "at": 1},
                {"axis": "column", "operation": "insert", "at": 1, "count": 1024},
                {"axis": "row", "operation": "insert", "at": 1, "count": 1024},
                *[
                    {"axis": "column", "operation": "insert", "at": 1, "count": 1024}
                    for _ in range(15)
                ],
            ),
        )
    assert openpyxl.load_workbook(io.BytesIO(source))["Data"]["A1"].value == "Original"


@pytest.mark.parametrize("kind", ["protected", "extension", "metadata"])
def test_unmodeled_target_features_are_reported_without_flattening(kind):
    source = drawing_workbook()
    sheet = root(source)
    if kind == "metadata":
        sheet.find("s:sheetData/s:row/s:c", NS).set("cm", "1")
    else:
        etree.SubElement(
            sheet, tag("sheetProtection" if kind == "protected" else "extLst")
        )
    with pytest.raises(ValueError, match=r"Protected|extLst|metadata"):
        NativeWorkbookGrid().update(
            _replace(source, {SHEET_PART: etree.tostring(sheet)}), request()
        )


def test_exact_worksheet_key_is_required():
    bad = request()
    bad.worksheet.sheet_id = "2"
    with pytest.raises(ValueError, match="exact workbook revision"):
        NativeWorkbookGrid().update(build_workbook(), bad)


def test_serialized_package_tampering_is_detected(monkeypatch):
    from src.infrastructure import native_workbook_plan

    real = native_workbook_plan.extend_package

    def corrupt(*args, **kwargs):
        data = real(*args, **kwargs)
        sheet = root(data)
        sheet.find('s:sheetData/s:row/s:c[@r="A3"]/s:v', NS).text = "999"
        return _replace(data, {SHEET_PART: etree.tostring(sheet)})

    monkeypatch.setattr(native_workbook_plan, "extend_package", corrupt)
    with pytest.raises(ValueError, match="exact XML read-back"):
        NativeWorkbookGrid().update(build_workbook(), request())
