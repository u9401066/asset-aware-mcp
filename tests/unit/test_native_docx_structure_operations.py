"""Managed Word creation, evidence-bound structure, historical wiki and CAS."""

from __future__ import annotations

import hashlib
import io

import pytest
from docx import Document

from src.infrastructure.native_docx_structure import NativeDocxStructure
from tests.native_docx_helpers import call, read_dfm
from tests.native_docx_helpers import native_docx as native_docx
from tests.native_docx_structure_helpers import creation, paragraph


@pytest.fixture
def structured(native_docx):
    service, asset, source = native_docx
    service.docx_structure = NativeDocxStructure()
    service.docx_operations.structure = service.docx_structure
    return service, asset, source


def references(service, asset):
    return call(
        service,
        op="read_docx",
        asset_id=asset["asset_id"],
        revision=asset["revision"],
        limit=100,
    )["blocks"]


def test_create_edit_and_evidence_wiki_without_source(structured, tmp_path):
    service, _, _ = structured
    request = creation()
    result = call(service, op="create_docx", docx_create=request.model_dump())
    asset = result["asset"]
    assert asset["source"] is None and result["source_written"] is False
    contract = call(service, op="contract", for_op="create_docx")
    assert (
        contract["docx_structure_enabled"]
        and "create_docx" in contract["formats"]["docx"]
    )
    original = service.repository.read(asset["asset_id"], asset["revision"])
    assert Document(io.BytesIO(original)).tables[0].cell(1, 0).text == "007"
    table = next(
        block for block in references(service, asset) if block["type"] == "table"
    )
    assert call(service, op="verify", reference=table["evidence"])["valid"]
    text = read_dfm(service, asset)
    changed = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={"dfm_text": text.replace("1,234.50", "1,234.75")},
    )["asset"]
    data = service.repository.read(changed["asset_id"], changed["revision"])
    assert Document(io.BytesIO(data)).tables[0].cell(2, 1).text == "1,234.75"
    exported = call(
        service,
        op="export_wiki",
        asset_id=asset["asset_id"],
        revision=changed["revision"],
        output_dir=str(tmp_path / "wiki"),
    )
    assert exported["success"]
    assert call(service, op="verify", reference=table["evidence"])["valid"]


@pytest.mark.parametrize("position", ["start", "end", "before", "after"])
def test_insert_delete_body_preserves_source_and_historical_proof(structured, position):
    service, asset, source = structured
    original = source.read_bytes()
    anchor = references(service, asset)[1]["evidence"]
    insert = {"position": position, "blocks": [paragraph("TEMP 007", bold=True)]}
    if position in {"before", "after"}:
        insert["anchor"] = anchor
    added = call(
        service,
        op="add_docx_blocks",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_insert=insert,
    )["asset"]
    block = next(b for b in references(service, added) if b["preview"] == "TEMP 007")
    removed = call(
        service,
        op="delete_docx_blocks",
        asset_id=asset["asset_id"],
        expected_revision=added["revision"],
        docx_block_refs=[block["evidence"]],
    )["asset"]
    assert all(b["preview"] != "TEMP 007" for b in references(service, removed))
    assert call(service, op="verify", reference=block["evidence"])["valid"]
    assert source.read_bytes() == original


@pytest.mark.parametrize(
    "fault", ["stale", "reference", "archive", "header", "picture", "duplicate"]
)
def test_invalid_structural_targets_never_commit(structured, fault):
    service, asset, source = structured
    blocks = references(service, asset)
    reference = blocks[0]["evidence"]
    revision = asset["revision"]
    if fault == "stale":
        revision = "0" * 64
    elif fault == "reference":
        reference = {**reference, "value_sha256": "0" * 64}
    elif fault == "archive":
        call(
            service,
            op="archive",
            asset_id=asset["asset_id"],
            expected_revision=revision,
        )
    elif fault == "header":
        reference = next(
            b["evidence"]
            for b in blocks
            if "header" in b["evidence"]["locator"]["part"]
        )
    elif fault == "picture":
        reference = next(
            b["evidence"]
            for b in blocks
            if b["metadata"]["source_element"] == "w:drawing"
        )
    refs = [reference, reference] if fault == "duplicate" else [reference]
    with pytest.raises(ValueError):
        call(
            service,
            op="delete_docx_blocks",
            asset_id=asset["asset_id"],
            expected_revision=revision,
            docx_block_refs=refs,
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == asset["revision"]


def test_concurrent_structure_commit_preserves_winner(structured, monkeypatch):
    service, asset, _ = structured
    insert = service.docx_structure.insert
    winner = None

    def race(data, blocks, position):
        nonlocal winner
        candidate = insert(data, blocks, position)
        other, checks = insert(data, creation().blocks[-1:], "start")
        winner = service.repository.commit(
            asset["asset_id"], asset["revision"], other, checks
        )
        return candidate

    monkeypatch.setattr(service.docx_structure, "insert", race)
    with pytest.raises(ValueError, match=r"revision|concurrent|changed|stale"):
        call(
            service,
            op="add_docx_blocks",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_insert={"position": "end", "blocks": [paragraph("LOSER")]},
        )
    assert service.repository.load(asset["asset_id"]).revision == winner.revision
    assert "LOSER" not in read_dfm(
        service, {"asset_id": asset["asset_id"], "revision": winner.revision}
    )
