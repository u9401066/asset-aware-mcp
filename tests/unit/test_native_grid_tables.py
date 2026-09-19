"""Real XLSX table ranges, formulas, column identities, headers and native bytes."""

import io

import openpyxl
import pytest
from lxml import etree

from src.domain.native_grid import GridTransform, NativeGridEdit
from src.infrastructure.native_grid_filters import shift_filters
from src.infrastructure.native_grid_xml import tag
from src.infrastructure.native_ooxml import DOC_REL_NS, xml_bytes
from src.infrastructure.native_spreadsheet_reader import NS, NativeSpreadsheetReader
from tests.native_grid_helpers import apply_table_grid, table_workbook
from tests.native_workbook_helpers import _parts, _replace


def edit(axis="row", operation="insert", at=3, count=1, **kw):
    return GridTransform(
        NativeGridEdit(axis=axis, operation=operation, at=at, count=count, **kw)
    )


def table(data):
    return etree.fromstring(_parts(data)["xl/tables/table1.xml"])


def read(data, cell, sheet="Data"):
    return NativeSpreadsheetReader(data).read_cell(sheet, cell)


def test_insert_column_retains_ids_and_existing_styles_and_generates_unique_header():
    source = table_workbook()
    updated, result = apply_table_grid(source, edit(axis="column", at=2))
    root = table(updated)
    columns = root.findall("s:tableColumns/s:tableColumn", NS)
    assert root.get("ref") == "A2:D5"
    assert [(col.get("id"), col.get("name")) for col in columns] == [
        ("1", "Input"),
        ("4", "Column2"),
        ("2", "Calc"),
        ("3", "Label"),
    ]
    assert read(updated, "B2")["value"] == "Column2"
    assert read(updated, "B3")["kind"] == "blank"
    assert read(updated, "C3")["value"] == read(source, "B3")["value"]
    assert read(updated, "C3")["style_index"] == read(source, "B3")["style_index"]
    assert [
        node.get("colId") for node in root.findall("s:autoFilter/s:filterColumn", NS)
    ] == ["0", "3"]
    assert _parts(source)["xl/styles.xml"] == _parts(updated)["xl/styles.xml"]
    assert _parts(source)["docProps/app.xml"] == _parts(updated)["docProps/app.xml"]
    assert (
        _parts(source)["customXml/unchanged.xml"]
        == _parts(updated)["customXml/unchanged.xml"]
    )
    parsed = openpyxl.load_workbook(io.BytesIO(updated), data_only=False)
    assert parsed["Data"].tables["Table1"].ref == "A2:D5"
    assert result.changes[0]["tables"][0]["generated_headers"] == 1


def test_delete_first_column_shrinks_intervals_and_invalidates_exact_removed_column():
    source = table_workbook()
    updated, _ = apply_table_grid(source, edit(axis="column", operation="delete", at=1))
    assert [
        (node.get("id"), node.get("name"))
        for node in table(updated).findall("s:tableColumns/s:tableColumn", NS)
    ] == [("2", "Calc"), ("3", "Label")]
    assert read(updated, "A1", "Other")["value"] == "=SUM(Table1[[Calc]:[Label]])"
    assert read(updated, "A2", "Other")["value"] == "=SUM(#REF!)"
    assert read(updated, "A3")["value"] == "=#REF!*2"
    assert table(updated).find("s:autoFilter/s:sortState", NS) is None
    assert [
        node.get("colId")
        for node in table(updated).findall("s:autoFilter/s:filterColumn", NS)
    ] == ["1"]


@pytest.mark.parametrize(
    "formula,expected", [("=A3*2", "A3*2"), ("=[@Input]*2", "[[#This Row],Input]*2")]
)
def test_insert_first_data_row_generates_formula_at_correct_origin(formula, expected):
    source = table_workbook(formula=formula)
    updated, result = apply_table_grid(
        source, edit(at=3, count=2, inherit_format="after")
    )
    assert table(updated).get("ref") == "A2:C7"
    assert (
        table(updated).findtext(
            's:tableColumns/s:tableColumn[@id="2"]/s:calculatedColumnFormula',
            namespaces=NS,
        )
        == expected
    )
    assert read(updated, "B3")["value"] == "=" + expected
    assert read(updated, "B4")["value"] == "=" + (expected.replace("A3", "A4"))
    assert read(updated, "A5")["value"] == 10
    assert read(updated, "B3")["style_index"] == read(source, "B3")["style_index"]
    assert result.changes[0]["tables"][0]["generated_calculated_cells"] == 2


def test_delete_first_data_row_rebases_template_to_surviving_data():
    source = table_workbook(formula="=A3*2")
    updated, _ = apply_table_grid(source, edit(operation="delete", at=3))
    assert read(updated, "A3")["value"] == 20
    assert read(updated, "B3")["value"] == "=A3*2"
    assert (
        table(updated).findtext(
            's:tableColumns/s:tableColumn[@id="2"]/s:calculatedColumnFormula',
            namespaces=NS,
        )
        == "A3*2"
    )


def test_literal_repeated_a1_formulas_retain_actual_source_dependency_semantics():
    source = table_workbook(formula="=A3*2", expand_a1=False)
    assert read(source, "B4")["value"] == "=A3*2"
    updated, _ = apply_table_grid(source, edit(operation="delete", at=3))
    assert read(updated, "B3")["value"] == "=#REF!*2"


def test_header_deletion_keeps_data_and_names_with_hidden_header():
    source = table_workbook()
    updated, result = apply_table_grid(source, edit(operation="delete", at=2))
    root = table(updated)
    assert root.get("headerRowCount") == "0" and root.get("ref") == "A2:C4"
    assert root.find("s:autoFilter", NS) is None
    assert read(updated, "A2")["value"] == 10 and read(updated, "C2")["value"] == "x"
    assert root.find("s:tableColumns/s:tableColumn", NS).get("name") == "Input"
    assert result.changes[0]["tables"][0]["header_hidden_after_deletion"] is True
    parsed = openpyxl.load_workbook(io.BytesIO(updated))
    assert parsed["Data"].tables["Table1"].headerRowCount == 0


def test_totals_row_deletion_removes_only_totals_definitions():
    source = table_workbook(totals=True)
    updated, _ = apply_table_grid(source, edit(operation="delete", at=6))
    root = table(updated)
    assert root.get("totalsRowCount") == "0" and root.get("totalsRowShown") == "0"
    assert all(
        node.get("totalsRowFunction") is None and node.get("totalsRowLabel") is None
        for node in root.findall("s:tableColumns/s:tableColumn", NS)
    )
    assert read(updated, "A5")["value"] == 30
    assert read(updated, "A2")["value"] == "Input"


def test_whole_table_deletion_detaches_parts_and_repairs_surviving_references():
    source = table_workbook()
    updated, result = apply_table_grid(source, edit(operation="delete", at=2, count=4))
    parts = _parts(updated)
    assert parts["xl/tables/table1.xml"] == _parts(source)["xl/tables/table1.xml"]
    assert (
        etree.fromstring(parts["xl/worksheets/sheet1.xml"]).find("s:tableParts", NS)
        is None
    )
    assert not any(
        node.get("Type") == f"{DOC_REL_NS}/table"
        for node in etree.fromstring(parts["xl/worksheets/_rels/sheet1.xml.rels"])
    )
    assert read(updated, "A1", "Other")["value"] == "=SUM(#REF!)"
    assert (
        NativeSpreadsheetReader(updated).workbook.findtext(
            "s:definedNames/s:definedName", namespaces=NS
        )
        == "#REF!"
    )
    assert result.changes[0]["tables"][0]["after"] is None


def test_insert_before_table_moves_range_without_inventing_a_table_column():
    source = table_workbook()
    updated, _ = apply_table_grid(source, edit(axis="column", at=1))
    assert table(updated).get("ref") == "B2:D5"
    assert [
        node.get("id")
        for node in table(updated).findall("s:tableColumns/s:tableColumn", NS)
    ] == ["1", "2", "3"]
    assert read(updated, "B2")["value"] == "Input"


def test_duplicate_filter_columns_fail_and_deleted_whole_filter_is_removed():
    root = table(table_workbook())
    auto = root.find("s:autoFilter", NS)
    etree.SubElement(auto, tag("filterColumn"), colId="0")
    with pytest.raises(ValueError, match="duplicate filter"):
        shift_filters(root, edit(axis="column"))
    root = table(table_workbook())
    shift_filters(root, edit(operation="delete", at=2, count=4))
    assert root.find("s:autoFilter", NS) is None


def test_mapped_table_column_structure_requires_coordinated_field_mapping():
    source = table_workbook()
    root = table(source)
    root.set("tableType", "xml")
    source = _replace(source, {"xl/tables/table1.xml": xml_bytes(root)})
    with pytest.raises(ValueError, match="source field identities"):
        apply_table_grid(source, edit(axis="column", at=2))
    updated, _ = apply_table_grid(source, edit(at=1))
    assert table(updated).get("ref") == "A3:C6"
