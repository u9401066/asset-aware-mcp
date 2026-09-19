"""Standalone Word creation and byte-preserving body structure mutations."""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.domain.native_docx_structure import NativeDocxCreate
from src.infrastructure.native_docx_structure import (
    NativeDocxStructure,
    body_of,
    canonical,
)
from src.infrastructure.native_docx_workspace import MAIN_PART, checked_docx
from tests.native_docx_helpers import build_docx, replace_parts
from tests.native_docx_structure_helpers import creation


def test_create_rich_literal_content_merges_and_page_setup():
    request = creation()
    data = NativeDocxStructure().create(request)
    document = Document(io.BytesIO(data))
    assert document.paragraphs[0].text == "研究 007\tµg\n測試"
    run = document.paragraphs[0].runs[0]
    assert run.font.size.pt == 11.5 and run.bold and run.underline
    assert run.font.cs_bold and run._r.rPr.find(qn("w:szCs")).get(qn("w:val")) == "23"
    table = document.tables[0]
    assert table.cell(0, 2).text == "Header"
    assert table.cell(2, 0).text == "007"
    assert table.cell(1, 1).text == "-0.50\nSecond"
    assert table.cell(2, 1).text == "1,234.50"
    assert [column.width.twips for column in table.columns] == [1200, 1800, 2400]
    assert document.sections[0].left_margin.twips == 1440
    assert document.core_properties.author == "u9401066"


@pytest.mark.parametrize("position", ["start", "end", 1, 3])
def test_insert_delete_keeps_source_headers_media_and_foreign_xml(position):
    original = replace_parts(
        build_docx(), {"customXml/foreign.xml": b"<foreign> unchanged </foreign>"}
    )
    before = checked_docx(original)
    adapter = NativeDocxStructure()
    updated, report = adapter.insert(original, creation().blocks, position)
    after = checked_docx(updated)
    assert report.changed_parts == [MAIN_PART]
    assert all(
        after.parts[name] == content
        for name, content in before.parts.items()
        if name != MAIN_PART
    )
    restored, _ = adapter.delete(updated, report.changes[0]["body_positions"])
    assert canonical(checked_docx(restored).xml(MAIN_PART)) == canonical(
        before.xml(MAIN_PART)
    )
    assert checked_docx(restored).parts.keys() == before.parts.keys()


def test_empty_document_delete_all_and_reinsert():
    adapter = NativeDocxStructure()
    empty = adapter.create(NativeDocxCreate(name="empty.docx"))
    assert len(body_of(checked_docx(empty).xml(MAIN_PART))) == 1
    inserted, _ = adapter.insert(empty, creation().blocks, "end")
    restored, _ = adapter.delete(inserted, [2, 0, 1])
    assert canonical(checked_docx(empty).xml(MAIN_PART)) == canonical(
        checked_docx(restored).xml(MAIN_PART)
    )


def test_landscape_page_and_no_borders():
    request = creation().model_dump()
    request.update(page_width_twips=15840, page_height_twips=12240)
    request["blocks"][1]["borders"] = False
    data = NativeDocxStructure().create(NativeDocxCreate.model_validate(request))
    document = Document(io.BytesIO(data))
    assert document.sections[0].page_width.twips == 15840
    assert (
        document.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))[0].get(qn("w:val"))
        == "nil"
    )
