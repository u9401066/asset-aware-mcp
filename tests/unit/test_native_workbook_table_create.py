"""Native Table creation preserves data, style, package identity and evidence."""

import io

import openpyxl
import pytest
from lxml import etree
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, PatternFill

from src.domain.native_assets import NativeCellEdit
from src.domain.native_table_create import NativeTableCreate
from src.domain.native_table_edit import NativeTableUpdate
from src.infrastructure.native_grid_xml import cell_map, tag
from src.infrastructure.native_spreadsheet import NativeSpreadsheet
from src.infrastructure.native_spreadsheet_reader import NS
from src.infrastructure.native_workbook_structure import NativeWorkbookStructure
from src.infrastructure.native_workbook_table_create import NativeWorkbookTableCreate
from src.infrastructure.native_workbook_table_edit import NativeWorkbookTableEdit
from tests.native_workbook_helpers import _parts, _replace
from tests.unit.test_native_workbook_table_edit import mutate

SHEET = "xl/worksheets/sheet1.xml"
KEY = {"sheet_id": "1", "part": SHEET}
ADAPTER = NativeWorkbookTableCreate()


def source(*, rich=True):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Code", "Count", "Double"])
    sheet.append(["007", 3, None])
    sheet.append(["0010", 4, None])
    sheet["A2"].number_format = "@"
    sheet["B2"].number_format = "0.00"
    sheet["C2"].fill = PatternFill("solid", fgColor="FFCC00")
    sheet["C2"].font = Font(bold=True)
    if rich:
        sheet["A1"] = CellRichText(
            TextBlock(InlineFont(b=True), "Co"), TextBlock(InlineFont(i=True), "de")
        )
    sheet["G1"] = "Outside"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return _replace(
        output.getvalue(), {"customXml/preserved.xml": b"<keep>  exact </keep>"}
    )


def request(**changes):
    return NativeTableCreate.model_validate(
        {
            "worksheet": KEY,
            "ref": "A1:C3",
            "name": "Records",
            "columns": [{"name": n} for n in ("Code", "Count", "Double")],
            **changes,
        }
    )


def test_openpyxl_default_unlocked_workbook_creates_table_preserving_cells():
    original = source()
    assert b"workbookProtection" in _parts(original)["xl/workbook.xml"]
    output, result = ADAPTER.create(original, request())
    before, after = _parts(original), _parts(output)
    old_cells = cell_map(etree.fromstring(before[SHEET]))
    new_cells = cell_map(etree.fromstring(after[SHEET]))
    assert {k: etree.tostring(v, method="c14n") for k, v in old_cells.items()} == {
        k: etree.tostring(v, method="c14n") for k, v in new_cells.items()
    }
    with_book = openpyxl.load_workbook(io.BytesIO(output), rich_text=True)
    try:
        table = with_book["Data"].tables["Records"]
        assert table.ref == table.autoFilter.ref == "A1:C3"
        assert [(c.id, c.name) for c in table.tableColumns] == [
            (1, "Code"),
            (2, "Count"),
            (3, "Double"),
        ]
        assert table.tableStyleInfo.name == "TableStyleMedium2"
        assert with_book["Data"]["A1"].value[0].font.b
        assert with_book["Data"]["A1"].value[1].font.i
        assert with_book["Data"]["A2"].value == "007"
        assert with_book["Data"]["A2"].data_type == "s"
    finally:
        with_book.close()
    assert result.changes[0]["cells"] == {}
    for part in before:
        if part not in result.changed_parts:
            assert before[part] == after[part]
    assert before["xl/styles.xml"] == after["xl/styles.xml"]
    table = NativeWorkbookStructure().read(output, references=True)["tables"][0]
    assert table["header_cells"][0]["value"]["value"] == "Code"


def test_calculated_totals_and_styles_read_back_and_can_be_edited():
    original = source()
    intent = request(
        ref="A1:C4",
        totals_row=True,
        columns=[
            {"name": "Code", "totals": {"kind": "label", "value": "合計"}},
            {"name": "Count"},
            {
                "name": "Double",
                "calculated": {"formula": "=B2*2", "policy": "require_matching"},
                "totals": {"kind": "function", "value": "sum"},
            },
        ],
    )
    output, result = ADAPTER.create(original, intent)
    reader = NativeSpreadsheet(output)
    for cell, expected in {
        "A4": "合計",
        "C2": "=B2*2",
        "C3": "=B3*2",
        "C4": "=SUBTOTAL(109,[Double])",
    }.items():
        assert reader.read_cell("Data", cell)["value"] == expected
    assert (
        reader.read_cell("Data", "C2")["style_index"]
        == NativeSpreadsheet(original).read_cell("Data", "C2")["style_index"]
    )
    table = NativeWorkbookStructure().read(output)["tables"][0]
    assert 'autoFilter ref="A1:C3"' in table["xml"]
    assert result.changes[0]["cells"]["C2"]["before"]["kind"] == "blank"
    updated, _ = NativeWorkbookTableEdit().update(
        output,
        NativeTableUpdate(
            worksheet=KEY,
            part=table["part"],
            expected_ref="A1:C4",
            columns=[{"column_id": 3, "expected_name": "Double", "name": "Twice"}],
        ),
    )
    assert (
        NativeSpreadsheet(updated).read_cell("Data", "C4")["value"]
        == "=SUBTOTAL(109,[Twice])"
    )
    with pytest.raises(ValueError, match="table-aware"):
        NativeSpreadsheet(output).edit(
            [NativeCellEdit(sheet="Data", cell="C2", kind="number", value=3)]
        )


def test_blank_header_policy_never_coerces_or_renames_existing_values():
    original = source()
    original = mutate(
        original,
        SHEET,
        lambda root: root.find("s:sheetData/s:row", NS).remove(
            root.find("s:sheetData/s:row/s:c", NS)
        ),
    )
    with pytest.raises(ValueError, match="header must match"):
        ADAPTER.create(original, request())
    output, result = ADAPTER.create(original, request(header_policy="fill_blank"))
    assert NativeSpreadsheet(output).read_cell("Data", "A1")["value"] == "Code"
    assert result.changes[0]["cells"]["A1"]["before"]["kind"] == "blank"
    with pytest.raises(ValueError, match="header must match"):
        ADAPTER.create(
            source(),
            request(
                header_policy="fill_blank",
                columns=[{"name": "ID"}, {"name": "Count"}, {"name": "Double"}],
            ),
        )


def test_explicit_headerless_table_preserves_first_data_row():
    original = source()
    output, result = ADAPTER.create(
        original, request(ref="A2:C3", header_row=False, autofilter=False)
    )
    table = NativeWorkbookStructure().read(output)["tables"][0]
    assert table["attributes"]["headerRowCount"] == "0"
    assert table["header_cells"] == [] and "autoFilter" not in table["xml"]
    assert result.changes[0]["cells"] == {}
    assert NativeSpreadsheet(output).read_cell("Data", "A2")["value"] == "007"


@pytest.mark.parametrize(
    "name", ["R", "c", "A1", "r1c1", "2Records", "two words", "Records!", ""]
)
def test_invalid_table_names_rejected_by_typed_contract(name):
    with pytest.raises(ValueError):
        request(name=name)


@pytest.mark.parametrize(
    "name", ["紀錄", "_Records", "\\Records", "Records.v2", "XFE1"]
)
def test_valid_unicode_and_non_reference_names(name):
    output, _ = ADAPTER.create(source(), request(name=name))
    assert (
        NativeWorkbookStructure().read(output)["tables"][0]["attributes"]["name"]
        == name
    )


@pytest.mark.parametrize(
    "changes",
    [
        {"ref": "A1:B3"},
        {"ref": "a1:c3"},
        {"ref": "A1:C1"},
        {"ref": "A1:C2", "totals_row": True},
        {"style": {"name": "MissingStyle"}},
    ],
)
def test_range_and_style_guards(changes):
    with pytest.raises(ValueError):
        ADAPTER.create(source(), request(**changes))


def test_totals_never_consume_the_last_data_row():
    with pytest.raises(ValueError, match="totals row must be blank"):
        ADAPTER.create(source(), request(totals_row=True))


def test_existing_calculated_data_needs_explicit_replacement():
    original, _ = NativeSpreadsheet(source()).edit(
        [NativeCellEdit(sheet="Data", cell="C2", kind="number", value=900)]
    )
    columns = [{"name": n} for n in ("Code", "Count", "Double")]
    columns[-1]["calculated"] = {"formula": "=[@Count]*3", "policy": "require_matching"}
    with pytest.raises(ValueError, match="exception at C2"):
        ADAPTER.create(original, request(columns=columns))
    columns[-1]["calculated"]["policy"] = "replace_all"
    output, result = ADAPTER.create(original, request(columns=columns))
    assert result.changes[0]["cells"]["C2"]["before"]["value"] == 900
    assert NativeSpreadsheet(output).read_cell("Data", "C3")["value"] == "=[@Count]*3"


@pytest.mark.parametrize("kind", ["merge", "filter", "array", "shared", "protected"])
def test_overlapping_structures_and_protection_rejected(kind):
    def change(root):
        if kind == "merge":
            etree.SubElement(
                etree.SubElement(root, tag("mergeCells")), tag("mergeCell"), ref="B3:D5"
            )
        elif kind == "filter":
            etree.SubElement(root, tag("autoFilter"), ref="A1:C3")
        elif kind == "protected":
            etree.SubElement(root, tag("sheetProtection"), sheet="1")
        else:
            cell = root.find("s:sheetData/s:row/s:c", NS)
            etree.SubElement(cell, tag("f"), t=kind, ref="A1:B4").text = "1"

    with pytest.raises(ValueError, match=r"overlaps|Protected"):
        ADAPTER.create(mutate(source(), SHEET, change), request())


def test_table_and_defined_names_are_unique_ignoring_case():
    output, _ = ADAPTER.create(source(), request())
    with pytest.raises(ValueError, match="overlaps"):
        ADAPTER.create(output, request(name="Second"))
    with pytest.raises(ValueError, match="name conflicts"):
        ADAPTER.create(
            output, request(ref="G1:G2", name="records", columns=[{"name": "Outside"}])
        )
    original = mutate(
        source(),
        "xl/workbook.xml",
        lambda root: etree.SubElement(
            root.find("s:definedNames", NS), tag("definedName"), name="records"
        ).__setattr__("text", "Data!$A$2"),
    )
    with pytest.raises(ValueError, match="name conflicts"):
        ADAPTER.create(original, request())


@pytest.mark.parametrize(
    "flags", [{}, {"lockStructure": "0"}, {"lockWindows": "false", "lockRevision": "0"}]
)
def test_empty_or_disabled_protection_preserved(flags):
    original = mutate(
        source(),
        "xl/workbook.xml",
        lambda root: root.find("s:workbookProtection", NS).attrib.update(flags),
    )
    output, _ = ADAPTER.create(original, request())
    assert (
        dict(
            etree.fromstring(_parts(output)["xl/workbook.xml"])
            .find("s:workbookProtection", NS)
            .attrib
        )
        == flags
    )


@pytest.mark.parametrize(
    "flags",
    [
        {"lockStructure": "1"},
        {"lockWindows": "true"},
        {"lockRevision": "1"},
        {"workbookPassword": "ABCD"},
        {"lockStructure": "invalid"},
    ],
)
def test_active_or_unknown_protection_stays_blocked(flags):
    original = mutate(
        source(),
        "xl/workbook.xml",
        lambda root: root.find("s:workbookProtection", NS).attrib.update(flags),
    )
    with pytest.raises(ValueError, match="Protected workbook"):
        ADAPTER.create(original, request())
