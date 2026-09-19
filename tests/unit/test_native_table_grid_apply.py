"""A2T structural edits must preserve native payloads and commit one exact result."""

import io
import json

import openpyxl
import pytest

from src.application.native_table_projection import (
    canonical_table,
    matching_grid,
    workspace_hash,
)
from src.domain.native_grid import NativeGridUpdate
from src.domain.native_table_grid import proposed_grid, validate_grid
from src.domain.table_state import table_from_state, table_state
from src.infrastructure.native_workbook_grid import NativeWorkbookGrid
from tests.native_table_helpers import project, read_workspace, service_at
from tests.native_workbook_helpers import _call, _parts
from tests.unit.test_native_workbook_operations import read as read_workbook


def grid_service(root):
    service, tables = service_at(root)
    from src.application.native_table_grid_apply import NativeTableGridApply

    service.table_operations.grid_apply = NativeTableGridApply(
        NativeWorkbookGrid(),
        service.spreadsheets,
        service.table_operations.ranges,
    )
    return service, tables


def apply_grid(service, asset, table_id):
    record, digest = read_workspace(service, table_id)
    return _call(
        service,
        op="apply_table_workspace",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        table_id=table_id,
        expected_table_sha256=digest,
        worksheet_grid=record["structural_plan"]["worksheet_grid"],
    )


def test_structural_apply_preserves_rich_payloads_and_relocates_formulas(tmp_path):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    original, digest = read_workspace(service, table_id)
    source = tmp_path / "source.xlsx"
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    historical = original["source_cells"][8]["source"]["evidence"]
    tables.delete_row(table_id, 0)
    tables.remove_column(table_id, "C")
    tables.rename_column(table_id, "D", "Rich")
    tables.add_column(
        table_id, "Tag", "native", default_value={"kind": "string", "value": "new"}
    )
    tables.add_rows(
        table_id,
        [
            {
                "A": {"kind": "number", "value": 30},
                "Tag": {"kind": "string", "value": "last"},
            }
        ],
    )
    current, digest = read_workspace(service, table_id)
    assert not current["source_correspondence_unchanged"]
    assert current["structural_plan"]["status"] == "required"
    result = apply_grid(service, asset, table_id)
    stored = service.repository.load(asset["asset_id"])
    assert len(stored.history) == 2
    data = service.repository.read(asset["asset_id"], stored.revision)
    workbook = openpyxl.load_workbook(io.BytesIO(data), rich_text=True)
    sheet = workbook["Data"]
    assert sheet["A1"].value == 10 and sheet["A2"].value == 20
    assert sheet["A3"].value == 30
    assert sheet["B1"].value == "=A1*2"
    assert str(sheet["C1"].value) == "Bold normal"
    assert type(sheet["C1"].value).__name__ == "CellRichText"
    assert sheet["A1"].number_format == "$0.00"
    assert [sheet[f"E{row}"].value for row in range(1, 4)] == ["new", "new", "last"]
    assert sheet["D2"].value == "=SUM(A2)" and sheet["D2"].data_type == "s"
    assert workbook["Other"]["B1"].value == "=Data!A1"
    assert "A5:C6" in {str(value) for value in sheet.merged_cells.ranges}
    assert _parts(data)["customXml/item1.xml"] == _parts(before)["customXml/item1.xml"]
    assert _parts(data)["xl/styles.xml"] == _parts(before)["xl/styles.xml"]
    receipt = read_workbook(service, result["asset"])["operation_result"]
    assert receipt["changes"][-1]["destination_projection"]["end_cell"] == "E3"
    assert "complete_structural_workspace_values_read_back" in receipt["checks"]
    assert receipt["preserved_parts"] == len(_parts(before)) - len(
        receipt["changed_parts"]
    )
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime
    assert _call(service, op="verify", reference=historical)["valid"]
    assert read_workspace(
        service, table_id, workspace_reference=result["workspace_reference"]
    ) == (current, digest)
    assert current["table"]["native_binding"]["source"]["revision"] == asset["revision"]


def test_rename_keeps_identity_but_recreated_same_name_requires_grid(tmp_path):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    context = tables.read_workspace(table_id)
    old = context.column_ids[-1]
    tables.rename_column(table_id, "E", "Code")
    context = tables.read_workspace(table_id)
    assert matching_grid(context) and context.column_ids[-1] == old
    tables.remove_column(table_id, "Code")
    tables.add_column(
        table_id,
        "E",
        "native",
        default_value={"kind": "string", "value": "replacement"},
    )
    context = tables.read_workspace(table_id)
    assert context.column_ids[-1] != old and not matching_grid(context)
    result = apply_grid(service, asset, table_id)
    book = openpyxl.load_workbook(
        io.BytesIO(
            service.repository.read(asset["asset_id"], result["asset"]["revision"])
        )
    )
    assert [book["Data"][f"E{row}"].value for row in range(1, 4)] == ["replacement"] * 3


def test_identical_new_row_never_resurrects_deleted_source_identity(tmp_path):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    before = tables.read_workspace(table_id)
    tables.delete_row(table_id, 2)
    tables.add_rows(table_id, [before.rows[2]])
    after = tables.read_workspace(table_id)
    assert after.row_ids[2] != before.row_ids[2]
    assert not matching_grid(after)
    request = proposed_grid(after)
    assert [edit.operation for edit in request.edits] == ["delete", "insert"]
    apply_grid(service, asset, table_id)


def test_legacy_snapshot_hash_survives_column_identity_upgrade(tmp_path):
    service, tables = grid_service(tmp_path)
    _, projection = project(service, tmp_path)
    context = tables.read_workspace(projection["table_id"])
    legacy = table_state(context)
    legacy.pop("column_ids")
    legacy["native_binding"].pop("column_ids")
    expected = canonical_table(legacy)
    loaded = table_from_state(json.loads(expected))
    assert canonical_table(table_state(loaded)) == expected
    assert loaded.column_ids == []
    loaded.rename_column("E", "Code")
    assert loaded.column_ids == loaded.native_binding.column_ids
    assert matching_grid(loaded)
    assert (
        canonical_table(table_state(table_from_state(json.loads(expected)))) == expected
    )
    legacy.pop("created_at")
    assert workspace_hash(table_from_state(legacy)) == workspace_hash(
        table_from_state(legacy)
    )
    assert table_state(table_from_state(legacy))["created_at"] == ""


@pytest.mark.parametrize("kind", ["wrong_axis", "outside", "delete_survivor"])
def test_mismatched_grid_is_rejected_before_native_commit(tmp_path, kind):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    tables.add_rows(table_id, [{"A": {"kind": "number", "value": 30}}])
    context = tables.read_workspace(table_id)
    grid = proposed_grid(context).model_dump(exclude_none=True)
    if kind == "wrong_axis":
        grid["edits"][0]["axis"] = "column"
    elif kind == "outside":
        grid["edits"][0]["at"] = 20
    else:
        grid["edits"].insert(0, {"axis": "row", "operation": "delete", "at": 1})
    with pytest.raises(ValueError, match=r"projected|identities"):
        validate_grid(context, NativeGridUpdate.model_validate(grid))
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


@pytest.mark.parametrize("axis", ["row", "column"])
def test_remove_complete_projected_axis_keeps_other_native_content(tmp_path, axis):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path, start="E1", end="E3")
    table_id = projection["table_id"]
    if axis == "row":
        for _ in range(3):
            tables.delete_row(table_id, 0)
    else:
        tables.remove_column(table_id, "E")
    result = apply_grid(service, asset, table_id)
    book = openpyxl.load_workbook(
        io.BytesIO(
            service.repository.read(asset["asset_id"], result["asset"]["revision"])
        )
    )
    assert book["Other"]["A1"].value == "Do not change"
    event = read_workbook(service, result["asset"])["operation_result"]["changes"][-1]
    assert event["destination_projection"] is None


def test_late_native_conflict_retains_concurrent_revision(tmp_path, monkeypatch):
    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    tables.add_rows(table_id, [{"A": {"kind": "number", "value": 30}}])
    adapter = service.table_operations.grid_apply.grid
    update = adapter.update
    concurrent = []

    def raced(data, request):
        result = update(data, request)
        written = _call(
            service,
            op="update",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            edits=[
                {"sheet": "Data", "cell": "E1", "kind": "string", "value": "concurrent"}
            ],
        )
        concurrent.append(written["asset"]["revision"])
        return result

    monkeypatch.setattr(adapter, "update", raced)
    with pytest.raises(ValueError, match=r"[Ss]tale|changed"):
        apply_grid(service, asset, table_id)
    current = service.repository.load(asset["asset_id"])
    assert current.revision == concurrent[0] and len(current.history) == 2
    assert (tmp_path / "source.xlsx").read_bytes() == service.repository.read(
        asset["asset_id"], asset["revision"]
    )


def test_composite_readback_detects_unedited_format_tampering(tmp_path, monkeypatch):
    from lxml import etree

    from tests.native_workbook_helpers import _replace

    service, tables = grid_service(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    tables.add_rows(table_id, [{"A": {"kind": "number", "value": 30}}])
    original_edit = service.spreadsheets.edit

    def damaged(data, edits):
        updated, result = original_edit(data, edits)
        part = "xl/worksheets/sheet1.xml"
        root = etree.fromstring(_parts(updated)[part])
        root.find(".//{*}c[@r='A2']").set("s", "0")
        return _replace(updated, {part: etree.tostring(root)}), result

    monkeypatch.setattr(service.spreadsheets, "edit", damaged)
    with pytest.raises(ValueError, match="representation changed"):
        apply_grid(service, asset, table_id)
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


def test_invalid_native_column_default_never_changes_schema(tmp_path):
    service, tables = grid_service(tmp_path)
    _, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    before = workspace_hash(tables.read_workspace(table_id))
    with pytest.raises(ValueError):
        tables.add_column(table_id, "Invalid", "native", default_value="007")
    assert workspace_hash(tables.read_workspace(table_id)) == before
