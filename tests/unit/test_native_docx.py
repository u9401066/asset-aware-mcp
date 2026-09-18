"""Native DOCX updates must reuse DFM guards and preserve the human source."""

from __future__ import annotations

import hashlib
import io
from pathlib import Path

import pytest
import yaml
from docx import Document
from lxml import etree

from src.application.docx_service import DocxService
from src.domain.native_docx import MAX_DFM_BYTES, NativeDocxEdit
from src.infrastructure.native_docx_workspace import WORD_NS
from src.infrastructure.native_ooxml import NativeOOXMLPackage
from tests.native_docx_helpers import call, read_dfm, replace_parts
from tests.native_docx_helpers import native_docx as native_docx


def test_native_docx_edit_preserves_unrelated_parts_and_human_source(
    native_docx,
) -> None:
    service, asset, source = native_docx
    before = source.read_bytes()
    mtime = source.stat().st_mtime_ns
    dfm = (
        read_dfm(service, asset)
        .replace("原始段落", "修訂段落")
        .replace("Old value", "New value")
    )
    result = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={"dfm_text": dfm},
    )
    assert result["success"] and result["source_written"] is False
    assert result["asset"]["revision"] != asset["revision"]
    updated = service.repository.read(asset["asset_id"])
    old, new = NativeOOXMLPackage(before), NativeOOXMLPackage(updated)
    assert old.parts.keys() == new.parts.keys()
    assert result["operation_result"]["changed_parts"] == ["word/document.xml"]
    for name in old.parts.keys() - {"word/document.xml"}:
        assert old.parts[name] == new.parts[name], name
    document = Document(io.BytesIO(updated))
    assert document.paragraphs[1].text == "修訂段落 with italic context"
    assert document.paragraphs[1].runs[0].bold is True
    assert document.paragraphs[1].runs[1].italic is True
    assert document.tables[0].cell(1, 1).text == "New value"
    assert document.tables[0].style.name == "Table Grid"
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime
    assert service.repository.read(asset["asset_id"], asset["revision"]) == before
    assert "rendered_layout" in result["operation_result"]["review_required"]
    assert "dfm_post_save" in result["operation_result"]["checks"]
    assert any("block_id" in item for item in result["operation_result"]["changes"])


def test_native_dfm_chunks_are_stable_and_bound_to_full_asset_revision(
    native_docx,
) -> None:
    service, asset, _ = native_docx
    expected = read_dfm(service, asset)
    assert expected == read_dfm(service, asset)
    assert asset["asset_id"] in expected and asset["revision"] in expected
    assert "created:" not in expected.split("---")[1]
    pieces = []
    offset = 0
    while True:
        result = call(
            service,
            op="read_docx",
            asset_id=asset["asset_id"],
            revision=asset["revision"],
            text_offset=offset,
            text_limit=700,
            limit=2,
        )
        dfm = result["dfm"]
        assert dfm["text_sha256"] == hashlib.sha256(expected.encode()).hexdigest()
        assert len(result["blocks"]) == 2 and result["next_offset"] == 2
        assert dfm["representation_complete"] is False
        pieces.append(dfm["text_excerpt"])
        offset = dfm["next_text_offset"]
        if offset is None:
            break
    assert "".join(pieces) == expected
    assert result["asset"]["capabilities"]["read_docx"] is True
    assert call(service, op="contract")["formats"]["docx"] == [
        "read_docx",
        "read_docx_block",
        "update_docx",
        "verify",
        "export_wiki",
    ]


@pytest.mark.parametrize(
    "field", ["native_asset_id", "native_revision", "checksum", "doc_id"]
)
def test_wrong_dfm_session_cannot_create_a_revision(native_docx, field: str) -> None:
    service, asset, source = native_docx
    import re

    dfm = re.sub(
        rf"^{field}:.*$", f"{field}: wrong", read_dfm(service, asset), flags=re.M
    )
    with pytest.raises(ValueError, match=r"binding mismatch|checksum|doc_id mismatch"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_edit={"dfm_text": dfm},
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == asset["revision"]


def test_equal_source_bytes_do_not_allow_another_assets_dfm(native_docx) -> None:
    service, asset, source = native_docx
    second = call(service, op="register", source_path=str(source))["asset"]
    assert second["revision"] == asset["revision"]
    with pytest.raises(ValueError, match="binding mismatch"):
        call(
            service,
            op="update_docx",
            asset_id=second["asset_id"],
            expected_revision=second["revision"],
            docx_edit={"dfm_text": read_dfm(service, asset)},
        )


def test_noop_retains_exact_file_and_does_not_add_history(native_docx) -> None:
    service, asset, _ = native_docx
    result = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={"dfm_text": read_dfm(service, asset)},
    )
    assert result["asset"]["revision"] == asset["revision"]
    assert result["asset"]["revision_count"] == 1
    assert result["operation_result"]["changed_parts"] == []


def test_stale_and_archived_writes_fail_but_old_revision_remains_readable(
    native_docx,
) -> None:
    service, asset, _ = native_docx
    dfm = read_dfm(service, asset)
    result = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={"dfm_text": dfm.replace("原始段落", "修訂段落")},
    )
    with pytest.raises(ValueError, match="stale"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_edit={"dfm_text": dfm},
        )
    current = result["asset"]
    call(
        service,
        op="archive",
        asset_id=asset["asset_id"],
        expected_revision=current["revision"],
    )
    with pytest.raises(ValueError, match="Archived"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=current["revision"],
            docx_edit={"dfm_text": read_dfm(service, current)},
        )
    assert read_dfm(service, asset) == dfm


@pytest.mark.parametrize("change", ["missing", "duplicate", "style"])
def test_incomplete_or_unsupported_dfm_edits_fail(native_docx, change: str) -> None:
    service, asset, _ = native_docx
    dfm = read_dfm(service, asset)
    if change == "missing":
        dfm = dfm[: dfm.index("<!-- @b:")]
    elif change == "duplicate":
        dfm += "\n<!-- @b:p999 -->\nUnexpected paragraph\n"
    else:
        prefix, tail = dfm.split("<!-- dfm:styles\n", 1)
        styles, suffix = tail.split("\n-->", 1)
        fields = yaml.safe_load(styles)
        fields["default_font"]["size"] += 1
        dfm = (
            prefix
            + "<!-- dfm:styles\n"
            + yaml.safe_dump(fields).strip()
            + "\n-->"
            + suffix
        )
    with pytest.raises(ValueError, match=r"marker|style|parse"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_edit={"dfm_text": dfm},
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


@pytest.mark.parametrize("protected", [False, True])
def test_signed_or_protected_docx_is_readable_but_not_editable(
    native_docx, protected: bool
) -> None:
    service, _, source = native_docx
    data = source.read_bytes()
    replacements = {"_xmlsignatures/sig1.xml": b"<signature/>"}
    if protected:
        root = NativeOOXMLPackage(data).xml("word/settings.xml")
        etree.SubElement(root, f"{{{WORD_NS}}}documentProtection")
        replacements = {"word/settings.xml": etree.tostring(root)}
    source.write_bytes(replace_parts(data, replacements))
    asset = call(service, op="register", source_path=str(source))["asset"]
    dfm = read_dfm(service, asset).replace("原始段落", "修訂段落")
    with pytest.raises(ValueError, match=r"signed|Protected"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_edit={"dfm_text": dfm},
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


def test_tracked_changes_retain_author_and_explicit_review(native_docx) -> None:
    service, asset, _ = native_docx
    result = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={
            "dfm_text": read_dfm(service, asset).replace("原始段落", "修訂段落"),
            "track_changes": True,
            "revision_author": "研究作者",
        },
    )
    package = NativeOOXMLPackage(service.repository.read(asset["asset_id"]))
    insertions = package.xml("word/document.xml").findall(f".//{{{WORD_NS}}}ins")
    assert insertions and all(
        item.get(f"{{{WORD_NS}}}author") == "研究作者" for item in insertions
    )
    assert "word/settings.xml" in result["operation_result"]["changed_parts"]
    assert "fields_and_revisions" in result["operation_result"]["review_required"]


def test_package_guard_rejects_unrelated_changes_even_after_dfm_save(
    native_docx, monkeypatch
) -> None:
    service, asset, _ = native_docx
    save = DocxService.save_docx

    async def corrupt(self, *args, **kwargs):
        result = await save(self, *args, **kwargs)
        assert result["success"], result
        path = Path(result["output_path"])
        package = NativeOOXMLPackage(path.read_bytes())
        path.write_bytes(
            replace_parts(
                path.read_bytes(),
                {"word/styles.xml": package.parts["word/styles.xml"] + b"\n"},
            )
        )
        return result

    monkeypatch.setattr(DocxService, "save_docx", corrupt)
    with pytest.raises(ValueError, match="unrelated"):
        call(
            service,
            op="update_docx",
            asset_id=asset["asset_id"],
            expected_revision=asset["revision"],
            docx_edit={
                "dfm_text": read_dfm(service, asset).replace("原始段落", "修訂段落")
            },
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


def test_dtd_in_an_auxiliary_xml_part_rejects_before_legacy_parsing(
    native_docx,
) -> None:
    service, _, source = native_docx
    source.write_bytes(
        replace_parts(
            source.read_bytes(),
            {
                "customXml/unsafe.xml": b'<!DOCTYPE x [<!ENTITY e "unsafe">]><x>&e;</x>',
            },
        )
    )
    asset = call(service, op="register", source_path=str(source))["asset"]
    with pytest.raises(ValueError, match="DTD"):
        read_dfm(service, asset)


def test_dfm_limit_is_utf8_bytes_and_force_is_unavailable() -> None:
    with pytest.raises(ValueError, match="4 MiB"):
        NativeDocxEdit(dfm_text="文" * (MAX_DFM_BYTES // 2))
    with pytest.raises(ValueError):
        NativeDocxEdit.model_validate({"dfm_text": "text", "force": True})
