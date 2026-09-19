"""One native CAS resolves explicit A2T generation intent while preserving evidence."""

import io

import openpyxl
import pytest

from src.domain.native_table_workspace import NativeTableCellValue
from tests.native_grid_helpers import table_workbook
from tests.native_table_helpers import read_workspace
from tests.native_workbook_helpers import _call
from tests.unit.test_native_table_grid_apply import grid_service
from tests.unit.test_native_workbook_operations import read as read_workbook

GENERATED = {"kind": "native_generated", "value": None}


def prepare(root):
    service, tables = grid_service(root)
    source = root / "table.xlsx"
    source.write_bytes(table_workbook(formula="=A3*2"))
    asset = _call(service, op="register", source_path=str(source))["asset"]
    projected = _call(
        service,
        op="project_workbook_table",
        asset_id=asset["asset_id"],
        revision=asset["revision"],
        table_projection={
            "worksheet": {"sheet_id": "1", "part": "xl/worksheets/sheet1.xml"},
            "start_cell": "A2",
            "end_cell": "C5",
        },
    )
    return service, tables, source, asset, projected["table_id"]


def application(service, asset, table_id):
    current, digest = read_workspace(service, table_id)
    grid = current["structural_plan"]["worksheet_grid"]
    grid["edits"][0]["expand_tables"] = [
        {"part": "xl/tables/table1.xml", "expected_ref": "A2:C5"}
    ]
    if len(grid["edits"]) > 1:
        grid["edits"][1]["expand_tables"] = [
            {"part": "xl/tables/table1.xml", "expected_ref": "A2:C6"}
        ]
    return {
        "op": "apply_table_workspace",
        "asset_id": asset["asset_id"],
        "expected_revision": asset["revision"],
        "table_id": table_id,
        "expected_table_sha256": digest,
        "worksheet_grid": grid,
    }


def test_a2t_appends_native_table_rows_and_columns_in_one_commit(tmp_path):
    service, tables, source, asset, table_id = prepare(tmp_path)
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    initial, _ = read_workspace(service, table_id)
    evidence = initial["source_cells"][3]["source"]["evidence"]
    tables.add_rows(
        table_id,
        [
            {
                "A": {"kind": "number", "value": 40},
                "B": GENERATED,
                "C": {"kind": "string", "value": "new"},
            }
        ],
    )
    tables.add_column(
        table_id,
        "Review",
        "native",
        default_value={"kind": "string", "value": "checked"},
    )
    tables.update_cell(table_id, 0, "Review", GENERATED)
    current, digest = read_workspace(service, table_id)
    result = _call(service, **application(service, asset, table_id))
    assert result["success"]
    record = read_workbook(service, result["asset"])
    assert len(service.repository.load(asset["asset_id"]).history) == 2
    data = service.repository.read(asset["asset_id"], result["asset"]["revision"])
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert book["Data"].tables["Table1"].ref == "A2:D6"
    assert book["Data"].tables["Table1"].autoFilter.ref == "A2:D6"
    assert book["Data"]["D2"].value == "Column4"
    assert book["Data"]["A6"].value == 40 and book["Data"]["B6"].value == "=A6*2"
    assert [book["Data"][f"D{r}"].value for r in range(3, 7)] == ["checked"] * 4
    resolved = next(
        c
        for c in record["operation_result"]["changes"]
        if c.get("operation") == "resolve_native_generated_values"
    )
    assert resolved["cells"] == [
        {"cell": "D2", "kind": "string", "value": "Column4"},
        {"cell": "B6", "kind": "formula", "value": "=A6*2"},
    ]
    assert read_workspace(
        service, table_id, workspace_reference=result["workspace_reference"]
    ) == (current, digest)
    assert _call(service, op="verify", reference=evidence)["valid"]
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime
    with pytest.raises(ValueError, match="stale"):
        _call(service, **application(service, asset, table_id))


@pytest.mark.parametrize(
    "kind", ["missing", "blank", "wrong_cell", "survivor", "stale_table"]
)
def test_unresolved_or_invalid_generated_intent_cannot_commit(tmp_path, kind):
    service, tables, source, asset, table_id = prepare(tmp_path)
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    row = {"A": {"kind": "number", "value": 40}, "B": GENERATED}
    if kind == "missing":
        row.pop("B")
    elif kind == "blank":
        row["B"] = {"kind": "blank", "value": None}
    elif kind == "wrong_cell":
        row["A"] = GENERATED
    elif kind == "survivor":
        tables.update_cell(table_id, 1, "B", GENERATED)
    tables.add_rows(table_id, [row])
    args = application(service, asset, table_id)
    if kind == "stale_table":
        args["worksheet_grid"]["edits"][0]["expand_tables"][0]["expected_ref"] = "A2:C4"
    with pytest.raises(ValueError, match=r"Calculated table|native_generated|stale"):
        _call(service, **args)
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    assert len(service.repository.load(asset["asset_id"]).history) == 1
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


def test_generated_value_is_intent_not_a_literal_or_independent_export(tmp_path):
    with pytest.raises(ValueError, match="null"):
        NativeTableCellValue.model_validate({"kind": "native_generated", "value": 4})
    service, tables, _, asset, table_id = prepare(tmp_path)
    tables.update_cell(table_id, 1, "B", GENERATED)
    _, digest = read_workspace(service, table_id)
    for options in [
        {
            "op": "apply_table_workspace",
            "asset_id": asset["asset_id"],
            "expected_revision": asset["revision"],
        },
        {"op": "create_workbook_from_table", "table_workbook": {"name": "new.xlsx"}},
    ]:
        with pytest.raises(ValueError, match="native_generated"):
            _call(service, table_id=table_id, expected_table_sha256=digest, **options)
