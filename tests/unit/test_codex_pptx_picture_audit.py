"""Independent live-run auditing must reject plausible but wrong picture results."""

from __future__ import annotations

import json

import pytest
from pptx import Presentation

from src.domain.native_pptx import NativePptxShapeLocator
from src.infrastructure.native_pptx_pictures import add_pictures
from tests.codex_native_pdf.trace import canonical, complete_records
from tests.codex_pptx_pictures.history import validate_references, validate_stage
from tests.native_pptx_helpers import build_presentation
from tests.native_pptx_picture_helpers import raster, request, sources
from tests.native_workbook_helpers import _replace
from tests.unit.test_codex_native_pdf_audit import call, readbacks


def shape_readbacks():
    calls, ref = readbacks()
    for item in calls:
        item["arguments"]["native_request"]["op"] = "read_pptx_shape"
        text = item["result"]["content"][0]
        result = json.loads(text["text"])
        result["shape"] = result.pop("page")
        text["text"] = json.dumps(result)
    return calls, ref


@pytest.mark.parametrize("failure", [None, "future", "missing_chunk", "wrong_proof"])
def test_shape_reference_must_be_complete_before_mutation(failure):
    calls, ref = shape_readbacks()
    mutation = call(
        "replace_pptx_pictures",
        {"success": True},
        pptx_picture_edits=[{"reference": ref}],
    )
    proof = call(
        "verify", {"valid": True, "is_current_managed_revision": False}, reference=ref
    )
    if failure == "future":
        calls = [mutation, *calls, proof]
    else:
        calls += [mutation, proof]
    if failure == "missing_chunk":
        calls.pop(0)
    if failure == "wrong_proof":
        proof["arguments"]["native_request"]["reference"] = {
            **ref,
            "asset_id": "file_" + "c" * 32,
        }

    def check():
        records = complete_records(
            calls, operation="read_pptx_shape", record_key="shape"
        )
        assert canonical(ref) in records
        validate_references(calls, records, {"asset_id": ref["asset_id"]})

    if failure:
        with pytest.raises(ValueError):
            check()
    else:
        check()


@pytest.mark.parametrize(
    "failure",
    [None, "original_shape", "added_mapping", "extra_shape", "notes", "wrong_image"],
)
def test_independent_presentation_audit(tmp_path, failure):
    data, image = build_presentation(), raster()
    updated, result = add_pictures(data, [request(data, image)], sources(image))
    source, target = tmp_path / "source.pptx", tmp_path / "target.pptx"
    source.write_bytes(data)
    target.write_bytes(updated)
    mapping = validate_stage(source, target, [image])
    presentation = Presentation(target)
    if failure == "original_shape":
        presentation.slides[0].shapes[0].left += 1
    elif failure == "added_mapping":
        presentation.slides[0].shapes[-1].rotation = 2
    elif failure == "extra_shape":
        presentation.slides[0].shapes.add_textbox(0, 0, 100, 100)
    elif failure == "wrong_image":
        image = raster(color="blue")
    if failure in {"original_shape", "added_mapping", "extra_shape"}:
        presentation.save(target)
    elif failure == "notes":
        from src.infrastructure.native_pptx_package import NativePptxPackage

        package = NativePptxPackage(updated)
        part = next(
            p
            for p in package.parts
            if p.startswith("ppt/notesSlides/notesSlide") and p.endswith(".xml")
        )
        target.write_bytes(
            _replace(
                updated,
                {part: package.parts[part].replace(b"</p:notes>", b" </p:notes>")},
            )
        )
    assert (
        NativePptxShapeLocator(**result.changes[0]["locator"]).part
        == "ppt/slides/slide1.xml"
    )
    if failure:
        with pytest.raises(ValueError):
            validate_stage(source, target, [image], mapping)
    else:
        validate_stage(source, target, [image], mapping)
