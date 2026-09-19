"""Native definition lifecycle, inheritance, relationship relocation and deletion."""

import hashlib
import io
import json

import pytest
from docx import Document

from src.domain.native_docx_story_lifecycle import DocxStoryStructure
from src.infrastructure.native_docx_stories import W, catalog, read_story
from src.infrastructure.native_docx_story_lifecycle import change_structure
from src.infrastructure.native_docx_workspace import checked_docx
from src.infrastructure.native_ooxml import (
    DOC_REL_NS,
    REL_NS,
    TYPE_NS,
    relationships_path,
    xml_bytes,
)
from tests.native_docx_stories_helpers import HEADER_PART, HEADER_TEXT, story_document

NEW = "word/independent/header.xml"


def apply(data, *edits):
    return change_structure(
        data,
        DocxStoryStructure.model_validate(
            {
                "scope": "sections_and_following_inheritors",
                "expected_catalog_sha256": hashlib.sha256(
                    json.dumps(
                        catalog(checked_docx(data)),
                        ensure_ascii=False,
                        sort_keys=True,
                        separators=(",", ":"),
                    ).encode()
                ).hexdigest(),
                "edits": list(edits),
            }
        ),
    )


def clone(data, part=HEADER_PART, target=NEW):
    return {
        "op": "clone",
        "source_part": part,
        "source_part_sha256": hashlib.sha256(
            checked_docx(data).parts[part]
        ).hexdigest(),
        "part": target,
    }


def binding(section, part, kind="header", variant="default"):
    return {
        "op": "bind",
        "section_index": section,
        "story_kind": kind,
        "variant": variant,
        "part": part,
    }


def test_clone_bind_unlink_delete_preserves_native_stories_and_body():
    original = story_document()
    before = checked_docx(original)
    separate, result = apply(original, clone(original), binding(1, NEW))
    after = checked_docx(separate)
    assert after.parts[NEW] == before.parts[HEADER_PART]
    assert read_story(after, NEW)["text"].startswith(HEADER_TEXT)
    assert catalog(after)["sections"][1]["bindings"][0]["part"] == NEW
    doc = Document(io.BytesIO(separate))
    assert not doc.sections[1].header.is_linked_to_previous
    assert [p.text for p in doc.paragraphs] == [
        p.text for p in Document(io.BytesIO(original)).paragraphs
    ]
    for name in before.parts.keys() - set(result.changed_parts):
        assert after.parts[name] == before.parts[name]
    with pytest.raises(ValueError, match="bound"):
        apply(
            separate,
            {
                "op": "delete",
                "part": NEW,
                "expected_part_sha256": hashlib.sha256(after.parts[NEW]).hexdigest(),
            },
        )
    restored, result = apply(
        separate,
        binding(1, None),
        {
            "op": "delete",
            "part": NEW,
            "expected_part_sha256": hashlib.sha256(after.parts[NEW]).hexdigest(),
        },
    )
    final = checked_docx(restored)
    assert NEW not in final.parts
    assert catalog(final) == catalog(before)
    assert Document(io.BytesIO(restored)).sections[1].header.is_linked_to_previous
    assert result.changes[-1]["retained_dependencies"] == []


def test_create_blank_is_distinct_from_inherit_and_settings_are_explicit():
    original = story_document()
    data, result = apply(
        original,
        {
            "op": "create",
            "part": NEW,
            "story_kind": "header",
            "blocks": [{"kind": "paragraph", "runs": []}],
        },
        binding(1, NEW),
        {"op": "first_page", "section_index": 1, "enabled": True},
        {"op": "even_pages", "enabled": True},
    )
    package = checked_docx(data)
    assert read_story(package, NEW)["text"] == ""
    assert catalog(package)["even_and_odd_headers"]
    assert catalog(package)["sections"][1]["different_first_page"]
    assert catalog(package)["sections"][1]["bindings"][0]["part"] == NEW
    assert result.changes[-1]["op"] == "even_pages"
    same, result = apply(
        data,
        binding(1, NEW),
        {"op": "first_page", "section_index": 1, "enabled": True},
        {"op": "even_pages", "enabled": True},
    )
    assert same == data and result.changed_parts == []


def test_wrong_hash_and_unsafe_part_reject_without_mutation():
    original = story_document()
    with pytest.raises(ValueError, match="hash"):
        apply(original, {**clone(original), "source_part_sha256": "b" * 64})
    with pytest.raises(ValueError, match=r"already|collision"):
        apply(original, clone(original, target=HEADER_PART))
    for part in [
        "../header.xml",
        "word/../header.xml",
        "word/_rels/header.xml",
        "word/header.xml/",
        "word/header?.xml",
    ]:
        with pytest.raises(ValueError):
            apply(
                original,
                {
                    "op": "create",
                    "part": part,
                    "story_kind": "header",
                    "blocks": [{"kind": "paragraph", "runs": []}],
                },
            )


def test_picture_clone_relocates_relationships_and_drawing_ids_without_copying_media():
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    from PIL import Image

    doc = Document(io.BytesIO(story_document()))
    png = io.BytesIO()
    Image.new("RGB", (4, 4), "blue").save(png, format="PNG")
    doc.sections[0].header.paragraphs[0].add_run().add_picture(
        io.BytesIO(png.getvalue()), width=Inches(0.1)
    )
    paragraph = doc.sections[0].header.paragraphs[0]._p
    paragraph.set(
        "{http://schemas.microsoft.com/office/word/2010/wordml}paraId", "00000001"
    )
    inline = paragraph.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline"
    )
    for key, value in [("anchorId", "00000002"), ("editId", "00000003")]:
        inline.set(
            "{http://schemas.microsoft.com/office/word/2010/wordprocessingDrawing}"
            + key,
            value,
        )
    header_part = str(doc.sections[0].header.part.partname).lstrip("/")
    rel_id = doc.sections[0].header.part.relate_to(
        "https://example.org/source", DOC_REL_NS + "/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set("{" + DOC_REL_NS + "}id", rel_id)
    doc.sections[0].header.paragraphs[0]._p.append(link)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    old = checked_docx(data)
    updated, result = apply(data, clone(data, header_part), binding(1, NEW))
    package = checked_docx(updated)
    assert package.relationships(NEW) == old.relationships(header_part)
    assert (
        package.parts[relationships_path(NEW)]
        != old.parts[relationships_path(header_part)]
    )
    image_parts = [n for n in old.parts if n.startswith("word/media/")]
    assert image_parts
    assert all(package.parts[n] == old.parts[n] for n in image_parts)
    repairs = result.changes[0]["identity_repairs"]
    assert repairs and repairs[0]["before"] != repairs[0]["after"]
    assert len(repairs) == 4 and all(r["before"] != r["after"] for r in repairs)
    wp = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
    ids = [
        n.get("id") for part in [header_part, NEW] for n in package.xml(part).iter(wp)
    ]
    assert len(ids) == len(set(ids))
    removed, receipt = apply(
        updated,
        binding(1, None),
        {
            "op": "delete",
            "part": NEW,
            "expected_part_sha256": hashlib.sha256(package.parts[NEW]).hexdigest(),
        },
    )
    final = checked_docx(removed)
    assert NEW not in final.parts and relationships_path(NEW) not in final.parts
    assert receipt.changes[-1]["retained_dependencies"] == image_parts
    assert final.parts[image_parts[0]] == old.parts[image_parts[0]]


def test_cloned_fields_keep_instruction_and_cached_value():
    data = story_document()
    package = checked_docx(data)
    source = next(
        x["locator"]["part"]
        for x in catalog(package)["stories"]
        if x["locator"]["story_kind"] == "footer"
    )
    updated, _ = apply(data, clone(data, source), binding(1, NEW, kind="footer"))
    assert checked_docx(updated).parts[NEW] == package.parts[source]


def test_delete_checks_historical_and_foreign_incoming_relationships():
    from lxml import etree

    from src.infrastructure.native_ooxml_additions import extend_package

    data, _ = apply(
        story_document(),
        {
            "op": "create",
            "part": NEW,
            "story_kind": "header",
            "blocks": [{"kind": "paragraph"}],
        },
    )
    p = checked_docx(data)
    rels = p.xml("word/_rels/document.xml.rels")
    etree.SubElement(
        rels,
        "{" + REL_NS + "}Relationship",
        Id="rIdHistorical",
        Type=DOC_REL_NS + "/header",
        Target="independent/header.xml",
    )
    main = p.xml("word/document.xml")
    old = etree.SubElement(
        etree.SubElement(main.find(W + "body")[-1], W + "sectPrChange"), W + "sectPr"
    )
    ref = etree.SubElement(old, W + "headerReference")
    ref.set(W + "type", "default")
    ref.set("{" + DOC_REL_NS + "}id", "rIdHistorical")
    historical = p.replace(
        {
            "word/document.xml": xml_bytes(main),
            "word/_rels/document.xml.rels": xml_bytes(rels),
        }
    )
    edit = {
        "op": "delete",
        "part": NEW,
        "expected_part_sha256": hashlib.sha256(p.parts[NEW]).hexdigest(),
    }
    with pytest.raises(ValueError, match="historical"):
        apply(historical, edit)
    extra = etree.Element("{" + REL_NS + "}Relationships")
    etree.SubElement(
        extra,
        "{" + REL_NS + "}Relationship",
        Id="link",
        Type=DOC_REL_NS + "/header",
        Target="independent/header.xml",
    )
    foreign = extend_package(p, {}, {"word/_rels/styles.xml.rels": xml_bytes(extra)})
    with pytest.raises(ValueError, match="incoming"):
        apply(foreign, edit)


def test_setting_creation_without_existing_settings_part_and_binding_cascade():
    from docx.enum.section import WD_SECTION_START

    doc = Document(io.BytesIO(story_document()))
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph("Third section")
    buf = io.BytesIO()
    doc.save(buf)
    package = checked_docx(buf.getvalue())
    rels = package.xml("word/_rels/document.xml.rels")
    for n in list(rels):
        if n.get("Type") == DOC_REL_NS + "/settings":
            rels.remove(n)
    types = package.xml("[Content_Types].xml")
    for n in list(types):
        if n.get("PartName") == "/word/settings.xml":
            types.remove(n)
    data = package.replace(
        {
            "word/_rels/document.xml.rels": xml_bytes(rels),
            "[Content_Types].xml": xml_bytes(types),
        },
        {"word/settings.xml"},
    )
    package = checked_docx(data)
    source = next(
        x["locator"]["part"]
        for x in catalog(package)["stories"]
        if x["locator"]["story_kind"] == "header"
        and "SOURCE" in read_story(package, x["locator"]["part"])["text"]
    )
    updated, _ = apply(
        data,
        clone(data, source),
        binding(1, NEW),
        {"op": "even_pages", "enabled": True},
    )
    final = checked_docx(updated)
    listing = catalog(final)
    assert listing["even_and_odd_headers"]
    assert listing["sections"][2]["bindings"][0]["part"] == NEW
    assert listing["sections"][2]["bindings"][0]["inherited_from"] == 1
    assert Document(io.BytesIO(updated)).settings.odd_and_even_pages_header_footer
    assert any(
        n.get("PartName") == "/word/settings.xml"
        for n in final.xml("[Content_Types].xml").findall("{" + TYPE_NS + "}Override")
    )


@pytest.mark.parametrize(
    "tag",
    ["bookmarkStart", "commentReference", "sdt", "footnoteReference", "ins", "object"],
)
def test_clone_rejects_identities_requiring_additional_remapping(tag):
    from lxml import etree

    data = story_document()
    p = checked_docx(data)
    root = p.xml(HEADER_PART)
    etree.SubElement(root[0], W + tag)
    data = p.replace({HEADER_PART: xml_bytes(root)})
    with pytest.raises(ValueError, match="identity-aware"):
        apply(data, clone(data))


def test_reversed_definition_batch_restores_exact_original_bytes():
    original = story_document()
    create = {
        "op": "create",
        "part": NEW,
        "story_kind": "header",
        "blocks": [{"kind": "paragraph", "runs": [{"text": "TEMP 007"}]}],
    }
    made, _ = apply(original, create)
    deletion = {
        "op": "delete",
        "part": NEW,
        "expected_part_sha256": hashlib.sha256(
            checked_docx(made).parts[NEW]
        ).hexdigest(),
    }
    restored, receipt = apply(original, create, deletion)
    assert restored == original and receipt.changed_parts == []


def test_requested_block_readback_rejects_a_corrupt_serializer(monkeypatch):
    from src.infrastructure import native_docx_story_lifecycle as module

    original = module.xml_bytes

    def corrupt(root):
        raw = original(root)
        return (
            raw.replace(b"REQUESTED 007", b"WRONG 999")
            if root.tag == W + "hdr"
            else raw
        )

    monkeypatch.setattr(module, "xml_bytes", corrupt)
    with pytest.raises(ValueError):
        apply(
            story_document(),
            {
                "op": "create",
                "part": NEW,
                "story_kind": "header",
                "blocks": [{"kind": "paragraph", "runs": [{"text": "REQUESTED 007"}]}],
            },
        )


def test_new_part_cannot_capture_an_existing_dangling_reference():
    from lxml import etree

    from src.infrastructure.native_ooxml_additions import extend_package

    data = story_document()
    package = checked_docx(data)
    rels = etree.Element("{" + REL_NS + "}Relationships")
    etree.SubElement(
        rels,
        "{" + REL_NS + "}Relationship",
        Id="dangling",
        Type=DOC_REL_NS + "/header",
        Target="independent/header.xml",
    )
    source = extend_package(
        package, {}, {"word/_rels/styles.xml.rels": xml_bytes(rels)}
    )
    with pytest.raises(ValueError, match="dangling"):
        apply(source, clone(source))
