"""Native/A2T typed round trips, preservation and immutable input lineage."""

import io

import openpyxl
import pytest
from lxml import etree

from tests.native_table_helpers import apply, project, read_workspace, service_at
from tests.native_workbook_helpers import _call, _parts
from tests.unit.test_native_workbook_operations import read as read_workbook


def test_native_a2t_roundtrip_preserves_package_and_historical_snapshots(tmp_path):
    service, tables = service_at(tmp_path)
    asset, projection = project(service, tmp_path)
    table_id = projection["table_id"]
    source = tmp_path / "source.xlsx"
    before, mtime = source.read_bytes(), source.stat().st_mtime_ns
    original, digest = read_workspace(service, table_id, limit=173)
    assert digest == projection["table_sha256"]
    assert original["source_correspondence_unchanged"]
    assert original["table"]["rows"][0]["E"] == {"kind": "string", "value": "007"}
    assert original["table"]["rows"][1]["E"] == {"kind": "boolean", "value": True}
    assert original["table"]["rows"][2]["E"] == {"kind": "string", "value": "=SUM(A2)"}
    assert original["table"]["rows"][1]["B"] == {"kind": "formula", "value": "=A2*2"}
    assert len(original["source_cells"]) == 15
    first = original["source_cells"][0]["source"]
    assert _call(service, op="verify", reference=first["evidence"])["valid"]
    for column, value in [("A", "Updated"), ("E", "008")]:
        tables.update_cell(table_id, 0, column, {"kind": "string", "value": value})
    with pytest.raises(ValueError, match="workspace changed"):
        apply(service, asset, table_id, digest)
    current, digest = read_workspace(service, table_id)
    result = apply(service, asset, table_id, digest)
    assert result["changed"] and not result["source_written"]
    assert result["asset"]["capabilities"]["project_workbook_table"]
    assert result["asset"]["capabilities"]["edit_worksheets"]
    changed = service.repository.read(asset["asset_id"], result["asset"]["revision"])
    old_parts, new_parts = _parts(before), _parts(changed)
    changed_parts = {name for name in old_parts if old_parts[name] != new_parts[name]}
    assert changed_parts <= {
        "xl/worksheets/sheet1.xml",
        "xl/workbook.xml",
        "xl/sharedStrings.xml",
    }
    # Count repair must retain every original shared/rich-text entry.
    shared = [
        etree.fromstring(parts["xl/sharedStrings.xml"])
        for parts in (old_parts, new_parts)
    ]
    assert [etree.tostring(node) for node in shared[0]] == [
        etree.tostring(node) for node in shared[1]
    ]
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    roots = [
        etree.fromstring(parts["xl/worksheets/sheet1.xml"])
        for parts in (old_parts, new_parts)
    ]
    for address in ("D2", "B2", "A5", "C3"):
        nodes = [root.find(f".//s:c[@r='{address}']", ns) for root in roots]
        assert etree.tostring(nodes[0]) == etree.tostring(nodes[1])
    assert roots[0].find(".//s:c[@r='A1']", ns).get("s") == roots[1].find(
        ".//s:c[@r='A1']", ns
    ).get("s")
    book = openpyxl.load_workbook(io.BytesIO(changed), rich_text=True)
    assert book["Data"]["A1"].value == "Updated"
    assert book["Data"]["E1"].value == "008"
    assert book["Data"]["E3"].data_type == "s"
    assert book["Data"]["B2"].data_type == "f"
    assert list(book["Data"].merged_cells.ranges) == list(
        openpyxl.load_workbook(io.BytesIO(before))["Data"].merged_cells.ranges
    )
    reference = result["workspace_reference"]
    assert reference["revision"] == digest
    assert _call(service, op="verify", reference=reference)["valid"]
    event = read_workbook(service, result["asset"])["operation_result"]["changes"][-1]
    assert event["workspace_reference"] == reference and event["edited_cell_count"] == 2
    reloaded, _ = service_at(tmp_path)
    assert read_workspace(reloaded, table_id) == (current, digest)
    tables.update_cell(table_id, 0, "E", {"kind": "string", "value": "009"})
    assert tables.delete_table(table_id)
    assert read_workspace(reloaded, table_id, workspace_reference=reference) == (
        current,
        digest,
    )
    proof = _call(service, op="verify", reference=first["evidence"])
    assert proof["valid"] and not proof["is_current_managed_revision"]
    assert source.read_bytes() == before and source.stat().st_mtime_ns == mtime


def test_noop_and_stale_apply_do_not_advance_binding(tmp_path):
    service, tables = service_at(tmp_path)
    asset, projection = project(service, tmp_path)
    count = len(service.repository.list_assets(0, 100))
    result = apply(service, asset, projection["table_id"], projection["table_sha256"])
    assert not result["changed"] and result["asset"]["revision"] == asset["revision"]
    assert len(service.repository.list_assets(0, 100)) == count
    tables.update_cell(
        projection["table_id"], 0, "E", {"kind": "string", "value": "008"}
    )
    current, digest = read_workspace(service, projection["table_id"])
    result = apply(service, asset, projection["table_id"], digest)
    with pytest.raises(ValueError, match="stale"):
        apply(service, asset, projection["table_id"], digest)
    with pytest.raises(ValueError, match="binding"):
        apply(service, result["asset"], projection["table_id"], digest)
    assert current["table"]["native_binding"]["source"]["revision"] == asset["revision"]


def test_structure_change_can_create_new_workbook_but_cannot_shift_source(tmp_path):
    service, tables = service_at(tmp_path)
    asset, projection = project(service, tmp_path, start="B2", end="E3")
    table_id = projection["table_id"]
    tables.add_rows(table_id, [{"E": {"kind": "string", "value": "new"}}])
    record, digest = read_workspace(service, table_id)
    assert not record["source_correspondence_unchanged"]
    assert record["source_cells"][0]["source"]["cell"] == "B2"
    with pytest.raises(ValueError, match="correspondence changed"):
        apply(service, asset, table_id, digest)
    result = _call(
        service,
        op="create_workbook_from_table",
        table_id=table_id,
        expected_table_sha256=digest,
        table_workbook={"include_headers": True},
    )
    book = openpyxl.load_workbook(
        io.BytesIO(
            service.repository.read(
                result["asset"]["asset_id"], result["asset"]["revision"]
            )
        )
    )
    assert next(iter(book["Data"].values)) == ("B", "C", "D", "E")
    assert (
        book["Data"]["A2"].value == "=A2*2"
    )  # Explicit formula text; no relocation claim.
    assert book["Data"]["D4"].value == "new"
    event = read_workbook(service, result["asset"])["operation_result"]["changes"][0]
    assert event["mapping"]["first_data_row"] == 2
    assert event["mapping"]["columns"][0] == {"name": "B", "native_column": "A"}
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]


def test_ordinary_table_export_retains_literal_values_and_types(tmp_path):
    service, tables = service_at(tmp_path)
    table_id = tables.create_table(
        "summary",
        "Types",
        [{"name": name, "type": "text", "required": False} for name in "ABCDE"],
    )
    tables.add_rows(
        table_id, [dict(zip("ABCDE", ["007", "=1+1", True, 1.25, None], strict=True))]
    )
    _, digest = read_workspace(service, table_id)
    result = _call(
        service,
        op="create_workbook_from_table",
        table_id=table_id,
        expected_table_sha256=digest,
        table_workbook={},
    )
    data = service.repository.read(
        result["asset"]["asset_id"], result["asset"]["revision"]
    )
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert [
        (book["Data"][cell].value, book["Data"][cell].data_type)
        for cell in ("A1", "B1", "C1", "D1")
    ] == [("007", "s"), ("=1+1", "s"), (True, "b"), (1.25, "n")]
    tables.update_cell(table_id, 0, "A", {"unexpected": "object"})
    _, digest = read_workspace(service, table_id)
    with pytest.raises(ValueError, match="scalar"):
        _call(
            service,
            op="create_workbook_from_table",
            table_id=table_id,
            expected_table_sha256=digest,
            table_workbook={},
        )
