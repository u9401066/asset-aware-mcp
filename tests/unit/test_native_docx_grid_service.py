"""Managed revisions, full references, paging, source preservation and rollback."""

from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy

import pytest
from docx import Document

from src.application.native_document_service import NativeDocumentService
from src.application.native_docx_bridge import NativeDocxBridge
from src.infrastructure.native_asset_store import FileNativeAssetRepository
from src.infrastructure.native_docx_structure import NativeDocxStructure
from src.infrastructure.native_docx_workspace import FileNativeDocxWorkspaces
from src.infrastructure.native_spreadsheet import SpreadsheetFileAdapter
from src.infrastructure.native_wiki_publisher import FileNativeWikiPublisher
from tests.native_docx_helpers import call
from tests.unit.test_native_docx_grid import source


@pytest.fixture
def stack(tmp_path):
    path = tmp_path / "original.docx"
    path.write_bytes(source())
    service = NativeDocumentService(
        FileNativeAssetRepository(tmp_path / "native"),
        SpreadsheetFileAdapter(),
        wiki_publisher=FileNativeWikiPublisher(),
        docx=NativeDocxBridge(FileNativeDocxWorkspaces()),
        docx_structure=NativeDocxStructure(),
    )
    asset = call(service, op="register", source_path=str(path))["asset"]
    blocks = call(
        service,
        op="read_docx",
        asset_id=asset["asset_id"],
        revision=asset["revision"],
        limit=100,
    )["blocks"]
    reference = next(b["evidence"] for b in blocks if b["type"] == "table")
    return service, asset, reference, path


def read_complete(service, fields):
    parts = []
    offset = 0
    digest = None
    while True:
        page = call(service, **fields, text_offset=offset, text_limit=777)["table"]
        assert page["excerpt_char_range"][0] == offset
        assert digest is None or page["text_sha256"] == digest
        digest = page["text_sha256"]
        parts.append(page["text_excerpt"])
        offset = page["next_text_offset"]
        if offset is None:
            break
    text = "".join(parts)
    assert hashlib.sha256(text.encode()).hexdigest() == digest
    return json.loads(text)


def update(service, asset, reference, edits):
    return call(
        service,
        op="update_docx_table_grid",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_table_grid={"reference": reference, "edits": edits},
    )


def test_contract_paged_grid_update_history_and_full_wiki(stack, tmp_path):
    service, asset, reference, path = stack
    before, mtime = path.read_bytes(), path.stat().st_mtime_ns
    contract = call(service, op="contract", for_op="update_docx_table_grid")
    assert contract["docx_table_grid_enabled"]
    assert "read_docx_table" in contract["formats"]["docx"]
    original_fields = {
        "op": "read_docx_table",
        "asset_id": asset["asset_id"],
        "revision": asset["revision"],
        "docx_table_reference": reference,
    }
    original = read_complete(service, original_fields)
    result = update(
        service,
        asset,
        reference,
        [{"op": "insert", "axis": "row", "index": 3, "sizes_twips": [500]}],
    )
    assert (
        not result["source_written"]
        and result["asset"]["revision"] != asset["revision"]
    )
    current = read_complete(service, result["review_request"])
    assert current["rows"] == 4 and current["evidence"] != reference
    assert read_complete(service, original_fields) == original
    assert path.read_bytes() == before and path.stat().st_mtime_ns == mtime
    assert service.repository.read(asset["asset_id"], asset["revision"]) == before
    verified = call(service, op="verify", reference=reference)
    assert verified["valid"] and not verified["is_current_managed_revision"]
    wiki = call(
        service,
        op="export_wiki",
        asset_id=asset["asset_id"],
        revision=result["asset"]["revision"],
        output_dir=str(tmp_path / "wiki"),
    )
    assert wiki["success"]
    document = Document(io.BytesIO(service.repository.read(asset["asset_id"])))
    assert document.tables[0].cell(2, 0).text == "007"
    assert len(service.repository.load(asset["asset_id"]).history) == 2


@pytest.mark.parametrize("field", ["value_sha256", "asset_id", "revision", "locator"])
def test_tampered_full_reference_cannot_commit(stack, field):
    service, asset, reference, _ = stack
    bad = deepcopy(reference)
    bad[field] = (
        {"block_id": "t999", "part": "word/document.xml"}
        if field == "locator"
        else ("file_" + "f" * 32 if field == "asset_id" else "f" * 64)
    )
    with pytest.raises(ValueError, match="reference"):
        update(service, asset, bad, [{"op": "split", "row": 1, "column": 0}])
    assert len(service.repository.load(asset["asset_id"]).history) == 1


def test_failed_second_edit_rolls_back_whole_batch_and_noop_skips_history(stack):
    service, asset, reference, _ = stack
    with pytest.raises(ValueError, match="last axis"):
        update(
            service,
            asset,
            reference,
            [
                {"op": "insert", "axis": "row", "index": 3, "sizes_twips": [500]},
                {"op": "delete", "axis": "row", "index": 0, "count": 4},
            ],
        )
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
    result = update(service, asset, reference, [{"op": "split", "row": 2, "column": 2}])
    assert result["asset"]["revision"] == asset["revision"]
    assert len(service.repository.load(asset["asset_id"]).history) == 1


def test_stale_revision_cannot_edit_current_table(stack):
    service, asset, reference, _ = stack
    edits = [{"op": "split", "row": 1, "column": 0}]
    update(service, asset, reference, edits)
    with pytest.raises(ValueError, match="stale"):
        update(service, asset, reference, edits)


def test_large_receipt_is_fully_readable_through_review_paging(stack):
    service, asset, reference, _ = stack
    values = ["甲" * 4000, "乙" * 4000, "丙" * 4000]
    edits = [
        {
            "op": "insert",
            "axis": "row",
            "index": 3,
            "sizes_twips": [500],
            "cells": [
                [{"paragraphs": [{"runs": [{"text": value}]}]} for value in values]
            ],
        }
    ]
    result = update(service, asset, reference, edits)
    complete = read_complete(service, result["review_request"])
    cells = complete["operation_result"]["changes"][0]["edits"][0]["request"]["cells"]
    assert [cell["paragraphs"][0]["runs"][0]["text"] for cell in cells[0]] == values
    assert len(json.dumps(result, ensure_ascii=False)) < 10000


def test_same_file_bytes_can_recur_with_distinct_explicit_receipts(stack):
    service, asset, reference, _ = stack
    first = update(
        service,
        asset,
        reference,
        [{"op": "resize", "axis": "row", "index": 2, "sizes_twips": [600]}],
    )
    first_read = read_complete(service, first["review_request"])
    second = update(
        service,
        first["asset"],
        first["review_request"]["docx_table_reference"],
        [{"op": "resize", "axis": "row", "index": 2, "sizes_twips": [700]}],
    )
    third = update(
        service,
        second["asset"],
        second["review_request"]["docx_table_reference"],
        [{"op": "resize", "axis": "row", "index": 2, "sizes_twips": [600]}],
    )
    assert first["asset"]["revision"] == third["asset"]["revision"]
    latest = read_complete(service, third["review_request"])
    assert latest["native_xml"] == first_read["native_xml"]
    assert latest["evidence"] == first_read["evidence"]
    assert latest["operation_result"] != first_read["operation_result"]
    assert "Latest history entry" in latest["operation_receipt_policy"]


def test_review_budget_failure_prevents_commit(stack, monkeypatch):
    from src.application import native_docx_table_operations as module

    service, asset, reference, _ = stack
    original = module.record_text

    def reject(record):
        if record.get("operation_result"):
            raise ValueError("simulated complete review budget")
        return original(record)

    monkeypatch.setattr(module, "record_text", reject)
    with pytest.raises(ValueError, match="review budget"):
        update(service, asset, reference, [{"op": "split", "row": 1, "column": 0}])
    assert service.repository.load(asset["asset_id"]).revision == asset["revision"]
