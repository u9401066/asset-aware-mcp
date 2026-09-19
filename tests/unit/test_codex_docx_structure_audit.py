"""The independent Word audit accepts extra source snapshots and rejects drift."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from tests.codex_docx_structure.audit import validate_docx, validate_wiki
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.run import execute
from tests.native_docx_helpers import call
from tests.native_docx_helpers import native_docx as native_docx
from tests.native_docx_structure_helpers import creation
from tests.unit.test_native_docx_structure_operations import structured as structured


def test_codex_prompt_reaches_process_as_utf8_under_windows_locale(
    tmp_path, monkeypatch
):
    popen = subprocess.Popen

    def windows_locale(*args, **kwargs):
        kwargs.setdefault("encoding", "cp1252")
        return popen(*args, **kwargs)

    monkeypatch.setattr(subprocess, "Popen", windows_locale)
    prompt = "研究 007 µg"
    command = [
        sys.executable,
        "-c",
        "import sys; sys.stdout.buffer.write(sys.stdin.buffer.read())",
    ]
    result = execute(command, prompt, tmp_path, 10)
    assert result["returncode"] == 0 and not result["timed_out"]
    assert (tmp_path / "events.jsonl").read_bytes() == prompt.encode("utf-8")


@pytest.mark.parametrize("fault", [None, "zero", "grid", "merge", "size", "heading"])
def test_independent_docx_table_audit(tmp_path, fault):
    document = Document()
    run = document.add_paragraph().add_run("研究 007 µg")
    run.bold, run.font.size, run.font.name = True, Pt(11.5), "Arial"
    run.font.color.rgb = RGBColor.from_string("1122AA")
    table = document.add_table(4, 5)
    table.autofit = False
    for column in table.columns:
        column.width = Twips(1500)
    for row in table.rows:
        row.height = Twips(300)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if fault != "merge":
        table.cell(0, 0).merge(table.cell(0, 4))
    rows = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
    ]
    for r, values in enumerate(rows):
        for c, text in enumerate(values):
            if r == 0 and c > 0:
                continue
            run = table.cell(r, c).paragraphs[0].add_run(text)
            run.font.size, run.font.name = Pt(11.5), "Arial"
    if fault == "zero":
        table.cell(2, 1).paragraphs[0].runs[0].text = "7"
    elif fault == "grid":
        table.columns[0].width = Twips(1499)
    elif fault == "size":
        table.cell(2, 1).paragraphs[0].runs[0].font.size = Pt(11)
    elif fault == "heading":
        document.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("000000")
    path = tmp_path / "table.docx"
    document.save(path)
    if fault:
        with pytest.raises(ValueError):
            validate_docx(path)
    else:
        assert validate_docx(path) == ("007", False)


@pytest.mark.parametrize("fault", [None, "part", "projection"])
@pytest.mark.parametrize("locale_encoding", ["utf-8", "cp1252"])
def test_wiki_audit_selects_docx_with_extra_source_snapshot(
    structured, tmp_path, fault, locale_encoding, monkeypatch
):
    service, _, _ = structured
    request = creation().model_dump()
    request["blocks"] = request["blocks"][:2]
    asset = call(service, op="create_docx", docx_create=request)["asset"]
    data = service.repository.read(asset["asset_id"], asset["revision"])
    (tmp_path / "verified.docx").write_bytes(data)
    result = call(
        service,
        op="export_wiki",
        asset_id=asset["asset_id"],
        output_dir=str(tmp_path / "native-wiki"),
    )
    root = Path(result["output_dir"])
    unrelated = tmp_path / "native-wiki" / "source-pdf"
    unrelated.mkdir()
    (unrelated / "manifest.json").write_text(
        json.dumps({"asset_id": "source"}), encoding="utf-8"
    )
    manifest_path = root / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if fault == "part":
        part = root / manifest["part_attachments"]["word/document.xml"]["attachment"]
        part.write_bytes(part.read_bytes().replace(qn("w:t").encode(), b"wrong") + b" ")
    elif fault == "projection":
        manifest["projection"] = "pdf-pages-v1"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    read_text = Path.read_text

    def locale_read(path, encoding=None, errors=None, **kwargs):
        return read_text(
            path, encoding=encoding or locale_encoding, errors=errors, **kwargs
        )

    # Reproduce a non-UTF-8 Windows locale even when the tests run on Linux.
    monkeypatch.setattr(Path, "read_text", locale_read)
    if fault:
        with pytest.raises(ValueError):
            validate_wiki(tmp_path, asset)
    else:
        validate_wiki(tmp_path, asset)
