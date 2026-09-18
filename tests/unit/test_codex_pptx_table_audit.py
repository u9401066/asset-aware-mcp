"""The independent Codex table auditor rejects plausible but nonidentical outputs."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Pt

from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pptx_tables.audit import validate_table


@pytest.mark.parametrize(
    "failure", [None, "leading_zero", "merge", "grid", "extra_shape", "changed_shared"]
)
def test_independent_scanned_table_auditor(tmp_path, failure):
    presentation = make_presentation(failure)
    output = tmp_path / "table.pptx"
    presentation.save(output)
    if failure:
        with pytest.raises(ValueError):
            validate_table(
                output, expected_count=2, changed=failure == "changed_shared"
            )
    else:
        validate_table(output, expected_count=2)


def make_presentation(failure=None):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    expected = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
    ]
    for _index in range(2):
        shape = slide.shapes.add_table(4, 5, 100000, 100000, 6000000, 2200000)
        table = shape.table
        if failure != "merge":
            table.cell(0, 0).merge(table.cell(0, 4))
        for row, values in enumerate(expected):
            for column, text in enumerate(values):
                if row == 0 and column > 0:
                    continue
                table.cell(row, column).text = text
                table.cell(row, column).text_frame.paragraphs[0].runs[0].font.size = Pt(
                    14
                )
        if failure == "leading_zero":
            table.cell(2, 1).text = "7"
        if failure == "grid":
            table.columns[0].width += 1
        if failure == "changed_shared":
            table.cell(2, 1).text = "008"
    if failure == "extra_shape":
        slide.shapes.add_textbox(0, 0, 100, 100)
    return presentation


@pytest.mark.parametrize("case", ["exact", "recovered", "missing_restore"])
def test_history_audit_records_recovery_and_requires_restore(tmp_path, case):
    from tests.codex_pptx_tables.tables import validate_revisions

    history = []
    for index, value in enumerate(
        [
            "7" if case == "recovered" else "007",
            "008",
            "008" if case == "missing_restore" else "007",
            "007",
        ]
    ):
        presentation = make_presentation()
        shapes = presentation.slides[0].shapes
        shapes[0].table.cell(2, 1).text = value
        if index == 3:
            shapes[-1].element.getparent().remove(shapes[-1].element)
        name = str(index)
        presentation.save(tmp_path / name)
        history.append(
            {
                "sha256": name,
                "result": {
                    "changes": [{"operation": "add_table" if index == 0 else "edit"}]
                },
            }
        )
    if case == "missing_restore":
        with pytest.raises(ValueError, match="restoration"):
            validate_revisions(tmp_path, {"history": history})
    else:
        assert validate_revisions(tmp_path, {"history": history}) is (case == "exact")
