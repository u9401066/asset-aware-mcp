"""DOCX evidence binds full parsed blocks to immutable native revisions."""

from __future__ import annotations

import copy
import hashlib
import io
import json

import pytest
from docx import Document

from tests.native_docx_helpers import call, read_dfm
from tests.native_docx_helpers import native_docx as native_docx


def test_block_references_are_stable_and_hash_complete_formatting(native_docx):
    service, asset, source = native_docx
    parsed = service.docx.decompose(
        source.read_bytes(), asset["asset_id"], asset["revision"]
    )
    summaries = call(service, op="read_docx", asset_id=asset["asset_id"], limit=100)[
        "blocks"
    ]
    assert {b["id"]: b["evidence"] for b in summaries} == {
        r["block_id"]: r["evidence"] for r in parsed.blocks
    }
    assert {r["locator"]["part"] for r in parsed.blocks} >= {
        "word/document.xml",
        "word/header1.xml",
        "word/footer1.xml",
    }
    assert any(r["representation"].get("runs") for r in parsed.blocks)
    assert any(r["representation"].get("cell_formats") for r in parsed.blocks)
    for record in parsed.blocks:
        canonical = {k: v for k, v in record.items() if k != "evidence"}
        digest = hashlib.sha256(
            json.dumps(
                canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":")
            ).encode()
        ).hexdigest()
        assert digest == record["evidence"]["value_sha256"]
        excerpt = call(
            service,
            op="read_docx_block",
            asset_id=asset["asset_id"],
            block_id=record["block_id"],
        )["block"]
        assert excerpt["evidence"] == record["evidence"]
        assert excerpt["representation_complete"] is False
        verified = call(service, op="verify", reference=record["evidence"])
        assert verified["valid"] and all(verified["checks"].values())
        assert verified["source_freshness"].startswith("not_checked")


def test_long_unicode_block_is_reassembled_without_hashing_its_preview(native_docx):
    service, _, source = native_docx
    document = Document(io.BytesIO(source.read_bytes()))
    expected = "證據😀\n[[unsafe]] <script>x</script> " * 250
    document.paragraphs[1].text = expected
    document.save(source)
    asset = call(service, op="register", source_path=str(source))["asset"]
    summaries = call(service, op="read_docx", asset_id=asset["asset_id"], limit=100)[
        "blocks"
    ]
    summary = next(b for b in summaries if b["preview"].startswith("證據"))
    assert len(summary["preview"]) == 80
    pieces, offset = [], 0
    while True:
        block = call(
            service,
            op="read_docx_block",
            asset_id=asset["asset_id"],
            revision=asset["revision"],
            block_id=summary["id"],
            text_offset=offset,
            text_limit=377,
        )["block"]
        assert len(block["text_excerpt"]) <= 377
        assert block["evidence"] == summary["evidence"]
        assert block["text_sha256"] == hashlib.sha256(expected.encode()).hexdigest()
        assert "runs" not in block and "representation" not in block
        pieces.append(block["text_excerpt"])
        offset = block["next_text_offset"]
        if offset is None:
            break
    assert "".join(pieces) == expected
    past = call(
        service,
        op="read_docx_block",
        asset_id=asset["asset_id"],
        block_id=summary["id"],
        text_offset=len(expected) + 5,
    )["block"]
    assert past["text_excerpt"] == "" and past["next_text_offset"] is None
    assert past["text_complete"] is False


@pytest.mark.parametrize("mutation", ["hash", "part", "block", "revision"])
def test_forged_block_reference_never_verifies(native_docx, mutation):
    service, asset, _ = native_docx
    reference = copy.deepcopy(
        call(service, op="read_docx", asset_id=asset["asset_id"])["blocks"][0][
            "evidence"
        ]
    )
    if mutation == "hash":
        reference["value_sha256"] = "0" * 64
        result = call(service, op="verify", reference=reference)
        assert result["valid"] is False
        assert result["checks"]["block_representation_hash"] is False
        return
    if mutation == "part":
        reference["locator"]["part"] = "word/other.xml"
    elif mutation == "block":
        reference["locator"]["block_id"] = "p999999"
    else:
        reference["revision"] = "0" * 64
    with pytest.raises(ValueError):
        call(service, op="verify", reference=reference)


def test_old_references_survive_edit_and_archive_but_report_revision_freshness(
    native_docx,
):
    service, asset, source = native_docx
    references = [
        b["evidence"]
        for b in call(service, op="read_docx", asset_id=asset["asset_id"])["blocks"]
    ]
    current = call(
        service,
        op="update_docx",
        asset_id=asset["asset_id"],
        expected_revision=asset["revision"],
        docx_edit={"dfm_text": read_dfm(service, asset).replace("原始段落", "新段落")},
    )["asset"]
    call(
        service,
        op="archive",
        asset_id=asset["asset_id"],
        expected_revision=current["revision"],
    )
    for ref in references:
        result = call(service, op="verify", reference=ref)
        assert result["valid"] and result["archived"]
        assert result["is_current_managed_revision"] is False
    assert hashlib.sha256(source.read_bytes()).hexdigest() == asset["revision"]


def test_equal_text_in_other_blocks_does_not_validate_a_forged_locator(native_docx):
    service, _, source = native_docx
    document = Document(io.BytesIO(source.read_bytes()))
    document.paragraphs[1].text = "Repeated evidence"
    document.paragraphs[2].text = "Repeated evidence"
    document.save(source)
    asset = call(service, op="register", source_path=str(source))["asset"]
    blocks = call(service, op="read_docx", asset_id=asset["asset_id"])["blocks"]
    repeated = [b for b in blocks if b["preview"] == "Repeated evidence"]
    assert len(repeated) == 2
    forged = copy.deepcopy(repeated[0]["evidence"])
    forged["locator"] = repeated[1]["evidence"]["locator"]
    assert call(service, op="verify", reference=forged)["valid"] is False


def test_corrupt_immutable_docx_blob_fails_before_block_verification(native_docx):
    service, asset, _ = native_docx
    ref = call(service, op="read_docx", asset_id=asset["asset_id"])["blocks"][0][
        "evidence"
    ]
    blob = service.repository.root / asset["asset_id"] / "revisions" / asset["revision"]
    blob.write_bytes(b"corrupted native source")
    with pytest.raises(ValueError, match="hash verification"):
        call(service, op="verify", reference=ref)
