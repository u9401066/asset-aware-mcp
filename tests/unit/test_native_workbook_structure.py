"""Read native structures with an independent reader and preserve rich parts."""

from __future__ import annotations

import io

import openpyxl
import pytest
from lxml import etree

from src.domain.native_workbook import (
    NativeWorksheetInsert,
    NativeWorksheetKey,
    NativeWorksheetRename,
)
from src.infrastructure.native_ooxml import (
    DOC_REL_NS,
    REL_NS,
    SHEET_NS,
    TYPE_NS,
    xml_bytes,
)
from src.infrastructure.native_workbook_structure import NativeWorkbookStructure
from tests.native_workbook_helpers import _parts, _replace, build_workbook

ADAPTER = NativeWorkbookStructure()
NS = {"s": SHEET_NS}


def keys(data: bytes) -> list[NativeWorksheetKey]:
    return [
        NativeWorksheetKey.model_validate(item["key"])
        for item in ADAPTER.read(data)["worksheets"]
    ]


def test_rich_workbook_sheet_lifecycle_and_exact_unchanged_parts():
    source = build_workbook()
    original = _parts(source)
    data, added = ADAPTER.add(
        source, NativeWorksheetInsert(index=0, names=["暫存", "Notes"])
    )
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == [
        "暫存",
        "Notes",
        "Data",
        "Other",
    ]
    assert (
        _parts(data)["xl/worksheets/sheet1.xml"] == original["xl/worksheets/sheet1.xml"]
    )
    current_keys = keys(data)
    data, renamed = ADAPTER.rename(
        data, NativeWorksheetRename(key=current_keys[2], name="研究 O'Brien")
    )
    renamed_parts = _parts(data)
    independent = openpyxl.load_workbook(io.BytesIO(data))
    assert independent["Other"]["B1"].value == "='研究 O''Brien'!A2"
    sheet = independent["研究 O'Brien"]
    assert sheet["A1"].value == "Original" and sheet["A1"].font.bold
    assert sheet["A1"].comment.text == "Keep this comment"
    assert str(sheet.merged_cells) == "A5:D6"
    assert sheet["A2"].number_format == "$0.00"
    assert sheet["C3"].hyperlink.target == "https://example.com"
    assert len(sheet.tables) == 1 and len(sheet._charts) == 1
    assert b"'" in renamed_parts["xl/charts/chart1.xml"]
    assert "研究 O''Brien" in renamed_parts["xl/charts/chart1.xml"].decode()
    assert (
        renamed_parts["xl/worksheets/sheet1.xml"]
        == original["xl/worksheets/sheet1.xml"]
    )
    for name in set(original) - set(added.changed_parts) - set(renamed.changed_parts):
        assert renamed_parts[name] == original[name], name
    data, _ = ADAPTER.reorder(data, list(reversed(keys(data))))
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == [
        "Other",
        "研究 O'Brien",
        "Notes",
        "暫存",
    ]
    removed_key = keys(data)[-1]
    data, deleted = ADAPTER.delete(data, [removed_key])
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == [
        "Other",
        "研究 O'Brien",
        "Notes",
    ]
    assert removed_key.part in _parts(data)
    assert deleted.changes[0]["secure_erasure"] is False
    assert _parts(data)["customXml/item1.xml"] == original["customXml/item1.xml"]


def test_surviving_formula_blocks_deleting_data():
    source = build_workbook()
    with pytest.raises(ValueError, match="surviving explicit"):
        ADAPTER.delete(source, [keys(source)[0]])
    # Other has no dependants; its protected worksheet is retained as detached bytes.
    updated, _ = ADAPTER.delete(source, [keys(source)[1]])
    assert openpyxl.load_workbook(io.BytesIO(updated)).sheetnames == ["Data"]
    assert (
        _parts(updated)["xl/worksheets/sheet2.xml"]
        == _parts(source)["xl/worksheets/sheet2.xml"]
    )


def test_name_scope_and_views_follow_sheet_identity():
    data = build_workbook()
    root = etree.fromstring(_parts(data)["xl/workbook.xml"])
    names = etree.Element(f"{{{SHEET_NS}}}definedNames")
    etree.SubElement(
        names, f"{{{SHEET_NS}}}definedName", name="Scoped", localSheetId="0"
    ).text = "Data!$A$1"
    etree.SubElement(
        names, f"{{{SHEET_NS}}}definedName", name="Removed", localSheetId="1"
    ).text = "Other!$A$1"
    root.insert(list(root).index(root.find("s:calcPr", NS)), names)
    view = root.find("s:bookViews/s:workbookView", NS)
    view.set("activeTab", "1")
    customs = etree.SubElement(root, f"{{{SHEET_NS}}}customWorkbookViews")
    etree.SubElement(
        customs,
        f"{{{SHEET_NS}}}customWorkbookView",
        name="Custom",
        guid="{12345678-1234-1234-1234-123456789012}",
        activeSheetId="2",
    )
    data = _replace(data, {"xl/workbook.xml": xml_bytes(root)})
    data, _ = ADAPTER.reorder(data, list(reversed(keys(data))))
    metadata = ADAPTER.read(data)
    assert metadata["views"][0]["activeTab"] == "0"
    assert metadata["views"][0]["firstSheet"] == "1"
    assert metadata["defined_names"][0]["attributes"]["localSheetId"] == "1"
    assert metadata["custom_views"][0]["activeSheetId"] == "2"
    data, _ = ADAPTER.delete(data, [keys(data)[0]])
    metadata = ADAPTER.read(data)
    assert [item["attributes"]["name"] for item in metadata["defined_names"]] == [
        "Scoped"
    ]
    assert metadata["defined_names"][0]["attributes"]["localSheetId"] == "0"
    assert metadata["custom_views"][0]["activeSheetId"] == "1"


def test_3d_membership_requires_explicit_change():
    data = build_workbook()
    part = etree.fromstring(_parts(data)["xl/worksheets/sheet2.xml"])
    part.find(".//s:f", NS).text = "SUM(Data:Other!A2)"
    data = _replace(data, {"xl/worksheets/sheet2.xml": xml_bytes(part)})
    request = NativeWorksheetInsert(index=1, names=["Inside"])
    with pytest.raises(ValueError, match="3D-reference membership"):
        ADAPTER.add(data, request)
    data, result = ADAPTER.add(data, request, True)
    assert result.changes[0]["changed_3d_memberships"] == 1
    order = keys(data)
    with pytest.raises(ValueError, match="3D-reference membership"):
        ADAPTER.reorder(data, [order[1], order[0], order[2]])
    with pytest.raises(ValueError, match="surviving explicit"):
        ADAPTER.delete(data, [order[1]])


@pytest.mark.parametrize("guard", ["workbookProtection", "extLst", "fileSharing"])
def test_structure_guards_allow_read_but_not_mutation(guard):
    source = build_workbook()
    root = etree.fromstring(_parts(source)["xl/workbook.xml"])
    etree.SubElement(root, f"{{{SHEET_NS}}}{guard}")
    source = _replace(source, {"xl/workbook.xml": xml_bytes(root)})
    assert len(ADAPTER.read(source)["worksheets"]) == 2
    with pytest.raises(ValueError, match="workflow"):
        ADAPTER.add(source, NativeWorksheetInsert(index=0, names=["New"]))


def test_calc_chain_detaches_without_losing_new_relationships_or_old_parts():
    source = build_workbook()
    parts = _parts(source)
    rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
    etree.SubElement(
        rels,
        f"{{{REL_NS}}}Relationship",
        Id="rId777",
        Type=f"{DOC_REL_NS}/calcChain",
        Target="calcChain.xml",
    )
    types = etree.fromstring(parts["[Content_Types].xml"])
    etree.SubElement(
        types,
        f"{{{TYPE_NS}}}Override",
        PartName="/xl/calcChain.xml",
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
    )
    chain = f'<calcChain xmlns="{SHEET_NS}"><c r="B2" i="1"/></calcChain>'.encode()
    source = _replace(
        source,
        {
            "xl/_rels/workbook.xml.rels": xml_bytes(rels),
            "[Content_Types].xml": xml_bytes(types),
            "xl/calcChain.xml": chain,
        },
    )
    data, result = ADAPTER.add(source, NativeWorksheetInsert(index=0, names=["Added"]))
    assert "stale_calculation_chain_detached" in result.repairs
    assert _parts(data)["xl/calcChain.xml"] == chain
    assert b"rId777" not in _parts(data)["xl/_rels/workbook.xml.rels"]
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames[0] == "Added"
