"""Independent merge audit catches lost leading zeros, rich format and split drift."""

import io

import pytest
from pptx import Presentation
from pptx.util import Pt

from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pptx_tables.merges import TEMP, check_merge_delta, check_merge_stage


def save(presentation):
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def setup():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(5, 5, 100000, 100000, 6000000, 2750000).table
    table.cell(0, 0).merge(table.cell(0, 4))
    values = [
        ["Scanned inventory", "", "", "", ""],
        list(COLUMNS),
        *map(list, PAGE_ROWS[0]),
        TEMP,
    ]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            if r == 0 and c:
                continue
            cell = table.cell(r, c)
            cell.text = value
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    table.cell(4, 1).text_frame.paragraphs[0].runs[0].font.bold = True
    table.cell(4, 2).text_frame.paragraphs[0].runs[0].font.italic = True
    return presentation, table


@pytest.mark.parametrize(
    "failure", [None, "zero", "style", "merge", "source", "split_redistribution"]
)
def test_independent_merge_and_split_audit(failure):
    presentation, table = setup()
    before = save(presentation)
    check_merge_stage(before, 1)
    table.cell(4, 0).merge(table.cell(4, 3 if failure == "merge" else 4))
    after_merge = save(presentation)
    if failure == "zero":
        table.cell(4, 0).text_frame.paragraphs[1].runs[0].text = "0"
    elif failure == "style":
        table.cell(4, 0).text_frame.paragraphs[1].runs[0].font.bold = False
    elif failure == "source":
        table.cell(2, 1).margin_left += 1
    elif failure == "split_redistribution":
        table.cell(4, 0).split()
        for c, text in enumerate(TEMP):
            table.cell(4, c).text = text
    result = save(presentation)
    if failure:
        with pytest.raises(ValueError):
            if failure in {"style", "source"}:
                check_merge_stage(result, 2)
                check_merge_delta(before, result, 2)
            else:
                check_merge_stage(result, 3 if failure == "split_redistribution" else 2)
    else:
        check_merge_stage(result, 2)
        check_merge_delta(before, result, 2)
        table.cell(4, 0).split()
        split = save(presentation)
        check_merge_stage(split, 3)
        check_merge_delta(after_merge, split, 3)
