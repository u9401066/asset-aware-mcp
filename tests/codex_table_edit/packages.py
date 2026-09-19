"""Independent openpyxl/package verification, including rich header run formats."""

import io
import zipfile

from openpyxl import load_workbook

from tests.codex_native_pdf.trace import digest
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require


def check_packages(workspace, target, expected):
    history = target["history"]
    require(
        len(history) == 3, "Expected registration, transcription and one Table edit"
    )
    template = workspace / "template.xlsx"
    require(
        digest(template.read_bytes())
        == expected["template_sha256"]
        == history[0]["sha256"],
        "Template bytes differ",
    )
    require(
        template.stat().st_mtime_ns == expected["template_mtime_ns"],
        "Template mtime changed",
    )
    original_parts, original_fonts, previous_parts = None, None, None
    for index, entry in enumerate(history):
        data = (
            workspace
            / "data/native-assets"
            / target["asset_id"]
            / "revisions"
            / entry["sha256"]
        ).read_bytes()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        original_parts = original_parts or parts
        require(
            parts["xl/styles.xml"] == original_parts["xl/styles.xml"],
            "Native styles changed",
        )
        if previous_parts is not None:
            changed = set(entry["result"]["changed_parts"])
            require(
                all(
                    parts[name] == value
                    for name, value in previous_parts.items()
                    if name not in changed
                ),
                "Untouched native part changed",
            )
        previous_parts = parts
        book = load_workbook(io.BytesIO(data), rich_text=True)
        try:
            sheet, native = book["Sheet1"], book["Sheet1"].tables["Inventory"]
            require(
                native.ref == "A1:F4" and native.autoFilter.ref == "A1:F3",
                "Table/filter boundaries changed",
            )
            require(
                [c.id for c in native.tableColumns] == list(range(1, 7)),
                "Native column identities changed",
            )
            headers = [*COLUMNS, "CountLength"]
            if index == 2:
                headers[1] = "Quantity"
            require(
                [str(sheet.cell(1, c).value) for c in range(1, 7)] == headers,
                "Header cell text mismatch",
            )
            require(
                [col.name for col in native.tableColumns] == headers,
                "Header metadata mismatch",
            )
            runs = sheet["B1"].value
            require(
                len(runs) == 2 and all(hasattr(run, "font") for run in runs),
                "Rich header run structure lost",
            )
            fonts = [run.font for run in runs]
            original_fonts = original_fonts or fonts
            require(fonts == original_fonts, "Rich header run formatting changed")
            require(
                [run.text for run in runs]
                == (["Quan", "tity"] if index == 2 else ["Co", "unt"]),
                "Rich header run text differs",
            )
            if index:
                for row, truth in enumerate(PAGE_ROWS[0], 2):
                    require(
                        [sheet.cell(row, col).value for col in range(1, 6)]
                        == list(truth),
                        "Scan data characters changed",
                    )
                    require(
                        all(
                            sheet.cell(row, col).data_type == "s" for col in range(1, 6)
                        ),
                        "Scan data types changed",
                    )
            for row in (2, 3):
                require(
                    sheet[f"F{row}"].value
                    == (
                        "=LEN([@Quantity])+1"
                        if index == 2
                        else "=LEN([[#This Row],Count])"
                    ),
                    "Calculated formula mismatch",
                )
            require(
                sheet["F4"].value == "=SUBTOTAL(109,[CountLength])",
                "Totals formula differs",
            )
            require(
                sheet["A4"].value == ("Reviewed" if index == 2 else None),
                "Totals label differs",
            )
            require(
                book["Summary"]["A1"].value
                == f"=COUNTA(Inventory[{'Quantity' if index == 2 else 'Count'}])",
                "Dependent reference mismatch",
            )
            require(
                book.defined_names["OriginalCount"].attr_text
                == f"Inventory[{'Quantity' if index == 2 else 'Count'}]",
                "Defined-name reference mismatch",
            )
        finally:
            book.close()
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Published workbook differs",
    )
