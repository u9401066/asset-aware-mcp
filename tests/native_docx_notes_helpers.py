"""Independent OOXML fixture with note identities, rich text and body references."""

import hashlib
import io
import json
import zipfile

from docx import Document
from docx.shared import Pt
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "http://schemas.openxmlformats.org/package/2006/relationships"
TYPE = "http://schemas.openxmlformats.org/package/2006/content-types"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
FOOT = "word/annotations/source-footnotes.xml"
END = "word/annotations/source-endnotes.xml"


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def canonical_record(record):
    return json.dumps(
        record, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode()


def parts(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def package(values):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in values.items():
            archive.writestr(name, value)
    return output.getvalue()


def xml(node):
    return etree.tostring(node, encoding="UTF-8", xml_declaration=True, standalone=True)


def source_document():
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for text, refs in [
        ("Source 007 µg and preserved run.", [("footnote", 2)]),
        ("Second value -0.50 mg/L.", [("footnote", 8), ("endnote", 5)]),
    ]:
        p = doc.add_paragraph()
        p.add_run(text).bold = True
        for kind, identity in refs:
            run = p.add_run()._r
            pr = etree.SubElement(run, W + "rPr")
            etree.SubElement(pr, W + "vertAlign", {W + "val": "superscript"})
            etree.SubElement(run, W + kind + "Reference", {W + "id": str(identity)})
        if len(doc.paragraphs) == 1:
            doc.add_page_break()
    raw = io.BytesIO()
    doc.save(raw)
    values = parts(raw.getvalue())
    rels = etree.fromstring(values["word/_rels/document.xml.rels"])
    types = etree.fromstring(values["[Content_Types].xml"])
    for kind, part, entries in [
        (
            "footnote",
            FOOT,
            [(2, "FOOTNOTE 007 µg -0.50 mg/L"), (8, "REMOVE OLD FOOTNOTE")],
        ),
        ("endnote", END, [(5, "ENDNOTE 1,234.50 — EVIDENCE TAIL")]),
    ]:
        root = etree.Element(W + kind + "s", nsmap={"w": W[1:-1]})
        for identity, role in [(-1, "separator"), (0, "continuationSeparator")]:
            note = etree.SubElement(
                root, W + kind, {W + "id": str(identity), W + "type": role}
            )
            run = etree.SubElement(etree.SubElement(note, W + "p"), W + "r")
            etree.SubElement(run, W + role)
        for identity, text in entries:
            note = etree.SubElement(root, W + kind, {W + "id": str(identity)})
            para = etree.SubElement(note, W + "p")
            marker = etree.SubElement(para, W + "r")
            etree.SubElement(marker, W + kind + "Ref")
            run = etree.SubElement(para, W + "r")
            pr = etree.SubElement(run, W + "rPr")
            etree.SubElement(pr, W + "b")
            etree.SubElement(pr, W + "sz", {W + "val": "18"})
            etree.SubElement(
                pr, W + "rFonts", {W + "ascii": "Arial", W + "hAnsi": "Arial"}
            )
            etree.SubElement(
                run,
                W + "t",
                {"{http://www.w3.org/XML/1998/namespace}space": "preserve"},
            ).text = " " + text
            extra = etree.SubElement(etree.SubElement(note, W + "p"), W + "r")
            etree.SubElement(extra, W + "t").text = "KEEP SECOND PARAGRAPH"
        values[part] = xml(root)
        etree.SubElement(
            rels,
            "{" + R + "}Relationship",
            Id="rId" + kind,
            Type=REL + kind + "s",
            Target=part.removeprefix("word/"),
        )
        etree.SubElement(
            types,
            "{" + TYPE + "}Override",
            PartName="/" + part,
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml."
            + kind
            + "s+xml",
        )
    values["word/_rels/document.xml.rels"] = xml(rels)
    values["[Content_Types].xml"] = xml(types)
    return package(values)


def locator(kind="footnote", identity=2):
    return {
        "part": FOOT if kind == "footnote" else END,
        "note_kind": kind,
        "note_id": identity,
    }
