"""Audit scan-derived structural A2T writeback independently of production code."""

import io
import json
import zipfile

from openpyxl import load_workbook

from tests.codex_native_pdf.artifacts import load_asset
from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require
from tests.codex_table_grid.reads import read_events


def expected_values(final):
    values = [list(COLUMNS), *[list(row) for row in PAGE_ROWS[0]]]
    if final:
        values[0][-1] = "Review"
        values[1][1] = "008"
        values[1][-1] = values[2][-1] = "checked"
        values.append(["Added", "000", "0.0", "mg/L", "checked"])
    return values


def _check_native(target, directory):
    require(len(target["history"]) == 2, "Structural application must commit once")
    parts = []
    for index, event in enumerate(target["history"]):
        data = (directory / "revisions" / event["sha256"]).read_bytes()
        require(digest(data) == event["sha256"], "Workbook blob hash mismatch")
        book = load_workbook(io.BytesIO(data))
        try:
            require(book.sheetnames == ["Sheet1"], "Wrong structural worksheet")
            sheet = book["Sheet1"]
            cells = list(sheet.iter_rows())
            require(
                [[cell.value for cell in row] for row in cells]
                == expected_values(bool(index)),
                "Structural native values differ",
            )
            require(
                all(cell.data_type == "s" for row in cells for cell in row),
                "Literal types changed",
            )
        finally:
            book.close()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            parts.append({name: archive.read(name) for name in archive.namelist()})
    changed = target["history"][-1]["result"]["changed_parts"]
    require(
        all(
            parts[0][name] == parts[1].get(name)
            for name in parts[0]
            if name not in changed
        ),
        "Unchanged native members differ",
    )
    require(
        parts[0]["xl/styles.xml"] == parts[1]["xl/styles.xml"],
        "Structural workbook styles changed",
    )


def _identity_check(initial, current):
    require(
        len(initial["column_ids"]) == len(set(initial["column_ids"])) == 5,
        "Wrong original column identities",
    )
    require(
        [[row.get(name) for name in "ABCDE"] for row in initial["rows"]]
        == [
            [{"kind": "string", "value": value} for value in row]
            for row in expected_values(False)
        ],
        "Original projected values differ from the scan",
    )
    require(
        initial["native_binding"] == current["native_binding"],
        "Source binding migrated",
    )
    require(
        len(initial["row_ids"]) == 3 and len(current["row_ids"]) == 4,
        "Wrong structural row counts",
    )
    require(
        current["row_ids"][:2] == initial["row_ids"][:2]
        and initial["row_ids"][2] not in current["row_ids"]
        and len(set(current["row_ids"])) == 4,
        "Deleted row identity reused or surviving identity lost",
    )
    require(
        current["column_ids"][:4] == initial["column_ids"][:4]
        and len(set(current["column_ids"])) == 5
        and initial["column_ids"][4] not in current["column_ids"],
        "Deleted column identity reused or surviving identity lost",
    )
    names = [column["name"] for column in current["columns"]]
    require(names == ["A", "B", "C", "D", "Review"], "Wrong workspace columns")
    values = [[row.get(name) for name in names] for row in current["rows"]]
    expected = [
        [{"kind": "string", "value": value} for value in row]
        for row in expected_values(True)
    ]
    require(values == expected, "Frozen A2T input differs from intended values")


def validate_table_grid_workflow(calls, workspace, target, directory):
    _check_native(target, directory)
    mutations = [
        (index, call)
        for index, call in enumerate(calls)
        if call["tool"] == "document"
        and call["arguments"]["native_request"]["op"] == "apply_table_workspace"
    ]
    require(len(mutations) == 1, "Missing or additional structural application")
    applied_at, call = mutations[0]
    args, result = call["arguments"]["native_request"], payload(call)
    reads = list(read_events(calls))
    before = [
        (index, record)
        for index, request, record, _ in reads
        if index < applied_at
        and request["op"] == "read_table_workspace"
        and request["table_id"] == args["table_id"]
    ]
    require(len(before) >= 2, "Missing complete pre-edit and pre-apply workspace reads")
    initial, current = before[0][1]["table"], before[-1][1]["table"]
    require(
        digest(canonical(current))
        == args["expected_table_sha256"]
        == result["table_sha256"],
        "Application lacks the current workspace hash",
    )
    _identity_check(initial, current)
    table_edits = [
        (index, item["arguments"])
        for index, item in enumerate(calls)
        if item["tool"] in {"table_data", "table_manage"}
    ]
    require(
        bool(table_edits)
        and before[0][0] < min(index for index, _ in table_edits)
        and max(index for index, _ in table_edits) < before[-1][0],
        "A2T mutations lack complete reads before and after",
    )
    require(
        all(request["table_id"] == args["table_id"] for _, request in table_edits),
        "Unrelated A2T table mutation",
    )
    require(
        {"delete_row", "add_rows", "remove_column", "add_column", "update_cell"}
        <= {request["operation"] for _, request in table_edits},
        "Missing structural A2T edits",
    )
    _check_application(args, result, initial, current, target)
    _check_receipts(reads, applied_at, target)
    _check_snapshot(calls, reads, applied_at, result, current, workspace)


def _check_application(args, result, initial, current, target):
    require(
        args["asset_id"] == target["asset_id"]
        and args["expected_revision"] == target["history"][0]["sha256"]
        and result["asset"]["revision"] == target["revision"],
        "Wrong structural native transition",
    )
    binding = initial["native_binding"]
    require(
        binding["projection"]["start_cell"] == "A1"
        and binding["projection"]["end_cell"] == "E3",
        "Wrong original projection bounds",
    )
    require(
        binding["source"]["asset_id"] == target["asset_id"]
        and binding["source"]["revision"] == args["expected_revision"],
        "Wrong structural source binding",
    )
    grid = args["worksheet_grid"]
    require(
        grid["worksheet"] == binding["projection"]["worksheet"], "Wrong grid worksheet"
    )
    mapping = {"row": list(initial["row_ids"]), "column": list(initial["column_ids"])}
    for edit in grid["edits"]:
        values = mapping[edit["axis"]]
        at, count = edit["at"] - 1, edit.get("count", 1)
        require(0 <= at <= len(values), "Grid is outside source range")
        if edit["operation"] == "insert":
            values[at:at] = [None] * count
        else:
            require(at + count <= len(values), "Grid deletes outside source range")
            del values[at : at + count]
    for axis, key in (("row", "row_ids"), ("column", "column_ids")):
        old = set(initial[key])
        require(
            mapping[axis]
            == [identity if identity in old else None for identity in current[key]],
            "Grid correspondence differs from stable identities",
        )


def _check_receipts(reads, applied_at, target):
    before = after = False
    for index, args, record, revision in reads:
        if (
            args["op"] != "read_workbook"
            or args["asset_id"] != target["asset_id"]
            or args.get("workbook_view") != "references"
        ):
            continue
        require(isinstance(record["references"], list), "Missing reference inventory")
        if index < applied_at and revision == target["history"][0]["sha256"]:
            before = True
        if index > applied_at and revision == target["revision"]:
            require(
                record["operation_result"] == target["history"][-1]["result"],
                "Read receipt differs from current native history",
            )
            after = True
    require(before and after, "Missing complete pre/post native references and receipt")


def _check_snapshot(calls, reads, applied_at, result, current, workspace):
    reference = result["workspace_reference"]
    require(
        reference["revision"] == digest(canonical(current)),
        "Wrong frozen reference hash",
    )
    require(
        any(
            index > applied_at
            and args.get("workspace_reference") == reference
            and record["table"] == current
            for index, args, record, _ in reads
        ),
        "Frozen workspace was not completely read",
    )
    require(
        any(
            call["tool"] == "document"
            and call["arguments"]["native_request"].get("op") == "verify"
            and call["arguments"]["native_request"].get("reference") == reference
            and payload(call).get("valid")
            for call in calls[applied_at + 1 :]
        ),
        "Frozen workspace proof missing",
    )
    asset = load_asset(workspace, reference["asset_id"])
    require(
        asset["format"] == "a2t" and asset["revision"] == reference["revision"],
        "Wrong frozen asset",
    )
    data = (
        workspace
        / "data/native-assets"
        / asset["asset_id"]
        / "revisions"
        / asset["revision"]
    ).read_bytes()
    require(
        digest(data) == reference["revision"] and json.loads(data) == current,
        "Stored frozen workspace differs",
    )
