"""Independent ZIP/openpyxl audit across five managed workbook history entries."""

import io

from openpyxl import load_workbook

from tests.codex_native_pdf.trace import digest
from tests.codex_pdf.fixtures import COLUMNS, PAGE_ROWS
from tests.codex_pdf.trace import require
from tests.codex_table_create.packages import check_table
from tests.native_workbook_helpers import _parts


def check_packages(workspace, target):
    require(
        len(target["history"]) == 5 and target["source"] is None,
        "Expected independent creation, one Table and three totals transitions",
    )
    previous = None
    for index, entry in enumerate(target["history"]):
        data = (
            workspace
            / "data/native-assets"
            / target["asset_id"]
            / "revisions"
            / entry["sha256"]
        ).read_bytes()
        require(digest(data) == entry["sha256"], "Revision bytes differ")
        parts = _parts(data)
        book = load_workbook(io.BytesIO(data))
        try:
            require(book.sheetnames == ["Sheet1"], "Wrong worksheet inventory")
            sheet = book["Sheet1"]
            for row, truth in zip(
                sheet.iter_rows(max_row=3, max_col=5),
                [COLUMNS, *PAGE_ROWS[0]],
                strict=True,
            ):
                require(
                    tuple(c.value for c in row) == truth
                    and all(c.data_type == "s" for c in row),
                    "Scan strings/types changed",
                )
            require(sheet["F1"].value == "CountLength", "Wrong calculated header")
            if index == 0:
                require(
                    not sheet.tables
                    and all(sheet[f"F{r}"].value is None for r in (2, 3, 4)),
                    "Baseline already has native Table/calculations",
                )
            elif index == 1:
                check_table(sheet, entry)
            else:
                check_transition(sheet, entry, index)
        finally:
            book.close()
        if previous is not None:
            require(
                parts["xl/styles.xml"] == previous["xl/styles.xml"], "Styles changed"
            )
            require(
                all(
                    parts[name] == raw
                    for name, raw in previous.items()
                    if name not in entry["result"]["changed_parts"]
                ),
                "Untouched package part changed",
            )
        previous = parts
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Published bytes differ from final revision",
    )


def check_transition(sheet, entry, index):
    require(set(sheet.tables) == {"Inventory"}, "Wrong native Table inventory")
    table = sheet.tables["Inventory"]
    totals = index == 3
    require(
        table.ref == ("A1:F4" if totals else "A1:F3")
        and table.autoFilter.ref == "A1:F3"
        and table.totalsRowCount == int(totals),
        "Incorrect Table/filter/totals state",
    )
    require(
        [c.id for c in table.tableColumns] == list(range(1, 7))
        and [c.name for c in table.tableColumns] == [*COLUMNS, "CountLength"],
        "Column identity changed",
    )
    require(
        table.tableStyleInfo.name == "TableStyleMedium2"
        and table.tableStyleInfo.showRowStripes,
        "Native Table style changed",
    )
    require(
        table.tableColumns[5].calculatedColumnFormula.text == "LEN([@Count])"
        and table.tableColumns[5].totalsRowFunction == "sum"
        and table.tableColumns[0].totalsRowLabel == "Reviewed",
        "Hidden definitions changed",
    )
    require(
        [sheet[f"F{r}"].value for r in (2, 3)] == ["=LEN([@Count])"] * 2,
        "Calculated data changed",
    )
    expected = {
        2: [None] * 6,
        3: ["Reviewed", None, None, None, None, "=SUBTOTAL(109,[CountLength])"],
        4: ["Reviewed", None, None, None, None, "=SUBTOTAL(109,'Sheet1'!$F$2:$F$3)"],
    }[index]
    require(
        [sheet.cell(4, col).value for col in range(1, 7)] == expected,
        "Totals/detached cell values differ",
    )
    if index >= 3:
        require(
            all(
                sheet.cell(4, c).style_id == sheet.cell(3, c).style_id
                for c in range(1, 7)
            ),
            "Explicit inherited direct cell style differs",
        )
    change = entry["result"]["changes"][0]
    intent = change["totals_transition"]
    require(
        change["operation"] == "update_workbook_table"
        and intent["action"] == ("add" if totals else "remove")
        and not intent["worksheet_rows_moved"],
        "Missing exact transition receipt",
    )
    require(
        intent["before_ref"] == ("A1:F3" if totals else "A1:F4")
        and intent["after_ref"] == table.ref,
        "Wrong transition bounds",
    )
    if index == 4:
        require(set(change["cells"]) == {"F4"}, "Keep unexpectedly rewrote values")
        require(
            change["cells"]["F4"]["before"]["value"] == "=SUBTOTAL(109,[CountLength])",
            "Detached formula receipt loses original value",
        )
