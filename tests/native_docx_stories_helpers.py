"""Independent human-source Word fixture with shared, nonstandard story names."""

import io

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from src.infrastructure.native_docx_workspace import checked_docx
from src.infrastructure.native_ooxml import xml_bytes

HEADER_PART = "word/running/report.xml"
HEADER_TEXT = "SOURCE 007 µg " + "Complete header evidence. " * 9 + "END-HEADER"


def story_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.4)
    section.header_distance = Inches(0.25)
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "FIRST PAGE ONLY"
    header = section.header
    p = header.paragraphs[0]
    run = p.add_run(HEADER_TEXT)
    run.bold = True
    run.font.size = Pt(9)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string("176B48")
    run = p.add_run(" KEEP-ITALIC")
    run.italic = True
    run.font.size = Pt(9)
    table = header.add_table(rows=1, cols=2, width=Inches(4))
    table.cell(0, 0).text = "007"
    table.cell(0, 1).text = "-0.50 mg/L"
    header.add_paragraph("REMOVE THIS PARAGRAPH")
    footer = section.footer.paragraphs[0]
    footer.add_run("Page ")
    for tag, value in [
        ("fldChar", "begin"),
        ("instrText", " PAGE "),
        ("fldChar", "separate"),
        ("t", "1"),
        ("fldChar", "end"),
    ]:
        r = OxmlElement("w:r")
        node = OxmlElement("w:" + tag)
        if tag == "fldChar":
            node.set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType",
                value,
            )
        else:
            node.text = value
        r.append(node)
        footer._p.append(r)
    doc.add_paragraph("Section A / first page")
    doc.add_page_break()
    doc.add_paragraph("Section A / second page")
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    second.different_first_page_header_footer = False
    doc.add_paragraph("Section B / inherited shared header")
    raw = io.BytesIO()
    doc.save(raw)
    package = checked_docx(raw.getvalue())
    rels = package.xml("word/_rels/document.xml.rels")
    old = next(node for node in rels if node.get("Target") == "header2.xml")
    old.set("Target", "running/report.xml")
    types = package.xml("[Content_Types].xml")
    next(node for node in types if node.get("PartName") == "/word/header2.xml").set(
        "PartName", "/" + HEADER_PART
    )
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, value in package.parts.items():
            if name == "word/header2.xml":
                name = HEADER_PART
            if name == "word/_rels/document.xml.rels":
                value = xml_bytes(rels)
            if name == "[Content_Types].xml":
                value = xml_bytes(types)
            z.writestr(name, value)
    return out.getvalue()


HEADER_CORRECTED = HEADER_TEXT.replace("SOURCE ", "VERIFIED ", 1)
ADDED_TEXT = "REVIEWED 1,234.50"


def correction(record):
    part = record["locator"]["part"]
    if part == HEADER_PART:
        node = next(n for n in record["text_nodes"] if n["text"] == HEADER_TEXT)
        return {
            "part": part,
            "shared_scope": "all_sections_using_part",
            "edits": [
                {"op": "set_text", "path": node["path"], "text": HEADER_CORRECTED},
                {
                    "op": "insert_blocks",
                    "index": 3,
                    "blocks": [
                        {
                            "kind": "paragraph",
                            "runs": [
                                {
                                    "text": ADDED_TEXT,
                                    "font_name": "Arial",
                                    "font_size_pt": 9.0,
                                }
                            ],
                        }
                    ],
                },
                {"op": "delete_blocks", "index": 2, "count": 1},
            ],
        }
    node = next(n for n in record["text_nodes"] if n["text"] == "Page ")
    return {
        "part": part,
        "shared_scope": "all_sections_using_part",
        "edits": [{"op": "set_text", "path": node["path"], "text": "Verified page "}],
    }


def inspect_story_pages(data, stage):
    import pymupdf

    from tests.native_docx_layout_helpers import independent_pdf

    pdf_data = independent_pdf(data)
    with pymupdf.open(stream=pdf_data, filetype="pdf") as pdf:
        texts = [p.get_text() for p in pdf]
    assert len(texts) == 3, texts
    assert "FIRST PAGE ONLY" in texts[0]
    assert "END-HEADER" not in texts[0]
    for text in texts[1:]:
        assert ("VERIFIED 007" if stage else "SOURCE 007") in text
        for token in ["µg", "END-HEADER", "KEEP-ITALIC", "-0.50 mg/L"]:
            assert token in text, (token, text)
        assert (ADDED_TEXT in text) == bool(stage)
        assert ("REMOVE THIS PARAGRAPH" in text) != bool(stage)
    assert "Page " not in texts[0] and "Verified page " not in texts[0]
    for index, text in enumerate(texts[1:], 2):
        assert (("Verified page " if stage == 2 else "Page ") + str(index)) in text
    return pdf_data, texts
