"""Check retained sources, exact TIFF stages, literal cells and portable Wikis."""

import io
import json
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from PIL import Image

from tests.codex_native_image.checks import (
    covers_row,
    crop_bounds,
    frames,
    same_pixels,
    source_bytes,
    validate_png,
    validate_stages,
)
from tests.codex_native_pdf.artifacts import load_asset, read_json
from tests.codex_native_pdf.trace import canonical, digest, payload
from tests.codex_pdf.trace import require
from tests.real_pdf.corpus import check_source


def source_and_stages(workspace, expected, final, trace):
    source = workspace / "source.png"
    require(
        digest(source.read_bytes()) == expected["source_sha256"]
        and source.stat().st_mtime_ns == expected["source_mtime_ns"],
        "Source image changed",
    )
    check_source(workspace / "source.pdf", expected["case"])
    require(
        (workspace / "source.pdf").stat().st_mtime_ns == expected["pdf_mtime_ns"],
        "Source PDF mtime changed",
    )
    with pymupdf.open(workspace / "source.pdf") as pdf:
        page = pdf[expected["derivation"]["page_index"]]
        scale = 1400 / max(page.rect.width, page.rect.height)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
        upright = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    require(
        digest(upright.tobytes()) == expected["derivation"]["upright_rgb_sha256"],
        "PDF raster oracle differs",
    )
    with Image.open(source) as image:
        require(
            image.getexif()[274] == 6 and image.size == (upright.height, upright.width),
            "Fixture EXIF/stored geometry differs",
        )
    same_pixels(frames(source.read_bytes())[0], upright)
    source_asset = load_asset(workspace, final["source_asset_id"])
    require(
        source_asset["revision"] == expected["source_sha256"]
        and len(source_asset["history"]) == 1,
        "Managed original image changed",
    )
    target = load_asset(workspace, final["image_asset_id"])
    book = load_asset(workspace, final["workbook_asset_id"])
    ledger = read_json(
        workspace / "data/native-assets" / book["asset_id"] / "derivations.json"
    )
    require(
        len(ledger["events"]) == 1 and digest(canonical(ledger)) in trace["ledgers"],
        "Final ledger missing or unread",
    )
    claim = ledger["events"][0]["derivation"]
    require(len(claim["sources"]) == 1, "Unexpected derivation sources")
    region = claim["sources"][0]
    require(
        region["schema_version"] == "native-image-region-ref-v1"
        and (region["asset_id"], region["revision"])
        == (source_asset["asset_id"], source_asset["revision"]),
        "Wrong original source region",
    )
    require(
        canonical(region) in trace["regions"]
        and canonical(region) in trace["verified"],
        "Region not viewed and verified",
    )
    require(
        canonical(claim["target"]) in trace["cells"]
        and claim["target"]["locator"]["cell"] == "B2"
        and claim["target"]["revision"] == book["revision"],
        "Wrong target evidence",
    )
    require(claim["review"]["notes"], "Missing explicit Agent review")
    covers_row(region["selector"]["rect"], expected)
    crop = upright.crop(crop_bounds(region["selector"]["rect"], upright.size))
    revisions = [s["sha256"] for s in target["history"]]
    data = [source_bytes(workspace, target["asset_id"], r) for r in revisions]
    validate_stages(data, upright, crop)
    require(
        (workspace / "verified.tif").read_bytes() == data[-1], "Published TIFF differs"
    )
    historical = [
        r["evidence"]
        for r in trace["records"].values()
        if r["evidence"]["asset_id"] == target["asset_id"]
        and (
            r["evidence"]["revision"] == revisions[0]
            or (
                r["evidence"]["revision"] == revisions[1]
                and r["locator"]["frame_index"] == 2
            )
        )
    ]
    require(
        len(historical) == 3
        and all(canonical(r) in trace["verified"] for r in historical),
        "Original/deleted frame references not verified",
    )
    selections = [
        json.loads(r)
        for r in trace["verified"]
        if json.loads(r).get("schema_version") == "native-selection-ref-v1"
    ]
    require(
        any(
            r["parent"] == region and r["selector"]["pointer"] == "/source_pixel_bounds"
            for r in selections
        ),
        "Region selection not verified",
    )
    for rev in revisions:
        require((target["asset_id"], rev) in trace["catalogs"], "Missing stage read")
    raw_book = source_bytes(workspace, book["asset_id"], book["revision"])
    require(
        (workspace / "verified.xlsx").read_bytes() == raw_book,
        "Published workbook differs",
    )
    wb = load_workbook(io.BytesIO(raw_book), data_only=False)
    require(wb.sheetnames == ["Evidence"], "Unexpected workbook sheets")
    truth = [expected["case"]["columns"], expected["case"]["rows"][0]]
    sheet = wb["Evidence"]
    require((sheet.max_row, sheet.max_column) == (2, 3), "Unexpected table geometry")
    for index, values in enumerate(truth, 1):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column)
            require(
                cell.data_type == "s" and cell.value == value,
                "Literal transcription cell differs",
            )
    require(len(trace["cells"]) == 6, "All literal cells were not read")
    seen_cells = set()
    for cell in trace["cells"].values():
        ref = cell["evidence"]
        require(
            (ref["asset_id"], ref["revision"]) == (book["asset_id"], book["revision"]),
            "Cell read identity differs",
        )
        require(
            cell["sheet"] == "Evidence"
            and cell["value_excerpt"] == sheet[cell["cell"]].value
            and cell["kind"] == "string",
            "Read cell differs from native workbook",
        )
        seen_cells.add(cell["cell"])
    require(
        seen_cells == {"A1", "B1", "C1", "A2", "B2", "C2"},
        "Literal cell read coverage differs",
    )
    require(
        final["transcription"] == expected["case"]["rows"][0],
        "Reported transcription differs",
    )
    return target, book, ledger, region


def validate_wikis(workspace, trace, target, book, ledger, region):
    seen = set()
    for export in trace["exports"]:
        root = Path(export["output_dir"])
        require(root.is_relative_to(workspace / "native-wiki"), "Unexpected Wiki path")
        manifest = read_json(root / "manifest.json")
        key = (manifest["asset_id"], manifest["revision"])
        seen.add(key)
        require(
            (root / manifest["source_attachment"]).read_bytes()
            == source_bytes(workspace, *key),
            "Wiki source differs",
        )
        for name, metadata in manifest["files"].items():
            require(Path(name).name == name, "Unsafe Wiki filename")
            data = (root / name).read_bytes()
            require(
                digest(data) == metadata["sha256"]
                and len(data) == metadata["size_bytes"],
                "Wiki artifact changed",
            )
        records = [
            json.loads(line)
            for line in (root / "records.jsonl").read_text().splitlines()
        ]
        if key[0] == target["asset_id"]:
            require(
                manifest["projection"] == "image-frames-v1"
                and manifest["frame_count"] == 2,
                "Wrong image Wiki projection",
            )
            for record in records:
                require(
                    canonical(record["evidence"]) in trace["records"],
                    "Wiki frame evidence differs",
                )
                png = (root / record["preview_attachment"]).read_bytes()
                require(
                    digest(png) == record["preview_sha256"], "Wiki preview hash differs"
                )
                validate_png(
                    source_bytes(workspace, *key),
                    png,
                    record["locator"]["frame_index"],
                    768,
                )
            for item in manifest["image_inputs"]["attachments"].values():
                require(
                    (root / item["attachment"]).read_bytes()
                    == source_bytes(workspace, item["asset_id"], item["revision"]),
                    "Wiki operation input differs",
                )
            receipt = read_json(root / "image-operation.json")
            expected_receipt = next(
                h["result"] for h in target["history"] if h["sha256"] == key[1]
            )
            require(receipt == expected_receipt, "Wiki receipt differs")
        else:
            require(
                key == (book["asset_id"], book["revision"]), "Unexpected Wiki target"
            )
            extra = manifest["derivations"]
            require(
                extra["active_ids_for_revision"]
                == [ledger["events"][0]["derivation_id"]],
                "Wiki derivation missing",
            )
            require(len(extra["image_records"]) == 1, "Wiki region evidence missing")
            item = next(iter(extra["image_records"].values()))
            record = read_json(root / item["record_file"])
            require(record["evidence"] == region, "Wiki source region differs")
            require(
                (root / record["source_attachment"]).read_bytes()
                == (workspace / "source.png").read_bytes(),
                "Wiki original image missing",
            )
            validate_png(
                (workspace / "source.png").read_bytes(),
                (root / record["preview_attachment"]).read_bytes(),
                0,
                768,
                region["selector"]["rect"],
            )
            require(
                record["citation_presentation"]["inline"].startswith("Review ["),
                "Custom region citation missing",
            )
        if key[1] in {target["revision"], book["revision"]}:
            require(
                manifest["citation_contract"]["name"] == "image-review",
                "Custom citation contract missing",
            )
            require(
                all(
                    r["citation_presentation"]["inline"].startswith("Review [")
                    for r in records
                ),
                "Custom citation display missing",
            )
    require(
        seen
        == {
            (target["asset_id"], target["history"][0]["sha256"]),
            (target["asset_id"], target["revision"]),
            (book["asset_id"], book["revision"]),
        },
        "Missing original/final/workbook Wiki",
    )
    return len(seen)


def check_claimed_review(final, trace, calls):
    review = final["visual_review"]
    require(
        review["scope"] == "static_rgba8_frame_and_region_previews"
        and review["findings"]
        and final["limitations"],
        "Missing visual review scope/findings/limits",
    )
    require(
        {
            (r["asset_id"], r["revision"], r["frame_index"])
            for r in review["reviewed_frames"]
        }
        == trace["images"],
        "Reported review differs from actual PNG delivery",
    )
    require(
        any(
            c["arguments"]["native_request"]["op"] == "verify_derivation"
            and payload(c).get("references_valid")
            and payload(c).get("active")
            for c in calls
        ),
        "Derivation not verified",
    )
