"""Independent exact values, complete raster crops and region ledger/Wiki audit."""

import argparse
import base64
import io
import json

import pymupdf
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageStat

from tests.codex_native_pdf.artifacts import load_asset, read_json
from tests.codex_native_pdf.trace import canonical, complete_records, digest, payload
from tests.codex_pdf.trace import call_failed, require, tool_errors
from tests.codex_pptx_tables.derivations import complete_page


def region_image(workspace, reference, result, size, png):
    path = (
        workspace
        / "data"
        / "native-assets"
        / reference["asset_id"]
        / "revisions"
        / reference["revision"]
    )
    require(
        digest(path.read_bytes()) == reference["revision"], "Wrong region source bytes"
    )
    with pymupdf.open(path) as pdf:
        page = pdf[reference["parent"]["locator"]["page_index"]]
        x0, y0, x1, y1 = reference["selector"]["rect"]
        clip = pymupdf.Rect(
            x0 * page.rect.width,
            y0 * page.rect.height,
            x1 * page.rect.width,
            y1 * page.rect.height,
        )
        clip.transform(pymupdf.Identity)
        scale = size / max(clip.width, clip.height)
        require(
            max(page.rect.width, page.rect.height) * scale < 16000,
            "Audit raster budget exceeded",
        )
        direct = page.get_pixmap(
            matrix=pymupdf.Matrix(scale, scale),
            clip=clip,
            colorspace=pymupdf.csRGB,
            alpha=False,
            annots=True,
        )
        delivered = Image.open(io.BytesIO(png)).convert("RGB")
        require(
            delivered.size == (direct.width, direct.height)
            and delivered.tobytes() == direct.samples,
            "Delivered region differs from independent direct source rendering",
        )
        # Full-page render, then independent PIL pixel crop. No production region code.
        pix = page.get_pixmap(
            matrix=pymupdf.Matrix(scale, scale), alpha=False, annots=True
        )
        bounds = (clip * pymupdf.Matrix(scale, scale)).irect
        expected = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).crop(
            tuple(bounds)
        )
    actual = Image.open(io.BytesIO(png)).convert("RGB")
    require(actual.size == expected.size, "Region crop size mismatch")
    difference = ImageChops.difference(actual, expected)
    mean = max(ImageStat.Stat(difference).mean)
    require(
        mean < 1,
        f"Region differs materially from full-page raster crop: mean={mean}",
    )
    require(digest(png) == result["image_sha256"], "Region PNG digest mismatch")


def audit(output):
    expected, final = (
        read_json(output / "expected.json"),
        read_json(output / "last-message.txt"),
    )
    workspace = output / "workspace"
    events = [
        json.loads(line) for line in (output / "events.jsonl").read_text().splitlines()
    ]
    require(any(e["type"] == "turn.completed" for e in events), "Incomplete model turn")
    calls = []
    for event in events:
        if event["type"] != "item.completed":
            continue
        item = event["item"]
        require(
            item["type"] in {"agent_message", "reasoning", "plan", "mcp_tool_call"},
            "Non-MCP action",
        )
        if item["type"] == "mcp_tool_call":
            require(
                item["server"] == "asset_aware_under_test"
                and item["tool"] == "document"
                and item["arguments"]["op"] == "native",
                "Unexpected tool",
            )
            if not call_failed(item):
                calls.append(item)
    ops = {c["arguments"]["native_request"]["op"] for c in calls}
    require(
        {
            "read_pdf_region",
            "create",
            "update",
            "update_pdf",
            "verify",
            "record_derivation",
            "verify_derivation",
            "export_wiki",
            "publish",
            "history",
        }
        <= ops
        and "writeback" not in ops,
        "Incomplete/unsafe workflow",
    )
    source, target = (
        load_asset(workspace, final[k]) for k in ("source_asset_id", "table_asset_id")
    )
    source_path = workspace / "source.pdf"
    require(
        digest(source_path.read_bytes()) == expected["source_sha256"]
        and source_path.stat().st_mtime_ns == expected["source_mtime_ns"],
        "Human PDF changed",
    )
    require(
        len(source["history"]) == len(target["history"]) == 2,
        "Wrong source/target history",
    )
    with pymupdf.open(source_path) as pdf:
        require(all(not p.get_text() for p in pdf), "Source contains hidden text")
    pages = complete_records(calls)
    require(
        any(
            r["evidence"]["revision"] == source["revision"]
            and r["locator"]["page_index"] == 0
            for r in pages.values()
        ),
        "Current rotated page was not read completely",
    )
    with pymupdf.open(
        workspace
        / "data"
        / "native-assets"
        / source["asset_id"]
        / "revisions"
        / source["revision"]
    ) as pdf:
        require(pdf[0].rotation == 90 and len(pdf) == 3, "Wrong managed PDF correction")
    target_dir = workspace / "data" / "native-assets" / target["asset_id"]
    original_revision = target["history"][0]["sha256"]
    truth = {
        "A1": "Count",
        "B1": "Reading",
        "C1": "Unit",
        "A2": "007",
        "B2": "-0.50",
        "C2": "mg/L",
    }
    for revision in (original_revision, target["revision"]):
        book = load_workbook(
            io.BytesIO((target_dir / "revisions" / revision).read_bytes())
        )
        require(book.sheetnames == ["Sheet1"], "Unexpected workbook structure")
        for address, value in truth.items():
            cell = book.active[address]
            require(
                cell.value
                == (
                    "008"
                    if address == "A2" and revision == target["revision"]
                    else value
                )
                and cell.data_type == "s",
                "Transcription/type mismatch",
            )
        book.close()
    require(
        digest((workspace / "verified.xlsx").read_bytes()) == target["revision"],
        "Publication differs",
    )
    regions, image_count, sizes, ledger_buffers, complete_ledgers, cells = (
        {},
        0,
        {},
        {},
        set(),
        {},
    )
    ledger = read_json(target_dir / "derivations.json")
    events_by_id = {e["derivation_id"]: e for e in ledger["events"]}
    require(len(events_by_id) == 3, "Expected three assertions")
    observed_events = []
    current_ledger = {**ledger, "events": observed_events}
    for call in calls:
        args, result = call["arguments"]["native_request"], payload(call)
        op = args["op"]
        if op == "read_cell":
            cell = result["cell"]
            require(cell["next_text_offset"] is None, "Incomplete cell readback")
            cells[canonical(cell["evidence"])] = cell
        elif op == "read_pdf_region":
            region = result["region"]
            ref = region["evidence"]
            require(
                ref["parent"] == region["parent"]
                and ref["selector"] == region["selector"],
                "Region reference mismatch",
            )
            require(canonical(ref["parent"]) in pages, "Missing complete parent read")
            require(
                digest(canonical({k: v for k, v in region.items() if k != "evidence"}))
                == ref["value_sha256"],
                "Region representation hash mismatch",
            )
            images = [c for c in call["result"]["content"] if c["type"] == "image"]
            require(len(images) == 1, "No actual region PNG")
            png = base64.b64decode(images[0]["data"], validate=True)
            region_image(workspace, ref, result, args.get("render_size", 1024), png)
            key = canonical(ref)
            regions[key] = region
            sizes.setdefault(key, set()).add(args.get("render_size", 1024))
            image_count += 1
        elif op == "read_derivations":
            sha = digest(canonical(current_ledger))
            require(result["derivations_sha256"] == sha, "Ledger trace mismatch")
            if complete_page(
                result, sha, ledger_buffers, digest_field="derivations_sha256"
            ):
                complete_ledgers.add(sha)
        elif op == "record_derivation":
            sha = digest(canonical(current_ledger))
            require(
                sha in complete_ledgers and args["expected_derivations_sha256"] == sha,
                "Append without full pinned ledger",
            )
            event = events_by_id[result["derivation_id"]]
            require(
                canonical(event["derivation"]["target"]) in cells, "Target not read"
            )
            require(
                canonical(event["derivation"]["sources"][0]) in regions,
                "Source region not viewed",
            )
            observed_events.append(event)
    require(
        len(regions) == 3 and any({256, 384} <= v for v in sizes.values()),
        "Wrong region coverage/resolution workflow",
    )
    require(len(cells) == 7, "Missing original or updated complete cell reads")
    derivation_proofs = [
        payload(c)
        for c in calls
        if c["arguments"]["native_request"]["op"] == "verify_derivation"
    ]
    require(
        {
            p["derivation_id"]
            for p in derivation_proofs
            if p["active"] and p["references_valid"]
        }
        == set(events_by_id),
        "Missing valid derivation proofs",
    )
    require(
        digest(canonical(ledger)) in complete_ledgers and not ledger_buffers,
        "Incomplete final ledger",
    )
    # Full glyph coverage plus distinct source cells, independent of the model's chosen boxes.
    glyphs = {
        "A2": (153, 202, 181, 216),
        "B2": (233, 202, 266, 216),
        "C2": (363, 202, 393, 216),
    }
    for event in ledger["events"]:
        claim = event["derivation"]
        require(len(claim["sources"]) == 1, "Ambiguous source")
        ref = claim["sources"][0]
        cell = cells[canonical(claim["target"])]["cell"]
        require(
            cell in glyphs
            and claim["target"]["revision"] == original_revision
            and ref["revision"] == expected["source_sha256"],
            "Migrated assertion",
        )
        x0, y0, x1, y1 = ref["selector"]["rect"]
        gx0, gy0, gx1, gy1 = glyphs[cell]
        require(
            x0 * 612 <= gx0 and y0 * 792 <= gy0 and x1 * 612 >= gx1 and y1 * 792 >= gy1,
            "Region clips expected glyphs",
        )
        require(
            x1 * 612 - x0 * 612 < 145 and y1 * 792 - y0 * 792 < 70,
            "Region is not a distinct cell",
        )
    proofs = [
        payload(c) for c in calls if c["arguments"]["native_request"]["op"] == "verify"
    ]
    require(
        sum(p["valid"] and not p["is_current_managed_revision"] for p in proofs) >= 2,
        "No historical proofs",
    )
    manifests = list((workspace / "native-wiki").rglob("manifest.json"))
    require(len(manifests) == 2, "Missing wiki snapshots")
    for path in manifests:
        manifest = read_json(path)
        info = manifest["derivations"]
        current = manifest["revision"] == target["revision"]
        require(
            len(info["active_ids_for_revision"]) == (0 if current else 3),
            "Assertions inherited/lost",
        )
        if current:
            require("region_records" not in info, "Inherited region records")
            continue
        require(len(info["region_records"]) == 3, "Missing Wiki region records")
        for entry in info["region_records"].values():
            stored = read_json(path.parent / entry["record_file"])
            require(
                all(
                    stored[k] == v
                    for k, v in regions[canonical(entry["reference"])].items()
                ),
                "Wiki source record changed",
            )
            require(
                digest((path.parent / stored["source_attachment"]).read_bytes())
                == expected["source_sha256"],
                "Wiki source attachment changed",
            )
            png = (path.parent / entry["preview_attachment"]).read_bytes()
            region_image(
                workspace,
                entry["reference"],
                {"image_sha256": entry["preview_sha256"]},
                768,
                png,
            )
            require(
                "displayed CropBox fractions"
                in stored["citation_presentation"]["inline"],
                "Citation lost region locator",
            )
    return {
        "passed": True,
        "successful_native_calls": len(calls),
        "actual_region_images": image_count,
        "distinct_regions": len(regions),
        "derivations": 3,
        "errors": tool_errors(events),
        "limitations": [
            "Synthetic scan; no arbitrary-PDF or Excel fidelity claim.",
            "Semantic review fields are caller assertions; independent audit checks this fixture's known values and glyph coverage.",
        ],
    }


def write_audit(output):
    try:
        report = audit(output)
    except Exception as exc:
        report = {"passed": False, "error": str(exc)}
    (output / "audit.json").write_text(json.dumps(report, indent=2) + "\n")
    return report


if __name__ == "__main__":
    from pathlib import Path

    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    result = write_audit(parser.parse_args().output.resolve())
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result["passed"] else 1)
