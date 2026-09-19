"""Resolve static named source rectangles without evaluating worksheet formulas."""

from __future__ import annotations

from dataclasses import dataclass
from itertools import pairwise
from typing import TYPE_CHECKING

from openpyxl.formula.tokenizer import Token

from src.infrastructure.native_grid_formulas import _names
from src.infrastructure.native_grid_selectors import selectors
from src.infrastructure.native_grid_xml import Rectangle
from src.infrastructure.native_spreadsheet_reader import NS, _decode_text
from src.infrastructure.native_workbook_formulas import _qualifier_spans, formula_tokens

if TYPE_CHECKING:
    from lxml import etree

    from src.infrastructure.native_grid_table_state import GridTableState
    from src.infrastructure.native_workbook_plan import WorkbookPlan


@dataclass(frozen=True)
class GridSourceRange:
    sheet: str
    bounds: Rectangle


def table_bounds(
    state: GridTableState, selector: str, *, source_name: bool
) -> Rectangle:
    bounds = state.bounds
    values = selectors(selector) if selector else []
    items = {
        selector[value.start : value.end].casefold()
        for value in values
        if value.name is None
    }
    outside = list(selector)
    for value in values:
        outside[value.start : value.end] = " " * (value.end - value.start)
    if "@" in outside or "#this row" in items:
        raise ValueError("Current-row source selectors require an evaluation context")
    rows = []
    if "#all" in items or (source_name and not selector):
        rows = [(bounds.first_row, bounds.last_row)]
    else:
        if "#headers" in items and state.headers:
            rows.append((bounds.first_row, bounds.first_row))
        if "#data" in items or not items:
            rows.append(
                (bounds.first_row + state.headers, bounds.last_row - state.totals)
            )
        if "#totals" in items and state.totals:
            rows.append((bounds.last_row, bounds.last_row))
    if not rows or any(first > last for first, last in rows):
        raise ValueError("Named table source has no surviving rows")
    ordered = sorted(rows)
    if any(left[1] + 1 < right[0] for left, right in pairwise(ordered)):
        raise ValueError(
            "Disjoint table row selectors require a multi-range source workflow"
        )
    columns = [value for value in values if value.name is not None]
    lookup = {
        column.name.casefold(): index
        for index, column in enumerate(state.columns, bounds.first_column)
    }
    if len(columns) > 2 or (
        len(columns) == 2 and ":" not in selector[columns[0].end : columns[1].start]
    ):
        raise ValueError("Disjoint table columns require a multi-range source workflow")
    try:
        positions = [
            lookup[value.name.casefold()] for value in columns if value.name is not None
        ]
    except KeyError as exc:
        raise ValueError("Named table source has an unknown column") from exc
    first, last = (
        (min(positions), max(positions))
        if positions
        else (bounds.first_column, bounds.last_column)
    )
    return Rectangle(ordered[0][0], first, max(last_row for _, last_row in rows), last)


class GridSourceResolver:
    def __init__(self, plan: WorkbookPlan, states: list[GridTableState]):
        self.sheets = {
            entry["name"].casefold(): entry["name"] for entry in plan.book.entries
        }
        self.tables: dict[str, GridTableState] = {}
        self.table_owners: dict[str, str] = {}
        for state in states:
            if state.deleted:
                continue
            owner = next(
                entry["name"]
                for entry in plan.book.entries
                if entry["key"]["part"] == state.worksheet
            )
            for name in {state.name, state.root.get("name", "")} - {""}:
                if (
                    name.casefold() in self.tables
                    and self.tables[name.casefold()] is not state
                ):
                    raise ValueError("Ambiguous table source name")
                self.tables[name.casefold()] = state
                self.table_owners[name.casefold()] = owner
        self.names: dict[tuple[str | None, str], etree._Element] = {}
        total = 0
        for node in plan.book.workbook.findall("s:definedNames/s:definedName", NS):
            index = node.get("localSheetId")
            if index is not None and not 0 <= int(index) < len(plan.book.entries):
                raise ValueError("Invalid named-source worksheet scope")
            owner = (
                plan.book.entries[int(index)]["name"].casefold()
                if index is not None
                else None
            )
            key = owner, _decode_text(node.get("name", "")).casefold()
            total += len((node.text or "").encode("utf-8"))
            if (
                not key[1]
                or key in self.names
                or len(self.names) >= 4096
                or total > 8 * 1024 * 1024
            ):
                raise ValueError(
                    "Ambiguous defined names or source resolution budget exceeded"
                )
            self.names[key] = node

    def resolve(
        self,
        expression: str,
        owner: str | None = None,
        *,
        source_name: bool = True,
        stack: tuple[tuple[str | None, str], ...] = (),
    ) -> GridSourceRange | None:
        text = expression.strip().removeprefix("=").strip()
        tokens = [token for token in formula_tokens(text) if token[3] != Token.WSPACE]
        while len(tokens) >= 3 and tokens[0][2] == "(" and tokens[-1][2] == ")":
            tokens = tokens[1:-1]
        if len(tokens) != 1 or tokens[0][3:] != (Token.OPERAND, Token.RANGE):
            raise ValueError(
                "Named source requires formula evaluation or explicit static range resolution"
            )
        operand = tokens[0][2]
        spans = _qualifier_spans(operand)
        if len(spans) > 1:
            raise ValueError("Named source has ambiguous worksheet qualifiers")
        if spans:
            start, end = spans[0]
            names = _names(operand[start:end])
            if names is None:
                return None  # External workbook coordinates are not transformed.
            if len(names) != 1:
                raise ValueError(
                    "Three-dimensional sources require a dedicated pivot workflow"
                )
            owner, operand = names[0], operand[end + 1 :]
        scope = owner.casefold() if owner is not None else None
        if scope is not None and scope not in self.sheets:
            raise ValueError("Named source worksheet is absent from the workbook")
        key = (scope, operand.casefold())
        if key not in self.names:
            key = None, operand.casefold()
        if key in self.names:
            if key in stack or len(stack) >= 32:
                raise ValueError(
                    "Cyclic named source or name-chain resolution budget exceeded"
                )
            return self.resolve(
                self.names[key].text or "",
                key[0],
                source_name=False,
                stack=(*stack, key),
            )
        bracket = operand.find("[")
        table_name = operand[:bracket] if bracket >= 0 else operand
        table = self.tables.get(table_name.casefold())
        if table is not None:
            sheet = self.table_owners[table_name.casefold()]
            if scope is not None and scope != sheet.casefold():
                raise ValueError(
                    "Named table source has conflicting worksheet identity"
                )
            return GridSourceRange(
                sheet,
                table_bounds(
                    table,
                    operand[bracket:] if bracket >= 0 else "",
                    source_name=source_name,
                ),
            )
        if owner is None:
            raise ValueError(
                "Named source has no exact worksheet or defined-name identity"
            )
        try:
            bounds = Rectangle.parse(operand)
        except ValueError as exc:
            raise ValueError(
                "Named source requires an explicit static rectangular range"
            ) from exc
        assert scope is not None
        return GridSourceRange(self.sheets[scope], bounds)
