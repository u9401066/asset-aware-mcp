"""Generate only new standalone textbox XML using public python-pptx APIs."""

from __future__ import annotations

from typing import TYPE_CHECKING

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

if TYPE_CHECKING:
    from src.domain.native_pptx import NativePptxShapeCreate


def build_textboxes(additions: list[NativePptxShapeCreate]) -> list[etree._Element]:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    result = []
    for item in additions:
        box = item.textbox
        shape = slide.shapes.add_textbox(
            Emu(box.left), Emu(box.top), Emu(box.width), Emu(box.height)
        )
        for index, runs in enumerate(box.paragraphs):
            paragraph = (
                shape.text_frame.paragraphs[0]
                if index == 0
                else shape.text_frame.add_paragraph()
            )
            for source in runs:
                run = paragraph.add_run()
                run.text = source.text
                run.font.bold, run.font.italic = source.bold, source.italic
                run.font.size = Pt(source.font_size_pt)
        # Detach python-pptx's XML subclass (its xpath signature differs from lxml).
        result.append(
            etree.fromstring(
                etree.tostring(shape.element),
                etree.XMLParser(
                    resolve_entities=False, no_network=True, huge_tree=False
                ),
            )
        )
    return result
