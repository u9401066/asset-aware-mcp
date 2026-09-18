"""Create presentations and edit only checked existing DrawingML text runs."""

from __future__ import annotations

import hashlib
import io
from typing import TYPE_CHECKING, Any

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

from src.domain.native_assets import NativeEditResult
from src.infrastructure.native_ooxml import xml_bytes
from src.infrastructure.native_pptx_package import NS, NativePptxPackage
from src.infrastructure.native_pptx_pictures import (
    add_pictures,
    read_picture,
    replace_pictures,
)
from src.infrastructure.native_pptx_records import shape_record, text_body
from src.infrastructure.native_pptx_shape_edit import add_shapes, delete_shapes

if TYPE_CHECKING:
    from collections.abc import Iterator

    from src.domain.native_pptx import (
        NativePptxShapeLocator,
        NativePptxTextEdit,
        NativePresentationCreate,
    )


def _run_text(package: NativePptxPackage, edit: NativePptxTextEdit) -> etree._Element:
    locator = edit.locator
    shape = package.locate(locator)
    body = text_body(shape, locator.row, locator.column)
    paragraphs = body.findall("a:p", NS)
    if locator.paragraph >= len(paragraphs):
        raise ValueError("Presentation paragraph is outside the existing structure")
    runs = paragraphs[locator.paragraph].findall("a:r", NS)
    if locator.run >= len(runs):
        raise ValueError("Presentation run is outside the existing structure")
    texts = runs[locator.run].findall("a:t", NS)
    if len(texts) != 1 or len(texts[0]):
        raise ValueError("Presentation run must contain exactly one plain text node")
    return texts[0]


def _plan(
    package: NativePptxPackage, edits: list[NativePptxTextEdit]
) -> list[tuple[NativePptxTextEdit, etree._Element, str | None]]:
    if not edits or len(edits) > 1000:
        raise ValueError("Presentation updates require 1 to 1,000 text edits")
    targets = set()
    result = []
    for edit in edits:
        text = _run_text(package, edit)
        target = (edit.locator.part, text.getroottree().getpath(text))
        if target in targets:
            raise ValueError("Duplicate presentation run edit")
        targets.add(target)
        digest = hashlib.sha256((text.text or "").encode("utf-8")).hexdigest()
        if digest != edit.expected_text_sha256:
            raise ValueError("Stale presentation run text; read the component again")
        result.append((edit, text, text.text))
    return result


def _verify_parts(
    before: NativePptxPackage,
    after: NativePptxPackage,
    plan: list[tuple[NativePptxTextEdit, etree._Element, str | None]],
    changed: list[str],
) -> None:
    for edit, _, original in plan:
        actual = _run_text(after, edit)
        if (actual.text or "") != edit.text:
            raise ValueError("Presentation run read-back failed")
        actual.text = original
    for part in changed:
        original = before.xml(part)
        if etree.tostring(original, method="c14n") != etree.tostring(
            after.roots[part], method="c14n"
        ):
            raise ValueError("Presentation changed outside the requested text nodes")


class NativePresentation:
    add_pictures = staticmethod(add_pictures)
    replace_pictures = staticmethod(replace_pictures)
    read_picture = staticmethod(read_picture)
    add_shapes = staticmethod(add_shapes)
    delete_shapes = staticmethod(delete_shapes)

    def package_parts(self, data: bytes) -> dict[str, bytes]:
        return NativePptxPackage(data).parts

    def inspect(self, data: bytes) -> dict[str, Any]:
        package = NativePptxPackage(data)
        shapes = list(package.shapes())
        size = package.presentation.find("p:sldSz", NS)
        return {
            "representation": "pptx_package",
            "main_part": package.main_part,
            "slide_count": len(package.slides),
            "shape_count": len(shapes),
            "slide_size_emu": dict(size.attrib) if size is not None else None,
            "slides": package.slides,
            "part_count": len(package.parts),
            "coverage": "Direct slide/notes shape trees; unknown XML and package parts retained. No inherited or rendered interpretation.",
        }

    def iter_shapes(
        self, data: bytes, *, offset: int = 0, limit: int | None = None
    ) -> Iterator[dict[str, Any]]:
        package = NativePptxPackage(data)
        for index, (locator, shape, parents) in enumerate(package.shapes()):
            if index < offset:
                continue
            if limit is not None and index >= offset + limit:
                break
            yield shape_record(locator, shape, parents)

    def read_shape(
        self, data: bytes, locator: NativePptxShapeLocator
    ) -> dict[str, Any]:
        package = NativePptxPackage(data)
        for candidate, shape, parents in package.shapes():
            if candidate == locator.model_dump():
                return shape_record(candidate, shape, parents)
        raise ValueError("Presentation shape locator does not resolve in this revision")

    def create(self, request: NativePresentationCreate) -> bytes:
        presentation = Presentation()
        presentation.slide_width = Emu(request.width)
        presentation.slide_height = Emu(request.height)
        for item in request.slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            for box in item.textboxes:
                shape = slide.shapes.add_textbox(
                    Emu(box.left), Emu(box.top), Emu(box.width), Emu(box.height)
                )
                for index, runs in enumerate(box.paragraphs):
                    paragraph = (
                        shape.text_frame.paragraphs[0]
                        if index == 0
                        else shape.text_frame.add_paragraph()
                    )
                    for source in runs:
                        run = paragraph.add_run()
                        run.text = source.text
                        run.font.bold, run.font.italic = source.bold, source.italic
                        run.font.size = Pt(source.font_size_pt)
            if item.notes:
                slide.notes_slide.notes_text_frame.text = item.notes
        output = io.BytesIO()
        presentation.save(output)
        data = output.getvalue()
        self.inspect(data)
        return data

    def edit(
        self, data: bytes, edits: list[NativePptxTextEdit]
    ) -> tuple[bytes, NativeEditResult]:
        package = NativePptxPackage(data)
        package.check_editable()
        plan = _plan(package, edits)
        changed = sorted(
            {edit.locator.part for edit, _, old in plan if (old or "") != edit.text}
        )
        for edit, text, _ in plan:
            text.text = edit.text
        replacements = {part: xml_bytes(package.roots[part]) for part in changed}
        updated = package.replace(replacements)
        checked = NativePptxPackage(updated)
        _verify_parts(package, checked, plan, changed)
        return updated, NativeEditResult(
            changed_parts=changed,
            preserved_parts=len(package.parts) - len(changed),
            changes=[
                {
                    "locator": edit.locator.model_dump(),
                    "before_text_sha256": edit.expected_text_sha256,
                    "after_text_sha256": hashlib.sha256(
                        edit.text.encode("utf-8")
                    ).hexdigest(),
                }
                for edit, _, _ in plan
            ],
            checks=[
                "run_text_preconditions",
                "exact_package_inventory",
                "untouched_part_bytes",
                "run_text_read_back",
                "unchanged_xml_outside_text_nodes",
            ],
            review_required=[
                "semantic_accuracy",
                "rendered_layout",
                "text_overflow",
                "inherited_formatting",
            ],
        )
