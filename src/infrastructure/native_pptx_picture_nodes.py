"""Build isolated picture XML and resolve exact, unambiguous embedded rasters."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from lxml import etree
from pptx import Presentation

from src.infrastructure.native_ooxml import DOC_REL_NS, TYPE_NS
from src.infrastructure.native_pptx_package import (
    NS,
    P_NS,
    NativePptxPackage,
    shape_identity,
)
from src.infrastructure.native_raster_image import inspect_image

if TYPE_CHECKING:
    from src.domain.native_pptx import NativePptxShapeLocator
    from src.domain.native_pptx_picture import NativePptxPictureCreate


def picture_image(
    package: NativePptxPackage, locator: NativePptxShapeLocator
) -> tuple[etree._Element, etree._Element, str, bytes, dict[str, Any]]:
    node = package.locate(locator)
    blips = node.findall(".//a:blip", NS)
    blip = node.find("p:blipFill/a:blip", NS)
    if node.tag != f"{{{P_NS}}}pic" or blip is None or blips != [blip]:
        raise ValueError("Target must be one picture with one embedded raster")
    if (
        blip.get(f"{{{DOC_REL_NS}}}link") is not None
        or blip.find("a:extLst", NS) is not None
    ):
        raise ValueError(
            "Linked or alternate picture representations require a separate workflow"
        )
    identity = blip.get(f"{{{DOC_REL_NS}}}embed", "")
    kind, path = package.relationships(locator.part).get(identity, ("", ""))
    if kind != f"{DOC_REL_NS}/image" or not path or path not in package.parts:
        raise ValueError("Picture image relationship does not resolve internally")
    data = package.parts[path]
    metadata = inspect_image(data)
    if image_content_type(package, path) != metadata["media_type"]:
        raise ValueError("Picture image content type does not match its bytes")
    return node, blip, path, data, metadata


def image_content_type(package: NativePptxPackage, path: str) -> str:
    root = package.xml("[Content_Types].xml")
    matches = [
        e.get("ContentType", "")
        for e in root
        if e.tag == f"{{{TYPE_NS}}}Override" and e.get("PartName") == "/" + path
    ]
    if not matches:
        extension = path.rsplit(".", 1)[-1].lower()
        matches = [
            e.get("ContentType", "")
            for e in root
            if e.tag == f"{{{TYPE_NS}}}Default"
            and e.get("Extension", "").lower() == extension
        ]
    if len(matches) != 1:
        raise ValueError("Picture needs an unambiguous image content type")
    return str(matches[0])


def build_picture(
    item: NativePptxPictureCreate, data: bytes, metadata: dict[str, Any]
) -> etree._Element:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    left, top, width, height = item.left, item.top, item.width, item.height
    iw, ih = metadata["width_px"], metadata["height_px"]
    if item.fit == "contain":
        scale = min(width / iw, height / ih)
        width, height = max(1, round(iw * scale)), max(1, round(ih * scale))
        left += (item.width - width) // 2
        top += (item.height - height) // 2
    shape = slide.shapes.add_picture(io.BytesIO(data), left, top, width, height)
    if item.fit == "cover":
        scale = max(width / iw, height / ih)
        shape.crop_left = shape.crop_right = max(
            0, (iw * scale - width) / (2 * iw * scale)
        )
        shape.crop_top = shape.crop_bottom = max(
            0, (ih * scale - height) / (2 * ih * scale)
        )
    if shape.image.blob != data:
        raise ValueError("Picture builder changed the source image bytes")
    node = etree.fromstring(
        etree.tostring(shape.element),
        etree.XMLParser(resolve_entities=False, no_network=True),
    )
    props = shape_identity(node)
    props.set("name", item.name)
    props.set("descr", item.description)
    return node
