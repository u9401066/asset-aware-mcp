"""Build only explicitly requested Word blocks with python-docx."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from lxml import etree

from src.domain.native_docx_structure import NativeDocxParagraph, merged_cells

if TYPE_CHECKING:
    from docx.document import Document as DocumentObject
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    from src.domain.native_docx_structure import (
        NativeDocxBlock,
        NativeDocxCreate,
        NativeDocxTable,
    )

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
TABLE_ALIGN = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}
VERTICAL = {
    "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
    "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
}


def fill_paragraph(paragraph: Paragraph, request: NativeDocxParagraph) -> None:
    if request.alignment is not None:
        paragraph.alignment = ALIGN[request.alignment]
    fmt = paragraph.paragraph_format
    fmt.keep_with_next, fmt.page_break_before = (
        request.keep_with_next,
        request.page_break_before,
    )
    if request.space_before_twips is not None:
        fmt.space_before = Twips(request.space_before_twips)
    if request.space_after_twips is not None:
        fmt.space_after = Twips(request.space_after_twips)
    if request.outline_level is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(request.outline_level))
        paragraph._p.get_or_add_pPr().append(outline)
    for item in request.runs:
        run = paragraph.add_run(item.text)
        run.bold, run.italic, run.underline = item.bold, item.italic, item.underline
        run.font.cs_bold, run.font.cs_italic = item.bold, item.italic
        if item.font_size_pt is not None:
            run.font.size = Pt(item.font_size_pt)
            size = OxmlElement("w:szCs")
            size.set(qn("w:val"), str(int(item.font_size_pt * 2)))
            run._r.get_or_add_rPr().insert_element_before(
                size,
                "w:highlight",
                "w:u",
                "w:effect",
                "w:bdr",
                "w:shd",
                "w:fitText",
                "w:vertAlign",
                "w:rtl",
                "w:cs",
                "w:em",
                "w:lang",
            )
        if item.font_name is not None:
            run.font.name = item.font_name
            fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), item.font_name)
            fonts.set(qn("w:cs"), item.font_name)
        if item.color_rgb is not None:
            run.font.color.rgb = RGBColor.from_string(item.color_rgb.upper())


def fill_table(table: Table, request: NativeDocxTable) -> None:
    table.autofit = False
    table.alignment = TABLE_ALIGN[request.alignment]
    props = table._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    props.find(qn("w:tblW")).set(qn("w:w"), str(sum(request.column_widths_twips)))
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single" if request.borders else "nil")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "auto")
        borders.append(border)
    props.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblPrChange"
    )
    for column, width in zip(table.columns, request.column_widths_twips, strict=True):
        column.width = Twips(width)
    grid = [list(row.cells) for row in table.rows]
    for cell_row in grid:
        for cell, width in zip(cell_row, request.column_widths_twips, strict=True):
            cell.width = Twips(width)
    if request.row_heights_twips is not None:
        for row, height in zip(table.rows, request.row_heights_twips, strict=True):
            row.height, row.height_rule = Twips(height), WD_ROW_HEIGHT_RULE.AT_LEAST
    if request.repeat_header:
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
    for merge in request.merges:
        grid[merge.row][merge.column].merge(grid[merge.end_row][merge.end_column])
    covered = merged_cells(request)
    grid = [list(row.cells) for row in table.rows]
    for row, items in enumerate(request.cells):
        for column, item in enumerate(items):
            if covered.get((row, column), (row, column)) != (row, column):
                continue
            cell = grid[row][column]
            cell.text = ""
            cell.paragraphs[0].clear()
            cell.vertical_alignment = VERTICAL[item.vertical_alignment]
            if item.fill_rgb:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), item.fill_rgb.upper())
                cell._tc.get_or_add_tcPr().insert_element_before(
                    shade,
                    "w:noWrap",
                    "w:tcMar",
                    "w:textDirection",
                    "w:tcFitText",
                    "w:vAlign",
                    "w:hideMark",
                    "w:headers",
                    "w:tcPrChange",
                )
            for index, paragraph in enumerate(item.paragraphs):
                target = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
                fill_paragraph(target, paragraph)


def populate(document: DocumentObject, blocks: list[NativeDocxBlock]) -> None:
    for block in blocks:
        if isinstance(block, NativeDocxParagraph):
            fill_paragraph(document.add_paragraph(), block)
        else:
            table = document.add_table(
                rows=len(block.cells), cols=len(block.column_widths_twips)
            )
            fill_table(table, block)


def build_blocks(blocks: list[NativeDocxBlock]) -> list[etree._Element]:
    scratch = Document()
    populate(scratch, blocks)
    return [
        etree.fromstring(etree.tostring(node))
        for node in scratch.element.body
        if node.tag != qn("w:sectPr")
    ]


def build_document(request: NativeDocxCreate) -> bytes:
    document = Document()
    section = document.sections[0]
    section.orientation = (
        WD_ORIENT.LANDSCAPE
        if request.page_width_twips > request.page_height_twips
        else WD_ORIENT.PORTRAIT
    )
    section.page_width, section.page_height = (
        Twips(request.page_width_twips),
        Twips(request.page_height_twips),
    )
    for name in ("top", "bottom", "left", "right"):
        setattr(
            section, f"{name}_margin", Twips(getattr(request, f"margin_{name}_twips"))
        )
    document.core_properties.author = request.author
    document.core_properties.last_modified_by = request.author
    document.core_properties.title = request.name
    document.core_properties.comments = ""
    document.core_properties.created = document.core_properties.modified = datetime.now(
        timezone.utc
    )
    populate(document, request.blocks)
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()
