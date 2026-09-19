"""Independent serialized readback of new paragraphs, grids and direct formatting."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from src.domain.native_docx_structure import NativeDocxParagraph, merged_cells
from src.infrastructure.native_docx_builder import ALIGN, TABLE_ALIGN, VERTICAL

if TYPE_CHECKING:
    from src.domain.native_docx_structure import NativeDocxBlock, NativeDocxTable


def require(value: bool, message: str) -> None:
    if not value:
        raise ValueError("DOCX readback differs: " + message)


def check_paragraph(actual: Paragraph, expected: NativeDocxParagraph) -> None:
    require(len(actual.runs) == len(expected.runs), "run count")
    require(actual.text == "".join(run.text for run in expected.runs), "paragraph text")
    for run, item in zip(actual.runs, expected.runs, strict=True):
        require(run.text == item.text, "exact literal run text")
        require(
            (run.bold, run.italic, run.underline)
            == (item.bold, item.italic, item.underline),
            "run flags",
        )
        size = run.font.size.pt if run.font.size is not None else None
        require(size == item.font_size_pt and run.font.name == item.font_name, "font")
        require(
            (run.font.cs_bold, run.font.cs_italic) == (item.bold, item.italic),
            "complex script flags",
        )
        cs_size = run._r.find("./" + qn("w:rPr") + "/" + qn("w:szCs"))
        require(
            (float(cs_size.get(qn("w:val"))) / 2 if cs_size is not None else None)
            == item.font_size_pt,
            "complex script size",
        )
        color = str(run.font.color.rgb) if run.font.color.rgb is not None else None
        require(
            color == (item.color_rgb.upper() if item.color_rgb else None), "run color"
        )
        if item.font_name is not None:
            if run._r.rPr is None or run._r.rPr.rFonts is None:
                raise ValueError("DOCX readback differs: missing font properties")
            fonts = run._r.rPr.rFonts
            require(
                all(
                    fonts.get(qn("w:" + key)) == item.font_name
                    for key in ("ascii", "hAnsi", "eastAsia", "cs")
                ),
                "script fonts",
            )
    require(
        actual.alignment == (ALIGN[expected.alignment] if expected.alignment else None),
        "paragraph alignment",
    )
    fmt = actual.paragraph_format
    require(
        (fmt.keep_with_next, fmt.page_break_before)
        == (expected.keep_with_next, expected.page_break_before),
        "paragraph flow flags",
    )
    for name in ("before", "after"):
        actual_spacing = getattr(fmt, "space_" + name)
        require(
            (actual_spacing.twips if actual_spacing is not None else None)
            == getattr(expected, f"space_{name}_twips"),
            "paragraph spacing",
        )
    outline = actual._p.find("./" + qn("w:pPr") + "/" + qn("w:outlineLvl"))
    require(
        (int(outline.get(qn("w:val"))) if outline is not None else None)
        == expected.outline_level,
        "outline level",
    )


def check_table(actual: Table, expected: NativeDocxTable) -> None:
    widths = [column.width.twips for column in actual.columns]
    require(widths == expected.column_widths_twips, "column grid")
    require(len(actual.rows) == len(expected.cells), "row count")
    require(
        actual.autofit is False and actual.alignment == TABLE_ALIGN[expected.alignment],
        "table layout",
    )
    total = actual._tbl.tblPr.find(qn("w:tblW"))
    require(
        total.get(qn("w:type")) == "dxa" and int(total.get(qn("w:w"))) == sum(widths),
        "table width",
    )
    borders = actual._tbl.tblPr.find(qn("w:tblBorders"))
    require(borders is not None and len(borders) == 6, "table borders")
    for border in borders:
        require(
            border.get(qn("w:val")) == ("single" if expected.borders else "nil"),
            "border style",
        )
        require(
            border.get(qn("w:sz")) == "4" and border.get(qn("w:color")) == "auto",
            "border format",
        )
    covered = merged_cells(expected)
    grid = [list(row.cells) for row in actual.rows]
    anchors = {}
    for row, items in enumerate(expected.cells):
        height = actual.rows[row].height
        require(
            actual.rows[row].height_rule
            == (WD_ROW_HEIGHT_RULE.AT_LEAST if expected.row_heights_twips else None),
            "row height rule",
        )
        require(
            (height.twips if height is not None else None)
            == (
                expected.row_heights_twips[row] if expected.row_heights_twips else None
            ),
            "row height",
        )
        header = actual.rows[row]._tr.find(
            "./" + qn("w:trPr") + "/" + qn("w:tblHeader")
        )
        require(
            (header is not None) == (expected.repeat_header and row == 0),
            "repeating header",
        )
        for column, item in enumerate(items):
            cell = grid[row][column]
            origin = covered.get((row, column), (row, column))
            require(cell._tc is grid[origin[0]][origin[1]]._tc, "merge coverage")
            if origin != (row, column):
                continue
            require(cell._tc not in anchors, "unexpected cell merge")
            anchors[cell._tc] = origin
            span = next(
                (
                    m.end_column - m.column + 1
                    for m in expected.merges
                    if (m.row, m.column) == origin
                ),
                1,
            )
            require(
                cell.width is not None
                and cell.width.twips == sum(widths[column : column + span]),
                "cell width",
            )
            require(
                len(cell.paragraphs) == len(item.paragraphs), "cell paragraph count"
            )
            for paragraph, wanted in zip(cell.paragraphs, item.paragraphs, strict=True):
                check_paragraph(paragraph, wanted)
            require(
                cell.vertical_alignment == VERTICAL[item.vertical_alignment],
                "cell alignment",
            )
            shade = cell._tc.tcPr.find(qn("w:shd"))
            require(
                (shade.get(qn("w:fill")) if shade is not None else None)
                == (item.fill_rgb.upper() if item.fill_rgb else None),
                "cell shading",
            )


def check_blocks(nodes: list[etree._Element], blocks: list[NativeDocxBlock]) -> None:
    require(len(nodes) == len(blocks), "block count")
    parent = Document()
    for node, block in zip(nodes, blocks, strict=True):
        root = parse_xml(etree.tostring(node))
        if isinstance(block, NativeDocxParagraph):
            require(root.tag == qn("w:p"), "paragraph kind")
            check_paragraph(Paragraph(root, parent), block)
        else:
            require(root.tag == qn("w:tbl"), "table kind")
            check_table(Table(root, parent), block)
