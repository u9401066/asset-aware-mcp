"""Build new table XML in a scratch presentation using public python-pptx APIs."""

from __future__ import annotations

from typing import TYPE_CHECKING

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from src.domain.native_pptx_table import NativePptxTableCreate, merge_map
from src.infrastructure.native_pptx_package import NS, shape_identity

if TYPE_CHECKING:
    from pptx.table import Table, _Cell

    from src.domain.native_pptx_table import NativePptxTableCellCreate

ALIGNMENTS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHORS = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def fill_cell(cell: _Cell, request: NativePptxTableCellCreate) -> None:
    cell.margin_left = cell.margin_right = Emu(request.margin)
    cell.margin_top = cell.margin_bottom = Emu(request.margin)
    cell.vertical_anchor = ANCHORS[request.vertical_anchor]
    if request.fill_rgb:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(request.fill_rgb.upper())
    frame = cell.text_frame
    frame.clear()
    for index, runs in enumerate(request.paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = ALIGNMENTS[request.alignment]
        for item in runs:
            run = paragraph.add_run()
            run.text = item.text
            run.font.bold, run.font.italic = item.bold, item.italic
            run.font.size = Pt(item.font_size_pt)
            if request.text_rgb:
                run.font.color.rgb = RGBColor.from_string(request.text_rgb.upper())


def populate_table(table: Table, request: NativePptxTableCreate) -> None:
    for target, source in [
        ("first_row", "first_row"),
        ("last_row", "last_row"),
        ("first_col", "first_column"),
        ("last_col", "last_column"),
        ("horz_banding", "row_banding"),
        ("vert_banding", "column_banding"),
    ]:
        setattr(table, target, getattr(request, source))
    for merge in request.merges:
        table.cell(merge.row, merge.column).merge(
            table.cell(merge.end_row, merge.end_column)
        )
    covered = merge_map(request)
    for row, cells in enumerate(request.cells):
        for column, item in enumerate(cells):
            if covered.get((row, column), (row, column)) == (row, column):
                fill_cell(table.cell(row, column), item)


def build_table(
    request: NativePptxTableCreate, style_id: str | None = None
) -> etree._Element:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(
        len(request.row_heights),
        len(request.column_widths),
        request.left,
        request.top,
        sum(request.column_widths),
        sum(request.row_heights),
    )
    table = shape.table
    for index, width in enumerate(request.column_widths):
        table.columns[index].width = width
    for index, height in enumerate(request.row_heights):
        table.rows[index].height = height
    populate_table(table, request)
    node = etree.fromstring(
        etree.tostring(shape.element),
        etree.XMLParser(resolve_entities=False, no_network=True),
    )
    style = node.find("a:graphic/a:graphicData/a:tbl/a:tblPr/a:tableStyleId", NS)
    assert style is not None
    if style_id is None:
        style.getparent().remove(style)
    else:
        style.text = style_id
    props = shape_identity(node)
    props.set("name", request.name)
    props.set("descr", request.description)
    return node
