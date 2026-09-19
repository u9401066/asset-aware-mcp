"""Build only new slide XML using destination layouts and public python-pptx APIs."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

from src.infrastructure.native_pptx_layouts import placeholders
from src.infrastructure.native_pptx_package import NS

if TYPE_CHECKING:
    from src.domain.native_pptx_slides import NativePptxSlideInsert


def build_slides(
    data: bytes, request: NativePptxSlideInsert, layout_roots: dict[str, etree._Element]
) -> list[etree._Element]:
    presentation = Presentation(io.BytesIO(data))
    available = {
        str(layout.part.partname).lstrip("/"): layout
        for master in presentation.slide_masters
        for layout in master.slide_layouts
    }
    result = []
    for item in request.slides:
        if item.layout_part not in available or item.layout_part not in layout_roots:
            raise ValueError("Requested layout is not owned by this presentation")
        expected_count = len(placeholders(layout_roots[item.layout_part]))
        slide = presentation.slides.add_slide(available[item.layout_part])
        if len(slide.placeholders) != expected_count:
            raise ValueError(
                "Layout placeholder cloning differs from the supported plan"
            )
        for box in item.textboxes:
            shape = slide.shapes.add_textbox(
                Emu(box.left), Emu(box.top), Emu(box.width), Emu(box.height)
            )
            for index, runs in enumerate(box.paragraphs):
                paragraph = (
                    shape.text_frame.paragraphs[0]
                    if index == 0
                    else shape.text_frame.add_paragraph()
                )
                for source in runs:
                    run = paragraph.add_run()
                    run.text = source.text
                    run.font.bold, run.font.italic = source.bold, source.italic
                    run.font.size = Pt(source.font_size_pt)
        root = etree.fromstring(etree.tostring(slide.element))
        if any(
            k.startswith("{" + NS["r"] + "}")
            for node in root.iter()
            for k in node.attrib
        ):
            raise ValueError(
                "Generated slide unexpectedly requires additional relationships"
            )
        result.append(root)
    # Never save the loaded source presentation: only new slide nodes are returned.
    return result
